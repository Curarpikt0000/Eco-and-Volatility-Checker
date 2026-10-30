"""external_data.py — 联动读取 KOL 和 Economic Dashboard 的数据。

- KOL by_day DB：取最近 N 天 KOL 言论，识别「状态变化」(主导方向 vs 该 KOL 上次)
- Economic Dashboard DB：取 Fed/中日流动性关键点(reserves/RRP/TGA/DR007/USDJPY 等)

只读，不写。供每日 report + dashboard 底部板块用。
时区 JST。所有数字来自真实 Notion，绝不编造。
"""
import sys, os, datetime, json
sys.path.insert(0, os.path.dirname(__file__))
from notion_writer import _req

# ─────────── KOL ───────────
KOL_BY_DAY_DB = "32347eb5fd3c8087b9c0f409f95f664e"
KOL_REGISTRY = os.path.join(os.path.dirname(__file__), "..", "data", "kol_registry.json")
KOL_INDEPENDENT = os.path.join(os.path.dirname(__file__), "..", "data", "kol_independent.json")
# Eco 自己的每日 KOL 全量方向快照仓库(独立副本, 进 git, 供周度对比用, 不依赖任何 agent 的 Notion DB)
KOL_DAILY_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "kol", "daily")
KOL_BACKFILL_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "kol", "backfill")


def _tgs(v, limit=None):
    """targets 归一为字符串(上游可能是 str / list / None)。

    ★2026-08-24: 深度重抓时部分记录的 targets 被写成<b>数组</b>。
      对 list 做 [:150] 不报错但会静默切成子列表, 比直接崩更难发现;
      对 list 做 .strip() 则直接 AttributeError(线上 build 实崩)。
      故所有取用点统一走本函数。
    """
    if v is None:
        s = ""
    elif isinstance(v, (list, tuple)):
        s = "、".join(str(x).strip() for x in v if str(x).strip())
    else:
        s = str(v).strip()
    return s[:limit] if limit else s


def load_registry():
    """读 KOL 名册(full list)。返回 [{id,display_name,domain,sector,institution},...]。"""
    import json
    try:
        d = json.load(open(KOL_REGISTRY))
        return d.get("kols", [])
    except Exception:
        return []


def load_independent_kol():
    """读 Eco 自己独立抓的 KOL 数据(cron agent 每日 web_search 写入)。
    格式: {"date":"YYYY-MM-DD", "changes":[{kol,sector,prev_dir,new_dir,date,comments,targets},...],
           "all":[{kol,sector,direction,date,comments,targets},...]}
    返回 (changes_list, meta) 或 ([], {})。"""
    import json
    if os.path.exists(KOL_INDEPENDENT):
        try:
            d = json.load(open(KOL_INDEPENDENT))
            return d.get("changes", []), d
        except Exception:
            pass
    return [], {}


def _rt(prop):
    """rich_text / title → 纯文本。"""
    arr = prop.get("rich_text") or prop.get("title") or []
    return "".join(x.get("plain_text", x.get("text", {}).get("content", "")) for x in arr)


def _sel(prop):
    s = prop.get("select")
    return s.get("name") if s else ""


def fetch_kol_recent(days=10, limit=200):
    """取最近 N 天 KOL by_day 记录。返回 [dict,...] 按日期降序。"""
    since = (datetime.date.today() - datetime.timedelta(days=days)).strftime("%Y-%m-%d")
    out = []
    cursor = None
    while True:
        payload = {
            "filter": {"property": "Date", "date": {"on_or_after": since}},
            "sorts": [{"property": "Date", "direction": "descending"}],
            "page_size": 100,
        }
        if cursor:
            payload["start_cursor"] = cursor
        st, b = _req("POST", f"/databases/{KOL_BY_DAY_DB}/query", payload)
        if st != 200:
            break
        for r in b.get("results", []):
            p = r["properties"]
            out.append({
                "date": (p.get("Date", {}).get("date") or {}).get("start", "")[:10],
                "kol": _sel(p.get("Name of KOL", {})),
                "type": _sel(p.get("KOL or IB View", {})),
                "direction": _sel(p.get("主导方向", {})),
                "sector": _sel(p.get("Sector", {})),
                "detail_sector": _sel(p.get("Detail Sector", {})),
                "suggestion": _rt(p.get("Suggestion", {})),
                "comments": _rt(p.get("Comments", {})),
                "dir_detail": _rt(p.get("方向明细", {})),
                "targets": _rt(p.get("多空标的", {})),
            })
        if b.get("has_more"):
            cursor = b["next_cursor"]
        else:
            break
        if len(out) >= limit:
            break
    return out


def kol_stance_changes(days=10):
    """识别最近 N 天内 KOL 的「状态变化」：某 KOL 最新主导方向 != 其之前一条。
    返回 [{kol, sector, prev_dir, new_dir, date, comments, targets},...]。
    """
    recs = fetch_kol_recent(days=days)
    # 按 KOL 分组，按日期升序
    by_kol = {}
    for r in recs:
        if not r["kol"] or not r["direction"]:
            continue
        by_kol.setdefault(r["kol"], []).append(r)
    changes = []
    for kol, rs in by_kol.items():
        rs.sort(key=lambda x: x["date"])
        for i in range(1, len(rs)):
            if rs[i]["direction"] != rs[i - 1]["direction"]:
                changes.append({
                    "kol": kol,
                    "sector": rs[i]["sector"],
                    "prev_dir": rs[i - 1]["direction"],
                    "new_dir": rs[i]["direction"],
                    "date": rs[i]["date"],
                    "comments": rs[i]["comments"][:300],
                    "targets": rs[i]["targets"][:150],
                })
    # 只保留最近的变化(去重同 KOL 取最新)
    latest = {}
    for ch in sorted(changes, key=lambda x: x["date"]):
        latest[ch["kol"]] = ch
    return sorted(latest.values(), key=lambda x: x["date"], reverse=True)


def _days_since_last_monday():
    """今天回溯到『上周一』的天数(含今天区间)。
    上周一 = 本周一 - 7天。用于 KOL 变化的回看窗口。"""
    import datetime as _dt
    today = _dt.date.today()
    this_monday = today - _dt.timedelta(days=today.weekday())  # 本周一
    last_monday = this_monday - _dt.timedelta(days=7)
    return (today - last_monday).days + 1


# ═══════ Eco 自己的 KOL 每日快照仓库(独立副本, 供周度对比, 不依赖任何 agent) ═══════
def save_kol_daily_snapshot(date_str=None):
    """把当日 kol_independent.json 的全量方向(all)落盘为 data/kol/daily/YYYY-MM-DD.json。
    这是 Eco 自己的独立数据仓库副本(进 git), 供周度对比用, 不依赖任何 agent 的 Notion DB。
    幂等: 同日重复调用覆盖当日快照。返回写入路径或 None(无独立数据时跳过)。"""
    import json
    if not os.path.exists(KOL_INDEPENDENT):
        return None
    try:
        ind = json.load(open(KOL_INDEPENDENT))
    except Exception:
        return None
    all_dirs = ind.get("all") or []
    if not all_dirs:
        return None
    ds = date_str or ind.get("date") or datetime.date.today().strftime("%Y-%m-%d")
    os.makedirs(KOL_DAILY_DIR, exist_ok=True)
    snap = {
        "date": ds,
        "count": len(all_dirs),
        "source": "Eco independent daily crawl (web_search 全量 KOL)",
        "kols": [
            {
                "kol": x.get("kol", ""),
                "sector": x.get("sector", ""),
                "direction": x.get("direction", ""),
                "targets": _tgs(x.get("targets"), 150),
                "comments": (x.get("comments") or "")[:300],
                # ★2026-08-22(Chao): 快照必须带 detail/sources, 否则 dashboard 从快照
                #   渲染详情层时又会退回「只有一句话」。detail 是 100-300 字结构化摘要。
                "detail": (x.get("detail") or "")[:1200],
                "detail_status": x.get("detail_status") or "",
                "sources": (x.get("sources") or [])[:5],
            }
            for x in all_dirs if x.get("kol")
        ],
    }
    path = os.path.join(KOL_DAILY_DIR, f"{ds}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(snap, f, ensure_ascii=False, indent=1)
    return path


def load_kol_daily_snapshots():
    """读 data/kol/daily/ 下全部快照。返回 {YYYY-MM-DD: {kol: {direction,sector,targets,comments}}} 按日期。
    无快照返回 {}。"""
    import json
    import glob
    out = {}
    if not os.path.isdir(KOL_DAILY_DIR):
        return out
    for p in glob.glob(os.path.join(KOL_DAILY_DIR, "*.json")):
        try:
            d = json.load(open(p))
        except Exception:
            continue
        ds = d.get("date") or os.path.basename(p)[:10]
        by_kol = {}
        for x in d.get("kols", []):
            if x.get("kol"):
                by_kol[x["kol"]] = x
        if by_kol:
            out[ds] = by_kol
    return out


def kol_full_history():
    """★每个 KOL 的【全量历史观点】, 按时间倒序, 供 dashboard 两层展开钻取。

    ★2026-08-26 架构收敛(Chao 指示「统一成一个数据库」):
      唯一真相源已改为 data/kol_store.sqlite (src/kol_store.py)。
      三条历史链路 —— 每日快照 / 历史回填 / 新增 KOL 一次性回填 ——
      全部经 kol_store.upsert() 汇入同一张 opinion 表, 不再于渲染时临时合并。
      本函数优先读 store; store 不可用时自动退回下方旧的双路合并实现,
      保证 dashboard 永不因 store 故障而白屏。

    返回 {kol_name: [{first_date,last_date,direction,comments,targets,source,origin}, ...]}
    列表按 last_date 倒序(最新在前)。无数据返回 {}。
    """
    import json
    import glob

    # ── 主路径: 统一 SSOT ──
    try:
        import kol_store as _ks
        if os.path.exists(_ks.STORE_PATH):
            _h = _ks.full_history()
            if _h:
                return _h
    except Exception as _e:      # noqa: BLE001  故障降级, 不让 dashboard 挂掉
        print(f"[kol_full_history] store 不可用, 退回旧双路合并: {_e}")

    # ── 降级路径: 旧的 daily + backfill 双路内存合并 ──
    out = {}

    # ── 路 1: 每日快照 → 连续同内容合并成区间 ──
    snaps = load_kol_daily_snapshots()
    for ds in sorted(snaps.keys()):                       # 升序遍历便于合并连续段
        for kol, x in snaps[ds].items():
            cmt = (x.get("comments") or "").strip()
            if not cmt:
                continue
            direction = (x.get("direction") or "").strip()
            _tg = x.get("targets") or ""
            targets = ", ".join(_tg) if isinstance(_tg, list) else str(_tg).strip()
            lst = out.setdefault(kol, [])
            if lst and lst[-1]["direction"] == direction and lst[-1]["comments"] == cmt:
                lst[-1]["last_date"] = ds                 # 延长区间
                # 区间延长时用最新一天的 detail(若有), 保证展开层拿到的是最新版本
                if x.get("detail"):
                    lst[-1]["detail"] = x.get("detail")
                    lst[-1]["sources"] = x.get("sources") or lst[-1].get("sources") or []
            else:
                lst.append({"first_date": ds, "last_date": ds,
                            "direction": direction, "comments": cmt,
                            "targets": targets, "source": "", "origin": "snapshot",
                            # ★2026-08-23: 快照已带深度摘要, 必须透传给 dashboard 展开层
                            "detail": (x.get("detail") or "").strip(),
                            "sources": x.get("sources") or []})

    # ── 路 2: 历史回填(带原文链接) ──
    if os.path.isdir(KOL_BACKFILL_DIR):
        for p in glob.glob(os.path.join(KOL_BACKFILL_DIR, "*.json")):
            try:
                d = json.load(open(p))
            except Exception:
                continue
            kol = (d.get("kol") or "").strip()
            if not kol:
                continue
            for h in d.get("history", []):
                cmt = (h.get("comments") or "").strip()
                dt = (h.get("date") or "").strip()
                if not cmt or not dt:
                    continue
                out.setdefault(kol, []).append({
                    "first_date": dt, "last_date": dt,
                    "direction": (h.get("direction") or "").strip(),
                    "comments": cmt,
                    "targets": _tgs(h.get("targets")),
                    "source": (h.get("source") or "").strip(),
                    "origin": "backfill",
                    # ★2026-08-23 全量回填(369/369): 历史观点也有四段式深度摘要
                    "detail": (h.get("detail") or "").strip(),
                    "sources": h.get("sources") or [],
                })

    # ── 跨路去重(同日同言论可能两路都有) + 倒序 ──
    for kol, lst in out.items():
        seen = set()
        dedup = []
        # backfill 优先(带 source), 故先排 origin
        for r in sorted(lst, key=lambda r: (r["last_date"], r["origin"] != "backfill")):
            key = (r["last_date"], r["comments"][:80])
            if key in seen:
                continue
            seen.add(key)
            dedup.append(r)
        dedup.sort(key=lambda r: r["last_date"], reverse=True)
        out[kol] = dedup
    return out


def kol_weekly_changes(period="week"):
    """窗口内 KOL 方向转向, 全部基于 Eco 自己的每日快照(data/kol/daily/), 不碰任何 agent DB。
    比较逻辑: 对每个 KOL, 取『窗口内最新一天快照的方向』vs『窗口起点之前最后一天快照的方向』, 不同则记为转向。
    period: day=本日(今天 vs 前一有快照日) / week=本周(默认, 本周最新 vs 上周最后) / month=本月(最新 vs 上月最后)。
    若快照不足(尚未攒够跨窗口数据), 回退到"最早快照 vs 最新快照"对比。
    返回 [{kol, sector, prev_dir, new_dir, date, comments, targets},...] (供 grouped 分组)。"""
    snaps = load_kol_daily_snapshots()
    if len(snaps) < 2:
        return []
    dates = sorted(snaps.keys())
    win_start, win_end, _ = _kol_period_window(period)
    # 窗口内快照 / 窗口之前的快照(取最后一天作对比基准)
    in_win = [d for d in dates if win_start <= d <= win_end]
    before = [d for d in dates if d < win_start]
    if in_win and before:
        cur_date, prev_date = in_win[-1], before[-1]
    else:
        # 跨窗口数据不足 → 用最新 vs 最早(诚实回退, 攒够后自动切回)
        cur_date, prev_date = dates[-1], dates[0]
    cur, prev = snaps.get(cur_date, {}), snaps.get(prev_date, {})
    # 无效方向(空/未找到/数据滞后等非真实立场)不参与转向判定
    _invalid_dir = {"", "未找到", "无", "n/a", "未知", "数据滞后", "-", "—"}
    changes = []
    for kol, c in cur.items():
        p = prev.get(kol)
        if not p:
            continue  # 上周无记录 → 不算"转向"(新覆盖的KOL不误报)
        cd = (c.get("direction") or "").strip()
        pd = (p.get("direction") or "").strip()
        if cd in _invalid_dir or pd in _invalid_dir:
            continue  # 任一端无效 → 不是真转向(避免"看空→未找到"伪转向)
        if cd and pd and cd != pd:
            changes.append({
                "kol": kol,
                "sector": c.get("sector", ""),
                "prev_dir": pd,
                "new_dir": cd,
                "date": cur_date,
                "comments": (c.get("comments") or "")[:300],
                "targets": _tgs(c.get("targets"), 150),
            })
    return sorted(changes, key=lambda x: x["kol"])


def _kol_period_window(period="week"):
    """把 本日/本周/本月 换算成 (窗口起点, 窗口终点, 新观点基准日) 三个 YYYY-MM-DD。

    ★Chao 2026-08-31 要求三档 filter。口径:
      day   = 今天(无今日快照则回退最近一天), 新观点基准=当天
      week  = 上周一 → 今天 滚动窗口(维持原有口径不变), 新观点基准=本周一
      month = 本月 1 号 → 今天, 新观点基准=本月 1 号
    """
    import datetime as _dt
    today = _dt.date.today()
    this_monday = today - _dt.timedelta(days=today.weekday())
    if period == "day":
        start = today
        new_base = today
    elif period == "month":
        start = today.replace(day=1)
        new_base = today.replace(day=1)
    else:                                    # week(默认, 与历史口径一致)
        start = this_monday - _dt.timedelta(days=7)
        new_base = this_monday
    return (start.strftime("%Y-%m-%d"), today.strftime("%Y-%m-%d"),
            new_base.strftime("%Y-%m-%d"))


def kol_weekly_views(min_comment_len=10, period="week"):
    """指定窗口内全部有实质观点的 KOL, 按 sector 分组(不只转向, 而是所有有意义的观点)。
    数据源: Eco 自己每日快照(优先窗口内最新一天; 无快照回退 kol_independent.json 的 all)。
    每个 KOL 取其窗口内最新一次观点(方向+言论+标的)。绝不编: 无言论/无数据则跳过。
    period: day=本日 / week=本周(默认) / month=本月 —— 见 _kol_period_window。
    返回 {"date": 观点日期, "period": 档位, "total": KOL数, "modules": [...]}。
    """
    import json
    import datetime as _dt
    # 优先: 滚动窗口聚合(独立数据仓库)
    # ★每个 KOL 取其窗口内【最近一次有实质观点】的记录 → 只要窗口内任意一天有观点
    #   就一直显示, 不会因"最新那天快照恰好空/稀疏/未跑"而整个模块空掉(Chao 2026-08 修正)。
    snaps = load_kol_daily_snapshots()
    rows = []
    src_date = ""
    _invalid = {"", "未找到", "无", "n/a", "未知", "数据滞后", "-", "—"}
    _win_start, _win_end, _new_base = _kol_period_window(period)
    since_map = {}   # kol -> 当前方向连续保持起始日
    if snaps:
        all_dates = sorted(snaps.keys())              # 升序
        # 滚动窗口内的快照日期
        window_dates = [d for d in all_dates if _win_start <= d <= _win_end]
        # 若窗口内一个快照都没有(极端: 长期没跑/当天还没跑), 诚实回退到最近有数据的一天
        if not window_dates:
            window_dates = all_dates[-1:]
        src_date = window_dates[-1] if window_dates else (all_dates[-1] if all_dates else "")
        # ── 每个 KOL 取窗口内【最近一次有实质观点】的记录(按日期从新到旧找) ──
        latest_view = {}   # kol -> (date, record)
        for dd in reversed(window_dates):             # 新→旧
            for kol, x in snaps.get(dd, {}).items():
                if kol in latest_view:
                    continue                          # 已有更新的记录
                cmt = (x.get("comments") or "").strip()
                direction = (x.get("direction") or "").strip()
                # 只收有实质言论的(方向可为分歧等, 但 comments 要有内容)
                if cmt:
                    latest_view[kol] = (dd, x)
        # ── 首现日期: 对每个当前观点, 回溯全历史算【当前方向连续保持起始日】 ──
        for kol, (vdate, x) in latest_view.items():
            cur_dir = (x.get("direction") or "").strip()
            since = vdate
            if cur_dir and cur_dir not in _invalid:
                # 从 vdate 往前回溯(全历史, 不限窗口), 方向相同则前推起始日
                earlier = [d for d in all_dates if d < vdate]
                for dd in reversed(earlier):
                    prev_rec = snaps.get(dd, {}).get(kol)
                    if not prev_rec:
                        break  # 断档 → 起始日停在此
                    pdir = (prev_rec.get("direction") or "").strip()
                    if pdir == cur_dir:
                        since = dd
                    else:
                        break  # 方向变了
            since_map[kol] = since
            # ★言论首现日(Chao 2026-08-31): since 记的是【方向】连续保持起始日,
            #   方向不变时它察觉不到言论内容已经换了 —— 实测最新快照 118 人里
            #   60 人言论与上一日逐字相同, 却因方向没变而全被当成"当前观点",
            #   导致本日/本周/本月三档都是 127 人、filter 形同虚设。
            #   故另算 cmt_since = 【这段言论文本】首次出现的日期, 用于窗口过滤。
            _cur_cmt = (x.get("comments") or "").strip()
            cmt_since = vdate
            if _cur_cmt:
                for dd in reversed([d for d in all_dates if d < vdate]):
                    prev_rec = snaps.get(dd, {}).get(kol)
                    if not prev_rec:
                        break                      # 断档 → 首现日停在此
                    if (prev_rec.get("comments") or "").strip() == _cur_cmt:
                        cmt_since = dd             # 同一段话, 继续前推
                    else:
                        break                      # 话变了 → 首现日确定
            rows.append({"kol": kol, "sector": x.get("sector", ""),
                         "direction": x.get("direction", ""),
                         "comments": x.get("comments", ""),
                         "targets": x.get("targets", ""), "date": vdate,
                         "since_date": since, "cmt_since": cmt_since})
    else:
        # 回退: kol_independent.json 的 all(当日全量, 无历史→首现日=当日)
        if os.path.exists(KOL_INDEPENDENT):
            try:
                d = json.load(open(KOL_INDEPENDENT))
                src_date = d.get("date", "")
                for x in d.get("all", []):
                    if x.get("kol"):
                        rows.append({"kol": x["kol"], "sector": x.get("sector", ""),
                                     "direction": x.get("direction", ""),
                                     "comments": x.get("comments", ""),
                                     "targets": x.get("targets", ""),
                                     "date": x.get("date", src_date),
                                     "since_date": x.get("date", src_date)})
            except Exception:
                pass
    # 新观点基准日(当前方向起始日 >= 基准日 = 本档窗口内新转成的)
    _this_monday_str = _new_base
    # 只保留有实质言论的
    rows = [r for r in rows if (r.get("comments") or "").strip()
            and len((r["comments"]).strip()) >= min_comment_len]
    # ★窗口过滤(Chao 2026-08-31): 只留【这段言论本身是窗口内新出现】的人。
    #   crawl 每天全量跑, 每人天天都有快照, 若不按言论首现日过滤,
    #   任何窗口都会覆盖全部人 → 三档都是 127 位, filter 等于没做。
    #   本日 = 言论首现日 >= 今天(无今日快照则回退最近一天, 即"当天或前一天有新观点");
    #   本周/本月同理按各自窗口起点。没有新观点就诚实留空, 绝不把所有人堆上去。
    _cut = _win_start
    if period == "day":
        # 本日: 以最新快照日为准(cron 可能还没跑当天), 取该日新出现的言论
        _cut = src_date or _win_start
    rows = [r for r in rows if (r.get("cmt_since") or r.get("date") or "") >= _cut]
    # 按归一 sector 分组
    groups = {}
    for r in rows:
        sec_raw = (r.get("sector") or "").strip()
        zh, en = _KOL_SECTOR_MAP.get(sec_raw, (sec_raw or "其他", sec_raw or "Other"))
        groups.setdefault(zh, {"sector": zh, "en": en,
                               "color": _KOL_SECTOR_COLOR.get(zh, "#8a8377"),
                               "views": []})
        since = r.get("since_date", r.get("date", src_date))
        _cs = r.get("cmt_since") or r.get("date") or src_date
        groups[zh]["views"].append({
            "kol": r.get("kol", ""),
            "direction": r.get("direction", ""),
            "comments": (r.get("comments") or "")[:400],
            "targets": (r.get("targets") or "")[:150],
            "date": r.get("date", src_date),
            "since_date": since,                    # 当前方向起始日
            "cmt_since": _cs,                       # 这段言论首次出现日
            "is_new": bool(_cs and _cs >= _this_monday_str),  # 窗口内新观点
        })
    # 方向强弱排序: 看多在前(更醒目), 模块内按方向强弱降序
    _rank = {"强烈看多": 2, "看多": 1, "分歧": 0, "中性": 0, "看空": -1, "强烈看空": -2}
    for g in groups.values():
        g["views"].sort(key=lambda v: (-_rank.get(v["direction"], 0), v["kol"]))
    # 模块按观点数降序
    modules = sorted(groups.values(), key=lambda g: len(g["views"]), reverse=True)
    total = sum(len(g["views"]) for g in modules)
    return {"date": src_date, "period": period, "win_start": _win_start,
            "total": total, "modules": modules}


# sector 中英归一(Notion DB 用英文, independent.json 用中文) → 统一中文模块名 + 英文副标 + 色
_KOL_SECTOR_MAP = {
    "Precious Metals": ("贵金属", "Precious Metals"),
    "贵金属": ("贵金属", "Precious Metals"),
    "贵金属与商品周期": ("贵金属", "Precious Metals & Commodity Cycle"),
    "Macro": ("宏观货币与金融体系", "Macro & Monetary"),
    "宏观货币与金融体系": ("宏观货币与金融体系", "Macro & Monetary"),
    "Equities": ("股权市场", "Equities"),
    "股权市场": ("股权市场", "Equities"),
    "Crypto": ("加密资产", "Crypto"),
    "加密资产": ("加密资产", "Crypto"),
    "Energy & Commodities": ("资源与能源安全", "Energy & Commodities"),
    "Energy": ("资源与能源安全", "Energy & Commodities"),
    "资源与能源安全": ("资源与能源安全", "Energy & Commodities"),
    "Government Debt": ("国债利率与债券市场", "Government Debt & Rates"),
    "国债利率与债券市场": ("国债利率与债券市场", "Government Debt & Rates"),
    "预测": ("预测与周期", "Forecast & Cycle"),
    "科技与未来趋势": ("科技与未来趋势", "Tech & Future"),
    "交易与市场微观结构": ("交易与市场微观结构", "Trading & Microstructure"),
    # ★2026-08-20 新增: 周期理论/金融占星/易经术数等"非常规"预测派系, 单独成一个 section
    "Cycles & Esoteric Forecasting": ("周期与术数预测", "Cycles & Esoteric Forecasting"),
    "周期与术数预测": ("周期与术数预测", "Cycles & Esoteric Forecasting"),
}
_KOL_SECTOR_COLOR = {
    "贵金属": "#bfa06a", "贵金属与商品周期": "#bfa06a",
    "宏观货币与金融体系": "#8ea1ad", "股权市场": "#a693a0",
    "加密资产": "#c9ac6b", "资源与能源安全": "#9aab97",
    "国债利率与债券市场": "#c08a7d", "预测与周期": "#8a8377",
    "科技与未来趋势": "#8ea1ad", "交易与市场微观结构": "#a693a0",
    "周期与术数预测": "#9d8bb0",
}


def kol_cumulative_changes(period="month"):
    """★月度口径(Chao 2026-08-31 采纳, 经实测两轮修正后定稿)。

    ── 为什么不用端点对比 ──
    端点对比(首尾两天)会把「月初看多→月中转空→月末又转多」判为『没变』。
    实测 2026-08: 端点口径本月仅 2 人转向, 而本周(端点)就有 75 人, 属系统性失真。

    ── 为什么也不逐次展示切换 ──
    实测发现快照的 direction 标注质量不足以支撑「月内第 N 次转向」这种精度:
    Jeffrey Gundlach 8/26-8/31 四天内容几乎是同一件事(不降息 / 押现金黄金 /
    AI芯片是顶部信号), 却被标成 看空→看多→看空; 且"25%黄金的防御姿态"被标看多,
    实为对黄金看多+对市场看空混在一个字段。逐次展示 = 把标注噪音当市场信号。
    (已试过按 多/空/中 三阵营归并, 387 次仅降到 253 次, 未解决根本问题。)

    ── 最终口径: 只输出「本月是否稳定」这一诚实结论 ──
    对每人统计其月内出现过的【阵营集合】:
      stable=True  → 全月只落在一个阵营, 立场稳定(这个结论数据撑得住)
      stable=False → 月内在多个阵营间出现过, 标为『反复/不稳定』, 不声称具体几次
    展示主体 = 该人月内最后一次有效观点; days_seen = 月内被抓到的天数(可信度参考)。
    返回 [{kol,sector,prev_dir,new_dir,date,comments,targets,
           stable,camps,days_seen,path}...]
    绝不编: 无有效方向不参与; 不输出无法支撑的切换次数。
    """
    snaps = load_kol_daily_snapshots()
    if len(snaps) < 2:
        return []
    dates = sorted(snaps.keys())
    win_start, win_end, _ = _kol_period_window(period)
    win = [d for d in dates if win_start <= d <= win_end]
    if len(win) < 2:
        return []

    _invalid = {"", "未找到", "无", "n/a", "未知", "数据滞后", "-", "—"}

    def _camp(v):
        """粗粒度阵营: 只分 多/空/中, 不认细档差异(细档漂移是标注噪音)。"""
        if not v:
            return None
        if "看多" in v:
            return "多"
        if "看空" in v:
            return "空"
        return "中"

    def _dir(day, kol):
        rec = snaps.get(day, {}).get(kol)
        if not rec:
            return None
        v = (rec.get("direction") or "").strip()
        return None if v in _invalid else v

    hist = {}
    for d in win:
        for kol in snaps.get(d, {}):
            v = _dir(d, kol)
            if v is None:
                continue
            hist.setdefault(kol, []).append((d, v, _camp(v)))

    out = []
    for kol, s in hist.items():
        if len(s) < 2:
            continue                        # 月内只出现一次, 无从判断稳定与否
        camps = sorted({c for _d, _v, c in s})
        _cs = set(camps)
        # ★2026-08-31 实测分级: 只看"阵营数>1"会把 Michael Saylor 这种
        #   强烈看多→看多 的细档漂移(掺了个中性档)也判成"立场反复", 误导交易员。
        #   故分两级: flip=真跨越多空(camps 同含 多与空, 实测 40 人)
        #            soft=仅多↔中 或 空↔中 摆动(未换阵营, 实测 32 人)
        _flip = ("多" in _cs and "空" in _cs)
        _stable = (len(camps) == 1)
        last_day, last_dir = s[-1][0], s[-1][1]
        rec = snaps.get(last_day, {}).get(kol) or {}
        # 展示用的 prev: 取最后一条之前、阵营不同的最近一条(没有则同向)
        prev_dir = last_dir
        for d0, v0, c0 in reversed(s[:-1]):
            if c0 != s[-1][2]:
                prev_dir = v0
                break
        out.append({
            "kol": kol,
            "sector": rec.get("sector", ""),
            "prev_dir": prev_dir,
            "new_dir": last_dir,
            "date": last_day,
            "comments": (rec.get("comments") or "")[:300],
            "targets": _tgs(rec.get("targets"), 150),
            "stable": _stable,
            "flip": _flip,
            "camps": "/".join(camps),
            "days_seen": len(s),
            "path": "→".join(v for _d, v, _c in s[-6:]),   # 仅末 6 条供人工核查
        })
    # 排序: 真跨越多空的排最前(交易含义最强), 其次仅多↔中/空↔中 摆动, 稳定的最后;
    #       同级按被抓到天数降序(样本足的更可信)
    return sorted(out, key=lambda x: (not x["flip"], x["stable"],
                                      -x["days_seen"], x["kol"]))


def kol_stance_changes_grouped(days=None, period="week"):
    """窗口内 KOL 方向转向, 按 sector 模块分组。
    ★数据源: Eco 自己的每日快照(data/kol/daily/, kol_weekly_changes), 完全独立, 不依赖任何 agent 的 Notion DB。
    period: day=本日 / week=本周(默认) / month=本月。窗口内最新快照 vs 窗口起点前最后快照对比转向;
    跨窗口数据不足时回退最新vs最早(诚实, 攒够自动切回)。
    返回 {"since": 对比基准日, "period": 档位, "days": N, "total": 转向总数,
          "modules": [{"sector":中文, "en":英文, "color":色, "changes":[{kol,prev_dir,new_dir,date,comments,targets},...]}]}
    绝不编: 无数据则 modules=[]。"""
    # ★主源: Eco 自己的每日快照做窗口对比(完全独立)
    # 口径分流(Chao 2026-08-31): month 用【累计】口径(逐日相邻切换汇总),
    #   day/week 用【端点】口径(窗口内最新 vs 窗口前最后)。
    #   原因: 端点口径在月级窗口会把中途来回摆动的人抵消成"没变",
    #   实测本月仅 2 人 vs 本周 75 人, 属系统性失真而非平滑。
    cumulative = (period == "month")
    if cumulative:
        # ★不过滤 stable: 月档头部文案本就分列「反复 / 稳定」两类, 卡片徽章也各自区分,
        #   全量展示才能回答「谁一直没动摇」。(2026-08-31 我一度误加过滤, 会让
        #   n_stable 恒为 0、头部文案自相矛盾, 已撤销。)
        raw = kol_cumulative_changes(period=period)
    else:
        raw = kol_weekly_changes(period=period)
    # 基准日: 用快照里实际对比的 prev 日期(窗口起点之前最后一天)
    snaps = load_kol_daily_snapshots()
    win_start, _win_end, _ = _kol_period_window(period)
    since = ""
    if snaps:
        dates = sorted(snaps.keys())
        before = [d for d in dates if d < win_start]
        since = before[-1] if before else dates[0]
    days = len(snaps)

    # 按归一 sector 分组
    groups = {}
    for ch in raw:
        sec_raw = (ch.get("sector") or "").strip()
        zh, en = _KOL_SECTOR_MAP.get(sec_raw, (sec_raw or "其他", sec_raw or "Other"))
        groups.setdefault(zh, {"sector": zh, "en": en,
                               "color": _KOL_SECTOR_COLOR.get(zh, "#8a8377"),
                               "changes": []})
        groups[zh]["changes"].append({
            "kol": ch.get("kol", ""),
            "prev_dir": ch.get("prev_dir", ""),
            "new_dir": ch.get("new_dir", ""),
            "date": ch.get("date", ""),
            "comments": (ch.get("comments") or "")[:300],
            "targets": _tgs(ch.get("targets"), 150),
            # 月度累计口径专有字段(端点口径下为默认值)
            "stable": ch.get("stable"),
            "flip": ch.get("flip"),
            "camps": ch.get("camps", ""),
            "days_seen": ch.get("days_seen", 0),
            "path": ch.get("path", ""),
        })
    # 模块内按日期降序; 模块按转向数降序
    for g in groups.values():
        g["changes"].sort(key=lambda x: x["date"], reverse=True)
    modules = sorted(groups.values(), key=lambda g: len(g["changes"]), reverse=True)
    total = sum(len(g["changes"]) for g in modules)
    # 月度口径附加统计: 立场稳定 / 跨多空反转 / 仅强度摆动 三级
    n_stable = sum(1 for g in modules for c in g["changes"] if c.get("stable"))
    n_flip = sum(1 for g in modules for c in g["changes"]
                 if c.get("stable") is False and c.get("flip"))
    n_soft = sum(1 for g in modules for c in g["changes"]
                 if c.get("stable") is False and not c.get("flip"))
    n_unstable = n_flip + n_soft
    return {"since": since, "period": period, "days": days,
            "cumulative": cumulative,
            "n_stable": n_stable, "n_unstable": n_unstable,
            "n_flip": n_flip, "n_soft": n_soft,
            "total": total, "modules": modules}


def kol_today(date_str=None):
    """取指定日期(默认最新有数据日)的 KOL 言论。返回 [dict,...]。"""
    recs = fetch_kol_recent(days=7)
    if not recs:
        return []
    if date_str is None:
        date_str = max(r["date"] for r in recs if r["date"])
    return [r for r in recs if r["date"] == date_str]


# ─────────── Economic Dashboard 流动性 ───────────
ECON_DB = {
    "fed_liquidity": "4df3fc720f0c4cdc902cb18c009e4802",   # A5: reserves/RRP/TGA
    "risk_report": "750f9b463e234537bc7c07213da4d8fb",     # A6: 每日风控报告
    "ust_yields": "76f9359788ef481faeba19e5b74d5dc4",      # A1: 收益率
    "pboc": "b54e379b6aea4062a15c891697af1230",            # B2: 中国流动性
    "boj": "286e24acda1242f9a7dfca9aa1c89a2e",             # B3: 日本流动性
}


def _num(prop):
    return prop.get("number")


def _latest_row(db_id, title_field="Date"):
    """取某 DB 最新一行(按 title 降序)。返回 properties dict 或 None。"""
    st, b = _req("POST", f"/databases/{db_id}/query",
                 {"sorts": [{"property": title_field, "direction": "descending"}], "page_size": 2})
    if st == 200 and b.get("results"):
        return [r["properties"] for r in b["results"]]
    return None


def _pct_change(series, key, days):
    """时序在 N 个自然日前 vs 最新 的变动。

    ★ 用<b>自然日回溯</b>而非"往前数 N 个点"：日频序列有周末与节假日缺口，
      数点会让「过去一周」实际跨越十几天，口径失真。
    返回 (最新值, 基准值, 绝对变动, 百分比变动) —— 找不到基准就返回 None，不外推。
    """
    import datetime as _dt
    if not series:
        return None
    try:
        last = series[-1]
        last_d = _dt.date.fromisoformat(last["date"])
    except Exception:
        return None
    target = last_d - _dt.timedelta(days=days)
    # 取 <= target 里最靠近的一条（宽容 10 天，超出就算没有基准）
    base = None
    for r in reversed(series[:-1]):
        try:
            d = _dt.date.fromisoformat(r["date"])
        except Exception:
            continue
        if d <= target:
            if (target - d).days <= 10:
                base = r
            break
    if not base or base.get(key) is None or last.get(key) is None:
        return None
    cur, prev = last[key], base[key]
    delta = cur - prev
    pct = (delta / prev * 100) if prev else None
    return {"cur": cur, "prev": prev, "prev_date": base["date"],
            "delta": delta, "pct": pct}


def _omo_weekly_sum(omo, days):
    """OMO 是流量不是存量 → 比「最新值变动」应看<b>区间投放总额</b>。

    ★ 存量指标(DR007/余额)看点位变动；流量指标(逆回购投放)看区间求和。
      混用会得出荒谬结论（比如"逆回购较上周-100%"其实只是当天没操作）。
    """
    import datetime as _dt
    if not omo:
        return None
    try:
        last_d = _dt.date.fromisoformat(omo[-1]["date"])
    except Exception:
        return None
    lo = last_d - _dt.timedelta(days=days)
    prev_lo = lo - _dt.timedelta(days=days)

    def _sum(a, b):
        s, n = 0.0, 0
        for r in omo:
            try:
                d = _dt.date.fromisoformat(r["date"])
            except Exception:
                continue
            if a < d <= b:
                s += sum(r.get(k) or 0 for k in ("on", "d7", "other"))
                n += 1
        return s, n
    cur, n1 = _sum(lo, last_d)
    prev, n2 = _sum(prev_lo, lo)
    return {"cur": cur, "prev": prev, "n_cur": n1, "n_prev": n2,
            "delta": cur - prev,
            "pct": ((cur - prev) / prev * 100) if prev else None}


def fetch_china_liquidity(days=400):
    """B2 中国流动性时序 → 四条曲线（Chao 2026-08-25 要求）。

    三层融资渠道 + 一个股市稳定工具：
      ① OMO   逆回购，日频，短期     → 隔夜 / 7天 分品种（<b>流量</b>）
      ② DR007 银行间质押回购利率，日频 → 短期资金<b>价格</b>（存量点位）
      ③ MLF   中期借贷便利，月频      → 中期投放（流量）
      ④ SFISF 证券基金保险互换便利     → <b>股市托底工具</b>，看有无剧增

    ★ PSL 已于 2026-08-25 下架（Chao：「已经是过时的了，或者说不太需要了」）：
      余额自 2024-02 峰值 34,022 亿腰斩到 2025-05 的 17,939 亿，官网停更 448 天，
      工具本身在清退。抓取器 `pboc_mlf_psl_scraper` 保留（数据仍在写 B2），
      <b>只是不再渲染到 dashboard</b>。

    ★ 频率混装的坑：B2 是<b>日频</b>表，但 MLF/SFISF 是<b>事件频</b>指标，
      只在公告日/操作日那一行有值，其余日子为 None。<b>这不是缺数</b>。
      下游必须 dropna 后再连线，否则折线会断成一堆孤点。
    """
    out = {"status": "未获取", "omo": [], "dr007": [], "mlf": [], "sfisf": [],
           "margin": []}
    rows, cursor = [], None
    for _ in range(14):                       # 上限 1400 行，够 3 年日频
        body = {"page_size": 100,
                "sorts": [{"property": "Date", "direction": "descending"}]}
        if cursor:
            body["start_cursor"] = cursor
        st, b = _req("POST", f"/databases/{ECON_DB['pboc']}/query", body)
        if st != 200:
            out["error"] = f"Notion HTTP {st}"
            return out
        rows += b.get("results", [])
        if not b.get("has_more"):
            break
        cursor = b.get("next_cursor")
    if not rows:
        out["error"] = "B2 无数据"
        return out

    def _t(p):
        ti = p.get("Date", {}).get("title") or []
        return ti[0]["plain_text"] if ti else None

    def _n(p, k):
        x = p.get(k)
        return x.get("number") if x and x.get("type") == "number" else None

    for r in rows:
        p = r["properties"]
        d = _t(p)
        if not d:
            continue
        ov, sv = _n(p, "隔夜逆回购_亿"), _n(p, "7天逆回购_亿")
        oth = _n(p, "其他期限逆回购_亿")
        if ov is not None or sv is not None or oth is not None:
            out["omo"].append({"date": d, "on": ov, "d7": sv, "other": oth,
                               "rate": _n(p, "逆回购利率_pct")})
        v = _n(p, "DR007_pct")
        if v is not None:
            out["dr007"].append({"date": d, "v": v})
        v = _n(p, "MLF_操作量_亿")
        if v is not None:
            out["mlf"].append({"date": d, "v": v, "bal": _n(p, "MLF_余额_亿")})
        v = _n(p, "SFISF_规模_亿")
        if v is not None:
            out["sfisf"].append({"date": d, "v": v,
                                 "rate_bp": _n(p, "SFISF_中标费率_bp")})
        # ★ A股两融余额（Chao 2026-08-25）：数据一直在 B2 同一张表里（543 条 / 2024-06 起），
        #   但从未渲染到 dashboard —— 「数据已有」≠「已在看板上」。
        #   用途：SFISF 停发独立公告后，两融是<b>日频</b>的杠杆资金活跃度代理，
        #   能回答「资金在进还是在退」，比季度/年度的 SFISF 口径灵敏得多。
        v = _n(p, "A股两融余额_万亿")
        if v is not None:
            out["margin"].append({"date": d, "v": v})

    for k in ("omo", "dr007", "mlf", "sfisf", "margin"):
        out[k].sort(key=lambda z: z["date"])
    out["status"] = "ok"
    out["as_of"] = max((s[-1]["date"] for s in
                        (out["omo"], out["dr007"], out["mlf"], out["sfisf"],
                         out["margin"]) if s),
                       default=None)

    # ── MLF 净投放（Chao 2026-08-25：「总投放量，但是有到期量，对吗？」）──
    # ★ 到期量不是估算，是<b>精确推算</b>：MLF 全部为 1 年期（实测 20/20 条 tenor 均为「1年」），
    #   所以本月到期量 ≡ 12 个月前的操作量。数据来自我们自己已入库的同一序列。
    #   净投放 = 本月操作量 − 12 个月前操作量 —— 这才是对银行体系流动性的真实影响：
    #   投放 5000 亿但到期 6000 亿，实为净回笼 1000 亿，只看投放会得出反向结论。
    # ★ PBoC 公告本身不含到期量（实抓最近 3 期确认无「到期/续作/净投放」字样），
    #   因此前 12 个月无法计算 → 显式留 None，不外推、不用 0 冒充。
    mlf_by_month = {r["date"][:7]: r["v"] for r in out["mlf"]}
    for r in out["mlf"]:
        y, m = int(r["date"][:4]), int(r["date"][5:7])
        mat = mlf_by_month.get(f"{y-1}-{m:02d}")
        r["matured"] = mat
        r["net"] = (r["v"] - mat) if mat is not None else None

    # ── 多周期变动（Chao：「过去一周、一月甚至半年的趋势」）──
    # 存量/价格类看点位变动；流量类(OMO)看区间投放总额。
    out["deltas"] = {
        "dr007": {f"{d}d": _pct_change(out["dr007"], "v", d)
                  for d in (7, 30, 180)},
        "omo": {f"{d}d": _omo_weekly_sum(out["omo"], d)
                for d in (7, 30, 180)},
        "mlf": {f"{d}d": _pct_change(out["mlf"], "v", d)
                for d in (30, 90, 180)},
        # 净投放单列一行：它与「操作量」是两个不同的口径，不能合并展示
        "mlf_net": {f"{d}d": _pct_change(
            [x for x in out["mlf"] if x.get("net") is not None], "net", d)
            for d in (30, 90, 180)},
        # 两融是<b>存量</b>（余额），比点位差，不是区间求和
        "margin": {f"{d}d": _pct_change(out["margin"], "v", d)
                   for d in (7, 30, 180)},
    }

    # ── 宽基 ETF 份额（Chao 2026-08-25，B 方案）──
    # ★ 只看<b>份额</b>不看规模/价格：规模=净值×份额会被指数涨跌污染
    #   （指数涨10%、份额没动，规模也涨10%，看着像流入其实没有）。
    #   份额增加 = 真有人申购 = 真金白银进场。这是判断国家队在不在场唯一干净的口径。
    # 数据由 src/fetchers/etf_share.py 采自<b>上交所官方</b>每日 ETF 规模表，落盘 data/etf_share.json。
    out["etf"] = _load_etf_share()
    if out["etf"].get("series"):
        out["deltas"]["etf"] = {
            code: {f"{d}d": _pct_change(s, "v", d) for d in (7, 30, 180)}
            for code, s in out["etf"]["series"].items()
        }
    return out


ETF_NAMES = {"510300": "沪深300ETF", "510050": "上证50ETF",
             "510500": "中证500ETF"}


def _load_etf_share():
    """读 data/etf_share.json → {series: {code: [{date, v}]}, names: {...}}。

    文件由 etf_share.py 采集器维护（键为 YYYYMMDD，值为 {code: 亿份}）。
    这里只做读取与格式转换，<b>不在渲染路径里发网络请求</b>——
    否则 dashboard 构建会被上游限流拖垮。

    ★★ 缺口护栏（2026-08-25 实测教训）：回填未跑完时序列会「两端有、中间空」，
       相邻两点跨越数月缺口，画出来就是一根 −72% 的悬崖 —— 看着像巨额赎回，
       实际只是没采到。这类假象比缺数据危险得多（会直接导致误判）。
       故：<b>相邻两点间隔 > 20 天即判定为缺口，从缺口处截断，只保留最近的连续段</b>，
       并回报 gap_trimmed 供上层标注。宁可少画一段，也不画出假崩盘。
    """
    path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                        "data", "etf_share.json")
    if not os.path.exists(path):
        return {"status": "未采集", "series": {}}
    try:
        with open(path, encoding="utf-8") as f:
            raw = json.load(f)
    except (OSError, ValueError) as e:
        # ★ 只吞「文件读不到 / JSON 坏了」这两类<b>数据问题</b>。
        #   绝不用裸 except：那样代码里的 NameError/TypeError 会被伪装成
        #   "读取失败"，看板照常渲染，bug 永远浮不上来（2026-08-25 实际踩到）。
        return {"status": f"读取失败: {type(e).__name__}", "series": {}}
    series = {}
    for ds in sorted(raw):
        iso = f"{ds[:4]}-{ds[4:6]}-{ds[6:8]}"
        for code, v in (raw[ds] or {}).items():
            if v is None:
                continue
            series.setdefault(code, []).append({"date": iso, "v": v})

    trimmed = 0
    for code, s in list(series.items()):
        if len(s) < 2:
            continue
        cut = 0
        for i in range(len(s) - 1, 0, -1):
            d1 = datetime.date.fromisoformat(s[i]["date"])
            d0 = datetime.date.fromisoformat(s[i - 1]["date"])
            if (d1 - d0).days > 20:      # 跨越缺口
                cut = i
                break
        if cut:
            trimmed += len(s) - len(s[cut:])
            series[code] = s[cut:]

    return {"status": "ok" if series else "空",
            "series": series, "names": ETF_NAMES,
            "gap_trimmed": trimmed,
            "as_of": max((s[-1]["date"] for s in series.values() if s),
                         default=None)}


def fetch_liquidity_points():
    """取 Economic Dashboard 流动性关键点。返回 dict(取到什么算什么，缺就 None)。"""
    out = {}
    # Fed 流动性 (A5)
    rows = _latest_row(ECON_DB["fed_liquidity"])
    if rows:
        cur = rows[0]
        prev = rows[1] if len(rows) > 1 else {}
        for label, fld in [("reserves_T", "Reserves_T"), ("on_rrp_B", "ON_RRP_B"), ("tga_B", "TGA_B")]:
            v = _num(cur.get(fld, {}))
            pv = _num(prev.get(fld, {})) if prev else None
            if v is not None:
                out[label] = {"value": v, "prev": pv,
                              "delta": (round(v - pv, 2) if pv is not None else None)}
    # 每日风控报告 (A6) — 取风控灯/结论文本
    rr = _latest_row(ECON_DB["risk_report"])
    if rr:
        cur = rr[0]
        for name, p in cur.items():
            if p.get("type") == "rich_text" and _rt(p):
                out.setdefault("risk_notes", {})[name] = _rt(p)[:400]
            elif p.get("type") == "select" and p.get("select"):
                out.setdefault("risk_lights", {})[name] = p["select"]["name"]
    # 收益率 (A1)
    yr = _latest_row(ECON_DB["ust_yields"])
    if yr:
        cur = yr[0]
        prev = yr[1] if len(yr) > 1 else {}
        for fld in ["10Y", "2Y", "30Y"]:
            v = _num(cur.get(fld, {}))
            pv = _num(prev.get(fld, {})) if prev else None
            if v is not None:
                out.setdefault("yields", {})[fld] = {"value": v, "delta": (round(v - pv, 3) if pv is not None else None)}
    return out


# ─────────── 三大央行资产负债表 (JP/CN/US) ───────────
# 数据源: Economic Dashboard 已维护的 B5(PBoC月度)/B6(BoJ旬报)/B7(Fed周度) Notion DB。
# 读最新两行, 每科目算环比。对比口径随频率不同(Fed=WoW真周环比 / BoJ=较上期旬报 / PBoC=较上月MoM)。
CB_BS_DB = {
    "US": "dea7e939bb394d538f3cea3ff0da5a6b",  # B7 Fed 周度
    "JP": "481f6e1960e4444383d66a1e1cafb49b",  # B6 BoJ 旬报
    "CN": "20b0eb37d69744109cca9fba0c61b929",  # B5 PBoC 月度
}

# 每国: (title字段, 单位, 对比口径标签, 资产科目[(显示名,字段)], 负债科目[(显示名,字段)], 总资产字段, 总负债字段)
CB_BS_SPEC = {
    "US": {
        "name": "美联储 Fed", "flag": "🇺🇸", "title": "Week", "unit": "$B", "period": "WoW",
        "assets": [
            ("美国国债 Treasuries", "资产_Treasuries_B"),
            ("短期国债 Bills（RMP 储备管理购买）", "资产_Bills_B", "sub"),
            ("MBS 抵押债", "资产_MBS_B"),
            # ★2026-09-25 口径更正(Chao:「FIMA 使用量我记得很小」): 真 FIMA = WORAL(资产侧,
            #   Fed 借美元给外国央行, 常态≈0)。此前这里显示的「FIMA 海外回购 $318B」读的是
            #   B7 字段 资产_FIMA_B = WLRRAFOIAL, 实为**负债侧**外国官方逆回购池 → 已移到负债侧并正名。
            ("FIMA 回购便利（Fed 借美元给外国央行）", "资产_FIMA真_B"),   # FRED H41RESPPALGTRFNWW(非 WORAL)
            ("SRF 常备回购", "资产_SRF_B"),
        ],
        "liabilities": [
            ("准备金 Reserves", "负债_Reserves_B"),
            ("隔夜逆回购 ON RRP", "负债_ON_RRP_B"),
            ("外国官方及国际账户逆回购（外国央行存放）", "资产_FIMA_B"),   # 字段名是上游历史遗留, 实为负债侧 WLRRAFOIAL
            ("财政部账户 TGA", "负债_TGA_B"),
            ("流通货币 Currency", "负债_Currency_B"),
        ],
        "total_a": "总资产_B", "total_l": None,
    },
    "JP": {
        "name": "日本央行 BoJ", "flag": "🇯🇵", "title": "Date", "unit": "兆¥", "period": "较上期",
        "assets": [
            ("日本国债 JGB", "资产_国债JGB_兆JPY"),
            ("公司债", "资产_公司债_兆JPY"),
            ("ETF", "资产_ETF_兆JPY"),
            ("J-REIT", "资产_J_REIT_兆JPY"),
        ],
        "liabilities": [
            ("银行券 Banknotes", "负债_银行券_兆JPY"),
            ("经常项目存款", "负债_经常项目存款_兆JPY"),
            ("政府存款", "负债_政府存款_兆JPY"),
        ],
        "total_a": "总资产_兆JPY", "total_l": None,
    },
    "CN": {
        "name": "中国央行 PBoC", "flag": "🇨🇳", "title": "Month", "unit": "亿¥", "period": "MoM",
        "assets": [
            ("外汇占款", "资产_外汇占款_亿"),
            ("对政府债权", "资产_对政府债权_亿"),
            ("对其他存款性公司债权", "资产_对其他存款性公司债权_亿"),
        ],
        "liabilities": [
            ("储备货币", "负债_储备货币_亿"),
            ("货币发行", "负债_货币发行_亿", "sub"),
            ("政府存款", "负债_政府存款_亿"),
        ],
        "total_a": "总资产_亿", "total_l": "总负债_亿",
    },
}


def _title_val(props, field):
    p = props.get(field, {})
    arr = p.get("title") or []
    return "".join(x.get("plain_text", "") for x in arr)


def _bs_line(cur, prev, disp, fld, sub=False):
    """一个科目: 取当前值 + 环比。返回 {name,value,delta,pct,sub} 或 None(值缺)。
    sub=True 表示该项是上一总项的子项(如"货币发行"⊂"储备货币"),渲染层缩进显示"其中·",不与总项同层相加。"""
    v = _num(cur.get(fld, {}))
    if v is None:
        return None
    pv = _num(prev.get(fld, {})) if prev else None
    delta = round(v - pv, 3) if pv is not None else None
    pct = round((v - pv) / pv * 100, 2) if (pv not in (None, 0)) else None
    return {"name": disp, "value": v, "delta": delta, "pct": pct, "sub": sub}


def _usd_convert(cc_data, factor, orig_unit):
    """把一个央行的资负表所有值 × factor 转成 $B。orig_unit 记进 note 保留原口径可溯源。"""
    if not cc_data or factor is None:
        return cc_data
    def conv(line):
        if not line:
            return line
        for k in ("value", "delta"):
            if line.get(k) is not None:
                line[k] = round(line[k] * factor, 1)
        return line
    for sec in ("assets", "liabilities"):
        cc_data[sec] = [conv(x) for x in cc_data.get(sec, [])]
    cc_data["total_assets"] = conv(cc_data.get("total_assets"))
    cc_data["total_liab"] = conv(cc_data.get("total_liab"))
    cc_data["unit"] = "$B"
    cc_data["orig_unit"] = orig_unit
    return cc_data


def fetch_ecb_balance_sheet(to_usd=True, usd_per_eur=None):
    """欧元体系 Eurosystem 资产负债表(完整分项版, 格式与另三国一致)。

    ★★2026-09-09 改名(Chao 追问口径时发现标签撒谎):
      原名「欧洲央行 ECB」是<b>错的</b> —— 数据源 ILM/W.U2.C 中
      U2=整个欧元区、C=Eurosystem, 即 <b>ECB 本部 + 20 家成员国央行的合并表</b>,
      不是 ECB 一家。实测合并总资产 €5,912B, 而德国央行一家就 €2,271B;
      卡片上的「黄金 1,432B」主要是德/意/法等国央行的金库, 不在法兰克福。
      看单国央行请用 fetch_euro_ncb_balance_sheets()。
    ★数据源: ECB 官方 Data Portal SDMX API 的 ILM 数据集(Eurosystem 周度合并财务报表),
      免费无key, wildcard 一次拉全部 43 分项(单条 series key 查询对维度取值敏感易 400, wildcard 最稳)。
      base: https://data-api.ecb.europa.eu/service/data/ILM/W.U2.C... 单位=百万欧元。
      政策利率 ECBMRRFR(MRO)/ECBDFR(存款便利)/ECBMLFR(边际贷款) 走 FRED。
    to_usd: 全部分项按当天 USD/EUR(DEXUSEU, 1欧元=X美元)折成 $B(乘)。
    返回与 CB_BS_SPEC 一致结构。绝不编: 抓不到的科目跳过。"""
    import requests
    import csv as _csv
    import io as _io

    # ── ILM wildcard 一次拉全部分项 (BS_ITEM.COUNT_AREA.CURRENCY_TRANS → (值,期次)) ──
    # ★拉最近 2 期(lastNObservations=2)以便算真实周环比 delta/pct。
    #   曾只拉 1 期导致全站 ECB 分项 delta/pct 恒为 None(渲染成 18 处 n/a)。
    #   实测 43/43 科目两期齐全; 若某科目缺上期, 该科目 delta 保持 None(诚实留空, 绝不编造)。
    got = {}          # tail -> 最新值
    got_prev = {}     # tail -> 上一期值
    period = None
    try:
        r = requests.get("https://data-api.ecb.europa.eu/service/data/ILM/W.U2.C...",
                         params={"lastNObservations": 2, "format": "csvdata"},
                         headers={"User-Agent": "Mozilla/5.0"}, timeout=35)
        if r.status_code == 200:
            _by_tail = {}
            for row in _csv.DictReader(_io.StringIO(r.text)):
                tail = f'{row["BS_ITEM"]}.{row["COUNT_AREA"]}.{row["CURRENCY_TRANS"]}'
                try:
                    _by_tail.setdefault(tail, {})[row["TIME_PERIOD"]] = float(row["OBS_VALUE"])
                except (ValueError, KeyError):
                    pass
            # 全局最新期次 = 所有科目中最大的 TIME_PERIOD(形如 2026-W33, 字典序即时间序)
            _all_p = sorted({p for v in _by_tail.values() for p in v})
            if _all_p:
                period = _all_p[-1]
                prev_p = _all_p[-2] if len(_all_p) > 1 else None
                for tail, obs in _by_tail.items():
                    if period in obs:
                        got[tail] = obs[period]
                    if prev_p and prev_p in obs:
                        got_prev[tail] = obs[prev_p]
    except Exception:
        pass

    # 科目映射(显示名 → ILM tail key), 已独立验证真实可抓, 单位百万€
    ASSETS = [
        ("黄金及黄金债权 Gold", "A010000.Z5.Z0Z"),
        ("外币资产 FX Assets", "A020000.U4.Z06"),
        ("货币政策证券 Securities(APP+PEPP)", "A070100.U2.EUR"),
        ("对信贷机构贷款 Lending to CIs", "A050000.U2.EUR"),
        ("其他资产 Other Assets", "A110000.Z5.Z01"),
    ]
    LIABS = [
        ("流通银行券 Banknotes", "L010000.Z5.EUR"),
        ("对信贷机构负债 Liab. to CIs", "L020000.U2.EUR"),
        ("存款便利 Deposit Facility", "L020200.U2.EUR", "sub"),
        ("对其他居民负债 Liab. to Others", "L050000.U2.EUR"),
        ("重估账户 Revaluation", "L140000.Z5.Z01"),
        ("资本与储备 Capital&Reserves", "L150000.Z5.Z01"),
        ("其他负债 Other Liab.", "L120000.Z5.Z01"),
    ]
    total_key = "T000000.Z5.Z01"

    # 折美元汇率: 1€ = usd_per_eur 美元
    if usd_per_eur is None:
        try:
            from fetchers.fred import fetch_fred_latest
            usd_per_eur, _ = fetch_fred_latest("DEXUSEU")
        except Exception:
            usd_per_eur = None
    if not usd_per_eur:
        usd_per_eur = 1.16
    # 百万€ → $B: /1000(百万→十亿€) × usd_per_eur; 不折美元则 /1000 得十亿€
    factor = (usd_per_eur / 1000.0) if to_usd else (1.0 / 1000.0)

    def _line(name, tail, sub=False):
        if tail not in got:
            return None
        val = round(got[tail] * factor, 1)
        # 周环比: 仅当上期真实存在才算, 否则诚实留 None(绝不用 0 或估算值冒充)
        delta = pct = None
        if tail in got_prev:
            prev = got_prev[tail] * factor
            delta = round(val - prev, 1)
            if prev:
                pct = round((val - prev) / abs(prev) * 100, 2)
        return {"name": name, "value": val,
                "delta": delta, "pct": pct, "sub": sub}

    # 政策利率(FRED)
    rates = {}
    for key, sid in (("mro", "ECBMRRFR"), ("dfr", "ECBDFR"), ("mlf", "ECBMLFR")):
        try:
            from fetchers.fred import fetch_fred_latest
            v, _ = fetch_fred_latest(sid)
            if v is not None:
                rates[key] = v
        except Exception:
            pass

    if not got:
        return {"name": "欧元体系 Eurosystem", "flag": "🇪🇺", "unit": "$B" if to_usd else "十亿€",
                "period": "周度", "date": None, "assets": [], "liabilities": [],
                "total_assets": None, "total_liab": None, "rates": rates, "status": "未找到"}

    assets = [x for x in (_line(n, k) for n, k in ASSETS) if x]
    liabs = [x for x in (_line(s[0], s[1], len(s) > 2 and s[2] == "sub") for s in LIABS) if x]
    ta = _line("总资产", total_key)
    return {
        "name": "欧元体系 Eurosystem", "flag": "🇪🇺",
        "unit": "$B" if to_usd else "十亿€", "orig_unit": "十亿€",
        "period": "周度", "date": period,       # 如 2026-W32
        "assets": assets, "liabilities": liabs,
        "total_assets": ta, "total_liab": None,
        "rates": rates,
        "usd_per_eur": round(usd_per_eur, 4),
        "source": "ECB Data Portal ILM (Eurosystem 周度财务报表)",
    }


def fetch_euro_ncb_balance_sheets(to_usd=True, usd_per_eur=None, months=78):
    """欧元区五国央行(德/意/荷/西/芬)各自的资产负债表总资产 —— 月度序列。

    ★2026-09-09 新增(Chao: 「把剩下 5 个国家单独做一个表, 并在子表下面再分别做
      这五个国家的央行的资产负债表」)。

    ★★口径警告(必须随数据一起显示, 否则会造成严重误读):
      本函数的数字与「TIC 德国持有美债 107.8B」<b>不是包含关系, 也不是同一批持有人</b>。
        - TIC「德国 107.8B」= 登记在德国的托管账户里的美债, 持有人是德国的
          保险/养老金/基金/企业等<b>所有外国持有人</b>, 央行只占极小部分。
        - 本函数 = 该国央行<b>自己</b>的资产负债表总资产。
      实测佐证: Bundesbank 的「对非欧元区外币债权」(TTA008) 仅 €0.004B ——
      <b>欧元区各国央行几乎不持有美债</b>(外汇储备已大部分上缴 ECB 统一管理)。
      所以「德国央行在增持美债」这类推论不能从这两张表推出来。

    数据源: ECB Data Portal SDMX, BSI 数据集 NCB 部门(BS_REP_SECTOR='N')。
      key 模板: BSI/M.{CC}.N.N.T00.A.1.Z5.0000.Z01.E
      标题实测 = "Total assets of the NCB, Stocks", 单位百万欧元, 月度。
      ★为什么用 ECB 统一口径而不是各国央行自己的 API: 意/荷/西/芬 四家 API
        端点各不相同且我实测的候选端点全部不可用(返回 HTML 页面而非数据);
        ECB BSI 一套 key 模板五国通吃, 且口径统一可横向比较。
      ★局限: BSI 的 NCB 部门只有 T00(总资产) 和 L40(发行债务证券) 两个科目,
        <b>没有黄金/证券等分项</b>。德国的详细分项可另走 Bundesbank API
        (BBBK11.D.TTA032 等), 但那样五国口径就不一致了, 故此处只做总资产。

    返回 {CC: {name, flag, series:[(YYYY-MM, 值),...], latest, first, delta, delta_pct,
               high, low, status, source, unit}}。抓不到标 status, 绝不编。
    """
    import requests

    NCB = {
        "DE": ("德国央行 Bundesbank", "🇩🇪"),
        "IT": ("意大利央行 Banca d'Italia", "🇮🇹"),
        "NL": ("荷兰央行 DNB", "🇳🇱"),
        "ES": ("西班牙央行 Banco de España", "🇪🇸"),
        "FI": ("芬兰央行 Suomen Pankki", "🇫🇮"),
    }
    out = {}

    rate = None
    if to_usd:
        # 与 fetch_ecb_balance_sheet 用同一口径: FRED DEXUSEU(1欧元=X美元), 兜底 1.16
        rate = usd_per_eur
        if rate is None:
            try:
                from fetchers.fred import fetch_fred_latest
                rate, _ = fetch_fred_latest("DEXUSEU")
            except Exception:
                rate = None
        if not rate:
            rate = 1.16
    unit = "$B" if (to_usd and rate) else "十亿€"

    for cc, (zh, flag) in NCB.items():
        rec = {"name": zh, "flag": flag, "series": [], "status": "未找到",
               "unit": unit,
               "source": "ECB Data Portal BSI (NCB 部门总资产, 月度)"}
        try:
            u = ("https://data-api.ecb.europa.eu/service/data/BSI/"
                 f"M.{cc}.N.N.T00.A.1.Z5.0000.Z01.E"
                 f"?format=jsondata&lastNObservations={int(months)}")
            r = requests.get(
                u, timeout=45,
                headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"})
            if r.status_code != 200:
                rec["status"] = f"HTTP {r.status_code}"
                out[cc] = rec
                continue
            d = r.json()
            ds = d.get("dataSets") or []
            ser = (ds[0].get("series") or {}) if ds else {}
            if not ser:
                rec["status"] = "数据不足"
                out[cc] = rec
                continue
            k0 = list(ser)[0]
            obs = ser[k0].get("observations") or {}
            tdim = ((d.get("structure") or {}).get("dimensions") or {}).get("observation") or []
            times = [v.get("id") for v in (tdim[0].get("values") or [])] if tdim else []
            pts = []
            for i, v in sorted(obs.items(), key=lambda z: int(z[0])):
                ti = int(i)
                if ti >= len(times) or not v or v[0] is None:
                    continue
                val = float(v[0]) / 1000.0            # 百万€ → 十亿€
                if to_usd and rate:
                    val *= rate
                pts.append((times[ti], round(val, 1)))
            if len(pts) < 2:
                rec["status"] = "数据不足"
                out[cc] = rec
                continue
            vals = [v for _, v in pts]
            rec.update({
                "series": pts, "status": "ok",
                "latest": pts[-1], "first": pts[0],
                "delta": round(pts[-1][1] - pts[0][1], 1),
                "delta_pct": (round((pts[-1][1] - pts[0][1]) / pts[0][1] * 100, 1)
                              if pts[0][1] else None),
                "high": round(max(vals), 1), "low": round(min(vals), 1),
                "cadence_note": "月度更新 · 滞后约1-2月",
            })
        except Exception as e:
            rec["status"] = f"抓取失败: {type(e).__name__}"
        out[cc] = rec
    return out


def fetch_cb_balance_sheets(to_usd=True):
    """读 JP/CN/US 三大央行资产负债表最新两期, 每科目算环比。
    to_usd=True: 统一换算成 $B(十亿美元)横向可比(汇率走 FRED DEXJPUS/DEXCHUS)。
    返回 {"US":{...}, "JP":{...}, "CN":{...}}, 每个:
      {name, flag, unit, period, date, assets:[{name,value,delta,pct}], liabilities:[...],
       total_assets:{value,delta,pct}, total_liab:{...}|None}
    取到什么算什么, 缺的科目跳过(绝不编)。"""
    out = {}
    for cc, spec in CB_BS_SPEC.items():
        rows = _latest_row(CB_BS_DB[cc], title_field=spec["title"])
        if not rows:
            out[cc] = {"name": spec["name"], "flag": spec["flag"], "unit": spec["unit"],
                       "period": spec["period"], "date": None, "assets": [], "liabilities": [],
                       "total_assets": None, "total_liab": None, "status": "未找到"}
            continue
        cur = rows[0]
        prev = rows[1] if len(rows) > 1 else {}
        assets = [x for x in (_bs_line(cur, prev, s[0], s[1], len(s) > 2 and s[2] == "sub") for s in spec["assets"]) if x]
        liabs = [x for x in (_bs_line(cur, prev, s[0], s[1], len(s) > 2 and s[2] == "sub") for s in spec["liabilities"]) if x]
        ta = _bs_line(cur, prev, "总资产", spec["total_a"]) if spec.get("total_a") else None
        tl = _bs_line(cur, prev, "总负债", spec["total_l"]) if spec.get("total_l") else None
        out[cc] = {
            "name": spec["name"], "flag": spec["flag"], "unit": spec["unit"],
            "period": spec["period"], "date": _title_val(cur, spec["title"]),
            "assets": assets, "liabilities": liabs,
            "total_assets": ta, "total_liab": tl,
        }
    # ── 统一换算成 $B(十亿美元), 全部用当天最新汇率(Chao 要求横向可比) ──
    usd_per_eur = None
    if to_usd:
        try:
            from fetchers.fred import fetch_fred_latest
            jpy, _ = fetch_fred_latest("DEXJPUS")   # 日元/美元(如 147.5)
            cny, _ = fetch_fred_latest("DEXCHUS")   # 人民币/美元(如 7.15)
            usd_per_eur, _ = fetch_fred_latest("DEXUSEU")  # 1欧元=X美元
            # BoJ: 兆¥(1e12¥) → $B: ×1000/jpy
            if "JP" in out and jpy:
                out["JP"] = _usd_convert(out["JP"], 1000.0 / jpy, "兆¥")
            # PBoC: 亿¥(1e8¥) → $B: ×0.1/cny
            if "CN" in out and cny:
                out["CN"] = _usd_convert(out["CN"], 0.1 / cny, "亿¥")
            # 记录换算汇率(可溯源)
            out["_fx"] = {"USDJPY": jpy, "USDCNY": cny, "USDEUR": usd_per_eur}
        except Exception as e:
            out["_fx_error"] = str(e)

    # ── ECB(欧洲央行): 独立源, 用同一当天 USD/EUR 折算 ──
    try:
        out["ECB"] = fetch_ecb_balance_sheet(to_usd=to_usd, usd_per_eur=usd_per_eur)
    except Exception as e:
        out["_ecb_error"] = str(e)

    # ── 总资产/总负债强制加总: 补"其他"差额项(Chao 要求, 如中国总负债7170≠5798+891.5) ──
    def _add_residual(blk):
        """若有总额且明细未加满, 补一行'其他(差额)'使明细总和=总额。子项(sub)不计入加和。"""
        for sec, tot_key in (("assets", "total_assets"), ("liabilities", "total_liab")):
            tot = blk.get(tot_key)
            lines = blk.get(sec) or []
            if not tot or tot.get("value") is None or not lines:
                continue
            summed = sum(x["value"] for x in lines if x.get("value") is not None and not x.get("sub"))
            resid = round(tot["value"] - summed, 1)
            # 只在差额显著(>总额0.5%且>1)时补, 避免舍入噪音
            if abs(resid) > max(1.0, abs(tot["value"]) * 0.005):
                lines.append({"name": "其他 Other（差额）", "value": resid,
                              "delta": None, "pct": None, "sub": False})
                blk[sec] = lines
    for cc in ("US", "JP", "CN", "ECB"):
        if cc in out and isinstance(out[cc], dict):
            _add_residual(out[cc])
    return out


# ─────────── 外国官方在纽约联储托管的美债 (Fed H.4.1) ───────────
# Chao 需求(2026-08, 附FT图): 外国央行/官方机构在NY Fed托管的可流通美债。
# 反映外国官方对美债的减持/去美元化趋势(FT图红箭头强调2025-26急跌)。
# 数据源: Fed H.4.1 weekly release HTML(每周四更新), "Securities held in custody
#   for foreign official and international accounts" 段的 "Marketable U.S. Treasury securities" 行。
#   ★FRED 的 WMTSECL 等 series 2012 已停更, 活跃真数据只在 Fed H.4.1 release 本身。
# 单位: 百万美元 → 转 $T(万亿)。绝不编: 抓不到标 status。
FED_H41_URL = "https://www.federalreserve.gov/releases/h41/current/h41.htm"


def fetch_cb_liquidity_swaps(years=1):
    """★2026-08-24(Chao 点名遗失指标): 央行美元流动性互换 (Central Bank Liquidity Swaps)。

    Chao 原话:「我刚刚发现有一个指标我们好像遗失掉了,叫"央行货币互换"。
              本周成交额是 1.21 亿美元, 欧洲央行在 8月19号成交了 1.19 亿美元。」
    实测两个数字与官方 API 完全吻合(W34=1.21亿 / 08-19 ECB=1.19亿), 已交叉验证。

    为什么重要: 美联储与外国央行的美元互换额度是<b>离岸美元荒的直接温度计</b>。
      平时接近零; 一旦某国央行大额动用, 说明该辖区银行体系在境外市场借不到美元。
      2008/2020 危机期间曾冲到数千亿美元。与 FIMA 回购、SOFR-IORB 利差互为佐证。

    数据源: 纽约联储官方 API (markets.newyorkfed.org), 免费无需 key, 可查历史。
      端点: /api/fxs/all/search.json?startDate=&endDate=
      ★注意 /latest.json 常返回空数组(只在当日有新操作时才有值), 不能用作主源;
        必须用 search.json 按日期区间拉取。
    单位: 原始 amount 为<b>美元</b>(非百万/十亿), 直接除 1e8 得"亿美元"。
    """
    import datetime as _dt
    out = {"status": "error", "as_of": None, "ops": [], "weekly": [],
           "by_cb": {}, "latest_week_usd": 0.0, "source": "NY Fed markets API",
           "source_url": "https://markets.newyorkfed.org/api/fxs/all/search.json"}
    try:
        end = _dt.date.today()
        start = end - _dt.timedelta(days=int(365 * years) + 10)
        url = ("https://markets.newyorkfed.org/api/fxs/all/search.json"
               f"?startDate={start.isoformat()}&endDate={end.isoformat()}")
        import urllib.request as _u
        import json as _j
        req = _u.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        raw = _j.load(_u.urlopen(req, timeout=45))
        ops = ((raw or {}).get("fxSwaps") or {}).get("operations") or []
    except Exception as e:
        out["error"] = f"NY Fed API 失败: {str(e)[:150]}"
        return out                      # 抓不到就如实返回 error, 绝不编造

    if not ops:
        out["status"] = "empty"
        out["error"] = "区间内无互换操作(平时可能确实为零)"
        return out

    rows = []
    for o in ops:
        try:
            amt = float(o.get("amount") or 0)
        except Exception:
            continue
        rows.append({
            "date": (o.get("tradeDate") or "")[:10],
            "cb": o.get("counterparty") or "",
            "usd": amt,
            "term": o.get("termInDays"),
            "rate": o.get("interestRate"),
            "maturity": (o.get("maturityDate") or "")[:10],
        })
    rows.sort(key=lambda r: r["date"], reverse=True)

    # 按 ISO 周汇总(便于画柱状图与"本周成交额")
    wk = {}
    for r in rows:
        try:
            y, w, _ = _dt.date.fromisoformat(r["date"]).isocalendar()
        except Exception:
            continue
        wk.setdefault((y, w), 0.0)
        wk[(y, w)] += r["usd"]
    weekly = [{"week": f"{y}-W{w:02d}", "usd": v} for (y, w), v in sorted(wk.items())]

    by_cb = {}
    for r in rows:
        by_cb.setdefault(r["cb"], {"n": 0, "usd": 0.0})
        by_cb[r["cb"]]["n"] += 1
        by_cb[r["cb"]]["usd"] += r["usd"]

    out.update({
        "status": "ok",
        "as_of": rows[0]["date"],
        "ops": rows[:60],               # 最近 60 笔明细
        "weekly": weekly,
        "by_cb": by_cb,
        "latest_week_usd": weekly[-1]["usd"] if weekly else 0.0,
        "latest_week": weekly[-1]["week"] if weekly else "",
        "n_ops": len(rows),
        "peak_week_usd": max((w["usd"] for w in weekly), default=0.0),
    })
    return out


def _custody_history_fred(start="2026-01-01"):
    """从 FRED WMTSECL1(外国官方托管可流通美债, 周度 as-of Wednesday, 2002至今活跃)
    拉历史序列。返回 [(date, value_$T), ...] 升序。空列表=拉取失败。"""
    try:
        from fetchers.fred import fetch_fred_history
    except Exception:
        from src.fetchers.fred import fetch_fred_history
    raw = fetch_fred_history("WMTSECL1", start=start)  # 单位: 百万美元
    return [(d, round(v / 1e6, 4)) for d, v in raw]     # → $T


def _mspd_maturing_within_1yr(record_date, rows):
    """给定某 record_date 的 MSPD table_3 明细行, 加总"距该日1年内到期"的可交易国债 outstanding。
    outstanding_amt 缺失时用 issued_amt+redeemed_amt(redeemed为负)兜底。单位: 百万美元 → 返回 $T。"""
    from datetime import datetime
    try:
        rec = datetime.strptime(record_date, "%Y-%m-%d")
    except Exception:
        return None
    total = 0.0
    for r in rows:
        md = r.get("maturity_date")
        if not md or md == "null":
            continue
        try:
            mdt = datetime.strptime(md, "%Y-%m-%d")
        except Exception:
            continue
        days = (mdt - rec).days
        if 0 <= days <= 366:
            oa = r.get("outstanding_amt")
            if oa in ("null", None, ""):
                try:
                    oa = float(r.get("issued_amt", 0) or 0) + float(r.get("redeemed_amt", 0) or 0)
                except Exception:
                    oa = 0
            else:
                try:
                    oa = float(oa)
                except Exception:
                    oa = 0
            total += oa
    return round(total / 1e6, 4)  # 百万 → 万亿($T)


# ── 日本一般会计当初予算总额(官方权威静态值, 兆円) ──
# 来源: 財務省 各年度予算 / 首相官邸「予算の概要」。当初予算(補正前)=年度启动时已确定盘子。
# FY2026=122.3兆(令和8年度, +7.1兆, 首相官邸 40621_ext_20_1.pdf 确认)。历年当初予算公开权威值。
JP_GENERAL_BUDGET = {
    2016: 96.72, 2017: 97.45, 2018: 97.71, 2019: 101.46, 2020: 102.66,
    2021: 106.61, 2022: 107.60, 2023: 114.38, 2024: 112.57, 2025: 115.20,
    2026: 122.30,
}


def fetch_fiscal_budget(cache_path=None):
    """日美年度财政花费(政府总支出/预算)柱状图数据。方案A: 每财年一柱, 双轴。
    - 美国: MTS table5 line5691 Total Outlays 财年末累计(9月)=全年实际支出; 当前财年至今=partial。单位 $T。
      美国财年 10/1 起。历史年=confirmed(已决算), 当前进行中财年=partial(至今累计)。
    - 日本: 一般会计当初予算总额(官方权威静态表 JP_GENERAL_BUDGET)。日本财年 4/1 起。单位 兆円。
      当初予算=年度启动已确定盘子(补正前); 全部标 confirmed(已公布的既定预算)。
    返回 {us:[{fy,value_t,status}], jp:[{fy,value_oku,status}], as_of, status, source}。
    绝不编: 美国 API 失败则 us 空并标 status; 日本用官方静态值。"""
    import requests
    import datetime
    out = {"us": [], "jp": [], "as_of": None, "status": "ok",
           "source": "美国 US Treasury MTS Table 5 (Total Outlays, 财年累计) · 日本 財務省 一般会計当初予算"}
    today = datetime.date.today()
    # 美国当前财年(10/1 起): 若当前月>=10 则财年=年+1
    cur_us_fy = today.year + 1 if today.month >= 10 else today.year

    # ── 美国: 财年末(9月)累计 = 全年实际支出 ──
    us_by_fy = {}
    try:
        r = requests.get(
            "https://api.fiscaldata.treasury.gov/services/api/fiscal_service/v1/accounting/mts/mts_table_5",
            params={"filter": "line_code_nbr:eq:5691,record_calendar_month:eq:09",
                    "fields": "record_date,record_fiscal_year,current_fytd_net_outly_amt",
                    "sort": "-record_date", "page[size]": "20"},
            timeout=45)
        if r.status_code == 200:
            for x in r.json().get("data", []):
                fy = x.get("record_fiscal_year")
                v = x.get("current_fytd_net_outly_amt")
                if fy and v not in (None, "", "null"):
                    try:
                        us_by_fy[int(fy)] = round(float(v) / 1e12, 3)  # → $T
                    except (ValueError, TypeError):
                        pass
        # 当前进行中财年: 拿最新月的 fytd 累计(partial)
        r2 = requests.get(
            "https://api.fiscaldata.treasury.gov/services/api/fiscal_service/v1/accounting/mts/mts_table_5",
            params={"filter": "line_code_nbr:eq:5691",
                    "fields": "record_date,record_fiscal_year,current_fytd_net_outly_amt",
                    "sort": "-record_date", "page[size]": "1"},
            timeout=45)
        latest_partial = None
        if r2.status_code == 200 and r2.json().get("data"):
            row = r2.json()["data"][0]
            fy = row.get("record_fiscal_year")
            v = row.get("current_fytd_net_outly_amt")
            rd = row.get("record_date")
            if fy and v not in (None, "", "null"):
                try:
                    ify = int(fy)
                    # 只有该财年还没有 9 月末完整值时才算 partial
                    if ify not in us_by_fy:
                        latest_partial = (ify, round(float(v) / 1e12, 3), rd)
                except (ValueError, TypeError):
                    pass
        # 组装美国(近10财年 + 当前 partial)
        us_list = []
        for fy in sorted(us_by_fy):
            us_list.append({"fy": fy, "value_t": us_by_fy[fy], "status": "confirmed"})
        if latest_partial:
            us_list.append({"fy": latest_partial[0], "value_t": latest_partial[1],
                            "status": "partial", "as_of": latest_partial[2]})
        # 只保留近 10 年
        out["us"] = us_list[-11:]
        if latest_partial:
            out["as_of"] = latest_partial[2]
        elif us_by_fy:
            out["as_of"] = f"FY{max(us_by_fy)}"
    except Exception as e:
        out["status"] = f"美国MTS错误:{e}"

    # ── 日本: 一般会计当初予算(官方静态权威值) ──
    for fy in sorted(JP_GENERAL_BUDGET):
        out["jp"].append({"fy": fy, "value_oku": JP_GENERAL_BUDGET[fy], "status": "confirmed"})
    out["jp"] = out["jp"][-11:]

    return out


def fetch_maturing_treasury(cache_path=None):
    """私营部门(含Fed)持有的、1年内到期需展期的【可交易国债】规模, 月末采样。
    ★口径: MSPD table_3(逐券明细) 按 maturity_date 筛"距 record_date ≤366天"的 Marketable 券加总 outstanding。
    ★注意: 该口径为总量(含 Fed SOMA + 私营), 未单独剔除 Fed 持仓(dashboard 明确标注)。
    返回 {history_long[(YYYY-MM,$T)] 2001起, history_recent[(YYYY-MM,$T)] 近两年, value, as_of, status, source}。
    有本地缓存(data/maturing_treasury.json)则增量更新, 只补最新缺失月, 避免每次全量拉。"""
    import os, json, requests
    from datetime import datetime, timedelta
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "maturing_treasury.json")
    BASE = ("https://api.fiscaldata.treasury.gov/services/api/fiscal_service"
            "/v1/debt/mspd/mspd_table_3_market")
    FIELDS = "record_date,maturity_date,outstanding_amt,issued_amt,redeemed_amt,security_type_desc"

    # 1) 读缓存
    cached = {}
    if os.path.exists(cache_path):
        try:
            cj = json.load(open(cache_path))
            cached = {d: v for d, v in cj.get("monthly", [])}
        except Exception:
            cached = {}

    # 2) 按年拉(增量: 已缓存年份的所有月都齐则跳过, 只拉当前年+缺失年)
    # ★铁律: 分页任一页失败 → 整年返回 None(调用方跳过该年, 保留缓存真值),
    #   绝不用残缺分页数据算出偏小值去覆盖真实历史(2024=10127行/2025=10372行必然翻页)。
    def _pull_year(year):
        out = {}
        page = 1
        got_rows = 0
        total_count = None
        while True:
            # ★Treasury API: filter 的 : 和 , 是语法字符不可编码; 仅 page[] 方括号需编码; page size 上限 10000。
            url = (f"{BASE}?filter=record_date:gte:{year}-01-01,record_date:lte:{year}-12-31"
                   f"&fields={FIELDS}&page%5Bsize%5D=10000&page%5Bnumber%5D={page}")
            try:
                resp = requests.get(url, timeout=40)
                payload = resp.json()
                data = payload.get("data", [])
                if total_count is None:
                    total_count = (payload.get("meta") or {}).get("total-count")
            except Exception as e:
                print(f"[maturing_treasury] {year} 第{page}页拉取失败, 放弃整年(保留缓存真值): {e}")
                return None
            if not data:
                break
            got_rows += len(data)
            # 按 record_date 分组
            by_date = {}
            for r in data:
                by_date.setdefault(r["record_date"], []).append(r)
            for rd, rows in by_date.items():
                out.setdefault(rd, []).extend(rows)
            if len(data) < 10000:
                break
            page += 1
            if page > 50:  # 安全阀: 防 API 异常导致无限翻页
                print(f"[maturing_treasury] {year} 分页超过50页, 异常中止, 放弃整年")
                return None
        # ★完整性校验: 实收行数须等于 API 声明的 total-count, 否则视为残缺, 放弃整年
        if total_count is not None:
            try:
                if int(total_count) != got_rows:
                    print(f"[maturing_treasury] {year} 行数不符 (API声明{total_count} 实收{got_rows}), "
                          f"判定残缺, 放弃整年(保留缓存真值)")
                    return None
            except (TypeError, ValueError):
                pass
        return out

    now = datetime.utcnow()
    start_year = 2001
    monthly = dict(cached)  # YYYY-MM-DD → $T
    # 需要拉的年份: 未缓存足月的年 + 当前年(总是刷新最新)
    cached_years = {}
    for d in cached:
        cached_years[d[:4]] = cached_years.get(d[:4], 0) + 1
    failed_years = []
    for year in range(start_year, now.year + 1):
        y = str(year)
        # 已有>=11个月且非当前年 → 跳过(历史不变)
        if cached_years.get(y, 0) >= 11 and year < now.year:
            continue
        yr_rows = _pull_year(year)
        if yr_rows is None:
            # ★该年拉取残缺/失败: 跳过, 保留缓存中该年的真实历史值, 绝不用残缺值覆盖
            failed_years.append(y)
            continue
        for rd, rows in yr_rows.items():
            val = _mspd_maturing_within_1yr(rd, rows)
            if val is not None:
                monthly[rd] = val

    if not monthly:
        return {"history_long": [], "history_recent": [], "value": None,
                "as_of": None, "status": "MSPD 无数据",
                "source": "US Treasury MSPD table 3 (marketable, maturity≤1yr)"}

    # 3) 排序 + 写缓存
    items = sorted(monthly.items())
    os.makedirs(os.path.dirname(cache_path), exist_ok=True)
    json.dump({"monthly": items, "updated": now.strftime("%Y-%m-%d"),
               "source": "US Treasury MSPD table 3 (marketable, maturity within 1yr, 含Fed+私营)"},
              open(cache_path, "w"), ensure_ascii=False, indent=1)

    # 4) 长图(全部, YYYY-MM) + 近两年图
    hist_long = [(d[:7], v) for d, v in items]
    cutoff = (now - timedelta(days=760)).strftime("%Y-%m")
    hist_recent = [(m, v) for m, v in hist_long if m >= cutoff]
    last_d, last_v = items[-1]
    _status = "ok"
    if failed_years:
        _status = f"ok(部分年份拉取失败, 已保留缓存真值: {','.join(failed_years)})"
    return {
        "history_long": hist_long,          # 2001至今月末
        "history_recent": hist_recent,      # 近两年月末
        "value": last_v, "as_of": last_d,
        "status": _status,
        "stale_years": failed_years,
        "source": "US Treasury MSPD table 3 (可交易国债·1年内到期·含Fed+私营)",
    }


def fetch_foreign_custody_ust():
    """抓外国官方托管可流通美债(当前值+周变动+历史序列)。
    ★口径统一走 FRED WMTSECL1(周度, 2002至今活跃, 单一可回溯口径), H.4.1 HTML 仅补总托管。
    返回 {value($T), as_of, wow_delta_bn, wow_pct, total_custody_tn, history[(date,$T)], status, source}。"""
    import re
    import requests
    # === 主口径: FRED WMTSECL1 历史序列 ===
    # 短期图=近12个月; 长期图=近10年(均周度点)
    from datetime import datetime, timedelta
    _start_short = (datetime.utcnow() - timedelta(days=380)).strftime("%Y-%m-%d")   # ~12个月前
    _start_long = (datetime.utcnow() - timedelta(days=3660)).strftime("%Y-%m-%d")   # ~10年前
    hist = _custody_history_fred(start=_start_short)          # 短期(12月)
    hist_long = _custody_history_fred(start=_start_long)       # 长期(10年)
    hist_2008 = _custody_history_fred(start="2007-07-04")      # 2007至今全周期(WMTSECL1序列2002-12建库但2007-07-04才第一个真实非零值, 之前237点为0=统计未开始, 起点取真实数据起点避免贴地假横线)
    fred_val = fred_as_of = fred_wow_bn = fred_wow_pct = None
    if len(hist) >= 2:
        fred_as_of, fred_val = hist[-1]
        prev = hist[-2][1]
        fred_wow_bn = round((fred_val - prev) * 1000, 1)          # $T→$B
        fred_wow_pct = round((fred_val - prev) / prev * 100, 2) if prev else None
    # === 补充: H.4.1 HTML 拿总托管(含机构债/MBS) ===
    total_custody_tn = None
    try:
        r = requests.get(FED_H41_URL,
                         headers={"User-Agent": "EcoVolChecker research (contact ANONYMIZED_EMAIL_ADDRESS_0_2)"},
                         timeout=30)
        if r.status_code == 200:
            txt = re.sub(r"<[^>]+>", " ", r.text)
            txt = re.sub(r"&#xa0;", " ", txt)
            txt = re.sub(r"\s+", " ", txt)
            i = txt.find("held in custody for foreign official")
            if i >= 0:
                seg = txt[i:i + 700]
                mt = re.search(r"international accounts\s+([\d,]+)", seg)
                if mt:
                    total_custody_tn = round(int(mt.group(1).replace(",", "")) / 1e6, 3)
    except Exception:
        pass  # 总托管拿不到不阻塞主口径

    # === 构造返回: FRED 为主口径(历史+当前+周环比), H.4.1 补总托管 ===
    if fred_val is not None:
        return {
            "value": round(fred_val, 3),                    # $T
            "as_of": fred_as_of,
            "wow_delta_bn": fred_wow_bn,                    # 周变动 十亿美元
            "wow_pct": fred_wow_pct,
            "total_custody_tn": total_custody_tn,
            "history": hist,                               # [(date,$T)] 升序, 近12个月周点(短期图)
            "history_long": hist_long,                     # [(date,$T)] 升序, 近10年周点(长期图)
            "history_2008": hist_2008,                     # [(date,$T)] 升序, 2008至今周点(全周期结构图)
            "status": "ok",
            "source": "FRED WMTSECL1 (Fed H.4.1 custody, weekly Wed)",
        }
    return {"value": None, "as_of": None, "history": hist, "history_long": hist_long,
            "history_2008": hist_2008,
            "total_custody_tn": total_custody_tn,
            "status": "FRED WMTSECL1 无数据"}


def fetch_custody_acceleration(weeks=26):
    """外国官方托管美债(Fed H.4.1 WMTSECL1)超短期【加速度】分析。
    ★数据颗粒度=周度(每周三 as-of), 无日度。故 trailing 7/14/28/56 天 ≡ 1/2/4/8 周(诚实标注)。
    加速度 = 二阶差分(变化的变化), 零轴上=流入加速/流出减速, 零轴下=流出加速/流入减速。
      7天(1周)  = v[i] - 2·v[i-1] + v[i-2]
      14天(2周) = v[i] - 2·v[i-2] + v[i-4]
      28天(4周) = v[i] - 2·v[i-4] + v[i-8]
      56天(8周) = v[i] - 2·v[i-8] + v[i-16]
    主平滑指标 = EMA(3周)斜率的差分(压制周度噪声, 拐点更干净)。
    ★fredgraph CSV 端点对 requests 偶发超时 → 优先带 key FRED API(fetch_fred_history)。
    返回 {points:[{date, a7, a14, a28, a56, ema_accel}], asof, unit, source, status, weeks}。
    绝不编: 拉不到→points=[] + status。单位=百万美元($M)。默认 weeks=26(过去约6个月)。
    """
    import datetime as _dt
    # 拉够长的历史(算 EMA + 4周二阶差分需预热, 多取 20 周余量)
    need = weeks + 20
    cosd = (_dt.date.today() - _dt.timedelta(weeks=need + 4)).strftime("%Y-%m-%d")
    pts = []
    # ★通道1: 项目带 key FRED API(fetch_fred_history) —— 现有 custody 板块一直用它, 最稳
    try:
        try:
            from fetchers.fred import fetch_fred_history
        except Exception:
            from src.fetchers.fred import fetch_fred_history
        raw = fetch_fred_history("WMTSECL1", start=cosd)  # 单位: 百万美元
        for d, v in raw:
            try:
                pts.append((str(d).strip(), float(v)))
            except (ValueError, TypeError):
                pass
    except Exception:
        pass
    # 通道2: curl 子进程回退(fredgraph CSV, 该端点在本环境偶发超时)
    if len(pts) < 10:
        import subprocess
        url = f"https://fred.stlouisfed.org/graph/fredgraph.csv?id=WMTSECL1&cosd={cosd}"
        text = ""
        for _attempt in range(3):
            try:
                r = subprocess.run(["curl", "-s", url, "--max-time", "30",
                                    "-H", "User-Agent: Mozilla/5.0"],
                                   capture_output=True, text=True, timeout=38)
                if r.returncode == 0 and r.stdout and "observation_date" in r.stdout:
                    text = r.stdout
                    break
            except Exception:
                pass
        if text:
            pts = []
            for ln in text.strip().splitlines()[1:]:
                if "," not in ln:
                    continue
                d, v = ln.split(",", 1)
                v = v.strip()
                if v and v != ".":
                    try:
                        pts.append((d.strip(), float(v)))
                    except ValueError:
                        pass
    if len(pts) < 10:
        return {"points": [], "asof": None, "unit": "$M",
                "source": "FRED WMTSECL1 (Fed H.4.1 custody, weekly Wed)",
                "status": "拉取失败(端点超时)", "weeks": weeks}
    dates = [p[0] for p in pts]
    v = [p[1] for p in pts]
    n = len(v)
    # EMA(3周) + 斜率(一阶差) + 斜率差分(加速度)
    span = 3
    alpha = 2.0 / (span + 1)
    ema = [v[0]]
    for i in range(1, n):
        ema.append(alpha * v[i] + (1 - alpha) * ema[-1])
    ema_slope = [None] + [ema[i] - ema[i - 1] for i in range(1, n)]
    out = []
    for i in range(n):
        rec = {"date": dates[i]}
        rec["a7"] = round(v[i] - 2 * v[i - 1] + v[i - 2], 0) if i >= 2 else None
        rec["a14"] = round(v[i] - 2 * v[i - 2] + v[i - 4], 0) if i >= 4 else None
        rec["a28"] = round(v[i] - 2 * v[i - 4] + v[i - 8], 0) if i >= 8 else None
        rec["a56"] = round(v[i] - 2 * v[i - 8] + v[i - 16], 0) if i >= 16 else None
        # EMA 斜率差分(平滑加速度)
        rec["ema_accel"] = (round(ema_slope[i] - ema_slope[i - 1], 0)
                            if i >= 2 and ema_slope[i] is not None
                            and ema_slope[i - 1] is not None else None)
        out.append(rec)
    # 只保留最近 weeks 周(有完整加速度的)
    out = [r for r in out if r["a28"] is not None][-weeks:]
    return {"points": out, "asof": dates[-1], "unit": "$M",
            "source": "FRED WMTSECL1 (Fed H.4.1 custody, 周度 as-of Wed)",
            "status": "ok" if out else "无有效加速度点", "weeks": weeks}


def write_custody_notion(cust=None):
    """把外国官方托管美债写入 Notion DB_CUSTODY(周度, 以 as_of 作 title 幂等 upsert)。
    返回写入的 page id 或 None(数据无效/无DB时跳过)。"""
    import config as c
    import notion_writer as nw
    if cust is None:
        cust = fetch_foreign_custody_ust()
    db = c.NOTION_DB.get("custody")
    if not db or cust.get("value") is None or not cust.get("as_of"):
        return None
    props = {
        "托管美债($T)": nw.prop_num(cust.get("value")),
        "周变动($B)": nw.prop_num(cust.get("wow_delta_bn")),
        "周变动(%)": nw.prop_num(cust.get("wow_pct")),
        "托管总额($T)": nw.prop_num(cust.get("total_custody_tn")),
        "数据源": {"rich_text": [{"type": "text", "text": {"content": cust.get("source", "Fed H.4.1")}}]},
    }
    return nw.upsert(db, cust["as_of"], props, title_field="Date")


# ─────────── 分国别持有美债 (US Treasury TIC — Major Foreign Holders) ───────────
# Chao 需求(2026-08): 在"外国官方托管美债"图下方加 日本/中国 各过去10年持有美债折线。
# ★口径区别: 上面 custody 是 NY Fed 托管账户"外国官方合计"; 这里是 TIC 分国别口径
#   (含官方+私人持有, 月度), 权威源=美国财政部 TIC Major Foreign Holders 历史文件。
# 数据源: ticdata.treasury.gov/Publish/mfhhis01.txt (官方历史 TSV, 2000-03 至今, 免key)。
#   ★注意 Publish/mfh.txt 是旧缓存(停在2023-01), 历史+最新都在 mfhhisNN.txt 系列。
# 单位: 十亿美元($B)。绝不编: 抓不到标 status。
TIC_MFH_HIST_URL = "https://ticdata.treasury.gov/Publish/mfhhis01.txt"
# ★2026-08-21 修复(Chao 报"图上说滞后233天,右上角却写每月更新"):
# mfhhis01.txt 是<b>历史存档表</b>, 只到上一自然年年末(2026-07 更新时内容仍止于 2025-12),
# 当年月度必须另取 slt_table5(13 个月滚动窗口)。只用存档表会让数据永远落后半年以上,
# 且 2026 全年真值缺失 —— 属"静默陈旧", 违反项目铁律。两源合并, 当期表优先。
TIC_SLT_T5_URL = ("https://ticdata.treasury.gov/resource-center/"
                  "data-chart-center/tic/Documents/slt_table5.html")

# ★★2026-09-09 重大修正(Chao 追问「欧盟到底哪个央行在增持美债」时挖出):
#
# 此前认定「Table 5 只列前20大, 德/意/荷/西/芬 落在 All Other, 故 EU9 只能退回
# 年度存档表, 滞后至上年末(251 天)」—— <b>这个结论是错的, 因为我只查了一张表</b>。
# 实测 Table 5 确实只有 4/9 国(法/比/卢/爱), 那部分判断没错;
# 但 TIC 另有 <b>Table 3</b>, 含 <b>113 个国家/地区</b>、欧元区 9 国全齐、
# 覆盖 2020-01 → 2026-06, 与日/中完全同频。
#   → EU9 全口径可做到滞后约 2 月, 既不必砍成 4 国, 也不必估算。
#   → 代价: Table 3 只回溯到 2020-01, 2008-2019 长历史仍需年度表拼接。
#
# ★Table 3 还独家提供<b>全球</b>「官方 vs 非官方」拆分(码 99990 / 99991):
#     2020-01  官方 4169.9B  私人 2857.4B  官方占比 59.3%
#     2026-06  官方 3778.1B  私人 5520.8B  官方占比 40.6%
#   全球央行六年净减持约 392B, 私人净增约 2663B —— 这是回答「央行在做什么」的
#   唯一官方口径。**注意: 官方/私人拆分只有全球总量, 113 个国家里没有任何一国
#   带 official 拆分**, 所以「日本央行持有多少美债」TIC 给不出, 不是抓不到而是
#   财政部不公布。要看单国央行行为必须去各国央行自己的资产负债表。
TIC_SLT_T3_URL = "https://ticdata.treasury.gov/Publish/slt_table3.txt"


def _parse_tic_table3(raw):
    """解析 TIC Table 3(全国别月度, 113 个国家/地区) → {name: {'YYYY-MM': $B}}。

    ★2026-09-09 新增。为什么需要它见 TIC_SLT_T3_URL 上方的长注释:
      Table 5 只列前 20 大, 欧元区只有 4/9 国; Table 3 有 113 个条目、9 国全齐。

    文件结构(tab 分隔):
      行7  表头 : Country | Country Code | Date | Holdings | Net U.S. Sales | ...
      行8  机器名: country | country_code | date | for_treas_pos | ...
      数据行    : Austria | 10189 | 2026-06 | 5728 | -39 | 3687 | ...
      第 4 列(idx 3) = for_treas_pos = 总持仓, 单位<b>百万美元</b>

    ★列名必须按表头认, 别按加总猜: 那 3 组数是 总额/长期/短期, <b>不是</b>官方/私人。
      我第一次就把 (1116683, 1046231, 70452) 误当成 总/官方/私人 —— 读了表头才发现
      是 总/长期/短期(1046231+70452=1116683 恰好也吻合, 加总验证救不了这个错)。

    两行全球汇总另存 key: 'OFFICIAL'(码99990) / 'PRIVATE'(码99991)。
    解析失败一律返回空 dict, 由调用方回退, 绝不编造。
    """
    import re
    out = {}
    try:
        for ln in raw.replace("\r", "").split("\n"):
            c = ln.split("\t")
            if len(c) < 4:
                continue
            nm = c[0].strip().strip('"')
            dt = c[2].strip()
            if not re.fullmatch(r"\d{4}-\d{2}", dt):
                continue
            try:
                v = float(c[3].replace(",", "")) / 1000.0     # 百万$ → $B
            except ValueError:
                continue
            if nm.startswith("Of Which: Foreign Official"):
                key = "OFFICIAL"
            elif nm.startswith("Of Which: Foreign Non-Official"):
                key = "PRIVATE"
            elif nm.startswith("Of Which") or nm.startswith("Grand Total"):
                continue
            else:
                key = nm
            out.setdefault(key, {})[dt] = v
    except Exception:
        return {}
    return out


def fetch_tic_official_vs_private():
    """全球「外国官方(央行/主权基金) vs 外国私人」持有美债月度序列。

    ★这是 TIC 唯一的官方/私人拆分, <b>只有全球总量, 不分国别</b>(码 99990/99991)。
      Chao 2026-09-09 问「日本/美国央行和其他机构持有美债有没有区别」——
      有区别且巨大, 但<b>分国拆分财政部不公布</b>, 只能看全球总量 + 各国央行自己的表。

    实测(2026-09-09):
      2020-01  官方 4169.9B  私人 2857.4B  官方占比 59.3%
      2026-06  官方 3778.1B  私人 5520.8B  官方占比 40.6%
      → 全球央行六年净减持 ~392B, 私人净增 ~2663B。
    """
    out = {"status": "未找到", "source": "US Treasury TIC Table 3 (monthly)",
           "official": [], "private": [], "share": []}
    import requests
    try:
        r = requests.get(
            TIC_SLT_T3_URL, timeout=60,
            headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"})
        if r.status_code != 200:
            out["status"] = f"HTTP {r.status_code}"
            return out
        d = _parse_tic_table3(r.text)
    except Exception as e:
        out["status"] = f"抓取失败: {type(e).__name__}"
        return out

    o, p = d.get("OFFICIAL", {}), d.get("PRIVATE", {})
    both = sorted(set(o) & set(p))
    if len(both) < 2:
        out["status"] = "数据不足"
        return out
    out["official"] = [(m, round(o[m], 1)) for m in both]
    out["private"] = [(m, round(p[m], 1)) for m in both]
    out["share"] = [(m, round(100 * o[m] / (o[m] + p[m]), 2)) for m in both]
    f, l = both[0], both[-1]
    out.update({
        "status": "ok",
        "latest_month": l,
        "official_latest": round(o[l], 1),
        "private_latest": round(p[l], 1),
        "share_latest": round(100 * o[l] / (o[l] + p[l]), 1),
        "share_first": round(100 * o[f] / (o[f] + p[f]), 1),
        "first_month": f,
        "official_delta": round(o[l] - o[f], 1),
        "private_delta": round(p[l] - p[f], 1),
        "cadence_note": "月度更新 · 滞后约2月",
    })
    return out


def _parse_tic_slt_table5(raw_html, key):
    """解析 TIC Table 5(当期 13 个月滚动窗口) -> {'YYYY-MM': float}。
    表头形如 ['Country','2026-06','2026-05',...]; 国家行首列为国名。
    解析失败一律返回 {} (由调用方回退到历史存档表), 绝不编造。"""
    import re as _re
    import html as _html

    def _cells(row):
        cs = _re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, _re.S | _re.I)
        return [_html.unescape(_re.sub(r"<[^>]+>", "", c)).replace("\xa0", " ").strip()
                for c in cs]

    rows = _re.findall(r"<tr[^>]*>(.*?)</tr>", raw_html, _re.S | _re.I)
    tab = [_cells(r) for r in rows]
    months, out = [], {}
    for t in tab:
        if t and t[0].strip().lower() == "country":
            months = [c.strip() for c in t[1:]]
            break
    if not months:
        return {}
    # 目标国家在 Table 5 中的行名
    aliases = {"Japan": ["japan"],
               "China": ["china, mainland", "china,mainland", "china mainland"]}
    # ★EU 在 TIC 无合计行, 与 _parse_tic_country 保持完全一致的 9 国口径逐月加总,
    #   否则两源(存档表/当期表)口径不一致会造成序列跳变。
    EU_MEMBERS = {"germany", "france", "italy", "netherlands", "belgium",
                  "luxembourg", "ireland", "spain", "finland"}
    if key == "EU":
        acc = {}
        seen = set()
        for t in tab:
            if not t:
                continue
            nm = t[0].strip().lower()
            if nm in EU_MEMBERS and nm not in seen:
                seen.add(nm)
                for m, v in zip(months, t[1:]):
                    if not _re.fullmatch(r"\d{4}-\d{2}", m or ""):
                        continue
                    v = (v or "").replace(",", "").strip()
                    try:
                        acc[m] = acc.get(m, 0.0) + float(v)
                    except ValueError:
                        continue
        # 成员国缺失过多则判定解析不可靠, 宁可回退存档表也不给半截加总
        if len(seen) < len(EU_MEMBERS):
            return {}
        return acc
    names = aliases.get(key)
    if not names:
        return {}
    for t in tab:
        if not t:
            continue
        if t[0].strip().lower() in names:
            for m, v in zip(months, t[1:]):
                if not _re.fullmatch(r"\d{4}-\d{2}", m or ""):
                    continue
                v = (v or "").replace(",", "").strip()
                try:
                    out[m] = float(v)
                except ValueError:
                    continue
            break
    return out
_TIC_MON = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6,
            "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12}


def _parse_tic_country(raw, country_key):
    """从 TIC mfhhis 文本解析某国月度序列。country_key: 'Japan' / 'China' / 'EU'。
    文件结构: 每个年份块有一"月份行"(Jan..Dec)+"年份行"(Country + 4位年份),
    随后 Japan / 'China, Mainland' 数据行, 按列对齐。返回 {'YYYY-MM': value_$B} 。
    ★'EU' = 欧元区主要成员国加总(德/法/意/荷/比/卢/爱/西/芬), 逐月求和(TIC 无欧盟合计行;
      注: Belgium 含 Euroclear 托管, 属 TIC 口径固有特性)。"""
    import re
    # 欧元区主要成员(精确行名匹配, 避免误伤脚注文本)
    EU_MEMBERS = {"Germany", "France", "Italy", "Netherlands", "Belgium",
                  "Luxembourg", "Ireland", "Spain", "Finland",
                  "Belgium-Luxembourg"}
    # ★2026-09-09 分组(交易员视角: 托管中心 vs 真实需求国信号相反, 见 fetch 内注释)
    #   Belgium-Luxembourg 是早年合并行名, 归入托管中心。
    EU_CUSTODY_M = {"Belgium", "Luxembourg", "Ireland", "Belgium-Luxembourg"}
    EU_REAL_M = {"Germany", "France", "Italy", "Netherlands", "Spain", "Finland"}
    _AGG = {"EU": EU_MEMBERS, "EU_CUSTODY": EU_CUSTODY_M, "EU_REAL": EU_REAL_M}
    lines = raw.replace("\r", "").split("\n")
    series = {}          # 单国: {月:值}
    eu_series = {}       # EU: {月: 累加值}
    cur_months = cur_years = None
    for l in lines:
        cells = [c.strip() for c in l.split("\t")]
        # 月份行: 至少含 6 个月份缩写
        monpos = {i: _TIC_MON[c] for i, c in enumerate(cells) if c in _TIC_MON}
        if len(monpos) >= 6:
            cur_months, cur_years = monpos, None
            continue
        # 年份行: 第0列 == 'Country', 其余为4位年份
        if cells and cells[0] == "Country":
            cur_years = {i: int(c) for i, c in enumerate(cells) if re.fullmatch(r"\d{4}", c)}
            continue
        name = cells[0].strip('"').strip() if cells else ""
        if country_key in _AGG:
            matched = name in _AGG[country_key]
        elif country_key == "Japan":
            matched = (name == "Japan")
        elif country_key == "China":
            matched = name.startswith("China")
        else:
            matched = False
        if matched and cur_months and cur_years:
            for i, mm in cur_months.items():
                yy = cur_years.get(i)
                if not yy or i >= len(cells):
                    continue
                v = cells[i].replace(",", "").strip()
                try:
                    val = float(v)
                except ValueError:
                    continue
                mk = f"{yy:04d}-{mm:02d}"
                if country_key in _AGG:
                    eu_series[mk] = eu_series.get(mk, 0.0) + val
                else:
                    series[mk] = val
    return eu_series if country_key in _AGG else series


def fetch_country_ust_holdings(years=10):
    """抓 日本 / 中国 / 欧盟(欧元区加总) 分国别持有美债的月度序列(近 N 年 + 2008起长历史)。
    源: 美国财政部 TIC Major Foreign Holders 历史文件(官方, 月度, $B, 覆盖2000-至今)。
    返回 {"Japan":{...}, "China":{...}, "EU":{...}} 每个:
      {name, flag, series:[(YYYY-MM, $B),...]升序(近N年), series_long:[...](2008起长历史),
       latest:(month,val), first:(month,val), delta_bn, delta_pct, high, low, status, source}。
    ★EU = 欧元区主要成员(德/法/意/荷/比/卢/爱/西/芬)逐月加总(TIC 无欧盟合计行; Belgium 含 Euroclear 托管)。
    绝不编: 解析失败留空 series + status。"""
    import requests
    import datetime
    out = {}
    meta = {"Japan": ("日本", "🇯🇵"), "China": ("中国大陆", "🇨🇳"),
            "EU": ("欧元区合计(9国)", "🇪🇺"),
            # ★分组见下方 _EU_GROUPS 注释: 托管中心与真实需求国的信号含义相反, 必须分开
            "EU_CUSTODY": ("欧元区托管中心", "🏦"),
            "EU_REAL": ("欧元区真实需求国", "🇪🇺")}
    cutoff = (datetime.date.today().replace(day=1) -
              datetime.timedelta(days=int(years * 365.25) + 45)).strftime("%Y-%m")
    long_cutoff = "2008-01"  # 2008 起长历史
    raw = None
    try:
        r = requests.get(TIC_MFH_HIST_URL,
                         headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"},
                         timeout=40)
        if r.status_code == 200 and "MAJOR FOREIGN HOLDERS" in r.text:
            raw = r.text
    except Exception:
        raw = None
    # ★当期表(slt_table5): 补 2026 年以来的月度, 覆盖存档表的滞后。抓不到就仅用存档表。
    cur_html = None
    try:
        r5 = requests.get(TIC_SLT_T5_URL,
                          headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"},
                          timeout=40)
        if r5.status_code == 200 and "Major Foreign Holders" in r5.text:
            cur_html = r5.text
    except Exception:
        cur_html = None
    # ★★Table 3(全国别月度, 113 个条目): 2026-09-09 新增, 优先级<b>高于</b> Table 5。
    #   Table 5 只列前20大 → 欧元区只有 4/9 国(法比卢爱), 逼得 EU 退回年度存档表,
    #   滞后 251 天。Table 3 九国全齐且到 2026-06 —— EU 全口径终于能与日/中同频。
    #   代价: Table 3 只回溯 2020-01, 更早的长历史仍由 mfhhis01 年度表提供(两者互补)。
    t3 = {}
    try:
        r3 = requests.get(
            TIC_SLT_T3_URL,
            headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"},
            timeout=60)
        if r3.status_code == 200 and "Country" in r3.text:
            t3 = _parse_tic_table3(r3.text)
    except Exception:
        t3 = {}
    # Table 3 里的行名 → 本函数的 key
    _T3_NAME = {
        "Japan": ["Japan"],
        "China": ["China, Mainland"],
        "EU": ["Germany", "France", "Italy", "Netherlands", "Belgium",
               "Luxembourg", "Ireland", "Spain", "Finland"],
        "EU_CUSTODY": ["Belgium", "Luxembourg", "Ireland"],
        "EU_REAL": ["Germany", "France", "Italy", "Netherlands", "Spain", "Finland"],
    }
    # ★★欧元区分组(Chao 2026-09-09: 「你把我当做一个国债交易员, 按照我需要关注的点,
    #   把欧盟的数据重新体现一下…现在滞后 251 天的数据全部拢在一起, 这是不可以接受的」)
    #
    #   交易员视角下, 欧元区 9 国<b>不是同一类持有人</b>, 拢成一个数会抹掉信号:
    #     - 托管中心(比/卢/爱): 全世界的钱寄存在 Euroclear/Clearstream, 持仓变化
    #       主要反映<b>第三国资金流向</b>(2014 年比利时暴涨被广泛认为是中国借道),
    #       与欧洲自身需求无关。卢森堡持仓 434B 是本国 GDP 的 4.3 倍即为明证。
    #     - 真实需求国(德/法/意/荷/西/芬): 持仓量级与经济体量匹配, 反映欧洲机构
    #       (保险/养老金/基金)的真实配置行为, 是美欧利差与对冲成本的函数。
    #   两组混在一起, 就无法区分「欧洲在买美债」和「别人借道欧洲在买美债」。
    _EU_GROUPS = {
        "EU_CUSTODY": (["Belgium", "Luxembourg", "Ireland"],
                       "欧元区托管中心", "🏦",
                       "比利时(Euroclear)/卢森堡(Clearstream)/爱尔兰 —— "
                       "第三国资金借道寄存, 变化多反映境外资金流向而非欧洲自身需求"),
        "EU_REAL": (["Germany", "France", "Italy", "Netherlands", "Spain", "Finland"],
                    "欧元区真实需求国", "🇪🇺",
                    "德/法/意/荷/西/芬 —— 持仓量级与经济体量匹配, "
                    "反映欧洲保险/养老金/基金的真实配置(美欧利差与对冲成本驱动)"),
    }

    def _t3_series(k):
        """从 Table 3 取某 key 的月度序列(EU 为 9 国逐月加总)。缺任一国的月份不计入,
        避免用不完整的加总冒充全口径 —— 宁可少一个月, 不出一个错的数。"""
        names = _T3_NAME.get(k) or []
        got = [t3.get(n) or {} for n in names]
        if not all(got):
            return {}
        if len(got) == 1:
            return dict(got[0])
        common = set(got[0])
        for g in got[1:]:
            common &= set(g)
        return {m: sum(g[m] for g in got) for m in common}

    for key, (zh, flag) in meta.items():
        if not raw:
            out[key] = {"name": zh, "flag": flag, "series": [], "series_long": [],
                        "status": "未找到", "source": "US Treasury TIC MFH"}
            continue
        s = _parse_tic_country(raw, key)
        cur_ok = False
        if cur_html:
            cur = _parse_tic_slt_table5(cur_html, key)
            if cur:
                s.update(cur)          # 当期表优先(含官方修订值)
                cur_ok = True
        s3 = _t3_series(key)
        if s3:
            s.update(s3)               # ★Table 3 最优先: 全国别 + 当期
            cur_ok = True
        pts = sorted((m, round(v, 1)) for m, v in s.items() if m >= cutoff)
        pts_long = sorted((m, round(v, 1)) for m, v in s.items() if m >= long_cutoff)
        if len(pts) < 2:
            out[key] = {"name": zh, "flag": flag, "series": pts, "series_long": pts_long,
                        "status": "数据不足", "source": "US Treasury TIC MFH"}
            continue
        first_m, first_v = pts[0]
        last_m, last_v = pts[-1]
        vals = [v for _, v in pts]
        out[key] = {
            "name": zh, "flag": flag, "series": pts, "series_long": pts_long,
            "latest": (last_m, last_v), "first": (first_m, first_v),
            "delta_bn": round(last_v - first_v, 1),
            "delta_pct": round((last_v - first_v) / first_v * 100, 1) if first_v else None,
            "high": round(max(vals), 1), "low": round(min(vals), 1),
            "status": "ok",
            "source": "US Treasury TIC Major Foreign Holders (monthly)",
            # ★2026-09-09 起改用 Table 3(113 国全国别月度), EU 不再受前20大限制。
            #   旧注释称「EU 只能用年度存档表滞后半年」——<b>已过时</b>, 保留在
            #   TIC_SLT_T3_URL 上方作为错误记录。
            "lag_reason": (None if cur_ok else
                           "TIC 当期表(Table 5)仅列前20大持有国，欧元区多数成员国归入 All Other，"
                           "无法还原9国加总口径；故欧盟数据只能取自年度存档表，滞后至上一年年末。"),
            "cadence_note": ("月度更新 · 滞后约2月" if cur_ok else "年度存档表 · 滞后至上年末"),
            # ★★溯源(Chao 2026-09-09: 「我所有的数据都需要可以回溯源头,
            #   以及你抓取到这些数据的路径」)—— 每个数据点都要能查到出处。
            "members": _T3_NAME.get(key) or [],
            "provenance": {
                "primary": TIC_SLT_T3_URL if s3 else TIC_SLT_T5_URL,
                "primary_desc": ("TIC Table 3「U.S. Treasury Securities Held by "
                                 "Foreign Residents」全国别月度表(113 个国家/地区)"
                                 if s3 else
                                 "TIC Table 5「Major Foreign Holders」当期表(前20大)"),
                "history": TIC_MFH_HIST_URL,
                "history_desc": "TIC mfhhis01「Major Foreign Holders」年度存档表(2000-至今)",
                "column": "for_treas_pos(第4列, 总持仓, 百万美元→$B)",
                "agg": ("逐月加总: " + " + ".join(_T3_NAME.get(key) or [])
                        if len(_T3_NAME.get(key) or []) > 1 else "单一国家行"),
                "note": ("TIC 为<b>托管口径</b>: 统计证券登记在哪国托管账户，"
                         "非最终受益人。美国财政部原文: 'cannot attribute holdings "
                         "with complete accuracy … held in a custodial account in a "
                         "third country, the true ownership will not be reflected'。"),
            },
        }
    return out


# ─────────── 美国国债拍卖 (Treasury Auctions) ───────────
TREASURY_AUCTION_API = ("https://api.fiscaldata.treasury.gov/services/api/"
                        "fiscal_service/v1/accounting/od/auctions_query")
# 关注的关键券种(Note/Bond, 反映市场对中长债需求; Bill 是货币工具口径不同, 不纳入)
AUCTION_KEY_TERMS = ["2-Year", "3-Year", "5-Year", "7-Year", "10-Year", "20-Year", "30-Year"]


def _auc_num(x):
    """把 API 字符串数字转 float, 'null'/空 → None。"""
    if x in (None, "", "null"):
        return None
    try:
        return float(x)
    except (ValueError, TypeError):
        return None


def fetch_treasury_auctions(history_days=420):
    """抓美国财政部国债拍卖数据(fiscaldata 官方 API, 免费无 key)。
    每个关键券种(2/3/5/7/10/20/30Y Note&Bond): 最新 + 过去3次(共4次) + 下次拍卖日程。
    返回 {terms: {term: {history:[...], next:{...}|None}}, upcoming:[...], as_of, status, source}。
    每次拍卖字段: auction_date, security_type, offering_bn, accepted_bn, bid_to_cover,
    high_yield, indirect_pct(间接投标占比=外国央行代理需求), issue_date。"""
    import requests
    import datetime
    from collections import defaultdict
    start = (datetime.date.today() - datetime.timedelta(days=history_days)).isoformat()
    today = datetime.date.today()
    try:
        r = requests.get(TREASURY_AUCTION_API,
                         params={"sort": "-auction_date", "page[size]": "800",
                                 "filter": f"auction_date:gte:{start}"},
                         timeout=50)
        if r.status_code != 200:
            return {"terms": {}, "upcoming": [], "as_of": None, "status": f"HTTP {r.status_code}"}
        rows = r.json().get("data", [])
    except Exception as e:
        return {"terms": {}, "upcoming": [], "as_of": None, "status": f"错误:{e}"}

    def pack(x):
        offer = _auc_num(x.get("offering_amt"))
        acc = _auc_num(x.get("total_accepted"))
        ind = _auc_num(x.get("indirect_bidder_accepted"))
        # 间接投标占比 = 间接中标 / 总中标(近似外国央行/官方代理需求份额)
        ind_pct = round(ind / acc * 100, 1) if (ind and acc) else None
        return {
            "auction_date": x.get("auction_date"),
            "security_type": x.get("security_type"),
            "security_term": x.get("security_term"),
            "offering_bn": round(offer / 1e9, 1) if offer else None,
            "accepted_bn": round(acc / 1e9, 1) if acc else None,
            "bid_to_cover": _auc_num(x.get("bid_to_cover_ratio")),
            "high_yield": _auc_num(x.get("high_yield")),
            "indirect_pct": ind_pct,
            "issue_date": x.get("issue_date"),
            "reopening": x.get("reopening"),
        }

    def _norm_term(t):
        """★2026-10-30 修: 官方 API 对「增发(reopening)」的期限写成 '9-Year 10-Month' /
        '29-Year 10-Month' / '19-Year 10-Month'(剩余期限), 而不是 '10-Year'/'30-Year'/'20-Year'。
        旧代码只认整数写法 → 10-07 的 10Y、10-08 的 30Y 增发全被过滤, 看板拍卖停在 10-06。
        归一: 'N-Year M-Month' → 四舍五入到最近整数年; 只对「整数年+1~11个月」且归一后落在关键券种时生效。"""
        import re as _re
        m = _re.match(r"^(\d+)-Year(?:\s+(\d+)-Month)?$", str(t or "").strip())
        if not m:
            return t
        yrs, mons = int(m.group(1)), int(m.group(2) or 0)
        yrs += 1 if mons >= 6 else 0
        return f"{yrs}-Year"

    done = defaultdict(list)     # term → 已完成(有中标结果)
    future = defaultdict(list)   # term → 未来(auction_date>=today 且未开标)
    for x in rows:
        st, term, ad = x.get("security_type"), _norm_term(x.get("security_term")), x.get("auction_date")
        # ★2026-09-07 修: 剔除 TIPS(通胀保值债券)。
        #   TIPS 的 security_type 同样是 Note/Bond、security_term 同样是 10-Year/30-Year,
        #   靠这两个字段过不掉。但它报的是<b>实际收益率</b>(扣通胀)且发行规模小得多,
        #   混进来会污染名义券曲线 —— 实测两条脏数据:
        #     10-Year 2026-07-23  规模 21.0B  收益率 2.438%  (邻月名义券 42.0B / 4.68%)
        #     30-Year 2026-02-19  规模  9.0B  收益率 2.473%  (邻月名义券 25.0B / 5.22%)
        #   官方 API 有专用标识字段 inflation_index_security = Yes/No, 用它精确过滤
        #   (旁证: TIPS 行才有 cpi_base_reference_period / index_ratio_on_issue_date)。
        if str(x.get("inflation_index_security") or "").strip().lower() in ("yes", "y", "true"):
            continue
        if st not in ("Note", "Bond") or term not in AUCTION_KEY_TERMS or not ad:
            continue
        d = datetime.date.fromisoformat(ad)
        has_result = x.get("high_yield") not in (None, "null", "")
        if d >= today and not has_result:
            future[term].append(pack(x))
        elif has_result:
            done[term].append(pack(x))

    terms = {}
    latest_date = None
    for term in AUCTION_KEY_TERMS:
        hist = sorted(done[term], key=lambda a: a["auction_date"], reverse=True)[:4]  # 最新+过去3
        nxt = sorted(future[term], key=lambda a: a["auction_date"])
        if not hist and not nxt:
            continue
        if hist:
            ld = hist[0]["auction_date"]
            if latest_date is None or ld > latest_date:
                latest_date = ld
        terms[term] = {"history": hist, "next": nxt[0] if nxt else None}

    # 全局最近一次拍卖日程(所有关键券种混排取最近未来)
    allf = [f for t in future for f in future[t]]
    upcoming = sorted(allf, key=lambda a: a["auction_date"])[:6]
    return {
        "terms": terms,
        "upcoming": upcoming,
        "as_of": latest_date,
        "status": "ok" if terms else "无拍卖数据",
        "source": "US Treasury fiscaldata.treasury.gov auctions_query",
    }


def write_auctions_notion(auc=None):
    """把国债拍卖写入 Notion DB_AUCTIONS(每次拍卖一行, cusip+日期作幂等 key)。
    以 '券种 拍卖日' 作 title。返回写入行数。"""
    import config as c
    import notion_writer as nw
    if auc is None:
        auc = fetch_treasury_auctions()
    db = c.NOTION_DB.get("auctions")
    if not db or auc.get("status") != "ok":
        return 0
    n = 0
    for term, blk in auc["terms"].items():
        for a in blk["history"]:
            title = f"{a['security_type']} {term} {a['auction_date']}"
            props = {
                "券种": {"select": {"name": f"{a['security_type']} {term}"}},
                "拍卖日": {"date": {"start": a["auction_date"]}},
                "发行规模($B)": nw.prop_num(a.get("offering_bn")),
                "中标额($B)": nw.prop_num(a.get("accepted_bn")),
                "中标率(BTC)": nw.prop_num(a.get("bid_to_cover")),
                "最高中标收益率(%)": nw.prop_num(a.get("high_yield")),
                "间接投标占比(%)": nw.prop_num(a.get("indirect_pct")),
            }
            try:
                nw.upsert(db, title, props, title_field="Auction")
                n += 1
            except Exception:
                pass
    return n


# ─────────── 货币供应量 M0/M1/M2 ───────────
MONEY_SUPPLY_OVERRIDE = os.path.join(os.path.dirname(__file__), "..", "data", "money_supply_override.json")
# FRED fredgraph.csv 直连(免 key), 取最后一行=最新
_FRED_MS = {"m1": "M1SL", "m2": "M2SL", "m0": "BOGMBASE"}


def _fred_latest(series_id):
    """curl FRED fredgraph.csv 取最新 (date, value)。失败返回 (None, None)。"""
    import requests
    try:
        r = requests.get(f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={series_id}", timeout=20)
        if r.status_code != 200:
            return None, None
        lines = [ln for ln in r.text.strip().splitlines() if "," in ln]
        # 从末尾找最后一个有数值的行(FRED 用 '.' 表示缺失)
        for ln in reversed(lines[1:]):
            d, v = ln.split(",", 1)
            v = v.strip()
            if v and v != ".":
                try:
                    return d.strip()[:7], round(float(v), 1)
                except ValueError:
                    continue
        return None, None
    except Exception:
        return None, None


def fetch_money_supply(to_usd=True):
    """三国货币供应量 M0/M1/M2(+日本 M3)。
    US: FRED 直连(自动最新); JP/CN: 无稳定免费 API→读 money_supply_override.json(weekly cron agent 用官方源更新)。
    to_usd=True(默认): 三国统一折算成 $B(十亿美元)横向可比,汇率走 FRED DEXJPUS/DEXCHUS(与央行资负表同源);
      本币原值保留在 orig_m0/orig_m1/orig_m2/orig_m3 + orig_unit 可溯源。
    返回 {"US":{...}, "JP":{...}, "CN":{...}}, 每国:
      {name, flag, unit, as_of, m0, m0_label, m1, m2, (m3), source, (orig_*), (_fx)}
    绝不编: US 抓不到该项留 None; JP/CN 缺 override 文件则该国留空 status。折不动(缺汇率)则保留本币并标 orig_unit。"""
    import json
    out = {}

    # 美国: FRED 直连. M0 无官方概念→用基础货币 BOGMBASE 代理. 本身就是 $B, 无需折算
    us = {"name": "美国 Fed", "flag": "🇺🇸", "unit": "$B",
          "m0_label": "基础货币 Monetary Base",
          "source": "FRED (M1SL/M2SL/BOGMBASE), 月度"}
    as_ofs = []
    for k, sid in _FRED_MS.items():
        d, v = _fred_latest(sid)
        us[k] = v
        if d:
            as_ofs.append(d)
    us["as_of"] = max(as_ofs) if as_ofs else None
    if us.get("m2") is None and us.get("m1") is None:
        us["status"] = "未找到"
    out["US"] = us

    # 日本 / 中国: 读 override(本币)
    try:
        ov = json.load(open(MONEY_SUPPLY_OVERRIDE))
    except Exception:
        ov = {}
    jp_meta = {"name": "日本 BoJ", "flag": "🇯🇵"}
    cn_meta = {"name": "中国 PBoC", "flag": "🇨🇳"}
    for cc, meta in (("JP", jp_meta), ("CN", cn_meta)):
        blk = ov.get(cc)
        if isinstance(blk, dict) and (blk.get("m2") is not None or blk.get("m1") is not None):
            row = dict(meta)
            row.update({k: blk.get(k) for k in ("unit", "as_of", "m0", "m0_label", "m1", "m2", "m3", "source")})
            out[cc] = row
        else:
            out[cc] = dict(meta, status="未找到")

    # ── 统一折算成 $B(十亿美元)横向可比 ──
    if to_usd:
        try:
            from fetchers.fred import fetch_fred_latest
            jpy, _ = fetch_fred_latest("DEXJPUS")   # 日元/美元(如 147.5)
            cny, _ = fetch_fred_latest("DEXCHUS")   # 人民币/美元(如 7.15)
            # 万亿本币 → $B: (值×1e12 本币) / 汇率 / 1e9 = 值×1000/汇率
            def _to_usd_row(row, rate):
                if not rate or row.get("status") == "未找到":
                    return
                factor = 1000.0 / rate
                orig_unit = row.get("unit")
                for k in ("m0", "m1", "m2", "m3"):
                    if row.get(k) is not None:
                        row[f"orig_{k}"] = row[k]           # 保留本币原值
                        row[k] = round(row[k] * factor, 1)  # 折算 $B
                if orig_unit:
                    row["orig_unit"] = orig_unit
                row["unit"] = "$B"
            if "JP" in out and jpy:
                _to_usd_row(out["JP"], jpy)
            if "CN" in out and cny:
                _to_usd_row(out["CN"], cny)
            out["_fx"] = {"USDJPY": jpy, "USDCNY": cny}
        except Exception as e:
            out["_fx_error"] = str(e)
    return out


def fetch_m2_history(months=126, to_usd=True):
    """三国 M2 月度历史序列(默认126月≈10.5年),用于折线图。
    ★折 $B 用【当天 crawl 的最新美元汇率】统一折算全序列(Chao 2026-08 要求):
      不再用各月历史汇率, 而是抓当日最新 DEXJPUS/DEXCHUS, 整条序列都乘同一汇率。
      → 曲线形状纯粹反映各国本币 M2 增长(剥离汇率波动), 反映真实"放水力度"横向对比。
    源: US=FRED M2SL / JP=BOJ stat-search CSV(第10列亿円,Shift_JIS) / CN=东方财富 datacenter API(亿元)。
    返回 {"US":{unit,points:[{date,value},...],orig_unit,latest,source,fx}, "JP":{...}, "CN":{...}}。
    绝不编: 某国抓不到→该国 points=[] + status。汇率抓不到→退回内置 fallback 并标 fx_note。"""
    import requests
    from datetime import datetime, timedelta
    cosd = (datetime.utcnow() - timedelta(days=int(months * 30.5) + 40)).strftime("%Y-%m-01")

    def _fred_series(sid, start):
        """拉 FRED 月度全序列 [(YYYY-MM, float)],升序。"""
        try:
            r = requests.get(f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={sid}&cosd={start}", timeout=25)
            if r.status_code != 200:
                return []
            out = []
            for ln in r.text.strip().splitlines()[1:]:
                if "," not in ln:
                    continue
                d, v = ln.split(",", 1)
                v = v.strip()
                if v and v != ".":
                    try:
                        out.append((d.strip()[:7], float(v)))
                    except ValueError:
                        pass
            return out
        except Exception:
            return []

    # ★当天最新汇率(全序列统一折算, 不用各月历史汇率)
    jpy_now = cny_now = None
    try:
        from fetchers.fred import fetch_fred_latest
        jpy_now, _ = fetch_fred_latest("DEXJPUS")   # 日元/美元
        cny_now, _ = fetch_fred_latest("DEXCHUS")   # 人民币/美元
    except Exception:
        pass
    # fallback: 若 fetch_fred_latest 不可用, 从 fredgraph 序列取最后一个有效值
    if not jpy_now:
        _js = _fred_series("DEXJPUS", cosd)
        jpy_now = _js[-1][1] if _js else 157.54
    if not cny_now:
        _cs = _fred_series("DEXCHUS", cosd)
        cny_now = _cs[-1][1] if _cs else 6.7474

    out = {}
    # ── 美国: FRED M2SL, 本身 $B ──
    us_pts = _fred_series("M2SL", cosd)
    out["US"] = {
        "name": "美国 Fed", "flag": "🇺🇸", "unit": "$B",
        "orig_unit": "$B",
        "points": [{"date": d, "value": round(v, 1)} for d, v in us_pts[-months:]],
        "source": "FRED M2SL 月度",
    }

    # ── 日本: BOJ stat-search CSV(Shift_JIS, 第10列 M2存量亿円) ──
    jp_pts = []
    try:
        r = requests.get("https://www.stat-search.boj.or.jp/ssi/mtshtml/csv/md02_m_1_en.csv", timeout=30)
        if r.status_code == 200:
            text = r.content.decode("shift_jis", errors="ignore")
            for ln in text.splitlines():
                cols = ln.split(",")
                if len(cols) < 10:
                    continue
                dt = cols[0].strip()
                if "/" not in dt:
                    continue
                y, m = dt.split("/")[:2]
                if not (y.isdigit() and m.isdigit()):
                    continue
                raw = cols[9].strip()
                try:
                    oku = float(raw)  # 亿円
                except ValueError:
                    continue
                mon = f"{y}-{int(m):02d}"
                trillion_yen = oku / 10000.0  # 亿円→万亿円
                jp_pts.append((mon, trillion_yen))
    except Exception:
        pass
    jp_pts = jp_pts[-months:]
    out["JP"] = {
        "name": "日本 BoJ", "flag": "🇯🇵", "unit": "$B" if to_usd else "万亿円",
        "orig_unit": "万亿円",
        "points": [], "source": "BOJ Time-Series (Money Stock M2 存量, 月度)",
    }
    for mon, tri in jp_pts:
        val = tri
        if to_usd:
            val = round(tri * 1000.0 / jpy_now, 1)  # 万亿円→$B, 全序列用当天最新汇率
        out["JP"]["points"].append({"date": mon, "value": round(val, 1), "orig": round(tri, 1)})

    # ── 中国: 东方财富 datacenter API(亿元) ──
    cn_pts = []
    try:
        url = ("https://datacenter-web.eastmoney.com/api/data/v1/get?columns=TIME,BASIC_CURRENCY"
               "&pageSize=400&pageNumber=1&sortColumns=REPORT_DATE&sortTypes=1&reportName=RPT_ECONOMY_CURRENCY_SUPPLY")
        r = requests.get(url, headers={"Referer": "https://data.eastmoney.com/", "User-Agent": "Mozilla/5.0"}, timeout=25)
        if r.status_code == 200:
            rows = (r.json().get("result") or {}).get("data") or []
            for row in rows:
                t = str(row.get("TIME", ""))  # "2026年06月份"
                m2 = row.get("BASIC_CURRENCY")
                if m2 is None:
                    continue
                import re as _re
                mt = _re.search(r"(\d{4})\D+(\d{1,2})", t)
                if not mt:
                    continue
                mon = f"{mt.group(1)}-{int(mt.group(2)):02d}"
                wan_yi = float(m2) / 10000.0  # 亿元→万亿元
                cn_pts.append((mon, wan_yi))
    except Exception:
        pass
    cn_pts = cn_pts[-months:]
    out["CN"] = {
        "name": "中国 PBoC", "flag": "🇨🇳", "unit": "$B" if to_usd else "万亿元",
        "orig_unit": "万亿元",
        "points": [], "source": "东方财富 datacenter (PBoC 口径 M2, 月度)",
    }
    for mon, wy in cn_pts:
        val = wy
        if to_usd:
            val = round(wy * 1000.0 / cny_now, 1)  # 万亿元→$B, 全序列用当天最新汇率
        out["CN"]["points"].append({"date": mon, "value": round(val, 1), "orig": round(wy, 1)})

    # 记录本次折算所用的当天汇率(可溯源, 供 dashboard 说明)
    if to_usd:
        out["_fx"] = {"USDJPY": round(jpy_now, 4), "USDCNY": round(cny_now, 4)}

    for cc in ("US", "JP", "CN"):
        blk = out[cc]
        if blk["points"]:
            blk["latest"] = blk["points"][-1]
        else:
            blk["status"] = "未找到"
    return out


def write_money_supply_notion(ms=None):
    """把三国货币供应量写入 Notion DB_MONEY_SUPPLY(每国一行, '国家 as_of' 作 title 幂等)。
    返回写入行数。DB id 缺失或无数据则返回 0。"""
    import config as c
    import notion_writer as nw
    db = c.NOTION_DB.get("money_supply")
    if not db:
        return 0
    ms = ms or fetch_money_supply()
    n = 0
    for cc in ("US", "JP", "CN"):
        d = ms.get(cc) or {}
        if d.get("status") == "未找到" or (d.get("m2") is None and d.get("m1") is None):
            continue
        title = f"{d.get('name', cc)} {d.get('as_of', '')}".strip()
        props = {
            "国家": {"select": {"name": d.get("name", cc)}},
            "口径日期": {"rich_text": [{"text": {"content": str(d.get("as_of", ""))}}]},
            "单位": {"select": {"name": d.get("unit", "")}},
            "M0/基础货币": nw.prop_num(d.get("m0")),
            "M1": nw.prop_num(d.get("m1")),
            "M2": nw.prop_num(d.get("m2")),
            "M3": nw.prop_num(d.get("m3")),
            "来源": {"rich_text": [{"text": {"content": str(d.get("source", ""))[:300]}}]},
        }
        try:
            nw.upsert(db, title, props, title_field="Country")
            n += 1
        except Exception:
            pass
    return n


def fetch_credit_impulse(years=8):
    """三国 Credit Impulse(信贷脉冲)——中期领先指标。

    定义(BIS/学界标准): Credit Impulse = [新增信贷流量的变化] / GDP
      = Δ(信贷流量)/GDP = 信贷存量的二阶差分 / GDP
      衡量的不是债务总量、也不是新增债务, 而是【新增信贷的加速度】。
      领先实体经济约 6-9 个月; 中国信贷脉冲是全球商品/风险资产最强领先指标之一。

    ★口径: 用 BIS 官方 credit-to-GDP ratio (Q*PAM770A, credit to private
      non-financial sector, % of GDP, 季度) 的【二阶差分】近似 Credit Impulse。
      因 GDP 分母变化远慢于信贷, ratio 二阶差分 ≈ Δ信贷流量/GDP(已验证与
      绝对额法方向/拐点完全一致, 数值差<0.2pp)。用 ratio 法的好处:
        ① 三国统一 BIS 口径, 跨国可比(BIS 专门保证国际可比性)
        ② 都到 2025-10, 避免中国季度 GDP 分母序列缺失(FRED 中国 GDP 停在 2023)
        ③ BIS 官方比值, 权威。
    ★频率: BIS 季度数据, 滞后约 1 季(最新到上一季末)。这是全球唯一跨国可比
      的信贷口径, 无周度版本。定位=季度更新的【中期指标】, 领先性足够。
    绝不编: 某国抓不到→该国 points=[] + status="未找到"。

    返回 {"US":{name,flag,points:[{date,ci}],latest,latest_date,signal,source},
          "CN":{...}, "EA":{...}}。signal: strong_pos/pos/neutral/neg/strong_neg。
    """
    import requests
    from datetime import datetime, timedelta
    cosd = (datetime.utcnow() - timedelta(days=int(years * 365) + 400)).strftime("%Y-%m-01")

    def _fred_ratio(sid):
        """拉 BIS credit-to-GDP ratio 季度序列 [(YYYY-MM-DD, float)] 升序。
        优先用项目带 key 的 FRED API 通道(稳); 失败回退免 key fredgraph CSV。"""
        # 通道1: 带 key API (fetchers.fred.fetch_fred_history) —— 更稳
        try:
            from fetchers.fred import fetch_fred_history
            h = fetch_fred_history(sid, start=cosd)
            if h:
                out = []
                for d, v in h:
                    try:
                        out.append((str(d).strip(), float(v)))
                    except (ValueError, TypeError):
                        pass
                if out:
                    return out
        except Exception:
            pass
        # 通道2: 免 key fredgraph CSV 回退
        try:
            r = requests.get(
                f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={sid}&cosd={cosd}",
                timeout=25, headers={"User-Agent": "Mozilla/5.0"})
            if r.status_code != 200:
                return []
            out = []
            for ln in r.text.strip().splitlines()[1:]:
                if "," not in ln:
                    continue
                d, v = ln.split(",", 1)
                v = v.strip()
                if v and v != ".":
                    try:
                        out.append((d.strip(), float(v)))
                    except ValueError:
                        pass
            return out
        except Exception:
            return []

    def _impulse(ratio):
        """ratio 二阶差分 → Credit Impulse(pp of GDP)。"""
        r = [(d, v) for d, v in ratio if v is not None]
        out = []
        for i in range(2, len(r)):
            d = r[i][0]
            ci = (r[i][1] - r[i - 1][1]) - (r[i - 1][1] - r[i - 2][1])
            out.append({"date": d, "ci": round(ci, 2)})
        return out

    def _signal(ci):
        if ci is None:
            return "unknown"
        if ci >= 0.8:
            return "strong_pos"
        if ci >= 0.15:
            return "pos"
        if ci <= -0.8:
            return "strong_neg"
        if ci <= -0.15:
            return "neg"
        return "neutral"

    spec = {
        "US": ("美国", "🇺🇸", "QUSPAM770A"),
        "CN": ("中国", "🇨🇳", "QCNPAM770A"),
        "EA": ("欧元区", "🇪🇺", "QXMPAM770A"),
        "JP": ("日本", "🇯🇵", "QJPPAM770A"),
    }
    # 长历史(2008至今)序列: 用于粗颗粒度参考图。cosd 提前到 2006(留二阶差分预热)
    cosd_long = "2006-01-01"

    def _fred_ratio_long(sid):
        """拉 2006 起的长序列(供 2008 至今参考图)。
        ★优先带 key FRED API(fetch_fred_history, 稳); fredgraph CSV 端点在本环境偶发超时作回退。"""
        # 通道1: 带 key API
        try:
            from fetchers.fred import fetch_fred_history
            h = fetch_fred_history(sid, start=cosd_long)
            if h:
                out = []
                for d, v in h:
                    try:
                        out.append((str(d).strip(), float(v)))
                    except (ValueError, TypeError):
                        pass
                if out:
                    return out
        except Exception:
            pass
        # 通道2: 免 key fredgraph CSV 回退(带 2 次重试)
        for _attempt in range(2):
            try:
                r = requests.get(
                    f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={sid}&cosd={cosd_long}",
                    timeout=18, headers={"User-Agent": "Mozilla/5.0"})
                if r.status_code != 200:
                    continue
                out = []
                for ln in r.text.strip().splitlines()[1:]:
                    if "," not in ln:
                        continue
                    d, v = ln.split(",", 1)
                    v = v.strip()
                    if v and v != ".":
                        try:
                            out.append((d.strip(), float(v)))
                        except ValueError:
                            pass
                if out:
                    return out
            except Exception:
                continue
        return []

    def _bis_ratio(cc):
        """★2026-10-30 新增: BIS 官方 API 直连 [(YYYY-MM-DD, float)] 升序。
        根因: FRED 镜像的 BIS 信贷序列(Q*PAM770A)停在 2025-Q3(date 2025-10-01), 滞后 394 天,
        而 BIS 官方 stats.bis.org 已有 2026-Q1。同一接口对照: US 总债务 N 口径取到 2026-Q1, 链路可信。
        季度 'YYYY-Qn' → 该季首月 1 日(与 FRED 日期约定一致, 使两通道可合并比较)。"""
        try:
            r = requests.get(
                f"https://stats.bis.org/api/v1/data/WS_TC/Q.{cc}.P.A.M.770.A?format=csv",
                timeout=40, headers={"User-Agent": "Mozilla/5.0"})
            if r.status_code != 200:
                return []
            import csv as _csv, io as _io
            out = []
            for row in _csv.DictReader(_io.StringIO(r.text)):
                tp, v = row.get("TIME_PERIOD", ""), row.get("OBS_VALUE", "")
                if len(tp) == 7 and tp[4:6] == "-Q" and v:
                    out.append((f"{tp[:4]}-{(int(tp[6]) - 1) * 3 + 1:02d}-01", float(v)))
            return sorted(out)
        except Exception:
            return []

    _bis_cc = {"US": "US", "CN": "CN", "EA": "XM", "JP": "JP"}

    def _merge_latest(fred_pts, bis_pts):
        """两通道取并集(同日以 FRED 为准), 让序列延伸到更新的那个。"""
        d = dict(bis_pts)
        d.update(dict(fred_pts))
        return sorted(d.items())

    out = {}
    for cc, (name, flag, sid) in spec.items():
        _bp = _bis_ratio(_bis_cc[cc])
        ratio = _merge_latest(_fred_ratio(sid), _bp)
        pts = _impulse(ratio) if len(ratio) >= 3 else []
        # 长历史(2008起)——粗颗粒度参考
        ratio_long = _merge_latest(_fred_ratio_long(sid), _bp)
        pts_long_all = _impulse(ratio_long) if len(ratio_long) >= 3 else []
        pts_long = [p for p in pts_long_all if p["date"] >= "2008-01-01"]
        if pts:
            latest = pts[-1]["ci"]
            out[cc] = {
                "name": name, "flag": flag, "points": pts,
                "points_long": pts_long,
                "latest": latest, "latest_date": pts[-1]["date"],
                "signal": _signal(latest),
                "source": f"BIS credit-to-GDP ratio (FRED {sid}), 二阶差分",
                "status": "ok",
            }
        else:
            out[cc] = {
                "name": name, "flag": flag, "points": [],
                "points_long": pts_long,
                "latest": None, "latest_date": None, "signal": "unknown",
                "source": f"BIS credit-to-GDP ratio (FRED {sid})",
                "status": "未找到",
            }
    return out


# ─────────── 国债市场压力三联图 (对齐 Morgan Stanley 三图, 过去3年真实公开数据) ───────────
# Chao 需求(2026-08): 参照 MS 报告三图(收益率+波动性/BrokerTec价差/BrokerTec DV01量),
#   用过去3年【真实公开数据】新增三个图, 竖向排列, 数据进 GitHub + Notion。
# ★诚实边界: MS 图2(BrokerTec日内bid-ask)/图3(BrokerTec周度DV01量)是 Morgan Stanley
#   BrokerTec 专有数据, 无免费公开源(已搜索确认只对付费机构客户开放)。故图2/图3 换成
#   主题对齐的免费公开【压力代理】, 并在 dashboard 明确标注"非BrokerTec原指标"。
# 三图逻辑(与 MS 一致): 波动性↑(MOVE) → 市场压力↑(期限溢价/IG OAS) → 风险传导(曲线利差/HY OAS)。
#   图1 收益率+波动性: DGS10/DGS2(FRED) + MOVE指数(yfinance ^MOVE, ICE专有但Yahoo有免费历史)
#   图2 国债市场压力代理: THREEFYTP10(10yr期限溢价,NY Fed ACM) + BAMLC0A0CM(IG企业债OAS)
#   图3 曲线/信用压力代理: T10Y2Y(10Y-2Y曲线利差) + BAMLH0A0HYM2(HY OAS)
# 全部 FRED 带 key API 优先(fredgraph CSV 本环境偶发超时); MOVE 走 yfinance。绝不编: 抓不到标 status。


def _fred_series_hist(sid, start):
    """拉单个 FRED 序列历史, 带 key API 优先, curl CSV 回退。返回 [(date,val),...] 只含有效点。"""
    pts = []
    try:
        try:
            from fetchers.fred import fetch_fred_history
        except Exception:
            from src.fetchers.fred import fetch_fred_history
        raw = fetch_fred_history(sid, start=start)
        for d, v in raw:
            try:
                if v is not None:
                    pts.append((str(d).strip(), float(v)))
            except (ValueError, TypeError):
                pass
    except Exception:
        pass
    if len(pts) < 20:
        import subprocess
        url = f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={sid}&cosd={start}"
        for _attempt in range(3):
            try:
                r = subprocess.run(["curl", "-s", url, "--max-time", "30",
                                    "-H", "User-Agent: Mozilla/5.0"],
                                   capture_output=True, text=True, timeout=38)
                if r.returncode == 0 and r.stdout and "observation_date" in r.stdout:
                    pts = []
                    for ln in r.stdout.strip().splitlines()[1:]:
                        if "," not in ln:
                            continue
                        d, v = ln.split(",", 1)
                        v = v.strip()
                        if v and v != ".":
                            try:
                                pts.append((d.strip(), float(v)))
                            except ValueError:
                                pass
                    if len(pts) >= 20:
                        break
            except Exception:
                pass
    return pts


def _fetch_move_history(start):
    """MOVE 指数(ICE BofA, 专有)历史 —— 优先 Yahoo chart API v8 直连(带 UA, 返完整历史),
    yfinance ^MOVE 作降级(2026-08 起 yfinance 只返当日1点, chart API 仍返完整序列)。
    返回 [(date,val),...]。抓不到返回 []（诚实, 不编）。"""
    import datetime as _dt
    # 主源: Yahoo chart API v8 直连
    try:
        import requests as _rq
        try:
            _sd = _dt.date.fromisoformat(start[:10])
            _rng_days = (_dt.date.today() - _sd).days
        except Exception:
            _rng_days = 1200
        _rng = "5y" if _rng_days > 1095 else ("3y" if _rng_days > 730 else "2y")
        url = ("https://query1.finance.yahoo.com/v8/finance/chart/%5EMOVE"
               f"?range={_rng}&interval=1d")
        resp = _rq.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=25)
        if resp.ok:
            j = resp.json()
            res = (j.get("chart", {}).get("result") or [None])[0]
            if res and res.get("timestamp"):
                ts = res["timestamp"]
                closes = res["indicators"]["quote"][0]["close"]
                out = []
                for t, c in zip(ts, closes):
                    if c is None:
                        continue
                    d = _dt.date.fromtimestamp(t).isoformat()
                    if d >= start[:10]:
                        out.append((d, round(float(c), 2)))
                if len(out) >= 2:
                    return out
    except Exception:
        pass
    # 降级: yfinance
    try:
        import yfinance as yf
    except Exception:
        return []
    try:
        df = yf.download("^MOVE", start=start, interval="1d",
                         progress=False, auto_adjust=False)
        if df is None or len(df) == 0:
            return []
        c = df["Close"]
        try:
            if hasattr(c, "columns"):
                c = c.iloc[:, 0]
        except Exception:
            pass
        c = c.dropna()
        out = []
        for idx, val in c.items():
            try:
                out.append((idx.date().isoformat(), round(float(val), 2)))
            except Exception:
                pass
        return out
    except Exception:
        return []


def _weekly_resample(pts, agg="last"):
    """把日度 [(date,val)] 降采样到周度, 减少折线噪声。agg: last|mean|max|peak。
    ★日期标签用该周【真实观测末日】而非名义周五 —— 名义周五会把当前未完成周
      (如周一 08-17)标成未来日期 08-21, 造成 as-of 显示未来、图表 X 轴外推。
      节假日周同理(周四收盘则标周四), 更贴合真实数据。
    ★agg="peak"(2026-09-25, Chao:「我要看的就是这些突出的抬升点, 确保数据正确」):
      每周保留周末值, 另外若周内日度最高/最低比周末值更极端, 把该极端日也作为独立点保留
      (用它的真实日期)。这样周中尖峰不会被削掉——此前 MOVE 日度高 141.67/139.88 被画成 137.26,
      OFR FSI 日度高 +2.743(2025-04-07)被画成 1.60。图上高/低/现 = 真实日度极值。
    返回 [(date, val)]。"""
    import datetime as _dt
    from collections import defaultdict
    buckets = defaultdict(list)
    for d, v in pts:
        try:
            dt = _dt.date.fromisoformat(d[:10])
        except Exception:
            continue
        # 归到该周的周五(仅作分桶键, 不作对外日期标签)
        friday = dt + _dt.timedelta(days=(4 - dt.weekday()))
        buckets[friday.isoformat()].append((dt.isoformat(), v))
    out = []
    for wk in sorted(buckets):
        rows = sorted(buckets[wk])          # 按真实日期排序
        vs = [v for _, v in rows]
        real_last_date = rows[-1][0]        # ★该周真实观测末日
        if agg == "peak":
            last_v = vs[-1]
            keep = {real_last_date: last_v}
            hi_d, hi_v = max(rows, key=lambda x: x[1])
            lo_d, lo_v = min(rows, key=lambda x: x[1])
            if hi_v > last_v:
                keep[hi_d] = hi_v
            if lo_v < last_v:
                keep[lo_d] = lo_v
            for d_, v_ in sorted(keep.items()):
                out.append((d_, round(v_, 4)))
            continue
        if agg == "mean":
            val = sum(vs) / len(vs)
        elif agg == "max":
            val = max(vs)
        else:
            val = vs[-1]
        out.append((real_last_date, round(val, 4)))
    return out


def fetch_treasury_stress_panels(years=3):
    """国债市场压力三联图(对齐 Morgan Stanley 三图, 过去 N 年真实公开数据, 竖向排列)。
    返回 {panels:[p1,p2,p3], asof, years, status}。每个 panel:
      {id, title, subtitle, note, unit_left, unit_right, series:[{name,color,axis,points:[{date,v}]}], source}
    全部真实公开数据; 图2/图3 是公开压力代理(非 BrokerTec 原指标, note 明确标注)。
    绝不编: 某序列抓不到→该 series.points=[] + 整体 status 反映缺失。周度降采样降噪。"""
    import datetime as _dt
    start = (_dt.date.today() - _dt.timedelta(days=int(365.3 * years) + 10)).isoformat()

    def _mk(pts, weekly=True, agg="last"):
        if not pts:
            return []
        wp = _weekly_resample(pts, agg=agg) if weekly else pts
        return [{"date": d, "v": v} for d, v in wp]

    # ── 图1: 收益率 + 波动性 ──
    dgs10 = _fred_series_hist("DGS10", start)
    dgs2 = _fred_series_hist("DGS2", start)
    move = _fetch_move_history(start)
    panel1 = {
        "id": "yields_vol",
        "title": "① 美债收益率与波动性",
        "subtitle": "US Treasury Yields & MOVE Volatility",
        "note": "MOVE 指数 = ICE BofA 债券市场隐含波动率(国债版 VIX)。收益率与 MOVE 同向飙升 = 利率风险定价剧烈。",
        "unit_left": "MOVE 指数", "unit_right": "收益率 %",
        "series": [
            {"name": "MOVE 指数", "color": "#c17d6a", "axis": "left",
             "dash": True, "points": _mk(move, agg="peak")},   # peak: 周中尖峰不削(Chao 要看抬升点)
            {"name": "10年收益率", "color": "#6b8fb5", "axis": "right",
             "points": _mk(dgs10, agg="last")},
            {"name": "2年收益率", "color": "#7fa085", "axis": "right",
             "points": _mk(dgs2, agg="last")},
        ],
        "source": "FRED DGS10/DGS2 · ICE BofA MOVE (Yahoo Finance ^MOVE)",
    }

    # ── 图2: 国债市场压力代理(对齐 MS 图2 流动性/成本恶化主题) ──
    tp10 = _fred_series_hist("THREEFYTP10", start)   # 10yr 期限溢价(NY Fed ACM), %
    igoas = _fred_series_hist("BAMLC0A0CM", start)    # IG 企业债 OAS, %
    panel2 = {
        "id": "mkt_stress",
        "title": "② 国债市场压力代理",
        "subtitle": "Treasury Market Stress Proxy (公开替代 BrokerTec 日内价差)",
        "note": "⚠ 非 BrokerTec 日内 bid-ask 价差(MS 专有数据无公开源)——改用免费公开压力代理，两条线都是「持有风险资产要求的补偿」，走高=市场压力↑。<br>"
                "<b>① 10年期限溢价(左轴, 橙线)</b>：投资者持有10年期国债、相比不断滚动短债，额外要求的年化补偿(%)。"
                "<b>怎么看</b>：走高=市场担心久期风险/供给过剩/通胀不确定，要求更高补偿；走低甚至转负=避险买盘涌入压低补偿。"
                "<b>怎么用</b>：期限溢价快速抬升往往先于长端收益率飙升，是国债承压的早期信号。<br>"
                "<b>② IG企业债利差 OAS(右轴, 紫线)</b>：<b>IG</b>=Investment Grade(投资级公司债，评级BBB-/Baa3及以上)；<b>OAS</b>=Option-Adjusted Spread(期权调整利差)。即投资级公司债相对同期限国债的信用利差(%，期权调整后)。"
                "<b>怎么看</b>：走阔=市场要求更高信用补偿=融资环境收紧/避险；收窄=信用环境宽松、风险偏好回升。"
                "<b>怎么用</b>：IG OAS 是「机构级」流动性/信用压力的温度计，比股市更早反映融资面紧张。<br>"
                "<b>两线同向走高 = 融资/流动性环境明显收紧</b>（对应 MS 原图「买卖价差扩大、成本上升」的主题）。",
        "unit_left": "期限溢价 %", "unit_right": "IG OAS %",
        "series": [
            {"name": "10年期限溢价", "color": "#b58a6a", "axis": "left",
             "points": _mk(tp10, agg="last")},
            {"name": "IG企业债利差(OAS)", "color": "#8a7fa8", "axis": "right",
             "points": _mk(igoas, agg="last")},
        ],
        "source": "FRED THREEFYTP10 (NY Fed ACM 10yr term premium) · BAMLC0A0CM (ICE BofA IG OAS)",
    }

    # ── 图3: 曲线利差 + 信用压力(对齐 MS 图3 活跃度/风险传导主题) ──
    t10y2y = _fred_series_hist("T10Y2Y", start)       # 10Y-2Y 曲线利差, %
    hyoas = _fred_series_hist("BAMLH0A0HYM2", start)   # HY OAS, %
    panel3 = {
        "id": "curve_credit",
        "title": "③ 收益率曲线与信用压力",
        "subtitle": "Yield Curve & Credit Stress (公开替代 BrokerTec DV01 成交量)",
        "note": "⚠ 非 BrokerTec 周度 DV01 成交量(MS 专有数据无公开源)——改用免费公开风险/曲线代理，反映利率预期与风险传导。<br>"
                "<b>① 10Y-2Y 曲线利差(左轴, 蓝线)</b>：10年期国债收益率 − 2年期国债收益率(%)。"
                "<b>怎么看</b>：<span style=\"color:#2e9e5b\">正值=正常向上倾斜</span>(长债利率高于短债)；"
                "<span style=\"color:#d64545\">负值=倒挂</span>(短债高于长债，历史上是衰退的经典领先信号)；由负转正的「陡峭化」常出现在降息预期升温/衰退临近。"
                "<b>怎么用</b>：看拐点——倒挂后重新转正(bear steepening)往往对应政策转向或risk-off。<br>"
                "<b>② 高收益债利差 HY OAS(右轴, 红线)</b>：<b>HY</b>=High Yield(高收益债/垃圾级，评级BB+/Ba1及以下)；<b>OAS</b>=Option-Adjusted Spread(期权调整利差)。即高收益公司债相对国债的信用利差(%)。"
                "<b>怎么看</b>：这是市场「风险偏好」最灵敏的指标——骤升=避险情绪爆发、风险资产承压；低位平稳=risk-on。"
                "<b>怎么用</b>：HY OAS 突破关键位(如>4.5%)是本系统的硬性卖出触发之一；它领先或同步于股市大跌。<br>"
                "<b>HY OAS 骤升 + 曲线快速变动 = 风险传导、市场剧烈调仓</b>（对应 MS 原图「成交量激增/风险交易」的主题）。",
        "unit_left": "10Y-2Y 利差 %", "unit_right": "HY OAS %",
        "series": [
            {"name": "10Y-2Y 曲线利差", "color": "#6b8fb5", "axis": "left",
             "points": _mk(t10y2y, agg="last")},
            {"name": "高收益债利差(HY OAS)", "color": "#c0757d", "axis": "right",
             "points": _mk(hyoas, agg="last")},
        ],
        "source": "FRED T10Y2Y (10Y-2Y spread) · BAMLH0A0HYM2 (ICE BofA HY OAS)",
    }

    panels = [panel1, panel2, panel3]
    # asof = 所有 series 里最新的日期
    all_last = []
    for p in panels:
        for s in p["series"]:
            if s["points"]:
                all_last.append(s["points"][-1]["date"])
    asof = max(all_last) if all_last else None
    # status: 统计缺失序列
    missing = [f'{p["id"]}/{s["name"]}' for p in panels for s in p["series"] if not s["points"]]
    status = "ok" if not missing else ("部分缺失: " + ", ".join(missing))
    return {"panels": panels, "asof": asof, "years": years, "status": status}


def fetch_ofr_fsi(years=3):
    """OFR 金融压力指数(Financial Stress Index) —— 对齐'国债市场压力'主题的官方权威总览指标。
    数据源: OFR 官方 CSV (financialresearch.gov/financial-stress-index/data/fsi.csv, 免key, 日频2000至今)。
    列: Date, OFR FSI(总), Credit, Equity valuation, Safe assets, Funding, Volatility, +3区域。
    OFR FSI: 0=金融压力处于历史正常水平; 正值=压力高于正常; 负值=低于正常。
    返回 {panel:{...同stress_panel结构...}, asof, years, status, latest:{分量最新值}}。绝不编: 抓不到→status。
    """
    import datetime as _dt
    import subprocess
    start_year = (_dt.date.today() - _dt.timedelta(days=int(365.3 * years) + 10))
    url = "https://www.financialresearch.gov/financial-stress-index/data/fsi.csv"
    text = ""
    # 通道1: urllib
    try:
        import urllib.request
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        text = urllib.request.urlopen(req, timeout=30).read().decode("utf-8", "replace")
    except Exception:
        text = ""
    # 通道2: curl 回退
    if "OFR FSI" not in text:
        for _attempt in range(3):
            try:
                r = subprocess.run(["curl", "-s", url, "--max-time", "30",
                                    "-H", "User-Agent: Mozilla/5.0"],
                                   capture_output=True, text=True, timeout=38)
                if r.returncode == 0 and "OFR FSI" in (r.stdout or ""):
                    text = r.stdout
                    break
            except Exception:
                pass
    if "OFR FSI" not in text:
        return {"panel": None, "asof": None, "years": years,
                "status": "拉取失败(OFR 端点超时)", "latest": {}}
    lines = text.strip().splitlines()
    header = [h.strip() for h in lines[0].split(",")]
    idx = {h: i for i, h in enumerate(header)}
    # 需要的列(官方 FSI 5 大分量: 信用/股票估值/融资/安全资产/波动性 + 总指数)
    want = {"OFR FSI": "总指数", "Credit": "信用", "Equity valuation": "股票估值",
            "Funding": "融资", "Safe assets": "安全资产", "Volatility": "波动性"}
    series_pts = {k: [] for k in want}
    for ln in lines[1:]:
        cells = ln.split(",")
        if not cells or len(cells) < len(header):
            continue
        d = cells[idx["Date"]].strip()
        try:
            dt = _dt.date.fromisoformat(d)
        except Exception:
            continue
        if dt < start_year:
            continue
        for col in want:
            j = idx.get(col)
            if j is None or j >= len(cells):
                continue
            v = cells[j].strip()
            if v and v not in (".", ""):
                try:
                    series_pts[col].append((d, round(float(v), 3)))
                except ValueError:
                    pass
    # 周度降采样降噪
    colors = {"OFR FSI": "#c0757d", "Credit": "#b58a6a", "Equity valuation": "#c9a86a",
              "Funding": "#8a7fa8", "Safe assets": "#7fa085", "Volatility": "#6b8fb5"}
    series = []
    for col, zh in want.items():
        wp = _weekly_resample(series_pts[col], agg="peak") if series_pts[col] else []   # peak: 保留周中尖峰
        series.append({
            "name": (f"OFR FSI 总指数" if col == "OFR FSI" else zh),
            "color": colors[col],
            "axis": "left",
            "width": (2.4 if col == "OFR FSI" else 1.3),
            "points": [{"date": d, "v": v} for d, v in wp],
        })
    all_last = [s["points"][-1]["date"] for s in series if s["points"]]
    asof = max(all_last) if all_last else None
    latest = {s["name"]: (s["points"][-1]["v"] if s["points"] else None) for s in series}
    panel = {
        "id": "ofr_fsi",
        "title": "④ OFR 金融压力指数",
        "subtitle": "OFR Financial Stress Index (官方权威·全市场压力总览)",
        "note": "OFR FSI = 美国金融研究办公室(财政部下属)编制的官方金融压力指数, 综合 33 个全球金融市场变量。"
                "<b>0 = 压力处于历史正常水平</b>; <span style=\"color:#d64545\">正值 = 压力高于正常</span>(承压); "
                "<span style=\"color:#2e9e5b\">负值 = 低于正常</span>(平静)。粗线为总指数, 细线为官方5大分量(信用/股票估值/融资/安全资产/波动性)——分量线用于定位压力来源(哪个市场在制造压力)。",
        "unit_left": "压力指数(0=正常)", "unit_right": "",
        "series": series,
        "source": "OFR (Office of Financial Research, US Treasury) FSI 官方 CSV",
        "single_axis": True,
    }
    return {"panel": panel, "asof": asof, "years": years,
            "status": "ok" if asof else "无有效数据", "latest": latest}


def write_stress_panels_notion(sp=None):
    """把国债市场压力三联图【最新值】写入 Notion DB_STRESS(以 asof 作 title 幂等 upsert)。
    存最新一期6个指标的当前读数(时序), 完整历史序列存 GitHub data/ 副本。
    返回写入 page id 或 None(无DB/无数据时跳过)。"""
    import config as c
    import notion_writer as nw
    if sp is None:
        sp = fetch_treasury_stress_panels()
    db = c.NOTION_DB.get("stress")
    if not db or not sp.get("asof"):
        return None
    # 从 panels 取每个 series 的最新值
    latest = {}
    for p in sp["panels"]:
        for s in p["series"]:
            latest[s["name"]] = s["points"][-1]["v"] if s["points"] else None
    props = {
        "MOVE指数": nw.prop_num(latest.get("MOVE 指数")),
        "10年收益率": nw.prop_num(latest.get("10年收益率")),
        "2年收益率": nw.prop_num(latest.get("2年收益率")),
        "10年期限溢价": nw.prop_num(latest.get("10年期限溢价")),
        "IG企业债OAS": nw.prop_num(latest.get("IG企业债利差(OAS)")),
        "10Y2Y曲线利差": nw.prop_num(latest.get("10Y-2Y 曲线利差")),
        "HY_OAS": nw.prop_num(latest.get("高收益债利差(HY OAS)")),
        "数据源": {"rich_text": [{"type": "text", "text": {"content":
                  "FRED + ICE BofA MOVE(Yahoo); 图2/图3为公开压力代理(非BrokerTec)"}}]},
    }
    return nw.upsert(db, sp["asof"], props, title_field="Date")


def write_ofr_notion(ofr=None):
    """把 OFR 金融压力指数【最新值】写入 Notion DB_OFR(以 asof 作 title 幂等 upsert)。
    存最新一期总指数+分量当前读数(时序), 完整3年序列存 GitHub data/ 副本。
    返回写入 page id 或 None(无DB/无数据时跳过)。"""
    import config as c
    import notion_writer as nw
    if ofr is None:
        ofr = fetch_ofr_fsi()
    db = c.NOTION_DB.get("ofr")
    if not db or not ofr.get("asof"):
        return None
    latest = ofr.get("latest", {})
    props = {
        "OFR_FSI总指数": nw.prop_num(latest.get("OFR FSI 总指数")),
        "信用": nw.prop_num(latest.get("信用")),
        "股票估值": nw.prop_num(latest.get("股票估值")),
        "融资": nw.prop_num(latest.get("融资")),
        "安全资产": nw.prop_num(latest.get("安全资产")),
        "波动性": nw.prop_num(latest.get("波动性")),
        "数据源": {"rich_text": [{"type": "text", "text": {"content":
                  "OFR (Office of Financial Research, US Treasury) FSI 官方 CSV"}}]},
    }
    return nw.upsert(db, ofr["asof"], props, title_field="Date")


def _fred_series(series_id, start=None):
    """curl FRED fredgraph.csv 取完整历史 [(date, value)] 升序。缺失('.')跳过。失败返回 []。"""
    import requests
    url = f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={series_id}"
    if start:
        url += f"&cosd={start}"
    try:
        r = requests.get(url, timeout=30)
        if r.status_code != 200:
            return []
        out = []
        for ln in r.text.strip().splitlines()[1:]:
            if "," not in ln:
                continue
            d, v = ln.split(",", 1)
            v = v.strip()
            if v and v != ".":
                try:
                    out.append((d.strip(), float(v)))
                except ValueError:
                    continue
        return out
    except Exception:
        return []


def _eia_series(series_id, start=None):
    """EIA API v2 petroleum stocks 周频序列。返回 [(date 'YYYY-MM-DD', value 千桶)] 升序。
    需 EIA_API_KEY(在 .env)。失败/无 key 返回 []。绝不编造。"""
    import requests
    key = os.environ.get("EIA_API_KEY")
    if not key:
        # 从 .env 兜底读
        envp = os.path.join(os.path.dirname(__file__), "..", ".env")
        if os.path.exists(envp):
            for line in open(envp):
                if line.startswith("EIA_API_KEY="):
                    key = line.strip().split("=", 1)[1]
                    break
    if not key:
        return []
    url = ("https://api.eia.gov/v2/petroleum/stoc/wstk/data/"
           f"?api_key={key}&frequency=weekly&data[0]=value"
           f"&facets[series][]={series_id}"
           "&sort[0][column]=period&sort[0][direction]=asc&length=5000")
    if start:
        url += f"&start={start}"
    try:
        r = requests.get(url, timeout=40)
        if r.status_code != 200:
            return []
        data = r.json().get("response", {}).get("data", [])
        out = []
        for x in data:
            per = x.get("period")
            val = x.get("value")
            if per and val is not None:
                try:
                    out.append((per, float(val)))
                except (ValueError, TypeError):
                    continue
        out.sort(key=lambda t: t[0])
        return out
    except Exception:
        return []


def _mof_jgb_yields(term_idx):
    """抓日本财务省(MOF)每日 JGB 收益率 CSV(历史 all + 当月), 取指定期限列。
    term_idx: CSV 中列索引(10Y=10, 30Y=14; 列0=Date)。
    返回 {date 'YYYY-MM-DD': yield%} 升序 dict。失败返回 {}。绝不编造。"""
    import csv, io, requests, datetime as _dt
    urls = [
        "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/historical/jgbcme_all.csv",
        "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/jgbcme.csv",
    ]
    out = {}
    for url in urls:
        try:
            r = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=40)
            if r.status_code != 200:
                continue
            rows = list(csv.reader(io.StringIO(r.content.decode("utf-8", errors="replace"))))
            for row in rows:
                if not row or len(row) <= term_idx:
                    continue
                d = row[0].strip()
                # 日期形如 2026/8/17
                if "/" not in d:
                    continue
                try:
                    dt = _dt.datetime.strptime(d, "%Y/%m/%d").date()
                except ValueError:
                    continue
                raw = row[term_idx].strip()
                if not raw or raw == "-":
                    continue
                try:
                    out[dt.isoformat()] = float(raw)
                except ValueError:
                    continue
        except Exception:
            continue
    return dict(sorted(out.items()))


def fetch_us_jp_yields(cache_path=None):
    """美日 10Y/30Y 国债收益率同图对比(过去一年, 日频)。绝不编造, 缺项标 status。
      美债: FRED DGS10/DGS30(日频, %)。日债: 日本财务省(MOF) 每日 JGB CSV 10Y/30Y列(日频, %)。
    返回 {status, as_of, series:{us_10y/us_30y/jp_10y/jp_30y: {name,color,dash,points:[(date,%)],latest,as_of}}}。
    """
    import json, datetime as _dt
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "us_jp_yields.json")
    today = _dt.date.today()
    cutoff = (today - _dt.timedelta(days=366)).isoformat()

    # 美债(FRED 日频) — 复用 _fred_series
    us10 = [(d, v) for d, v in _fred_series("DGS10", start=cutoff) if d >= cutoff]
    us30 = [(d, v) for d, v in _fred_series("DGS30", start=cutoff) if d >= cutoff]
    # 日债(MOF 日频)
    jp10_all = _mof_jgb_yields(10)
    jp30_all = _mof_jgb_yields(14)
    jp10 = [(d, v) for d, v in jp10_all.items() if d >= cutoff]
    jp30 = [(d, v) for d, v in jp30_all.items() if d >= cutoff]

    def pack(name, color, dash, pts, src):
        if len(pts) < 2:
            return {"status": "未获取", "name": name}
        return {"status": "ok", "name": name, "color": color, "dash": dash,
                "points": pts, "latest": pts[-1][1], "as_of": pts[-1][0], "source": src}

    series = {
        "us_10y": pack("美国 10Y", "#c0757d", "none", us10, "FRED DGS10 (日频)"),
        "us_30y": pack("美国 30Y", "#8a3f47", "none", us30, "FRED DGS30 (日频)"),
        "jp_10y": pack("日本 10Y", "#6b8fb5", "dash", jp10, "日本财务省 MOF JGB (日频)"),
        "jp_30y": pack("日本 30Y", "#3a5a7d", "dash", jp30, "日本财务省 MOF JGB (日频)"),
    }
    ok_any = any(s.get("status") == "ok" for s in series.values())
    asofs = [s["as_of"] for s in series.values() if s.get("status") == "ok"]
    out = {"status": "ok" if ok_any else "未获取",
           "as_of": max(asofs) if asofs else today.isoformat(),
           "series": series}
    try:
        with open(cache_path, "w") as f:
            json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def fetch_nikkei225(cache_path=None):
    """日经225指数(过去一年,日频)。FRED NIKKEI225。绝不编造,缺失跳过。
    返回 {status, as_of, latest, hi, lo, points:[(date,收盘)], source}。"""
    import json, datetime as _dt
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "nikkei225.json")
    today = _dt.date.today()
    cutoff = (today - _dt.timedelta(days=366)).isoformat()
    pts = [(d, v) for d, v in _fred_series("NIKKEI225", start=cutoff) if d >= cutoff]
    if len(pts) < 2:
        out = {"status": "未获取"}
    else:
        out = {"status": "ok", "as_of": pts[-1][0], "latest": pts[-1][1],
               "hi": max(v for _, v in pts), "lo": min(v for _, v in pts),
               "points": pts, "source": "FRED NIKKEI225 (日经225指数,日频)"}
    try:
        with open(cache_path, "w") as f:
            json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def _jpx_foreign_flow_weekly():
    """抓 JPX 投资部门别周报(东证Prime,金额), 解析海外投资家(Foreigners)每周净买卖(差引き Balance)。
    从 archive 页(00=本年,01=上年)解析 stock_val_1_*.xls 链接→逐个读:
      · 周结束日期从 Excel row3 '(M/D - M/D)' 解析(文件名code是YYMM週号非日期,不能当日期);
      · '海外投資家/Foreigners' 行的 差引き Balance(第6列 idx6)为净额, 字符串带逗号需去除。
    单位:日元→万亿日元(¥T)。返回 [(week_end 'YYYY-MM-DD', 净额¥T)] 升序。失败返回 []。绝不编造。"""
    import re, requests, datetime as _dt, io
    try:
        import pandas as pd
    except Exception:
        return []
    base = "https://www.jpx.co.jp"
    out = []
    links = []
    # archive-00=本年, archive-01=上年 (覆盖过去一年绰绰有余)
    for nn in (0, 1):
        u = base + f"/english/markets/statistics-equities/investor-type/00-00-archives-{nn:02d}.html"
        try:
            h = requests.get(u, headers={"User-Agent": "Mozilla/5.0"}, timeout=30).text
            links += re.findall(r'(/english/markets/statistics-equities/investor-type/[^"]*?stock_val_1_\d{6}\.xls)', h)
        except Exception:
            continue
    cutoff = _dt.date.today() - _dt.timedelta(days=380)
    seen = set()

    def _to_float(x):
        try:
            import pandas as _pd
            if _pd.isna(x):
                return None
        except Exception:
            pass
        if isinstance(x, (int, float)):
            try:
                f = float(x)
                return None if f != f else f  # NaN check
            except Exception:
                return None
        if isinstance(x, str):
            s = x.replace(",", "").replace("△", "-").replace("▲", "-").strip()
            if not s or s == "-":
                return None
            try:
                return float(s)
            except ValueError:
                return None
        return None

    for path in links:
        if path in seen:
            continue
        seen.add(path)
        try:
            raw = requests.get(base + path, headers={"User-Agent": "Mozilla/5.0"}, timeout=30).content
            df = pd.read_excel(io.BytesIO(raw), header=None, sheet_name=0)
        except Exception:
            continue
        # 1) 周结束日期: 扫前6行找 '( M/D - M/D )'
        wk_end = None
        for i in range(min(6, len(df))):
            cell = str(df.iloc[i, 0])
            m = re.search(r'(\d{4})年.*?(\d{1,2})/(\d{1,2})\s*[-–]\s*(\d{1,2})/(\d{1,2})', cell)
            if m:
                yr, _, _, em, ed = m.groups()
                try:
                    wk_end = _dt.date(int(yr), int(em), int(ed))
                except ValueError:
                    wk_end = None
                break
        if wk_end is None or wk_end < cutoff:
            continue
        # 2) 海外投資家 Balance(差引き, 第6列 idx6, 在'売り'行): 单位日元→亿日元(¥100M)
        bal = None
        for i in range(len(df)):
            c0 = str(df.iloc[i, 0])
            if "海外投資家" in c0 or "Foreign" in c0:
                # 差引き净额在'売り'行(本行)第6列; 兜底试下一行
                for j in (i, i + 1):
                    if j < len(df) and df.shape[1] > 6:
                        v = _to_float(df.iloc[j, 6])
                        if v is not None:
                            bal = v; break
                break
        if bal is not None:
            # Excel 单位=千円(1,000 yen): 千円→万亿日元(¥T) = /1e9
            out.append((wk_end.isoformat(), round(bal / 1e9, 3)))  # 千円→万亿日元(¥T)
    # ★★2026-10-06(Chao「这个图没更新」): JPX 自 2026 年 9 月第 3 周起换了发布格式 ——
    #   旧: 存档页 stock_val_1_YYMMWW.xls(海外投資家 差引 在 col6, 周范围在 row3 文字里)
    #   新: index 页 stock_1_w_YYYYMMDD_YYYYMMDD.xlsx(横向 14 个投资者块, 海外投資家 差引 在 col21,
    #       周范围在「文件名」里, 文件内只有 YYYYMMW 周标识)。新文件不出现在存档页 → 旧解析器抓不到,
    #       fetcher 不报错、缓存每天被重写但内容不变、日报仍标 ok(静默停更, 停在 09-11)。
    #   新格式单位已核: PDF 页眉「千円 1,000 yen」= 旧表同单位, 同样 /1e9 → 万亿日元。
    #   自洽校验: 买 − 卖 必须 == 差引(±2), 不符则丢弃该周(宁缺勿错)。
    have = {d for d, _ in out}
    try:
        _ih = requests.get(base + "/english/markets/statistics-equities/investor-type/index.html",
                           headers={"User-Agent": "Mozilla/5.0"}, timeout=30).text
        _new = sorted(set(re.findall(r'href="([^"]*stock_1_w_(\d{8})_(\d{8})\.xlsx)"', _ih)), key=lambda t: t[2])
    except Exception:
        _new = []
    for _href, _d0, _d1 in _new:
        try:
            wk_end = _dt.date(int(_d1[:4]), int(_d1[4:6]), int(_d1[6:8]))
        except ValueError:
            continue
        if wk_end < cutoff or wk_end.isoformat() in have:
            continue
        try:
            raw = requests.get(base + _href, headers={"User-Agent": "Mozilla/5.0"}, timeout=40).content
            df = pd.read_excel(io.BytesIO(raw), header=None, sheet_name=0)
        except Exception:
            continue
        # 定位「海外投資家」块: 第 3 行(idx3)起始列; 差引 Balance = 起始列+2 (売/買/差引/合計)
        jf = None
        for j in range(df.shape[1]):
            if "海外" in str(df.iloc[3, j]) if df.shape[0] > 3 else False:
                jf = j
                break
        if jf is None or df.shape[1] <= jf + 2:
            continue
        # 找「東証プライム」的「金額 Value」行(行内第 2 列含 Prime, 下一行是 Value)
        vrow = None
        for i in range(7, df.shape[0]):
            if "プライム" in str(df.iloc[i, 1]) or "Prime" in str(df.iloc[i, 1]):
                vrow = i + 1
                break
        if vrow is None or vrow >= df.shape[0] or "Value" not in str(df.iloc[vrow, 2]):
            continue
        sell, buy, bal = (_to_float(df.iloc[vrow, jf + k]) for k in (0, 1, 2))
        if None in (sell, buy, bal) or abs((buy - sell) - bal) > 2:
            continue   # 自洽失败: 列位置/口径不对, 丢弃, 不写错数
        out.append((wk_end.isoformat(), round(bal / 1e9, 3)))
    out.sort(key=lambda t: t[0])
    return out


def fetch_foreign_flow_japan(cache_path=None):
    """外资净买入日本股票(过去一年,周频)。JPX 投资部门别周报'海外投資家'差引き净额。
    绝不编造,抓不到留空 status。返回 {status,as_of,latest,points:[(week,¥T)],source}。"""
    import json
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "foreign_flow_japan.json")
    pts = _jpx_foreign_flow_weekly()
    if len(pts) < 2:
        out = {"status": "未获取"}
    else:
        out = {"status": "ok", "as_of": pts[-1][0], "latest": pts[-1][1],
               "hi": max(v for _, v in pts), "lo": min(v for _, v in pts),
               "points": pts,
               "source": "JPX 投资部门别交易(东证Prime,海外投資家净买卖,周频,万亿日元¥T,原表千円)"}
    try:
        with open(cache_path, "w") as f:
            json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def _imf_iip_series(country, entry, freq="A", start="2014"):
    """IMF SDMX 2.1 拉一国 IIP 资产/负债年度序列(USD)。
    country: USA/JPN/DEU/CHN; entry: A_P(资产)/L_P(负债)。返回 {year:USD} 。失败{}。绝不编造。"""
    import requests
    import xml.etree.ElementTree as ET
    url = f"https://api.imf.org/external/sdmx/2.1/data/IIP/{country}.{entry}.IIP.USD.{freq}?startPeriod={start}"
    try:
        r = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=60)
        if r.status_code != 200:
            return {}
        root = ET.fromstring(r.text)
        out = {}
        for el in root.iter():
            if el.tag.split("}")[-1] == "Obs":
                t = el.attrib.get("TIME_PERIOD"); v = el.attrib.get("OBS_VALUE")
                if t and v:
                    try:
                        out[t] = float(v)
                    except ValueError:
                        pass
        return out
    except Exception:
        return {}


def fetch_iip_four_countries(cache_path=None):
    """四国(美/日/德/中)国际投资头寸 IIP(过去约10年,年频): 对外总资产/总负债/净头寸。
    源: IMF SDMX 2.1 (api.imf.org) IIP 数据集, entry A_P(资产)/L_P(负债), INDICATOR=IIP, USD 年频。
    绝不编造, 缺国留 status。返回 {status,as_of,countries:{US/JP/DE/CN:{name,flag,color,
      assets:[(yr,$T)],liab:[(yr,$T)],net:[(yr,$T)],latest_*}}}。"""
    import json
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "iip_four.json")
    cmap = {"US": ("USA", "美国", "🇺🇸", "#c0757d"),
            "JP": ("JPN", "日本", "🇯🇵", "#6b8fb5"),
            "DE": ("DEU", "德国", "🇩🇪", "#7fa085"),
            "CN": ("CN_ISO_PLACEHOLDER", "中国", "🇨🇳", "#e0a92e")}
    # 中国 ISO 用 CHN
    cmap["CN"] = ("CHN", "中国", "🇨🇳", "#e0a92e")
    countries = {}
    all_years = []
    for key, (iso, name, flag, color) in cmap.items():
        a = _imf_iip_series(iso, "A_P")
        l = _imf_iip_series(iso, "L_P")
        yrs = sorted(set(a) & set(l))
        yrs = [y for y in yrs if y >= "2015"]
        if len(yrs) < 2:
            countries[key] = {"status": "未获取", "name": name, "flag": flag}
            continue
        assets = [(y, round(a[y] / 1e12, 2)) for y in yrs]
        liab = [(y, round(l[y] / 1e12, 2)) for y in yrs]
        net = [(y, round((a[y] - l[y]) / 1e12, 2)) for y in yrs]
        # ★2026-10-07 接季频(IMF IIP Q 口径与年频同一数据集, 实测美国 2025-Q4 季频值 == 2025 年频值)。
        #   只追加【晚于最新年末】的季度点(net_q), 年频序列与 latest_* 一律不动; 季度取不到就没有 net_q, 页面自动退回年频。
        qa = _imf_iip_series(iso, "A_P", freq="Q", start=str(int(yrs[-1])))
        ql = _imf_iip_series(iso, "L_P", freq="Q", start=str(int(yrs[-1])))
        qks = sorted(p for p in (set(qa) & set(ql)) if p > f"{yrs[-1]}-Q4")
        net_q = [(p, round((qa[p] - ql[p]) / 1e12, 2)) for p in qks]
        countries[key] = {
            "q_assets": round(qa[qks[-1]] / 1e12, 2) if qks else None,
            "q_liab": round(ql[qks[-1]] / 1e12, 2) if qks else None,
            "q_net": net_q[-1][1] if qks else None,
            "q_period": qks[-1] if qks else None,
            "net_q": net_q,
            "status": "ok", "name": name, "flag": flag, "color": color,
            "assets": assets, "liab": liab, "net": net,
            "latest_year": yrs[-1], "latest_assets": assets[-1][1],
            "latest_liab": liab[-1][1], "latest_net": net[-1][1],
        }
        all_years += yrs
    ok = any(c.get("status") == "ok" for c in countries.values())
    out = {"status": "ok" if ok else "未获取",
           "as_of": max(all_years) if all_years else "",
           "q_as_of": max((c.get("q_period") or "" for c in countries.values()), default=""),
           "source": "IMF International Investment Position (SDMX 2.1, api.imf.org), 年频 + 最新季频, USD",
           "countries": countries}
    try:
        with open(cache_path, "w") as f:
            json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def _ofr_hfm_series(mnemonic):
    """OFR Hedge Fund Monitor 免key JSON API 季度序列。返回 [(date 'YYYY-MM-DD', float)] 升序。失败返回 []。
    源: https://data.financialresearch.gov/hf/v1/series/timeseries?mnemonic=<code> (SEC Form PF 汇总, 无需key)。"""
    import requests
    url = f"https://data.financialresearch.gov/hf/v1/series/timeseries?mnemonic={mnemonic}"
    try:
        r = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=30)
        if r.status_code != 200:
            return []
        data = r.json()
        out = []
        for row in data:
            if isinstance(row, list) and len(row) >= 2 and row[1] is not None:
                try:
                    out.append((str(row[0]), float(row[1])))
                except (ValueError, TypeError):
                    continue
        return out
    except Exception:
        return []


def fetch_hf_leverage(cache_path=None):
    """对冲基金杠杆监测(OFR Hedge Fund Monitor, SEC Form PF 底层, 季度)。绝不编造。
      图A: 对冲基金美国主权债(美债)总名义敞口 GNE 占 US GDP % (2013Q1起, 季度)
           = FPF-ASSETCLASS_USGOV_GNE_SUM / (FRED GDP 十亿$×1e9) ×100
      图B: 对冲基金三类借款规模($万亿, 季度): Repo / Prime brokerage / Other secured
    源: OFR https://data.financialresearch.gov/hf/v1/series/timeseries?mnemonic=<code> (免key)
        GDP: FRED GDP (免key CSV)
    返回 {status, as_of, exposure:{points:[(q,%)], latest_pct, latest_usd_t}, borrow:{repo/prime/other:[(q,$T)], latest_*}}。
    """
    import json, requests
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "hf_leverage.json")

    def _ofr(m):
        try:
            r = requests.get(
                f"https://data.financialresearch.gov/hf/v1/series/timeseries?mnemonic={m}",
                timeout=30, headers={"User-Agent": "Mozilla/5.0"})
            d = r.json()
            s = d if isinstance(d, list) else d.get("timeseries", [])
            return [(str(x[0]), float(x[1])) for x in s if x and x[1] is not None]
        except Exception:
            return []

    def _fred_q(series):
        """FRED 季度序列 {date:val}(GDP 十亿$)。用 curl 子进程(该端点对 requests 偶发超时)。"""
        import subprocess
        url = f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={series}&cosd=2013-01-01"
        text = ""
        for _a in range(3):
            try:
                p = subprocess.run(["curl", "-s", "--max-time", "20", url],
                                   capture_output=True, text=True, timeout=25)
                if p.returncode == 0 and p.stdout and "," in p.stdout:
                    text = p.stdout
                    break
            except Exception:
                continue
        out = {}
        for ln in text.splitlines()[1:]:
            p = ln.split(",")
            if len(p) == 2 and p[1] not in ("", "."):
                out[p[0]] = float(p[1])
        return out

    repo = _ofr("FPF-BORROW_REPO_SUM")
    prime = _ofr("FPF-BORROW_PRIMEBROKER_SUM")
    other = _ofr("FPF-BORROW_OTHERSECURED_SUM")
    gne = _ofr("FPF-ASSETCLASS_USGOV_GNE_SUM")
    gdp = _fred_q("GDP")  # 十亿$ 年化, 季度点(月初日期如 2026-01-01)

    def _t(series):  # 转 $万亿
        return [(q, round(v / 1e12, 3)) for q, v in series]

    # 敞口/GDP%: OFR 季末(如 2026-03-31) 对 FRED 季初(2026-01-01)。按年季匹配。
    def _gdp_for_q(qdate):
        yr, mo = qdate[:4], qdate[5:7]
        qmap = {"03": "01", "06": "04", "09": "07", "12": "10"}
        key = f"{yr}-{qmap.get(mo, '01')}-01"
        return gdp.get(key)

    exp_pts = []
    for q, v in gne:
        g = _gdp_for_q(q)
        if g:
            exp_pts.append((q, round(v / (g * 1e9) * 100, 2)))

    ok = bool(repo and gne)
    out = {
        "status": "ok" if ok else "未获取",
        "as_of": (repo[-1][0] if repo else (gne[-1][0] if gne else "")),
        "source": "OFR Hedge Fund Monitor (SEC Form PF, 季度) · FRED GDP",
        "exposure": {
            "points": exp_pts,
            "latest_pct": exp_pts[-1][1] if exp_pts else None,
            "latest_usd_t": round(gne[-1][1] / 1e12, 2) if gne else None,
            "latest_q": gne[-1][0] if gne else "",
        },
        "borrow": {
            "repo": _t(repo), "prime": _t(prime), "other": _t(other),
            "latest_repo": round(repo[-1][1] / 1e12, 2) if repo else None,
            "latest_prime": round(prime[-1][1] / 1e12, 2) if prime else None,
            "latest_other": round(other[-1][1] / 1e12, 2) if other else None,
            "latest_q": repo[-1][0] if repo else "",
        },
    }
    try:
        with open(cache_path, "w") as f:
            json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def fetch_bis_gold_swaps(cache_path=None):
    """BIS 自营黄金掉期规模(吨)。BIS 机构自己那笔 gold swaps(GATA/Robert Lambourne 追踪)。

    数据由 data/bis_gold_swaps.json 驱动(真值锚点):
      - 年度值(3/31, kind=annual): BIS 年报明确确认的官方数字(2010 至今连续)。
      - 月度值(kind=monthly): GATA 顾问 Lambourne 从 BIS 官方月度 Statement of Account
        (bis.org/banking/balsheet/statofacc{YYMMDD}.pdf)推算, BIS 年报每年验证其准确。
    绝不编造: BIS 从不主动披露 swaps 行, 也不公开精确勾稽算法, 故本函数只消费已核实的
    公开真值(年报 + GATA 推算), 不做无法验证的自行换算。月度更新由 cron agent 读 GATA
    新文章补一个点(只增, 绝不编)。
    返回 {status, as_of, source, unit, note, latest_t, peak_t, peak_date, points:[{date,tonnes,kind,src}]}。
    """
    import json
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "bis_gold_swaps.json")
    try:
        with open(cache_path, encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        return {"status": "未获取", "as_of": "", "source": "BIS 自营黄金掉期(文件缺失)",
                "unit": "tonnes", "points": []}
    pts = [p for p in (data.get("points") or []) if p.get("date") and p.get("tonnes") is not None]
    pts.sort(key=lambda p: p.get("date", ""))
    if not pts:
        return {"status": "未获取", "as_of": data.get("as_of", ""),
                "source": data.get("source", "BIS 自营黄金掉期"), "unit": "tonnes", "points": []}
    peak = max(pts, key=lambda p: p["tonnes"])
    # ★2026-10-07: 「最新年报」必须取最后一个 kind=annual 点, 不能取最后一个点——补入月度点(2026-08 141t)后
    #   最后一点是 monthly, 页面会把 GATA 的月度推算误标成「年报确认」。无 annual 点才回退最后一点。
    _ann = [p for p in pts if p.get("kind") == "annual"]
    latest = _ann[-1] if _ann else pts[-1]
    return {
        "status": "ok",
        "as_of": data.get("as_of", pts[-1]["date"]),
        "gaps_note": data.get("gaps_note", ""),
        "source": data.get("source", "BIS 自营黄金掉期 (BIS own-account gold swaps)"),
        "unit": data.get("unit", "tonnes"),
        "note": data.get("note", ""),
        "latest_t": latest["tonnes"],
        "latest_date": latest["date"],
        "peak_t": peak["tonnes"],
        "peak_date": peak["date"],
        "points": pts,
    }


def fetch_silver_imports_data(json_path=None):
    """读印度白银月度进口落盘 JSON (src/fetch_silver_imports.py 生成, UN Comtrade)。
    绝不编造: 文件缺失 → status='未获取'。"""
    import json as _json
    if json_path is None:
        json_path = os.path.join(os.path.dirname(__file__), "..", "data", "silver_imports_india.json")
    try:
        with open(json_path, encoding="utf-8") as f:
            d = _json.load(f)
        if d.get("status") == "ok" and d.get("points"):
            return d
    except Exception:
        pass
    return {"status": "未获取", "as_of": "", "note": "印度白银进口 JSON 未找到"}


def _fetch_gold_premium_api():
    """★2026-09-08 新增: 从 WGC 官方 API 抓 premium 增量。

    背景(Chao 报「这张图没有每日更新啊」): 本指标原设计是<b>纯手工</b> ——
    「用户从 goldhub 下 xlsx → 人工重跑 parse 脚本」, 从未自动化过, 结果停在
    2026-08-14 断更 25 天, 页面却标着「每日更新」。

    修复过程中试过并失败的路(勿重走): xlsx 直链 403 / 带 Referer 403 /
    会话预热(cookie) 403 / ScraperAPI 5 种参数全 500 / playwright 真浏览器
    页内 fetch 仍 403 —— 说明是 IP 段被 CDN 拉黑, 不是指纹问题。
    自己猜的 API 路径(v13/charts/premium 等)全 404。

    真正的解法是<b>从页面 HTML 里挖出真实端点</b>(而不是猜):
        https://fsapi.gold.org/api/v11/charts/local-gold-price-premiums
    该域名未被封(页面自己在用), 直连 200, 返回
        chartData.india_premdisc / china_premdisc = [[epoch_ms, value], ...]
        chartData.asOfDate
    实测 260 点滚动窗口(近 1 年)、asOfDate=2026-09-04。

    ★注意只有近 1 年 —— 所以调用方必须与既有 xlsx 历史<b>合并</b>而非覆盖,
    否则会把印度 2012 起、中国 2003 起的长历史截断。
    """
    import urllib.request
    import datetime as _dt
    url = "https://fsapi.gold.org/api/v11/charts/local-gold-price-premiums"
    req = urllib.request.Request(url, headers={
        "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"),
        "Accept": "application/json,*/*",
        "Referer": "https://www.gold.org/goldhub/data/gold-premium",
    })
    with urllib.request.urlopen(req, timeout=60) as r:
        if r.status != 200:
            return None, None, None
        payload = json.loads(r.read().decode("utf-8", "ignore"))
    cd = payload.get("chartData") or {}

    def _conv(rows):
        out = []
        for row in (rows or []):
            if not isinstance(row, (list, tuple)) or len(row) < 2 or row[1] is None:
                continue
            ts = row[0]
            try:
                d = _dt.datetime.utcfromtimestamp(ts / 1000 if ts > 1e11 else ts).date()
            except Exception:
                continue
            out.append({"date": d.isoformat(), "premium": round(float(row[1]), 4)})
        return out

    return _conv(cd.get("india_premdisc")), _conv(cd.get("china_premdisc")), cd.get("asOfDate")


def fetch_gold_premium(india_path=None, china_path=None):
    """印度 + 中国黄金 domestic premium/discount (US$/oz, 日度)。
    源: World Gold Council goldhub (ICE Benchmark/MCX/RBI/NCDEX; Bloomberg/SGE)。
    数据由 scratch/parse_gold_premium.py 从 WGC 官方 xlsx 解析落盘 data/gold_premium_{india,china}.json。
    正=本地溢价(需求旺/供给紧, 看多实物), 负=折价(需求弱/进口过剩)。
    绝不编造: 文件缺失 → status='未获取'。更新方式: 用户从 WGC goldhub 下新 xlsx → 重跑 parse 脚本。
    返回 {status, india:{...}, china:{...}}。
    """
    import json as _json
    base = os.path.join(os.path.dirname(__file__), "..", "data")
    if india_path is None:
        india_path = os.path.join(base, "gold_premium_india.json")
    if china_path is None:
        china_path = os.path.join(base, "gold_premium_china.json")

    def _load(p):
        try:
            with open(p, encoding="utf-8") as f:
                d = _json.load(f)
            if d.get("status") == "ok" and d.get("points"):
                return d
        except Exception:
            pass
        return None

    ind = _load(india_path)
    chn = _load(china_path)

    # ★2026-09-08: 先用官方 API 抓增量, 与本地 xlsx 历史合并。
    #   API 只给近 1 年(260 点)滚动窗口, 本地 xlsx 有印度 2012 / 中国 2003 起长历史,
    #   所以是 merge 而非 replace —— 覆盖会截断十几年历史。
    #   同一天两边都有时<b>以 API 为准</b>(官方口径可能回溯修订, 如 2025-01 印度方法论重算)。
    api_note = ""
    try:
        api_ind, api_chn, api_asof = _fetch_gold_premium_api()
    except Exception as e:
        api_ind = api_chn = api_asof = None
        api_note = "API 增量失败(%s), 仅用本地历史" % type(e).__name__

    def _merge(local, incoming):
        """本地历史 + API 增量 → 按日期去重合并(API 优先), 返回 (dict, 是否有新增)。"""
        if not incoming:
            return local, 0
        base = dict((p["date"], p["premium"]) for p in (local or {}).get("points", []))
        before = len(base)
        for p in incoming:
            base[p["date"]] = p["premium"]        # API 覆盖同日旧值(官方修订优先)
        pts = [{"date": d, "premium": v} for d, v in sorted(base.items())]
        added = len(base) - before
        out = dict(local or {})
        out["status"] = "ok"
        out["points"] = pts
        out["as_of"] = pts[-1]["date"] if pts else out.get("as_of", "")
        out["latest"] = pts[-1]["premium"] if pts else out.get("latest")
        # ★同步 n —— 它是 xlsx 落盘时写的点数字段, 合并后若不更新会长期显示旧值,
        #   造成「日志说 3801、实际 3816」这种自相矛盾的假象(2026-09-08 实测踩到)。
        if "n" in out or pts:
            out["n"] = len(pts)
        return out, added

    n_add = 0
    if api_ind:
        ind, a = _merge(ind, api_ind)
        n_add += a
    if api_chn:
        chn, a = _merge(chn, api_chn)
        n_add += a
    if api_asof and n_add:
        api_note = "API 增量 +%d 点 (官方 asOfDate=%s)" % (n_add, api_asof)
    elif api_ind and not n_add:
        api_note = "API 已是最新, 无新增点"

    if not ind and not chn:
        return {"status": "未获取", "as_of": "",
                "note": "黄金 premium 数据未找到 (WGC xlsx 未解析?)"}
    return {
        "status": "ok",
        "as_of": (ind or chn).get("as_of", ""),
        "india": ind,
        "china": chn,
        "note": api_note,
    }


def fetch_ad_line_real(json_path=None, price_lookback=120, nh_window=20):
    """真正的 A/D 腾落线(NYSE 广度): 读 Economic-Dashboard 每日 cron 维护的 SP500 全成分股腾落数据。

    数据源: Curarpikt0000/Economic-Dashboard repo 的
      data/ad_line/ad_line_sp500_history.json (每日 cron job 334572d3e1a7 更新, 501 只成分股)。
    本地 repo 路径: ~/Projects/Economic-Dashboard (每日 git fetch 保持最新)。
    每条记录: {date, advance, decline, flat, total, ad_net, ad_ratio, advance_pct, cumulative}。
    cumulative = 累计 (advance-decline) = 真正的 A/D 腾落线。

    顶背离判定(可复现规则):
      在最近 nh_window 个交易日内, 若 advance_pct 的近期均值走弱 (广度恶化)
      但 cumulative 腾落线未同步创近 price_lookback 新高, 判为需警惕。
      更硬核: 比较 cumulative 近期高点 vs 全窗口高点 —— 若腾落线未创新高即为顶背离预警。

    返回结构兼容 market_breadth (供现有 section 复用),
      但 spy_points 装 cumulative 腾落线, ratio_points 装 advance_pct(参与率)。
    绝不编造: 文件缺失/为空 → status='未获取'。
    """
    import json as _json
    if json_path is None:
        json_path = os.path.expanduser(
            "~/Projects/Economic-Dashboard/data/ad_line/ad_line_sp500_history.json")

    try:
        with open(json_path, encoding="utf-8") as f:
            raw = _json.load(f)
    except Exception:
        return {"status": "未获取", "as_of": "",
                "note": f"A/D 腾落线数据文件未找到: {json_path} (Economic-Dashboard cron 未同步?)"}

    days = raw.get("days", [])
    # 过滤掉初始占位行 (total=0)
    days = [d for d in days if d.get("total", 0) > 0 and d.get("cumulative") is not None]
    if len(days) < 5:
        return {"status": "未获取", "as_of": "",
                "note": "A/D 腾落线有效数据点不足 5"}

    dates = [d["date"] for d in days]
    cum = [d["cumulative"] for d in days]
    adv_pct = [d.get("advance_pct") for d in days]
    n = len(days)
    as_of = dates[-1]

    # spy_points 复用为 cumulative 腾落线; ratio_points 复用为 advance_pct 参与率
    cum_points = list(zip(dates, [float(c) for c in cum]))
    pct_points = [(dates[i], float(adv_pct[i])) for i in range(n) if adv_pct[i] is not None]

    # 顶背离判定: 腾落线近 nh_window 是否创近 price_lookback 新高
    lb = min(price_lookback, n)
    nw = min(nh_window, n)
    lookback_high = max(cum[-lb:])
    recent_high = max(cum[-nw:])
    cum_made_new_high = recent_high >= lookback_high - 1e-9
    # 参与率(advance_pct)近 nw 均值 vs 前一段
    valid_pct = [p for p in adv_pct if p is not None]
    recent_pct = valid_pct[-nw:] if len(valid_pct) >= nw else valid_pct
    recent_pct_avg = round(sum(recent_pct) / len(recent_pct), 1) if recent_pct else None
    # 腾落线绝对趋势: 近 nw 净变化
    cum_delta_nw = cum[-1] - cum[-nw] if n >= nw else None

    # 顶背离 = 腾落线未创新高 且 参与率低于 50% (多数股走弱)
    divergence = None
    if cum_made_new_high:
        divergence = False  # 广度确认(腾落线跟随创新高)
    else:
        # 未创新高: 结合参与率判断是否恶化
        divergence = (recent_pct_avg is not None and recent_pct_avg < 50.0) or (
            cum_delta_nw is not None and cum_delta_nw < 0)

    gap_from_high = round((lookback_high - recent_high), 1)

    return {
        "status": "ok",
        "as_of": as_of,
        "source": "Economic-Dashboard A/D 腾落线 (SP500 全成分股, 每日 cron)",
        "tickers": raw.get("tickers"),
        "divergence": divergence,
        "spy_points": cum_points,      # 累计腾落线 (复用字段名保持 section 兼容)
        "ratio_points": pct_points,    # advance_pct 参与率
        "latest_cumulative": cum[-1],
        "latest_advance_pct": adv_pct[-1] if adv_pct else None,
        "latest_ad_net": days[-1].get("ad_net"),
        "evidence": {
            "lookback_days": lb,
            "nh_window": nw,
            "cum_lookback_high": round(lookback_high, 0),
            "cum_recent_high": round(recent_high, 0),
            "cum_made_new_high": cum_made_new_high,
            "recent_advance_pct_avg": recent_pct_avg,
            "cum_delta_recent": cum_delta_nw,
            "gap_from_high": gap_from_high,
            "breadth_confirmed": bool(cum_made_new_high),
        },
    }


def fetch_market_breadth(cache_path=None, lookback_days=250, nh_window=20):
    """美股市场广度: RSP(等权标普)/SPY(市值加权标普) 比值, 替代无数据源的 A/D 布尔判断。

    数据源: 东方财富 push2his 原生 API(免 key)。SPY=107.SPY, RSP=107.RSP, 日线复权。
    注: akshare 封装被东财挡, 但原生 push2his API 实测通, 故直接请求(带重试)。

    市场广度含义: RSP/SPY 比值下行 = 少数大权重股领涨、多数股走弱(广度恶化),
    与 NYSE A/D 腾落线「顶背离」同义, 但这是<b>连续可量化真数据</b>, 可画折线、可复现判定。

    数值化背离判定(可复现规则, 不靠 AI 主观):
      在最近 nh_window 个交易日内,
      若 S&P(用 SPY 代理)创下近 lookback_days 新高, 但 RSP/SPY 广度比 <b>未</b>同步创新高
      (广度比距其近 lookback 高点还差 > 2%), 则判定为 <b>顶背离(divergence=True)</b>。
      否则 divergence=False(广度确认)。

    返回 {status, as_of, divergence(bool), evidence{...}, spy_points:[(d,close)],
          ratio_points:[(d, RSP/SPY 归一化)], source}。
    绝不编造: 取不到 status='未获取'。
    """
    import json, requests, datetime as _dt, time as _time
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "market_breadth.json")

    def _load_cache():
        """读上次成功落盘的真数据(抓不到时兜底, 绝不覆盖成空)。"""
        try:
            with open(cache_path, encoding="utf-8") as f:
                c = json.load(f)
            if c.get("status") == "ok" and c.get("spy_points"):
                c["stale"] = True  # 标记为缓存(非当日新抓)
                return c
        except Exception:
            pass
        return None

    UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
          "Referer": "https://quote.eastmoney.com/"}
    _sess = requests.Session()
    _sess.headers.update(UA)

    def _kline(secid, beg="20180101"):
        end = _dt.date.today().strftime("%Y%m%d")
        url = "https://push2his.eastmoney.com/api/qt/stock/kline/get"
        params = {"secid": secid, "fields1": "f1,f2,f3,f4,f5,f6",
                  "fields2": "f51,f52,f53,f54,f55,f56,f57,f58",
                  "klt": 101, "fqt": 1, "beg": beg, "end": end}
        for attempt in range(6):
            try:
                r = _sess.get(url, params=params, timeout=25)
                d = r.json().get("data")
                if d and d.get("klines"):
                    # klines: "date,open,close,high,low,vol,amount,pct"
                    out = []
                    for ln in d["klines"]:
                        parts = ln.split(",")
                        out.append((parts[0], float(parts[2])))  # (date, close)
                    return out
            except Exception:
                pass
            _time.sleep(1.5 * (attempt + 1))  # 递增退避
        return None

    spy = _kline("107.SPY")
    _time.sleep(1.0)
    rsp = _kline("107.RSP")
    if not spy or not rsp:
        cached = _load_cache()
        if cached:
            return cached  # 东财限流/抖动 → 用上次真数据兜底(绝不覆盖成空)
        return {"status": "未获取", "as_of": "",
                "source": "美股市场广度 RSP/SPY(东方财富, 抓取失败且无缓存)", "divergence": None}

    spy_map = dict(spy)
    rsp_map = dict(rsp)
    common = sorted(set(spy_map) & set(rsp_map))
    if len(common) < nh_window + 5:
        cached = _load_cache()
        if cached:
            return cached
        return {"status": "未获取", "as_of": common[-1] if common else "",
                "source": "美股市场广度 RSP/SPY(数据不足)", "divergence": None}

    # 广度比 = RSP/SPY, 归一化到起点=100 便于观感
    base = rsp_map[common[0]] / spy_map[common[0]]
    ratio_pts = [(d, round((rsp_map[d] / spy_map[d]) / base * 100.0, 3)) for d in common]
    spy_pts = [(d, round(spy_map[d], 2)) for d in common]

    # 数值化背离判定
    recent = common[-nh_window:]
    lb = common[-lookback_days:] if len(common) >= lookback_days else common
    spy_vals_lb = [spy_map[d] for d in lb]
    ratio_vals_lb = [rsp_map[d] / spy_map[d] for d in lb]
    spy_hi = max(spy_vals_lb)
    ratio_hi = max(ratio_vals_lb)
    # 最近 nh_window 内 SPY 是否触及/接近 lookback 新高(>= 99.5% of hi)
    spy_recent_hi = max(spy_map[d] for d in recent)
    ratio_recent_hi = max(rsp_map[d] / spy_map[d] for d in recent)
    spy_made_high = spy_recent_hi >= spy_hi * 0.995
    # 广度比距其 lookback 高点差多少 %
    ratio_gap_pct = round((ratio_hi - ratio_recent_hi) / ratio_hi * 100.0, 2)
    breadth_confirmed = ratio_recent_hi >= ratio_hi * 0.98  # 广度也接近新高=确认
    divergence = bool(spy_made_high and not breadth_confirmed)

    last_d = common[-1]
    evidence = {
        "spy_recent_high": round(spy_recent_hi, 2),
        "spy_lookback_high": round(spy_hi, 2),
        "spy_made_new_high": spy_made_high,
        "ratio_gap_from_high_pct": ratio_gap_pct,
        "breadth_confirmed": breadth_confirmed,
        "lookback_days": len(lb),
        "nh_window": nh_window,
        "rule": "SPY创近lookback新高 且 RSP/SPY广度比距其高点>2% = 顶背离",
    }

    out = {
        "status": "ok",
        "as_of": last_d,
        "divergence": divergence,
        "evidence": evidence,
        "spy_points": spy_pts[-lookback_days:],
        "ratio_points": ratio_pts[-lookback_days:],
        "latest_spy": spy_pts[-1][1],
        "latest_ratio": ratio_pts[-1][1],
        "source": "美股市场广度 RSP(等权)/SPY(市值加权) · 东方财富 push2his API",
    }
    try:
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def fetch_silver_bank_positions(cache_path=None, weeks=520):
    """白银做市商(bullion banks = commercial)头寸方向 —— CFTC COT commercial 净持仓。

    ★这是「投行/做市商在 COMEX 白银上是接货还是压价」的<b>一手官方真数据等价物</b>,
    语义覆盖 Michael Lynch(@DtDS_WSS) 的 "cumulative & monthly issues and stops by bullion
    banks" 图想说的核心(做市商头寸方向), 但用 CFTC 官方 Socrata(免key/可回溯到1986/可每周更新),
    而非被 CME 封禁的 issues/stops 逐日抓取。

    含义: commercial(商业套保, 主体是 bullion banks) 净持仓通常为净空(对冲实物多头)。
      - 净空<b>扩大</b> = 做市商加空压价 / 卖压增强(与图里 "Issued" 交货压价同向)
      - 净空<b>收窄/转多</b> = 做市商减空 / 被逼平 (与图里 "Stopped" 接货 squeeze 同向)
    绝不编造: 抓不到读上次缓存真值, 绝不覆盖成空。

    返回 {status, as_of, source, unit, latest_net, latest_wow, peak_short_date,
          points:[{date, comm_net, comm_long, comm_short, open_interest}], inventory_note}。
    """
    import json
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data",
                                  "silver_bank_positions.json")

    def _load_cache():
        try:
            with open(cache_path, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None

    # ★2026-09-25 Chao「在这张图上叠加白银价格走势」: Yahoo SI=F 连续合约日线(chart API v8 直连)。
    # 抓不到读上次缓存价格, 绝不编造; 价格与 COT 抓取互不影响。
    def _silver_px():
        import datetime as _dt, urllib.request as _ur
        u = "https://query1.finance.yahoo.com/v8/finance/chart/SI=F?range=10y&interval=1d"
        try:
            r = _ur.urlopen(_ur.Request(u, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                                                    "AppleWebKit/537.36 Chrome/120 Safari/537.36"}), timeout=45)
            res = json.loads(r.read().decode("utf-8", "ignore"))["chart"]["result"][0]
            cl = res["indicators"]["quote"][0]["close"]
            return [[_dt.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d"), round(float(c), 3)]
                    for t, c in zip(res.get("timestamp") or [], cl) if c is not None]
        except Exception:
            return []
    price = _silver_px()
    if len(price) < 20:
        price = (_load_cache() or {}).get("price") or []
        price_status = "缓存" if price else "未获取"
    else:
        price_status = "ok"

    pts = []
    try:
        from fetchers import cot as _cot
        hist = _cot.fetch_history("silver", weeks=weeks)
        for r in hist:
            if not r.get("as_of"):
                continue
            pts.append({
                "date": r["as_of"],
                "comm_net": r.get("comm_net"),
                "comm_long": r.get("comm_long"),
                "comm_short": r.get("comm_short"),
                "open_interest": r.get("open_interest"),
            })
    except Exception:
        pts = []

    if not pts:
        cached = _load_cache()
        if cached and cached.get("points"):
            cached["status"] = "缓存"
            return cached
        return {"status": "未获取", "as_of": "", "unit": "contracts",
                "source": "CFTC COT 白银 commercial 净持仓(未获取)", "points": []}

    pts.sort(key=lambda p: p.get("date", ""))
    latest = pts[-1]
    prev = pts[-2] if len(pts) > 1 else None
    latest_net = latest.get("comm_net")
    latest_wow = (latest_net - prev["comm_net"]) if (prev and prev.get("comm_net") is not None
                                                      and latest_net is not None) else None
    # 净空最深点(comm_net 最小)
    valid = [p for p in pts if p.get("comm_net") is not None]
    peak_short = min(valid, key=lambda p: p["comm_net"]) if valid else None

    # COMEX 白银库存当前锚点(文字标注, 非每日折线 —— 无稳定免key每日源)
    inv_note = ("COMEX 白银库存约 335M oz(≈10,400 吨), 其中注册(registered)"
                "占比与做市商囤货状态需结合 CME 每日仓单报告(metalcharts.org 快照)。")

    out = {
        "status": "ok",
        "as_of": latest["date"],
        "unit": "contracts",
        "source": "CFTC COT · SILVER - COMMODITY EXCHANGE INC. · commercial 净持仓(官方 Socrata, 免key)",
        "latest_net": latest_net,
        "latest_long": latest.get("comm_long"),
        "latest_short": latest.get("comm_short"),
        "latest_oi": latest.get("open_interest"),
        "latest_wow": latest_wow,
        "peak_short_date": (peak_short or {}).get("date"),
        "peak_short_net": (peak_short or {}).get("comm_net"),
        "inventory_note": inv_note,
        "points": pts,
        "price": price,
        "price_status": price_status,
        "price_source": "Yahoo Finance · SI=F COMEX 白银连续合约 日收盘(美元/盎司)",
    }
    try:
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def fetch_comex_silver_issues_ref(cache_path=None):
    """C: COMEX 白银 issues/stops by bullion banks 静态参考(数据源=Michael Lynch @DtDS_WSS)。

    ★为什么静态不自动更新: 原始底层是 CME COMEX 每日 Issues&Stops by firm 报告, CME 官方
    明确封禁脚本抓取(IP block + Data Terms of Use), 无免key可回溯每日源。ANONYMIZED_PERSON_0_12 本人在
    Substack(econanalytics.substack.com)/X(@DtDS_WSS)周期性更新, 但**事件驱动、不规律、
    无 API、图为嵌入 PNG、不公开原始累加序列**。故本 section 只做诚实的静态参考锚点 + 明确
    标注来源与更新性质, 绝不伪造每日折线。真实可更新的做市商头寸方向见 fetch_silver_bank_positions
    (CFTC 官方一手数据)。

    数据由 data/comex_silver_issues_ref.json 驱动(从公开图手抄的关键锚点, 明确标注非实时)。
    返回 {status, as_of, source, source_url, note, points:[{date, cumulative_koz}], annotations}。
    """
    import json
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data",
                                  "comex_silver_issues_ref.json")
    try:
        with open(cache_path, encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        return {"status": "未获取", "as_of": "", "unit": "thousands_oz",
                "source": "COMEX 白银累计 issues/stops(文件缺失)", "points": []}
    pts = [p for p in (data.get("points") or []) if p.get("date")]
    pts.sort(key=lambda p: p.get("date", ""))
    return {
        "status": "ok" if pts else "未获取",
        "as_of": data.get("as_of", ""),
        "unit": data.get("unit", "thousands_oz"),
        "source": data.get("source", "Michael Lynch (@DtDS_WSS) · EconAnalytics"),
        "source_url": data.get("source_url", "https://econanalytics.substack.com/"),
        "update_nature": data.get("update_nature",
                                  "事件驱动·不规律更新·CME官方封禁脚本抓取·仅静态参考锚点"),
        "note": data.get("note", ""),
        "annotations": data.get("annotations", []),
        "points": pts,
    }


def fetch_gold_exports(cache_path=None):
    """美国黄金出口(Nonmonetary gold) —— FRED IEAXGG 季度真数据(1999→今, Mil USD)。

    ★对应用户图「各国黄金运回家·2025-26 美国出口货币黄金飙升」。图作者原标 Monetary gold,
    但 FRED 无此序列; 真正对应该暴涨现象的是 <b>Nonmonetary gold</b>(非货币黄金 = 民间/商业实物金
    出口, 非央行储备金)。本函数用 FRED 官方真序列 IEAXGG + 真标题, 诚实标注(不照抄不准的
    Monetary)。免 key(fredgraph.csv 直连), 可回溯, 每季更新。

    每季从 FRED 拉最新, 与本地缓存合并(只增新点), 抓不到读缓存真值绝不覆盖成空。
    返回 {status, as_of, source, unit, latest, peak, base_2024_avg, surge_x, points:[{date,value_musd}]}。
    """
    import json
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "gold_exports.json")

    def _load_cache():
        try:
            with open(cache_path, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None

    cached = _load_cache()
    # 从 FRED 拉最新(免key fredgraph.csv), 与缓存合并
    fresh = _fred_series("IEAXGG")  # [(date,val)] 升序
    pts = []
    if fresh:
        pts = [{"date": d, "value_musd": round(float(v), 1)} for d, v in fresh]
    elif cached and cached.get("points"):
        out = dict(cached)
        out["status"] = "缓存"
        return out
    if not pts:
        return {"status": "未获取", "as_of": "", "unit": "millions_usd",
                "source": "FRED IEAXGG 美国黄金出口(未获取)", "points": []}

    pts.sort(key=lambda p: p["date"])
    vals = [p["value_musd"] for p in pts]
    peak = max(vals); peak_date = pts[vals.index(peak)]["date"]
    base_2024 = [p["value_musd"] for p in pts if p["date"].startswith("2024")]
    base_avg = round(sum(base_2024) / len(base_2024), 0) if base_2024 else None
    latest = pts[-1]["value_musd"]
    surge_x = round(latest / base_avg, 1) if base_avg else None

    out = {
        "status": "ok",
        "as_of": pts[-1]["date"],
        "series_id": "IEAXGG",
        "unit": "millions_usd",
        "frequency": "quarterly",
        "source": "FRED / U.S. BEA · Exports of Goods: Nonmonetary gold (IEAXGG)",
        "source_url": "https://fred.stlouisfed.org/series/IEAXGG",
        "title_note": (cached or {}).get("title_note",
                       "图作者原标 Monetary gold, FRED 无此序列; 真正对应暴涨的是 Nonmonetary gold"
                       "(非货币黄金=民间/商业实物金出口, 非央行储备金)。本 section 用 FRED 官方真标题诚实标注。"),
        "latest": latest, "latest_date": pts[-1]["date"],
        "peak": peak, "peak_date": peak_date,
        "base_2024_avg": base_avg, "surge_x": surge_x,
        "points": pts,
    }
    try:
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def fetch_us_yield_century(cache_path=None):
    """图5: 美国国债收益率百年周期(4线月度: Fed Funds/3M/10Y/30Y, FRED)。

    数据由 data/us_yield_century.json 驱动(scratch/build_yield_century.py 从 FRED 生成)。
    每季 cron 增量刷新最新月。抓不到读缓存真值绝不覆盖成空。
    诚实: 各线起点为 FRED 收录起点(TB3M 1934/FEDFUNDS 1954/GS10 1953/GS30 1977),
    1920-1933 FRED 无月度序列, 不编造。周期锚点 1940大底/1981大顶/2020大底。
    返回 {status, as_of, source, unit, annotations, cycles, series:{key:{label,color,points}}}。
    """
    import json
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "us_yield_century.json")
    try:
        with open(cache_path, encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        return {"status": "未获取", "as_of": "", "unit": "%",
                "source": "美国国债收益率百年周期(文件缺失)", "series": {}}
    if not data.get("series"):
        return {"status": "未获取", "as_of": data.get("as_of", ""), "unit": "%",
                "source": data.get("source", ""), "series": {}}
    data["status"] = data.get("status", "ok")
    return data


# ★2026-08-22(Chao): COMEX 交割量「手 → 吨」换算常量。
# 合约规格来自 CME Group 官方规则手册, 已联网核对:
#   Gold     — Chapter 113 Gold Futures:     100 troy oz / 手
#   Silver   — Chapter 112 Silver Futures: 5,000 troy oz / 手
#   Platinum — NYMEX Platinum Futures (PL):    50 troy oz / 手
# 1 troy ounce = 31.1034768 g (国际标准, 精确值)
# → 吨 = 手数 × oz/手 × 31.1034768 / 1e6
_TROY_OZ_G = 31.1034768
_LOT_OZ = {"Gold": 100, "Silver": 5000, "Platinum": 50}


def fetch_comex_issue_stop_firms(cache_path=None):
    """需求B(Chao 2026-08-22): COMEX 三金属 per-firm top10 交货方/接货方。

    数据由 Comex-Daily-Report/data/comex_issue_stop_firms.json 驱动
    (src/build_issue_stop_firms.py 生成: archive PDF 133份 per-firm 精确解析
     + Notion「Daily auto tracking」增量, 覆盖 2026-01-06 起)。

    与 fetch_comex_issue_stop_weekly 的区别: 那个是大行汇总的周度净值曲线;
    这个保留 <b>每家机构</b> 明细(含 StoneX/Marex 等非银)且含 <b>铂金</b>,
    用于排 top10 榜单。

    ★排名口径(Chao 2026-08-22 定): 按 <b>总量</b> 排, 不按净额。
      「top10 发货方」= 该窗口内 issued 累计最大的十家;
      「top10 接货方」= stopped 累计最大的十家。
      净额会让大进大出的做市商掉榜, 不符合"谁交货最多"的直觉。

    返回 {status, as_of, coverage, bullion_banks,
          windows:[{key,label,days}],
          metals:{Gold:{window_key:{issuers:[...], stoppers:[...], total_i, total_s}}}}
    每条 = {firm, lots, share, is_bank}。抓不到 status='未获取', 绝不编造。
    """
    import json
    import datetime as _dt
    from collections import defaultdict

    if cache_path is None:
        cache_path = "/home/user/Projects/Comex-Daily-Report/data/comex_issue_stop_firms.json"
    try:
        with open(cache_path, encoding="utf-8") as f:
            raw = json.load(f)
    except Exception:
        return {"status": "未获取", "as_of": "",
                "source": "COMEX per-firm 交割明细(文件缺失)", "metals": {}}
    if raw.get("status") != "ok" or not raw.get("metals"):
        return {"status": "未获取", "as_of": raw.get("coverage", {}).get("end", ""),
                "source": raw.get("source", ""), "metals": {}}

    cov = raw.get("coverage", {})
    end = cov.get("end", "")
    try:
        last = _dt.date.fromisoformat(end)
    except Exception:
        last = None

    # 日/周/月三个粒度 + 全期。日=最后一个有交割的日子(非自然日, 交割不是每天都有)
    wins = [("d", "最新交割日", 1), ("w", "近一周", 7),
            ("m", "近一个月", 30), ("all", "全期", None)]
    banks = set(raw.get("bullion_banks", []))

    out_metals = {}
    for metal, bydate in raw["metals"].items():
        if not bydate:
            continue
        dates = sorted(bydate)
        per_win = {}
        for key, label, days in wins:
            if days is None or last is None:
                sel = dates
            elif key == "d":
                sel = dates[-1:]                     # 最后一个有数据的交割日
            else:
                cut = last - _dt.timedelta(days=days - 1)
                sel = [d for d in dates
                       if _dt.date.fromisoformat(d) >= cut] or dates[-1:]
            agg = defaultdict(lambda: {"i": 0, "s": 0})
            for d in sel:
                for firm, v in bydate[d].items():
                    agg[firm]["i"] += v.get("i", 0)
                    agg[firm]["s"] += v.get("s", 0)
                    # ★2026-08-22(Chao): 账户类型分项 C(客户)/H(自营)/U(未知)
                    for _k in ("i_c", "i_h", "i_u", "s_c", "s_h", "s_u"):
                        if v.get(_k):
                            agg[firm][_k] = agg[firm].get(_k, 0) + v[_k]
            ti = sum(v["i"] for v in agg.values())
            ts = sum(v["s"] for v in agg.values())

            def _rank(field, total):
                rows = [(f, v[field]) for f, v in agg.items() if v[field] > 0]
                rows.sort(key=lambda x: (-x[1], x[0]))
                out_rows = []
                for f, n in rows[:10]:
                    v = agg[f]
                    c = v.get(field + "_c", 0)
                    hh = v.get(field + "_h", 0)
                    u = v.get(field + "_u", 0)
                    out_rows.append({
                        "firm": f, "lots": n,
                        "tonnes": round(n * _LOT_OZ.get(metal, 0) * _TROY_OZ_G / 1e6, 3),
                        "share": round(n / total * 100, 1) if total else 0.0,
                        "is_bank": f in banks,
                        # ★账户类型: C=客户盘(清算通道) / H=自营盘(机构自有资金判断) / U=未知
                        "c": c, "h": hh, "u": u,
                        "h_pct": round(hh / n * 100, 1) if n else 0.0,
                    })
                return out_rows

            # 窗口级 C/H 汇总(供板块顶部一行概览)
            def _org_tot(field):
                return {k: sum(v.get(field + "_" + k, 0) for v in agg.values())
                        for k in ("c", "h", "u")}

            per_win[key] = {
                "label": label,
                "days": len(sel),
                "range": [sel[0], sel[-1]] if sel else ["", ""],
                "total_i": ti, "total_s": ts,
                "total_i_t": round(ti * _LOT_OZ.get(metal, 0) * _TROY_OZ_G / 1e6, 3),
                "total_s_t": round(ts * _LOT_OZ.get(metal, 0) * _TROY_OZ_G / 1e6, 3),
                "org_i": _org_tot("i"), "org_s": _org_tot("s"),
                "issuers": _rank("i", ti),
                "stoppers": _rank("s", ts),
            }
        out_metals[metal] = per_win

    return {
        "status": "ok" if out_metals else "未获取",
        "as_of": end,
        "source": raw.get("source", ""),
        "unit": raw.get("unit", "合约手数"),
        "coverage": cov,
        "bullion_banks": sorted(banks),
        "windows": [{"key": k, "label": l} for k, l, _ in wins],
        "metals": out_metals,
    }


def fetch_comex_issue_stop_weekly(cache_path=None):
    """图1: COMEX 做市商每周净 issue/stop(金+银, 两口径: 全投行/核心做市商)。

    数据由 Comex-Daily-Report/data/comex_issue_stop_weekly.json 驱动
    (scratch/build_issue_stop_weekly.py 扫 archive 133 PDF + Notion 6/6+ 生成)。
    净值=Σ(大行发货)−Σ(大行接货), 正=净发货(交货/压价), 负=净接货(囤货/看涨)。
    每日 cron 增量: 采集当日 PDF→per-firm 拆分→并入当周。抓不到读缓存绝不覆盖成空。
    返回 {status, as_of, gold:[{week,all_net,core_net,...}], silver:[...], banks_all, banks_core}。
    """
    import json
    if cache_path is None:
        # 数据在姊妹 Comex 项目, Eco dashboard 跨项目读
        cache_path = "/home/user/Projects/Comex-Daily-Report/data/comex_issue_stop_weekly.json"
    try:
        with open(cache_path, encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        return {"status": "未获取", "as_of": "",
                "source": "COMEX 做市商周净 issue/stop(文件缺失)", "gold": [], "silver": []}
    if not (data.get("gold") or data.get("silver")):
        return {"status": "未获取", "as_of": data.get("archive_range", ""),
                "source": data.get("note", ""), "gold": [], "silver": []}
    data["status"] = data.get("status", "ok")
    # as_of = 最新周
    last_weeks = [r["week"] for r in (data.get("silver") or []) + (data.get("gold") or [])]
    data["as_of"] = max(last_weeks) if last_weeks else data.get("archive_range", "")
    return data


def fetch_fiscal_news(cache_path=None, limit=20):
    """美日财政政策事件时间线(离散事件文本,非时序数字)。
    数据由 data/fiscal_news.json 驱动: 每日 cron agent 模式 web 检索权威源
    (US Treasury/Congress/路透/彭博 · 日本财务省/NHK/日经)动态写入最新美日财政举措
    (债务上限/持续决议CR/政府关门/补正预算/国债发行计划等)。
    绝不编造: 每条事件带真实 source_url。本函数只读文件+按日期倒序+截断到 limit 条。
    返回 {status, as_of, source, events:[{date,country,flag,category,title,summary,source_url,source_name}]}。
    """
    import json
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "fiscal_news.json")
    try:
        with open(cache_path, encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        return {"status": "未获取", "as_of": "", "source": "美日财政事件(文件缺失)", "events": []}
    events = data.get("events", []) or []
    # 只保留有必需字段的真实条目, 按日期倒序
    clean = [e for e in events if e.get("date") and e.get("title") and e.get("source_url")]
    clean.sort(key=lambda e: e.get("date", ""), reverse=True)
    clean = clean[:limit]
    return {
        "status": "ok" if clean else "未获取",
        "as_of": data.get("as_of", clean[0]["date"] if clean else ""),
        "source": data.get("source", "美日财政政策事件时间线"),
        "note": data.get("note", ""),
        "events": clean,
    }


def fetch_oil_inventory(cache_path=None):
    """美国石油库存运营红线三序列。绝不编造,取不到该项 status='未获取'。
      1) Brent-WTI 价差(过去一年,日频): FRED DCOILBRENTEU - DCOILWTICO(同日对齐,$/桶)
         价差转负(WTI>Brent)=Cushing 交割枢纽库存逼近 tank bottom 的市场信号。
      2) Cushing 原油库存(过去一年,周频): EIA W_EPC0_SAX_YCUOK_MBBL(千桶→百万桶)。运营红线~2000万桶(tank bottom)。
      3) SPR 战略石油储备(过去十年,周频): EIA WCSSTUS1(千桶→百万桶)。运营红线3亿桶(Amos Hochstein 披露)。
    返回 {status, as_of, spread:{...}, cushing:{...}, spr:{...}}。
    """
    import json, datetime as _dt
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "oil_inventory.json")
    today = _dt.date.today()
    out = {"status": "ok", "as_of": today.isoformat(),
           "spread": {"status": "未获取"}, "cushing": {"status": "未获取"}, "spr": {"status": "未获取"}}

    # 1) Brent-WTI 价差(过去一年,日频)
    start_1y = (today - _dt.timedelta(days=400)).isoformat()
    brent = dict(_fred_series("DCOILBRENTEU", start=start_1y))
    wti = dict(_fred_series("DCOILWTICO", start=start_1y))
    common = sorted(set(brent) & set(wti))
    # 只保留近约 366 天
    cutoff = (today - _dt.timedelta(days=366)).isoformat()
    spread_pts = [(d, round(brent[d] - wti[d], 2)) for d in common if d >= cutoff]
    if len(spread_pts) >= 2:
        last_d, last_v = spread_pts[-1]
        neg_days = sum(1 for _, v in spread_pts if v < 0)
        out["spread"] = {
            "status": "ok", "points": spread_pts, "as_of": last_d,
            "latest": last_v, "hi": max(v for _, v in spread_pts),
            "lo": min(v for _, v in spread_pts), "neg_days": neg_days,
            "source": "FRED (DCOILBRENTEU - DCOILWTICO), 日频",
        }

    # 2) Cushing 库存(过去一年,周频) — EIA 千桶→百万桶
    cush = _eia_series("W_EPC0_SAX_YCUOK_MBBL", start=(today - _dt.timedelta(days=400)).strftime("%Y-%m-%d"))
    cush = [(d, round(v / 1000.0, 2)) for d, v in cush if d >= cutoff]
    if len(cush) >= 2:
        cd, cv = cush[-1]
        out["cushing"] = {
            "status": "ok", "points": cush, "as_of": cd, "latest": cv,
            "hi": max(v for _, v in cush), "lo": min(v for _, v in cush),
            "redline": 20.0,  # 运营红线 2000 万桶(tank bottom)
            "source": "EIA Weekly Petroleum Status (W_EPC0_SAX_YCUOK_MBBL), 周频",
        }

    # 3) SPR(过去十年,周频) — EIA 千桶→百万桶
    spr_start = (today - _dt.timedelta(days=3670)).strftime("%Y-%m-%d")  # ~10年+buffer
    spr = _eia_series("WCSSTUS1", start=spr_start)
    spr = [(d, round(v / 1000.0, 2)) for d, v in spr]
    if len(spr) >= 2:
        sd, sv = spr[-1]
        out["spr"] = {
            "status": "ok", "points": spr, "as_of": sd, "latest": sv,
            "hi": max(v for _, v in spr), "lo": min(v for _, v in spr),
            "redline": 300.0,  # 运营红线 3 亿桶(Amos Hochstein 披露)
            "source": "EIA Weekly Ending Stocks Crude Oil in SPR (WCSSTUS1), 周频",
        }

    if all(out[k].get("status") != "ok" for k in ("spread", "cushing", "spr")):
        out["status"] = "未获取"
    try:
        with open(cache_path, "w") as f:
            json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def _fetch_tonar_history(months=4):
    """BoJ TONAR(无担保隔夜拆借加权平均利率, 日债隔夜融资成本基准, 对应美债 SOFR)日频时序。
    数据源: BoJ 官方每工作日 XLSX(md{YYYYMMDD}.xlsx, sheet 'コール', Average 行)。
    抓最近 N 个月的 final results 索引 → 逐个 XLSX 取 Average。绝不编: 抓不到的日跳过, 全失败返回 []。
    返回 [(date 'YYYY-MM-DD', rate%), ...] 升序。"""
    import datetime as _dt
    import re as _re
    import subprocess
    import io as _io
    try:
        import openpyxl
    except Exception:
        return []
    today = _dt.date.today()
    years_needed = sorted({today.year, (today - _dt.timedelta(days=months * 31)).year})
    # 1) 从各年索引页收集 XLSX 文件名(md{YYYYMMDD}.xlsx)
    files = []
    for y in years_needed:
        url = f"https://www.boj.or.jp/en/statistics/market/short/mutan/d_release/md/{y}/index.htm"
        try:
            r = subprocess.run(["curl", "-s", "--max-time", "25", "-H", "User-Agent: Mozilla/5.0", url],
                               capture_output=True, text=True, timeout=30)
            for m in _re.findall(r'md(\d{8})\.xlsx', r.stdout or ""):
                files.append((y, m))
        except Exception:
            pass
    if not files:
        return []
    # 去重 + 按日期取最近 months*22 个交易日(约 months 个月)
    files = sorted(set(files), key=lambda x: x[1])
    limit_n = int(months * 23)
    files = files[-limit_n:]
    # 2) 逐个 XLSX 取 Average(报告日 = 文件名日期的前一交易日, 但文件名 md{报告发布日}; 标题内含真实报告日)
    out = []
    for y, ymd in files:
        fu = f"https://www.boj.or.jp/en/statistics/market/short/mutan/d_release/md/{y}/md{ymd}.xlsx"
        try:
            r = subprocess.run(["curl", "-s", "--max-time", "20", "-H", "User-Agent: Mozilla/5.0", fu],
                               capture_output=True, timeout=25)
            if r.returncode != 0 or not r.stdout:
                continue
            wb = openpyxl.load_workbook(_io.BytesIO(r.stdout), data_only=True)
            ws = wb.worksheets[0]
            rep_date = None
            avg = None
            for row in ws.iter_rows(values_only=True):
                cells = [c for c in row if c is not None]
                for i, c in enumerate(cells):
                    cs = str(c)
                    # 报告日: 标题 "Uncollateralized Overnight Call Rate for August 18 ..." 或日文 "8月18日"
                    if rep_date is None:
                        mdate = _re.search(r'for\s+([A-Z][a-z]+)\s+(\d{1,2})', cs)
                        if mdate:
                            try:
                                mon = _dt.datetime.strptime(mdate.group(1), "%B").month
                                day = int(mdate.group(2))
                                rep_date = _dt.date(y, mon, day).isoformat()
                            except Exception:
                                pass
                    # Average 值
                    if "Average" in cs or cs.strip() == "平均":
                        for nxt in cells[i + 1:]:
                            try:
                                avg = float(nxt)
                                break
                            except (ValueError, TypeError):
                                continue
            if rep_date and avg is not None:
                out.append((rep_date, round(avg, 4)))
        except Exception:
            continue
    # 去重(同报告日取一次)+升序
    dd = {}
    for d, v in out:
        dd[d] = v
    return sorted(dd.items())


def fetch_basis_trade_monitor(cache_path=None, years=2):
    """基差套利去杠杆预警监控 —— 时序化你的「美债/日债基差套利 + SOFR 倒挂 + TONAR 倒挂」监控表。

    核心风险: 基差套利对冲基金(买现券/卖期货, repo 加 33-99x 杠杆)在
      「融资成本↑(SOFR 触顶 IORB) + carry 消失(收益率-SOFR 利差收窄/倒挂) + 波动↑(MOVE)」时
      被迫去杠杆强平 → 抛售美/日债 → 收益率跳升 → 踩踏(2020-03 / 反复出现的尾部风险)。

    产出 3 个双轴 panel(复用 _stress_panel_svg 渲染) + 1 个分期限状态矩阵:
      panel_funding : 融资成本压力(SOFR/IORB/EFFR/ON-RRP 利率 + 日本 TONAR), 看 SOFR 触顶 IORB
      panel_carry   : 套利 carry 空间(美债 2/5/10/30Y − SOFR 利差 + 日债 10/30Y − TONAR), 转负=倒挂=红区
      panel_trigger : 去杠杆触发器(MOVE 波动率 左轴 + Fed 准备金 右轴), 火药桶+火星
      matrix        : 分期限状态灯矩阵(美债 2/5/10/30Y + 日债 10/30Y), carry/信号灯/30日 sparkline

    全部真实公开数据(FRED / 日本 MOF / Economic-Dashboard Notion 准备金)。绝不编: 缺失序列跳过 + status 反映。
    周度降采样降噪。返回 {panels:{...}, matrix:[...], asof, status, lights:{...}}。
    """
    import json as _json
    import datetime as _dt
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "basis_trade_monitor.json")
    today = _dt.date.today()
    start = (today - _dt.timedelta(days=int(365.3 * years) + 15)).isoformat()

    def _wk(pts):
        """[(date,val)] → 周度降采样 [{date,v}]。"""
        if not pts:
            return []
        wp = _weekly_resample(pts, agg="last")
        return [{"date": d, "v": round(v, 4)} for d, v in wp]

    def _last(pts):
        return pts[-1] if pts else None

    # ── 拉核心序列(全 FRED, 日频) ──
    sofr = _fred_series_hist("SOFR", start)        # 隔夜融资成本(套利头寸的资金成本)
    iorb = _fred_series_hist("IORB", start)        # 准备金利率(SOFR 的政策上限/走廊顶)
    effr = _fred_series_hist("EFFR", start)        # 联邦基金有效利率
    rrp_rate = _fred_series_hist("RRPONTSYAWARD", start)  # ON RRP 利率(走廊底)
    dgs = {}
    for t, sid in [("2Y", "DGS2"), ("5Y", "DGS5"), ("10Y", "DGS10"), ("30Y", "DGS30")]:
        dgs[t] = _fred_series_hist(sid, start)

    # 日债(MOF 日频, 全期限用于 carry): 2Y/5Y/10Y/30Y
    jgb = {}
    for t, col in [("2Y", 6), ("5Y", 9), ("10Y", 10), ("30Y", 14)]:
        d = _mof_jgb_yields(col)
        jgb[t] = sorted(([k, v] for k, v in d.items() if k >= start), key=lambda x: x[0]) if d else []
    # 日本隔夜融资成本 TONAR(无担保隔夜拆借加权平均, 对应美债 SOFR): BoJ 官方每日 XLSX 拼近 4 月日频。
    # 抓不到则日债 carry 回退不做主判定(诚实标注), 绝不用常数近似冒充。
    try:
        tonar = _fetch_tonar_history(months=4)
    except Exception:
        tonar = []
    tonar_d = dict(tonar) if tonar else {}

    # ═══ Panel 1: 融资成本压力(SOFR 倒挂监控, 时序化表3) ═══
    funding_series = []
    if sofr:
        funding_series.append({"name": "SOFR 隔夜融资成本", "color": "#c0757d", "axis": "left",
                               "width": 2.2, "points": _wk(sofr)})
    if iorb:
        funding_series.append({"name": "IORB 准备金利率(走廊顶)", "color": "#8a3f47", "axis": "left",
                               "dash": True, "points": _wk(iorb)})
    if effr:
        funding_series.append({"name": "EFFR 联邦基金利率", "color": "#7fa085", "axis": "left",
                               "points": _wk(effr)})
    if rrp_rate:
        funding_series.append({"name": "ON RRP 利率(走廊底)", "color": "#b58a6a", "axis": "left",
                               "dash": True, "points": _wk(rrp_rate)})
    if tonar:
        funding_series.append({"name": "TONAR 日债隔夜融资(BoJ)", "color": "#c17d6a", "axis": "left",
                               "width": 2.0, "dash": True, "points": _wk(tonar)})
    panel_funding = {
        "id": "bt_funding",
        "title": "① 融资成本压力：SOFR 触顶 IORB + TONAR 走势（美 vs 日）",
        "subtitle": "Funding Cost Stress — US SOFR vs Corridor + JP TONAR",
        "unit_left": "%", "unit_right": "%", "single_axis": True,
        "series": funding_series,
        "source": "FRED SOFR / IORB / EFFR / RRPONTSYAWARD · BoJ TONAR（日频，周度降采样）",
    }

    # ═══ Panel 2: 套利 carry 空间(收益率 − SOFR 各期限利差) ═══
    # carry = 持有该期限国债收益率 − 隔夜融资成本(SOFR)。收窄/转负 = 套利利润消失 → 强平动机。
    sofr_d = dict(sofr) if sofr else {}
    carry_series = []
    carry_colors = {"2Y": "#6b8fb5", "5Y": "#7fa085", "10Y": "#c0757d", "30Y": "#8a3f47"}
    carry_latest = {}
    for t in ["2Y", "5Y", "10Y", "30Y"]:
        yld = dict(dgs.get(t, []))
        common = sorted(set(yld) & set(sofr_d))
        spread = [(d, round((yld[d] - sofr_d[d]) * 100, 1)) for d in common]  # bp
        if len(spread) >= 2:
            carry_series.append({"name": f"美债 {t}−SOFR carry", "color": carry_colors[t],
                                 "axis": "left", "points": _wk(spread)})
            carry_latest[f"US_{t}"] = spread[-1]
    # 日债 carry (JGB − TONAR): 日债侧真实融资成本基准(对应美债 SOFR)。虚线区分美日。
    jp_carry_colors = {"10Y": "#c17d6a", "30Y": "#8a5a52"}
    for t in ["10Y", "30Y"]:
        yld = dict(jgb.get(t, []))
        common = sorted(set(yld) & set(tonar_d))
        spread = [(d, round((yld[d] - tonar_d[d]) * 100, 1)) for d in common]  # bp
        if len(spread) >= 2:
            carry_series.append({"name": f"日债 {t}−TONAR carry", "color": jp_carry_colors[t],
                                 "axis": "left", "dash": True, "points": _wk(spread)})
            carry_latest[f"JP_{t}"] = spread[-1]
    panel_carry = {
        "id": "bt_carry",
        "title": "② 持券 Carry：国债收益率 − 隔夜融资成本（各期限利差，美债 vs 日债）",
        "subtitle": "Cash Carry — Yield minus Overnight Funding by Tenor (US SOFR / JP TONAR) · 注: 非期现 basis trade",
        "unit_left": "bp", "unit_right": "bp", "single_axis": True,
        "series": carry_series,
        "source": "FRED DGS2/5/10/30 − SOFR · 日本 MOF JGB − BoJ TONAR（日频，周度降采样，单位 bp；实线=美债，虚线=日债）",
    }

    # ═══ Panel 3: 去杠杆触发器(MOVE 波动率 + Fed 准备金) ═══
    move = _fetch_move_history(start)
    trigger_series = []
    if move:
        trigger_series.append({"name": "MOVE 债市波动率", "color": "#c17d6a", "axis": "left",
                               "width": 2.0,
                               "points": [{"date": d, "v": round(v, 4)} for d, v in _weekly_resample(move, agg="peak")]})
    # Fed 准备金(万亿$) — 从 FRED WRESBAL(周度, 百万$→万亿$)
    reserves = _fred_series_hist("WRESBAL", start)
    if reserves:
        res_t = [(d, round(v / 1_000_000.0, 3)) for d, v in reserves]  # 百万$ → $T
        trigger_series.append({"name": "Fed 准备金 $T（右轴，缓冲垫）", "color": "#6b8fb5", "axis": "right",
                               "points": _wk(res_t)})
    panel_trigger = {
        "id": "bt_trigger",
        "title": "③ 去杠杆触发器：波动率飙升 × 流动性缓冲枯竭",
        "subtitle": "Deleverage Trigger — MOVE Volatility × Fed Reserves",
        "unit_left": "MOVE", "unit_right": "$T",
        "series": trigger_series,
        "source": "FRED（^MOVE via Yahoo · WRESBAL 准备金余额）",
    }

    # ═══ 状态矩阵: 分期限 carry + 信号灯 + 30日 sparkline ═══
    # 信号灯规则(基于 carry, 越低越危险): >30bp 🟢 / 0~30bp 🟡 / <0bp(倒挂) 🔴
    def _light(carry_bp):
        if carry_bp is None:
            return "⚪"
        if carry_bp < 0:
            return "🔴"
        if carry_bp < 30:
            return "🟡"
        return "🟢"

    matrix = []
    for t in ["2Y", "5Y", "10Y", "30Y"]:
        yld = dict(dgs.get(t, []))
        common = sorted(set(yld) & set(sofr_d))
        spread = [(d, round((yld[d] - sofr_d[d]) * 100, 1)) for d in common]
        if not spread:
            matrix.append({"market": "美债", "tenor": t, "carry_bp": None, "light": "⚪",
                           "yield": None, "spark": [], "status": "未获取"})
            continue
        cur_c = spread[-1][1]
        spark = [v for _, v in spread[-30:]]
        spark_d = [d for d, _ in spread[-30:]]   # tooltip 用真实日期
        matrix.append({
            "market": "美债", "tenor": t, "carry_bp": cur_c, "light": _light(cur_c),
            "yield": (yld[spread[-1][0]] if spread[-1][0] in yld else None),
            "date": spread[-1][0], "spark": spark, "spark_d": spark_d, "status": "ok",
        })
    # 日债行(carry = JGB − TONAR, 真实融资成本基准, 对应美债 SOFR)。TONAR 抓不到才回退不判定。
    for t in ["2Y", "5Y", "10Y", "30Y"]:
        pts = jgb.get(t, [])
        yld = dict(pts)
        if not pts:
            matrix.append({"market": "日债", "tenor": t, "carry_bp": None, "light": "⚪",
                           "yield": None, "spark": [], "status": "未获取"})
            continue
        cur_y = pts[-1][1]
        if tonar_d:
            common = sorted(set(yld) & set(tonar_d))
            spread = [(d, round((yld[d] - tonar_d[d]) * 100, 1)) for d in common]
            if spread:
                cur_c = spread[-1][1]
                spark = [v for _, v in spread[-30:]]
                spark_d = [d for d, _ in spread[-30:]]   # tooltip 用真实日期
                matrix.append({
                    "market": "日债", "tenor": t, "carry_bp": cur_c, "light": _light(cur_c),
                    "yield": (yld[spread[-1][0]] if spread[-1][0] in yld else cur_y),
                    "date": spread[-1][0], "spark": spark, "spark_d": spark_d, "status": "ok",
                })
                continue
        # TONAR 抓不到 → 只展示收益率+走势, carry 诚实标 n/a
        spark = [v for _, v in pts[-30:]]
        spark_d = [d for d, _ in pts[-30:]]   # tooltip 用真实日期
        matrix.append({
            "market": "日债", "tenor": t, "carry_bp": None, "light": "⚪",
            "yield": cur_y, "date": pts[-1][0], "spark": spark, "spark_d": spark_d, "status": "ok",
            "note": "TONAR 隔夜融资序列本次未取到，carry 暂不判定",
        })

    # ── 合成风险灯(供 dashboard 顶部/简报) ──
    lights = {}
    # SOFR 触顶 IORB?
    s_last, i_last = _last(sofr), _last(iorb)
    if s_last and i_last:
        gap = (s_last[1] - i_last[1]) * 100  # bp, SOFR-IORB
        lights["sofr_iorb_gap_bp"] = round(gap, 1)
        lights["funding"] = "🔴" if gap >= 0 else ("🟡" if gap >= -5 else "🟢")
    # carry 最紧期限(美债 + 日债分开统计, 也给总最紧)
    us_carry = [m["carry_bp"] for m in matrix if m.get("market") == "美债" and m.get("carry_bp") is not None]
    jp_carry = [m["carry_bp"] for m in matrix if m.get("market") == "日债" and m.get("carry_bp") is not None]
    valid_carry = [m["carry_bp"] for m in matrix if m.get("carry_bp") is not None]
    if valid_carry:
        mn = min(valid_carry)
        lights["min_carry_bp"] = mn
        lights["carry"] = _light(mn)
    if us_carry:
        lights["us_min_carry_bp"] = min(us_carry)
    if jp_carry:
        lights["jp_min_carry_bp"] = min(jp_carry)
    # TONAR 倒挂灯(日债融资侧, 对应表6): 近端 carry 转负 = 日债去杠杆风险
    if tonar:
        lights["tonar"] = round(tonar[-1][1], 3)
        if jp_carry:
            jmn = min(jp_carry)
            lights["jp_funding"] = "🔴" if jmn < 0 else ("🟡" if jmn < 10 else "🟢")
    # 波动
    m_last = _last(move)
    if m_last:
        lights["move"] = round(m_last[1], 1)
        lights["vol"] = "🔴" if m_last[1] >= 120 else ("🟡" if m_last[1] >= 100 else "🟢")

    # ── asof / status ──
    all_last = []
    for p in (panel_funding, panel_carry, panel_trigger):
        for s in p["series"]:
            if s.get("points"):
                all_last.append(s["points"][-1]["date"])
    asof = max(all_last) if all_last else today.isoformat()
    missing = [f'{p["id"]}' for p in (panel_funding, panel_carry, panel_trigger) if not p["series"]]
    status = "ok" if not missing else ("部分缺失: " + ", ".join(missing))

    out = {
        "panels": {"funding": panel_funding, "carry": panel_carry, "trigger": panel_trigger},
        "matrix": matrix, "lights": lights, "asof": asof, "years": years, "status": status,
    }
    try:
        with open(cache_path, "w") as f:
            _json.dump(out, f, ensure_ascii=False, default=str)
    except Exception:
        pass
    return out


def fetch_liquidity_crisis_panel(start="2025-01-01", cache_path=None):
    """每日流动性危机监测盘 —— SOFR−EFFR + 信用利差 + 闪崩日标注。

    ★2026-09-18 新增(Chao: 「做一个去年1月份以来到今天的图, 尤其是闪崩的时间点帮我标注出来…
      同一张图上增加 oracle CDS 或者信贷差的曲线(跟 oracle 同级别的)…
      我需要这两个来观察每天的流动性危机」)。

    ★★为什么不是 Oracle 单名 CDS:
      单名 CDS(Oracle 5Y)属 Markit/ICE <b>付费</b>数据, 无免费合规源 —— 不硬凑、不伪造。
      Chao 已同意用「跟 Oracle 同级别」的信用利差替代。
      Oracle 真实评级(2026-09 实测核实):
        S&P <b>BBB−</b>(2026-07 下调, 距垃圾级仅一档) / Moody's Baa2 负面展望 / Fitch BBB
        债务 $129.5B ≈ 4.3× EBITDA (Reuters 2026-08-04)
      故主线取 <b>BBB 档 OAS</b>(最贴 Oracle), 并叠 <b>HY OAS</b> 作参照 ——
      Oracle 正卡在 BBB− 与垃圾级边界, 两线之间的距离就是市场给它的降级风险定价。

    ★三条线的分工(都是日频, 可每天更新):
      SOFR−EFFR  = 货币市场/回购的<b>资金面</b>压力(钱够不够)
      BBB OAS    = 投资级末端的<b>信用面</b>压力(愿不愿意借给高杠杆企业)
      HY OAS     = 垃圾级的信用面压力(风险偏好总开关)
      资金面与信用面同时恶化 = 真流动性危机; 只有一个动 = 局部事件。

    ★闪崩标注(Chao 选甲+乙两类都标):
      甲 = 金银(GC=F / SI=F), 乙 = 美股(^GSPC / ^IXIC)
      阈值<b>不拍脑袋</b>: 取各自 2025-01 以来日收益率分布的 <b>P1 分位</b>,
      即「本窗口内最极端的 1% 交易日」, 可解释且随数据自适应。

    返回 {series:{sofr_effr, bbb_oas, hy_oas, ig_oas}, crashes:[...],
          latest:{...}, provenance:{...}, status}
    """
    import datetime as _dt
    import json as _json
    import urllib.request as _ur

    def _px(sym, since):
        h = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                           "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120 Safari/537.36"}
        p1 = int(_dt.datetime(2024, 11, 1).timestamp())
        p2 = int(_dt.datetime.now().timestamp())
        u = (f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}"
             f"?period1={p1}&period2={p2}&interval=1d")
        try:
            d = _json.loads(_ur.urlopen(_ur.Request(u, headers=h), timeout=45)
                            .read().decode("utf-8", "ignore"))
            res = (d.get("chart") or {}).get("result") or []
            ts = res[0].get("timestamp") or []
            cl = ((res[0].get("indicators") or {}).get("quote") or [{}])[0].get("close") or []
            return [(_dt.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d"), float(c))
                    for t, c in zip(ts, cl) if c is not None]
        except Exception:
            return []

    out = {"series": {}, "crashes": [], "status": "ok", "start": start}

    # ── ① 资金面: SOFR − EFFR(bp) ──
    try:
        sofr = dict(_fred_series("SOFR", start="2024-11-01"))
        effr = dict(_fred_series("EFFR", start="2024-11-01"))
        pts = [(d, round((sofr[d] - effr[d]) * 100, 1))
               for d in sorted(set(sofr) & set(effr)) if d >= start]
        out["series"]["sofr_effr"] = {
            "name": "SOFR − EFFR", "unit": "bp", "axis": "left",
            "desc": "货币市场/回购资金面压力。有抵押本该比无抵押便宜，转正=回购市场缺钱。",
            "thresholds": {"warn": 5, "danger": 18},
            "data": pts, "status": "ok" if pts else "无数据",
        }
    except Exception as e:
        out["series"]["sofr_effr"] = {"name": "SOFR − EFFR", "data": [], "status": f"错误:{e}"}

    # ── ② 信用面: ICE BofA 分评级 OAS(%) ──
    for sid, key, nm, desc in [
        ("BAMLC0A4CBBB", "bbb_oas", "BBB 级 OAS",
         "投资级最末档信用利差。★Oracle 现为 S&P BBB−/Moody's Baa2，本线是其同级参照。"),
        ("BAMLH0A0HYM2", "hy_oas", "高收益 HY OAS",
         "垃圾级信用利差＝风险偏好总开关。Oracle 若再降一档即落入此区间。"),
        ("BAMLC0A0CM", "ig_oas", "投资级综合 OAS",
         "投资级整体信用利差，作基准对照。"),
    ]:
        try:
            s = _fred_series(sid, start="2024-11-01")
            pts = [(d, round(float(v), 3)) for d, v in s if d >= start]
            out["series"][key] = {
                "name": nm, "unit": "%", "axis": "right", "fred_id": sid,
                "desc": desc, "data": pts, "status": "ok" if pts else "无数据",
            }
        except Exception as e:
            out["series"][key] = {"name": nm, "data": [], "status": f"错误:{e}"}

    # ── ③ 闪崩日(甲 金银 / 乙 美股), 阈值 = 各自 P1 分位 ──
    CRASH = {
        "GC=F": ("黄金", "甲·贵金属"), "SI=F": ("白银", "甲·贵金属"),
        "^GSPC": ("标普500", "乙·美股"), "^IXIC": ("纳斯达克", "乙·美股"),
    }
    for sym, (zh, grp) in CRASH.items():
        px = _px(sym, start)
        if len(px) < 30:
            continue
        rets = []
        for i in range(1, len(px)):
            d, v = px[i]
            pv = px[i - 1][1]
            if pv:
                rets.append((d, 100.0 * (v / pv - 1)))
        win = [(d, r) for d, r in rets if d >= start]
        if not win:
            continue
        vals = sorted(r for _, r in win)
        p1 = vals[max(0, min(len(vals) - 1, int(len(vals) * 1 / 100)))]
        for d, r in win:
            if r <= p1:
                out["crashes"].append({"d": d, "asset": zh, "group": grp,
                                       "ret": round(r, 2), "threshold": round(p1, 2)})
    out["crashes"].sort(key=lambda z: (z["d"], z["asset"]))

    # ── ④ 黄金价格(COMEX 连续 GC=F, 美元/盎司) ──
    # ★2026-09-25 Chao:「黄金的曲线是不是也应该在这个图上」—— 闪崩标注本来就含金银,
    #   只有竖线没有价格线, 看不出崩之前涨了多少、崩完回到哪。与闪崩识别同源(同一次 Yahoo 拉取口径)。
    try:
        gpx = [(d, round(v, 2)) for d, v in _px("GC=F", start) if d >= start]
        out["series"]["gold"] = {
            "name": "黄金 COMEX 连续", "unit": "$/oz", "axis": "gold", "sym": "GC=F",
            "desc": "COMEX 黄金期货连续合约收盘价(Yahoo GC=F), 日频。",
            "data": gpx, "status": "ok" if gpx else "无数据",
        }
    except Exception as e:
        out["series"]["gold"] = {"name": "黄金 COMEX 连续", "data": [],
                                 "status": f"错误:{type(e).__name__}:{e}"}

    # 同日多资产合并(便于渲染成一条竖线)
    merged = {}
    for c in out["crashes"]:
        merged.setdefault(c["d"], []).append(c)
    out["crash_days"] = [
        {"d": d, "n": len(v), "groups": sorted({x["group"] for x in v}),
         "items": sorted(v, key=lambda z: z["ret"]),
         "worst": min(x["ret"] for x in v),
         "both": len({x["group"] for x in v}) > 1}
        for d, v in sorted(merged.items())
    ]

    for k, v in out["series"].items():
        if v.get("data"):
            v["latest"] = v["data"][-1]

    out["provenance"] = {
        "sofr_effr": "FRED SOFR & EFFR(纽约联储)，日频，差值×100 转 bp",
        "oas": "FRED ICE BofA US Corporate Index OAS：BAMLC0A4CBBB(BBB) / "
               "BAMLH0A0HYM2(HY) / BAMLC0A0CM(IG)，日频，单位 %",
        "oracle_note": "Oracle 单名 5Y CDS 属 Markit/ICE 付费数据，无免费合规源，"
                       "故以同评级 BBB 档 OAS 替代并叠 HY 作降级风险参照。"
                       "Oracle 评级：S&P BBB−(2026-07 下调，距垃圾级一档)、"
                       "Moody's Baa2 负面展望、Fitch BBB；债务 $129.5B≈4.3×EBITDA"
                       "（来源 Reuters 2026-08-04、S&P Global Ratings）。",
        "crash": "闪崩日 = 各资产 2025-01 起日收益率分布的 P1 分位(最极端 1% 交易日)，"
                 "阈值随数据自适应，非人为设定固定百分比。"
                 "甲=金银(GC=F/SI=F)，乙=美股(^GSPC/^IXIC)，数据源 Yahoo Finance。",
    }

    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data",
                                  "liquidity_crisis.json")
    try:
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False)
    except Exception:
        pass
    return out


def fetch_cot_managed_money(years=3, cache_path=None):
    """COMEX/NYMEX 金·银·铂 —— 对冲基金(Managed Money)持仓 + 未平仓量 + 日频价格。

    ★2026-09-18 新增(Chao: 「能否抓到对冲基金在 comex 上面的黄金白银和 pt 的
      开仓量和价格叠加图, 并且每天更新」)。

    ★★频率真相(必须随数据一起显示, 不许标成日更):
      CFTC COT 是<b>周度</b>: 每周二收盘持仓, 当周周五 15:30 ET 发布。
      「对冲基金持仓每天更新」在数据源层面不存在 —— 全世界都只有周度。
      本函数的做法: <b>持仓周度台阶 + 价格日度曲线</b>叠加, 各自标注真实频率。

    ★口径选择: 用 Disaggregated 而非 Legacy。
      Legacy 的 non-commercial 把对冲基金 + CTA + 指数基金混在一起;
      Disaggregated 的 <b>Managed Money</b> 才是纯对冲基金/CTA 这类投机资金,
      另有 Swap Dealer(做市商/银行)可作对手方参照。
      (项目原有 fetchers/cot.py 用的是 Legacy 口径, 两者并存、互不替代。)

    ★铂金在 **NYMEX** 不在 COMEX —— 渲染时不可笼统写成 COMEX。

    数据源:
      持仓 = CFTC Socrata Disaggregated futures-only (72hh-3qpy), 免 key
      价格 = Yahoo Finance 连续合约 GC=F / SI=F / PL=F, 日频

    返回 {metal: {name, exchange, cot:[...], price:[(d,close)],
                  latest:{...}, pctile:{...}, status, source}}
    """
    import datetime as _dt
    import json as _json
    import urllib.parse as _up
    import urllib.request as _ur

    UA = {"User-Agent": "EcoVolChecker research"}
    COT_BASE = "publicreporting.cftc.gov/resource/72hh-3qpy.json"
    METALS = {
        "gold": ("黄金 Gold", "COMEX", "GOLD - COMMODITY EXCHANGE INC.", "GC=F"),
        "silver": ("白银 Silver", "COMEX", "SILVER - COMMODITY EXCHANGE INC.", "SI=F"),
        "platinum": ("铂金 Platinum", "NYMEX", "PLATINUM - NEW YORK MERCANTILE EXCHANGE", "PL=F"),
    }
    start = (_dt.date.today() - _dt.timedelta(days=int(years * 365.25))).isoformat()

    def _cot(market):
        u = "https://" + COT_BASE + "?" + _up.urlencode({
            "$where": (f"market_and_exchange_names='{market}' "
                       f"AND report_date_as_yyyy_mm_dd >= '{start}'"),
            "$order": "report_date_as_yyyy_mm_dd", "$limit": 400})
        try:
            r = _ur.urlopen(_ur.Request(u, headers=UA), timeout=60)
            return _json.loads(r.read().decode("utf-8", "ignore"))
        except Exception:
            return []

    def _px(sym):
        rng = f"{max(int(years), 1)}y"
        u = (f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}"
             f"?range={rng}&interval=1d")
        h = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                           "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120 Safari/537.36"}
        try:
            r = _ur.urlopen(_ur.Request(u, headers=h), timeout=45)
            d = _json.loads(r.read().decode("utf-8", "ignore"))
            res = (d.get("chart") or {}).get("result") or []
            ts = res[0].get("timestamp") or []
            cl = ((res[0].get("indicators") or {}).get("quote") or [{}])[0].get("close") or []
            out = []
            for t, c in zip(ts, cl):
                if c is None:
                    continue
                out.append((_dt.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d"), round(float(c), 2)))
            return out
        except Exception:
            return []

    def _i(r, k):
        try:
            return int(float(r.get(k, 0) or 0))
        except Exception:
            return 0

    out = {}
    for key, (zh, exch, market, sym) in METALS.items():
        raw = _cot(market)
        rec = {"name": zh, "exchange": exch, "symbol": sym, "cot": [], "price": [],
               "status": "未找到",
               "source": ("CFTC Disaggregated COT (周度·周二持仓/周五发布) + "
                          "Yahoo Finance 连续合约 (日频)"),
               "cadence_note": "持仓周更 · 价格日更"}
        for r in raw:
            d = (r.get("report_date_as_yyyy_mm_dd") or "")[:10]
            if not d:
                continue
            ml, ms = _i(r, "m_money_positions_long_all"), _i(r, "m_money_positions_short_all")
            sl = _i(r, "swap_positions_long_all")
            ss = _i(r, "swap__positions_short_all") or _i(r, "swap_positions_short_all")
            rec["cot"].append({
                "d": d, "oi": _i(r, "open_interest_all"),
                "mm_long": ml, "mm_short": ms, "mm_net": ml - ms,
                "sd_long": sl, "sd_short": ss, "sd_net": sl - ss,
                "pm_long": _i(r, "prod_merc_positions_long"),
                "pm_short": _i(r, "prod_merc_positions_short"),
            })
        rec["price"] = _px(sym)
        if len(rec["cot"]) >= 2:
            last = rec["cot"][-1]
            nets = [c["mm_net"] for c in rec["cot"]]
            ois = [c["oi"] for c in rec["cot"]]
            rec["latest"] = last
            rec["prev"] = rec["cot"][-2]
            # 历史分位: 当前 MM 净 / OI 在本窗口里的位置
            rec["pctile"] = {
                "mm_net": round(100.0 * sum(1 for v in nets if v <= last["mm_net"]) / len(nets), 1),
                "oi": round(100.0 * sum(1 for v in ois if v <= last["oi"]) / len(ois), 1),
            }
            rec["mm_net_min"], rec["mm_net_max"] = min(nets), max(nets)
            rec["oi_min"], rec["oi_max"] = min(ois), max(ois)
            rec["status"] = "ok"
        if rec["price"]:
            rec["price_latest"] = rec["price"][-1]
        out[key] = rec

    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data",
                                  "cot_managed_money.json")
    try:
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False)
    except Exception:
        pass
    return out


def fetch_comex_inventory(cache_path=None):
    """COMEX & 上海(SHFE/SGE)贵金属库存 + GLD/SLV ETF 资金流。
    数据源: 公开 GitHub Pages comex-inventory-charts(每日自动更新, 内联 CHARTS JSON)。
    解析页面内 `const CHARTS = {...}` + `const SUMMARY = {...}`。绝不编: 抓不到返回 status='未获取'。
    返回 {status, as_of, summary, panels:{gold,silver,platinum,gld_flow_weekly,slv_flow_weekly}}。
    - 库存图: COMEX(左轴) vs 上海(右轴), 双轴折线, 单位吨。
    - ETF 资金流: 周净流量柱状(正=流入绿/负=流出红), 单位吨。
    """
    import json as _json
    import re as _re
    import datetime as _dt
    if cache_path is None:
        cache_path = os.path.join(os.path.dirname(__file__), "..", "data", "comex_inventory.json")

    def _fallback(note):
        """★铁律: 抓取/解析失败绝不返回空壳让 114 点曲线整块消失。
        有缓存 → 回退缓存真值 + status='ok(缓存)' + stale_since 标滞后天数; 无缓存才认输。"""
        try:
            with open(cache_path) as f:
                cached = _json.load(f)
            if cached.get("panels") or cached.get("flows"):
                cached = dict(cached)
                stale_since = cached.get("fetched_at") or cached.get("as_of")
                days = None
                try:
                    d0 = _dt.datetime.fromisoformat(str(stale_since)[:19]).date()
                    days = (_dt.date.today() - d0).days
                except Exception:
                    pass
                cached["status"] = "ok(缓存)"
                cached["stale_since"] = stale_since
                cached["stale_days"] = days
                cached["note"] = f"实时源不可用({note}), 展示缓存真值" + (f", 滞后 {days} 天" if days is not None else "")
                print(f"[comex_inventory] {note} → 回退缓存 (as_of={cached.get('as_of')}, 滞后{days}天)")
                return cached
        except Exception as e:
            print(f"[comex_inventory] 缓存回退也失败: {e}")
        return {"status": "未获取", "as_of": None, "note": note}

    # ★2026-09-19: 个人 GitHub 账号被反滥用系统限制匿名访问 → github.io 全线 404,
    #   本 fetcher 曾静默回退缓存(吃 3 天前旧数据)。改为<b>双源</b>: Cloudflare 主 + GitHub 备。
    #   Cloudflare 站是同一产物的镜像(实测含完整 const CHARTS), 不受账号限制影响。
    urls = ["https://comex-inventory-charts.pages.dev/",
            "https://curarpikt0000.github.io/comex-inventory-charts/"]
    html = None
    src_used = None
    tried = []
    import subprocess
    import sys as _sys
    # ★2026-10-06 Eco 接管(Chao「接入」): 先在本地自产(Notion 只读 + SPDR 官方 + SGE 官网周报), 不再依赖外站。
    #   外站 comex-inventory-charts 自 09-16 起停更(生产者在 Mac, 脚本写死 Mac 路径), 日报 cron 仍标 ok = 静默旧数据。
    #   自产失败才回落到外站, 外站也失败才回落缓存; 回落时 src_used/tried 会写进 note, 不再静默。
    try:
        _proj = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
        _tool = os.path.join(_proj, "tools", "comex_inv")
        _base = os.path.join(_proj, "data", "_comex_inv_work")      # 已加 .gitignore, 不进 git
        os.makedirs(os.path.join(_base, "data"), exist_ok=True)
        _seed = os.path.join(_proj, "data", "comex_inv", "sge_history.json")
        if os.path.exists(_seed):
            import shutil as _sh
            _sh.copy(_seed, os.path.join(_base, "data", "sge_history.json"))
        _out_html = os.path.join(_base, "inventory_charts.html")
        if os.path.exists(_out_html):
            os.remove(_out_html)
        _env = dict(os.environ, INV_BASE=_base)
        _ok = True
        for _s in ("build_inventory_data.py", "build_etf_data.py", "render_inventory_html.py"):
            _r = subprocess.run([_sys.executable, os.path.join(_tool, _s)], env=_env,
                                capture_output=True, text=True, timeout=240)
            if _r.returncode != 0:
                tried.append(f"自产:{_s} rc={_r.returncode} {(_r.stderr or '')[-100:].strip()}")
                _ok = False
                break
        if _ok and os.path.exists(_out_html):
            _b = open(_out_html, encoding="utf-8").read()
            if "const CHARTS" in _b:
                html, src_used = _b, "local:tools/comex_inv(Notion只读+SPDR+SGE官网)"
            else:
                tried.append("自产:产物无CHARTS")
    except Exception as e:  # noqa: BLE001
        tried.append(f"自产:{type(e).__name__} {str(e)[:80]}")
    for u in ([] if html else urls):
        try:
            r = subprocess.run(["curl", "-s", "-L", "--max-time", "30",
                                "-H", "User-Agent: Mozilla/5.0", u],
                               capture_output=True, text=True, timeout=35)
            body = r.stdout if r.returncode == 0 else ""
            # ★判据不是"拿到 200/非空", 而是<b>内容里真有 CHARTS</b>。
            #   GitHub 被限制后返回的是 404 页面, 非空但无数据 —— 光看非空会误判成功。
            if body and "const CHARTS" in body:
                html, src_used = body, u
                break
            tried.append(f"{u.split('/')[2]}:{'无CHARTS' if body else '空'}")
        except Exception as e:
            tried.append(f"{u.split('/')[2]}:{type(e).__name__}")
    if not html:
        return _fallback(f"comex-inventory-charts 全部源失败 ({'; '.join(tried)})")

    try:
        mc = _re.search(r'const CHARTS = (\{.*?\});\s*\nconst SUMMARY', html, _re.S)
        ms = _re.search(r'const SUMMARY = (\{.*?\});', html, _re.S)
        charts = _json.loads(mc.group(1)) if mc else {}
        summary = _json.loads(ms.group(1)) if ms else {}
    except Exception as e:
        return _fallback(f"解析失败: {e}")

    if not charts:
        return _fallback("CHARTS 数据为空")

    # ── 库存双轴 panel: COMEX(左) vs 上海(右) ──
    inv_colors = {"COMEX 库存 (吨)": "#c0757d", "SHFE 库存 (吨)": "#6b8fb5", "SGE 库存 (吨)": "#7fa085"}

    def _inv_panel(key, title, subtitle):
        c = charts.get(key)
        if not c:
            return None
        series = []
        for s in c.get("series", []):
            data = s.get("data", [])
            pts = [{"date": d, "v": v} for d, v in data if v is not None]
            if len(pts) < 2:
                continue
            nm = s.get("name", "")
            is_comex = "COMEX" in nm
            series.append({
                "name": nm, "color": inv_colors.get(nm, "#8a8578"),
                "axis": "left" if is_comex else "right",
                "width": 2.0 if is_comex else 1.6,
                "dash": not is_comex,
                "points": pts,
            })
        if not series:
            return None
        single = all(s["axis"] == "left" for s in series)
        return {"id": f"comex_inv_{key}", "title": title, "subtitle": subtitle,
                "unit_left": "COMEX 吨", "unit_right": "上海 吨", "single_axis": single,
                "series": series, "source": "comex-inventory-charts（COMEX 日频 · 上海 SHFE/SGE 周频）"}

    panels = {
        "gold": _inv_panel("gold", "黄金库存：COMEX vs 上海(SHFE)", "Gold Inventory — COMEX vs Shanghai"),
        "silver": _inv_panel("silver", "白银库存：COMEX vs 上海(SHFE/SGE)", "Silver Inventory — COMEX vs Shanghai"),
        "platinum": _inv_panel("platinum", "铂金库存：COMEX", "Platinum Inventory — COMEX"),
    }

    # ── ETF 资金流: 周净流量柱状(正绿负红) ──
    def _flow_bars(key, name):
        c = charts.get(key)
        if not c or not c.get("series"):
            return None
        data = c["series"][0].get("data", [])
        bars = [{"date": d, "v": v} for d, v in data if v is not None]
        if len(bars) < 2:
            return None
        return {"name": name, "bars": bars}

    flows = {
        "gld": _flow_bars("gld_flow_weekly", "GLD 黄金ETF 周净流量"),
        "slv": _flow_bars("slv_flow_weekly", "SLV 白银ETF 周净流量"),
    }

    # as_of
    asofs = []
    for p in panels.values():
        if p:
            for s in p["series"]:
                if s["points"]:
                    asofs.append(s["points"][-1]["date"])
    as_of = max(asofs) if asofs else None

    out = {"status": "ok", "as_of": as_of, "summary": summary,
           "panels": {k: v for k, v in panels.items() if v},
           "flows": {k: v for k, v in flows.items() if v},
           "source_url": src_used, "fetched_at": _dt.datetime.now().isoformat()}
    try:
        with open(cache_path, "w") as f:
            _json.dump(out, f, ensure_ascii=False, default=str)
    except Exception as e:
        print(f"[comex_inventory] 缓存写入失败(不影响本次渲染): {e}")
    return out


# ─────────── 世界前十经济体 政府债务/GDP (IMF WEO) ───────────
# 源: IMF DataMapper API (官方, 免 key)。指标 GGXWDG_NGDP =
#   General government gross debt, % of GDP。
# ★频率诚实说明: IMF WEO 一年发布两次(4月/10月), 本质是【年度】数据。
#   本 fetcher 每次运行都会重新拉取, 但值在两次 WEO 之间不会变化 ——
#   dashboard 用 stale 徽章显示"距上次发布 N 天", 绝不假装是周度/月度更新。
# ★同时含 IMF 【预测年份】(通常未来 5-6 年): 必须与实绩年份区分渲染,
#   否则会把预测当历史。用 is_forecast 标记。
IMF_DM_API = "https://www.imf.org/external/datamapper/api/v1"
# 世界前十大经济体(按名义 GDP, IMF 2025 口径)
TOP10_ECONOMIES = [
    ("USA", "美国", "🇺🇸"), ("CHN", "中国", "🇨🇳"), ("DEU", "德国", "🇩🇪"),
    ("JPN", "日本", "🇯🇵"), ("IND", "印度", "🇮🇳"), ("GBR", "英国", "🇬🇧"),
    ("FRA", "法国", "🇫🇷"), ("ITA", "意大利", "🇮🇹"), ("CAN", "加拿大", "🇨🇦"),
    ("BRA", "巴西", "🇧🇷"),
]


def fetch_debt_to_gdp(years=15):
    """世界前十经济体 政府债务/GDP (%)。IMF WEO 官方, 免 key。

    返回 {"countries":[{iso,name,flag,series:[(year,val)],latest,latest_year,
                        forecast:[(year,val)], chg_5y, high, low, status}],
          "as_of_year", "vintage", "status", "source"}
    ★series 只含【实绩年份】(<= 当前实绩年), forecast 单独放 —— 绝不混。
    ★抓不到 → status="未找到", countries=[] , 绝不编造。
    """
    import datetime as _dt
    import requests
    out = {"countries": [], "status": "未找到", "as_of_year": None,
           "vintage": None,
           "source": "IMF World Economic Outlook (DataMapper API) · GGXWDG_NGDP"}
    try:
        r = requests.get(f"{IMF_DM_API}/GGXWDG_NGDP", timeout=45,
                         headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"})
        if r.status_code != 200:
            out["note"] = f"IMF API HTTP {r.status_code}"
            return out
        vals = (r.json().get("values") or {}).get("GGXWDG_NGDP") or {}
    except Exception as e:
        out["note"] = f"IMF API 异常: {type(e).__name__}"
        return out
    if not vals:
        out["note"] = "IMF API 返回空"
        return out

    # ★实绩/预测分界: IMF 当年及以后为预测。用当前年份作界(保守: 当年也算预测,
    #   因为 WEO 当年数多为估计值 estimate)。
    this_year = _dt.date.today().year
    cutoff_year = this_year - 1          # <= cutoff 视为实绩
    min_year = this_year - years

    max_actual = None
    for iso, zh, flag in TOP10_ECONOMIES:
        s = vals.get(iso) or {}
        pts, fc = [], []
        for y, v in s.items():
            try:
                yi, vf = int(y), float(v)
            except (TypeError, ValueError):
                continue
            if vf is None:
                continue
            if yi <= cutoff_year:
                if yi >= min_year:
                    pts.append((yi, round(vf, 1)))
            else:
                fc.append((yi, round(vf, 1)))
        pts.sort()
        fc.sort()
        if not pts:
            out["countries"].append({
                "iso": iso, "name": zh, "flag": flag, "series": [],
                "forecast": fc, "status": "未找到"})
            continue
        ly, lv = pts[-1]
        max_actual = ly if max_actual is None else max(max_actual, ly)
        # 5 年变化(实绩内)
        chg5 = None
        base = [p for p in pts if p[0] <= ly - 5]
        if base:
            chg5 = round(lv - base[-1][1], 1)
        v_only = [v for _, v in pts]
        out["countries"].append({
            "iso": iso, "name": zh, "flag": flag,
            "series": pts, "forecast": fc,
            "latest": lv, "latest_year": ly,
            "chg_5y": chg5,
            "high": max(v_only), "low": min(v_only),
            "status": "ok",
        })
    ok = [c for c in out["countries"] if c["status"] == "ok"]
    if ok:
        out["status"] = "ok"
        out["as_of_year"] = max_actual
        # 债务率降序(最高的排前面)
        out["countries"].sort(key=lambda c: (c.get("latest") is None, -(c.get("latest") or 0)))
    return out


# ─────────── 美国分评级公司债: 收益率 / 利差 / 未偿总额 ───────────
# ★★口径陷阱(实测确认, 务必不要踩):
#   FRED 上 ICE BofA 带 "TRIV" 后缀的序列(如 BAMLCC0A0CMTRIV)标题是
#   "Total Return Index Value" = 【总回报指数】, 单位 Index —— 它随价格涨跌波动,
#   【不是】债券未偿总额。把它当"公司债总额"画图会得到完全错误的曲线。
#   真正的未偿总额只有 Fed Z.1 (Flow of Funds) 的【季度】序列。
#   免费源不存在"每日 + 分评级 + 未偿总额"(该数据属 ICE/Bloomberg 商业授权)。
# 因此本模块诚实分两层:
#   ① 日频: 各评级 有效收益率(EY) + 期权调整利差(OAS) —— 看利率与信用风险波动
#   ② 季频: Fed Z.1 非金融企业债务真实未偿额 —— 看总量是否增加
CORP_RATINGS = [
    # (标签, 有效收益率 series, OAS series, 投资级?)
    ("AAA", "BAMLC0A1CAAAEY", "BAMLC0A1CAAA", True),
    ("AA",  "BAMLC0A2CAAEY",  "BAMLC0A2CAA",  True),
    ("A",   "BAMLC0A3CAEY",   "BAMLC0A3CA",   True),
    ("BBB", "BAMLC0A4CBBBEY", "BAMLC0A4CBBB", True),
    ("BB",  "BAMLH0A1HYBBEY", "BAMLH0A1HYBB", False),
    ("B",   "BAMLH0A2HYBEY",  "BAMLH0A2HYB",  False),
    ("CCC及以下", "BAMLH0A3HYCEY", "BAMLH0A3HYC", False),
]
# 未偿总额(季度, Fed Z.1) —— 真实"总量"口径
CORP_OUTSTANDING = [
    ("非金融企业·债务证券", "NCBDBIQ027S"),
    ("非金融企业·债务证券+贷款", "BCNSDODNS"),
]


def _fred_hist(sid, start):
    """带 key FRED API 优先(项目铁律: CSV 端点在本 VM 偶发超时), 失败回退 CSV。
    返回 [(date, float)] 升序; 全失败返回 []。"""
    try:
        import sys as _sys
        _d = os.path.dirname(__file__)
        if _d not in _sys.path:
            _sys.path.insert(0, _d)
        from fetchers.fred import fetch_fred_history
        h = fetch_fred_history(sid, start=start)
        if h:
            out = []
            for d, v in h:
                try:
                    out.append((str(d).strip(), float(v)))
                except (TypeError, ValueError):
                    continue
            if out:
                return sorted(out)
    except Exception:
        pass
    return _fred_series(sid, start=start)


def fetch_corporate_credit(years=3):
    """美国分评级公司债: 日频收益率/利差 + 季频真实未偿总额。

    返回 {
      "ratings":[{label,ig,yield_latest,yield_date,yield_series:[(d,v)],
                  oas_latest,oas_date,oas_series:[(d,v)],
                  chg_1m_bp,chg_3m_bp,oas_chg_1m_bp,status}],
      "outstanding":[{label,latest,latest_date,series:[(d,v)],chg_yoy_pct,status}],
      "as_of","status","source","caveat"}
    ★绝不编: 任一序列抓不到 → 该项 status="未找到" + 空 series。
    ★单位: 收益率/OAS = %, 未偿额 = 十亿美元($B, 由 FRED 百万美元换算)。
    """
    import datetime as _dt
    start = (_dt.date.today() - _dt.timedelta(days=int(years * 365) + 40)).strftime("%Y-%m-%d")
    out = {"ratings": [], "outstanding": [], "status": "未找到", "as_of": None,
           "source": "FRED · ICE BofA US Corporate/High Yield Index (日频) + "
                     "Fed Z.1 Financial Accounts (季频未偿额)",
           "caveat": "免费源无“每日·分评级·未偿总额”(属 ICE/Bloomberg 商业授权)。"
                     "故总量用 Fed Z.1 季度真实未偿额, 日频仅收益率与利差。"}

    def _bp(series, days):
        """近 N 日变化(基点)。数据不足返回 None, 绝不外推。"""
        if len(series) < 2:
            return None
        last_d, last_v = series[-1]
        try:
            d0 = _dt.date.fromisoformat(last_d) - _dt.timedelta(days=days)
        except Exception:
            return None
        prior = [(d, v) for d, v in series if d <= d0.isoformat()]
        if not prior:
            return None
        return round((last_v - prior[-1][1]) * 100, 1)

    dates_seen = []
    for label, ey_sid, oas_sid, ig in CORP_RATINGS:
        ys = _fred_hist(ey_sid, start)
        os_ = _fred_hist(oas_sid, start)
        rec = {"label": label, "ig": ig,
               "yield_series": ys, "oas_series": os_,
               "yield_sid": ey_sid, "oas_sid": oas_sid}
        if ys:
            rec["yield_latest"], rec["yield_date"] = ys[-1][1], ys[-1][0]
            rec["chg_1m_bp"] = _bp(ys, 30)
            rec["chg_3m_bp"] = _bp(ys, 90)
            dates_seen.append(ys[-1][0])
        if os_:
            rec["oas_latest"], rec["oas_date"] = os_[-1][1], os_[-1][0]
            rec["oas_chg_1m_bp"] = _bp(os_, 30)
            dates_seen.append(os_[-1][0])
        rec["status"] = "ok" if (ys or os_) else "未找到"
        out["ratings"].append(rec)

    # 季度真实未偿额(FRED 单位: 百万美元 → 换算 $B)
    q_start = (_dt.date.today() - _dt.timedelta(days=365 * 12)).strftime("%Y-%m-%d")
    for label, sid in CORP_OUTSTANDING:
        s = _fred_hist(sid, q_start)
        rec = {"label": label, "sid": sid,
               "series": [(d, round(v / 1000.0, 1)) for d, v in s]}
        if s:
            rec["latest"] = round(s[-1][1] / 1000.0, 1)
            rec["latest_date"] = s[-1][0]
            # 同比(4 个季度前)
            if len(s) >= 5:
                prev = s[-5][1]
                if prev:
                    rec["chg_yoy_pct"] = round((s[-1][1] - prev) / prev * 100, 1)
            rec["status"] = "ok"
        else:
            rec["status"] = "未找到"
        out["outstanding"].append(rec)

    if dates_seen:
        out["as_of"] = max(dates_seen)
        out["status"] = "ok"
    return out


# ============ AI 产业链: 自由现金流 + 信用状况 ============
# 4 类 ≤20 家。CIK 为 SEC 官方标识; 空 CIK = 无 SEC 申报(未上市/外国私有), 显式排除并说明。
#
# ★第 4 个字段 = 中文名 (Chao 2026-08-25 要求图上 ticker 下方附中文)。
#   取名原则 (逐个联网核过, 不臆造):
#   - 有官方中文名的用官方: 维谛技术(Vertiv 中国官网自称) / 伊顿(Eaton 中国官网) /
#     美光科技 / 博通 / 微软 / 亚马逊 / 甲骨文 / 台积电 / 英伟达。
#   - 无官方中文名但中文财经媒体有稳定通行译名的用通行译名:
#     星座能源(CEG, 维基/雪球/国家核安全局均用) / 奇异维诺瓦(GEV, 中文维基条目名)。
#   - ★没有任何通行中文名的<b>保留英文原名</b>, 绝不自己编音译:
#     CoreWeave / Equinix / Digital Realty / Vistra / AMD / Meta / Alphabet。
#     中文维基与百度百科的条目名就是英文原文, 硬造"科尔维""亿速"之类反而误导。
#   ★★VST 的坑: 搜"Vistra 中文名"会命中「瑞致达」, 那是 Vistra Group Holdings
#     (香港信托服务商), 与本表的 Vistra Corp (德州电力/核电, NYSE:VST) 完全是两家公司。
#     曾差点写错, 保留英文原名最安全。
AI_UNIVERSE = [
    ("AI应用/模型层", [
        ("MSFT",  "Microsoft",        "0000789019", "微软"),
        ("GOOGL", "Alphabet",         "0001652044", "Alphabet（谷歌母公司）"),
        ("META",  "Meta",             "0001326801", "Meta（原 Facebook）"),
        ("AMZN",  "Amazon",           "0001018724", "亚马逊"),
        ("ORCL",  "Oracle",           "0001341439", "甲骨文"),
    ]),
    ("芯片/加速器", [
        ("NVDA",  "NVIDIA",           "0001045810", "英伟达"),
        ("AMD",   "AMD",              "0000002488", "AMD（超微半导体）"),
        ("AVGO",  "Broadcom",         "0001730168", "博通"),
        ("TSM",   "TSMC",             "0001046179", "台积电"),   # IFRS(20-F)
        ("MU",    "Micron",           "0000723125", "美光科技"),
    ]),
    ("算力/数据中心", [
        ("CRWV",  "CoreWeave",        "0001769628", "CoreWeave"),
        ("EQIX",  "Equinix",          "0001101239", "Equinix"),
        ("DLR",   "Digital Realty",   "0001297996", "Digital Realty"),
        ("VRT",   "Vertiv",           "0001674101", "维谛技术"),
    ]),
    ("电力/能源基建", [
        ("CEG",   "Constellation",    "0001868275", "星座能源"),
        ("VST",   "Vistra",           "0001692819", "Vistra"),
        ("GEV",   "GE Vernova",       "0001996810", "奇异维诺瓦"),
        ("ETN",   "Eaton",            "0001551182", "伊顿"),
    ]),
]

# 未上市 → 无 FCF/无公开债, 图上不静默省略而是显式列出
AI_EXCLUDED = [
    ("OpenAI",    "未上市, 无 SEC 申报, 无公开债"),
    ("Anthropic", "未上市, 无 SEC 申报, 无公开债"),
    ("xAI",       "未上市, 无 SEC 申报, 无公开债"),
    ("NextEra",   "公用事业, CapEx 记于公司自定义扩展标签, us-gaap 无标准字段可取"),
]

_OCF_TAGS = ["NetCashProvidedByUsedInOperatingActivities",
             "NetCashProvidedByUsedInOperatingActivitiesContinuingOperations"]
# CapEx 各公司标签不一, 同财年命中的多个科目求和(如 EQIX 拆 3 个)
_CAPEX_TAGS = ["PaymentsToAcquirePropertyPlantAndEquipment",
               "PaymentsToAcquireProductiveAssets",
               "PaymentsForCapitalImprovements",
               "PaymentsToAcquireOtherPropertyPlantAndEquipment",
               "PaymentsToAcquireMachineryAndEquipment",
               "PaymentsToAcquireBuildings",
               "PaymentsToDevelopRealEstateAssets",
               "PaymentsToAcquireRealEstate",
               "PaymentsForProceedsFromProductiveAssets"]
_IFRS_OCF = ["CashFlowsFromUsedInOperatingActivities"]
_IFRS_CAPEX = ["PurchaseOfPropertyPlantAndEquipmentClassifiedAsInvestingActivities"]
# 利息保障倍数 = EBIT / 利息支出
_EBIT_TAGS = ["OperatingIncomeLoss",
              "IncomeLossFromContinuingOperationsBeforeIncomeTaxes"
              "ExtraordinaryItemsNoncontrollingInterest"]
_INT_TAGS = ["InterestExpense", "InterestExpenseDebt",
             "InterestExpenseNonoperating", "InterestAndDebtExpense"]
_IFRS_EBIT = ["ProfitLossFromOperatingActivities"]
_IFRS_INT = ["FinanceCosts"]

# ── 图二 信用维度: 净债务/EBITDA + 利息保障倍数 ──
# ★锚定铁律: 绝不"每个标签各取最新值"—— 那会把 FY2012 的债务和 FY2026 的
#   EBIT 混算(实测 DLR 债务停更于 2011, ETN 营业利润停更于 2019)。
#   必须取【各构件共同存在的最近财年】。
_CR_EBIT = ["OperatingIncomeLoss", "ProfitLossFromOperatingActivities",
            "IncomeLossFromContinuingOperationsBeforeIncomeTaxes"
            "ExtraordinaryItemsNoncontrollingInterest"]
_CR_DA = ["DepreciationDepletionAndAmortization",
          "DepreciationAmortizationAndAccretionNet",
          "DepreciationAndAmortization", "DepreciationAmortizationAndImpairment",
          "Depreciation", "DepreciationNonproduction",
          "AmortizationOfIntangibleAssets", "DepreciationExpense"]
_CR_INT = ["InterestExpense", "InterestExpenseNonoperating",
           "InterestExpenseDebt", "InterestExpenseBorrowings",
           "InterestAndDebtExpense", "InterestExpenseNetOfCapitalizedInterest",
           "InterestCostsIncurred", "FinanceCosts"]
_CR_LTD = ["LongTermDebtNoncurrent", "LongTermDebt", "LongTermBorrowings",
           "DebtLongtermAndShorttermCombinedAmount",
           "LongTermDebtAndCapitalLeaseObligations",
           "NoncurrentPortionOfNoncurrentBondsIssued",
           "NoncurrentPortionOfLongtermBorrowings"]
_CR_STD = ["LongTermDebtCurrent", "DebtCurrent", "ShortTermBorrowings",
           "OtherShortTermBorrowings",
           "LongTermDebtAndCapitalLeaseObligationsCurrent",
           "CurrentPortionOfLongtermBorrowings"]
_CR_CASH = ["CashAndCashEquivalentsAtCarryingValue",
            "CashCashEquivalentsRestrictedCashAndRestrictedCashEquivalents",
            "CashAndCashEquivalents"]
_CR_STI = ["ShortTermInvestments", "MarketableSecuritiesCurrent",
           "AvailableForSaleSecuritiesDebtSecuritiesCurrent",
           "OtherShortTermInvestments"]
# 财报滞后容忍: 锚定财年早于「今年-2」视为过期, 标 n/a 而非照画
_CR_MAX_STALE_YEARS = 2


def _cr_series(facts, tags, ns, duration):
    """合并同类标签的年度序列: 靠前标签优先, 靠后的只补缺失年份。"""
    out = {}
    for t in tags:
        for k, v in (_xbrl_annual(facts, t, ns, duration) or {}).items():
            out.setdefault(k, v)
    return out


def fetch_ai_credit():
    """AI 产业链公司 信用维度: 净债务/EBITDA(杠杆) + 利息保障倍数。

    返回 {"rows":[{ticker,name,group,fy,leverage,coverage,ebitda,
                   net_debt,interest,status,note}],
          "as_of","status","source","caveat","lev_n","cov_n","total"}
    ★数据全部来自 10-K/20-F 官方申报值, 不做任何估算/代理/外推。
    ★拿不到 → 该维度留 None + note 写明原因, 绝不用 0 或行业均值填充。
    ★为何不用市场信用利差: 免费源全灭(FINRA TRACE 端点 404 /
      Boerse-Frankfurt 403 / WSJ 401), 用股价波动凑代理会失真, 故改用
      真实财报偿债能力指标。
    """
    import datetime as _dt
    out = {"rows": [], "as_of": None, "status": "未找到",
           "source": "SEC EDGAR XBRL companyfacts (10-K/20-F 年报官方申报值)",
           "caveat": "净债务 = 长期债务 + 短期债务 − 现金及等价物 − 短期投资; "
                     "EBITDA = 营业利润 + 折旧摊销(近似, 未加回股权激励等非现金项)。"
                     "各公司财年截止日不同, 图上已标注实际财年。"
                     "指标源于年报, 每年更新一次, 非日频。"}
    cur_year = _dt.date.today().year
    fys, lev_n, cov_n = [], 0, 0

    for gname, members in AI_UNIVERSE:
        for tk, cname, cik, zhname in members:
            rec = {"ticker": tk, "name": cname, "zh": zhname, "group": gname, "fy": None,
                   "leverage": None, "coverage": None, "ebitda": None,
                   "net_debt": None, "interest": None,
                   "status": "未找到", "note": ""}
            facts = _sec_facts(cik)
            if not facts:
                rec["note"] = "SEC 拉取失败"
                out["rows"].append(rec)
                continue
            ns = ("ifrs-full" if "ifrs-full" in (facts.get("facts") or {})
                  else "us-gaap")
            e_s = _cr_series(facts, _CR_EBIT, ns, True)
            d_s = _cr_series(facts, _CR_DA, ns, True)
            i_s = _cr_series(facts, _CR_INT, ns, True)
            l_s = _cr_series(facts, _CR_LTD, ns, False)
            s_s = _cr_series(facts, _CR_STD, ns, False)
            c_s = _cr_series(facts, _CR_CASH, ns, False)
            t_s = _cr_series(facts, _CR_STI, ns, False)
            if not e_s or not d_s:
                rec["note"] = "缺营业利润或折旧摊销年度申报"
                out["rows"].append(rec)
                continue

            notes = []
            lev_fy = max((y for y in e_s if y in d_s and y in l_s and y in c_s),
                         default=None)
            cov_fy = max((y for y in e_s if y in d_s and y in i_s and i_s[y]),
                         default=None)

            def _fresh(y):
                try:
                    return int(y[:4]) >= cur_year - _CR_MAX_STALE_YEARS
                except Exception:
                    return False

            if lev_fy and not _fresh(lev_fy):
                notes.append(f"债务科目停更于 FY{lev_fy[:4]}")
                lev_fy = None
            if cov_fy and not _fresh(cov_fy):
                notes.append(f"利息科目停更于 FY{cov_fy[:4]}")
                cov_fy = None

            if lev_fy:
                ebitda = e_s[lev_fy] + d_s[lev_fy]
                if ebitda > 0:
                    debt = l_s[lev_fy] + (s_s.get(lev_fy) or 0)
                    liq = c_s[lev_fy] + (t_s.get(lev_fy) or 0)
                    rec["ebitda"] = ebitda
                    rec["net_debt"] = debt - liq
                    rec["leverage"] = round((debt - liq) / ebitda, 2)
                    lev_n += 1
                else:
                    notes.append("EBITDA ≤ 0, 杠杆倍数无经济意义")
            elif not notes:
                notes.append("缺债务或现金年度申报")

            if cov_fy:
                itx = abs(i_s[cov_fy])
                rec["interest"] = itx
                rec["coverage"] = round((e_s[cov_fy] + d_s[cov_fy]) / itx, 1)
                cov_n += 1
            elif "利息" not in "".join(notes):
                notes.append("缺利息支出年度申报")

            fy = lev_fy or cov_fy
            if fy:
                rec["fy"] = fy
                fys.append(fy)
                rec["status"] = "ok"
            rec["note"] = "; ".join(notes)
            out["rows"].append(rec)

    if lev_n or cov_n:
        out["status"] = "ok"
        out["as_of"] = max(fys) if fys else None
    out["lev_n"], out["cov_n"] = lev_n, cov_n
    out["total"] = len(out["rows"])
    return out

# ★UA 必须含【真实邮箱】(2026-08-31 实测): SEC 对 data.sec.gov 宽松, 但
#   www.sec.gov/Archives/ 严格校验 —— 旧值 "EcoVolChecker research (contact: chao.jin)"
#   拉 Archives 一律 403, 导致读 filing 原始 XBRL 的补充源静默失败返回 None。
_SEC_UA = {"User-Agent": "Eco-Volatility-Checker chao.jin@uber.com"}
_FACTS_CACHE = {}


def _sec_facts(cik):
    """SEC companyfacts, 进程内缓存。失败返 None(绝不返空壳冒充成功)。"""
    if cik in _FACTS_CACHE:
        return _FACTS_CACHE[cik]
    import json
    import time
    import urllib.request
    url = f"https://data.sec.gov/api/xbrl/companyfacts/CIK{cik}.json"
    try:
        req = urllib.request.Request(url, headers=_SEC_UA)
        with urllib.request.urlopen(req, timeout=45) as r:
            d = json.loads(r.read().decode())
    except Exception as e:
        print(f"    [sec] CIK{cik} 拉取失败: {type(e).__name__} {str(e)[:60]}")
        d = None
    _FACTS_CACHE[cik] = d
    time.sleep(0.15)          # SEC 限速礼貌间隔
    return d


def _xbrl_annual(facts, tag, ns="us-gaap", duration=True):
    """年报(10-K/20-F)年度口径序列 {end: val}。只认 330~400 天区间, 排除季度。

    ★币种铁律: 同一标签常同时申报 USD 与本币(如 TSM 的 TWD, 差 ~32 倍)。
      必须【只取 USD】—— 早期实现遍历所有 unit 靠「后者覆盖」, 结果取决于
      SEC JSON 的 key 顺序, 一旦顺序变化就会静默产生 32 倍错误。
    """
    import datetime as _dt
    node = ((facts.get("facts") or {}).get(ns) or {}).get(tag)
    if not node:
        return {}
    units = node.get("units") or {}
    key = None
    for u in units:
        if u.upper() == "USD":
            key = u
            break
    if key is None:                     # 无 USD 申报 → 币种不可比, 宁缺勿错
        return {}
    out = {}
    for it in units[key]:
        if it.get("form") not in ("10-K", "10-K/A", "20-F", "20-F/A"):
            continue
        end, st = it.get("end"), it.get("start")
        if not end:
            continue
        if duration:
            if not st:
                continue
            try:
                d1 = _dt.date.fromisoformat(st)
                d2 = _dt.date.fromisoformat(end)
            except Exception:
                continue
            if not (330 <= (d2 - d1).days <= 400):
                continue
        out[end] = it["val"]       # 同 end 后出现者覆盖 = 取修订值
    return out


def _pick(facts, tags, ns="us-gaap", duration=True):
    for t in tags:
        s = _xbrl_annual(facts, t, ns, duration)
        if s:
            return s
    return {}


def fetch_ai_fcf():
    """AI 产业链公司 标准自由现金流(FCF = 经营现金流 − 资本开支)。

    返回 {"groups":[{name, members:[{ticker,name,fy,ocf,capex,fcf,
                                     ebit,interest,coverage,status}]}],
          "excluded":[...], "as_of", "status", "source", "caveat"}
    ★口径: 标准 FCF, 不做租赁调整(融资租赁本金 19 家中 7 家缺披露, 口径不可比)。
    ★绝不编: 任一公司缺 OCF 或 CapEx → status="未找到", 数值留 None, 不估算不近似。
    ★财年不同: 各公司 FY 截止日不同(如 MSFT 6月/NVDA 1月), 图上须标注实际 FY。
    """
    out = {"groups": [], "excluded": list(AI_EXCLUDED), "as_of": None,
           "status": "未找到",
           "source": "SEC EDGAR XBRL companyfacts (10-K/20-F 年报官方申报值)",
           "caveat": "标准 FCF = 经营现金流 − 资本开支, 未扣融资租赁本金"
                     "(19 家中 7 家未披露该字段, 强行调整会造成口径不可比)。"
                     "FCF 源于年报, 每年更新一次, 非日频。"
                     "各公司财年截止日不同, 卡片已标注实际财年。"}
    fys, ok_n, tot_n = [], 0, 0

    for gname, members in AI_UNIVERSE:
        grp = {"name": gname, "members": []}
        for tk, cname, cik, zhname in members:
            tot_n += 1
            rec = {"ticker": tk, "name": cname, "zh": zhname, "fy": None, "ocf": None,
                   "capex": None, "fcf": None, "ebit": None,
                   "interest": None, "coverage": None, "status": "未找到"}
            facts = _sec_facts(cik)
            if not facts:
                rec["status"] = "SEC 拉取失败"
                grp["members"].append(rec)
                continue

            ns_list = [("us-gaap", _OCF_TAGS, _CAPEX_TAGS, _EBIT_TAGS, _INT_TAGS)]
            if "ifrs-full" in (facts.get("facts") or {}):
                ns_list.insert(0, ("ifrs-full", _IFRS_OCF, _IFRS_CAPEX,
                                   _IFRS_EBIT, _IFRS_INT))
            ocf = capex_map = {}
            ns_used = None
            for ns, otags, ctags, etags, itags in ns_list:
                ocf = _pick(facts, otags, ns)
                if not ocf:
                    continue
                ns_used = (ns, etags, itags)
                fy = max(ocf)
                tot = 0.0
                hit = False
                for t in ctags:
                    s = _xbrl_annual(facts, t, ns)
                    if fy in s:
                        tot += s[fy]
                        hit = True
                if hit:
                    capex_map = {fy: tot}
                    break
                ocf = {}
            if not ocf or not capex_map:
                grp["members"].append(rec)
                continue

            fy = max(ocf)
            # ★补充源(Chao 2026-08-31 方案2): companyfacts 对 20-F 外国发行人滞后,
            #   若最新年度已落后当前年份 ≥1 年, 就去 EDGAR 找更新的年报直接读 XBRL。
            #   实测 TSM: companyfacts 停在 2024-12-31, 而 2025 年 20-F 早已申报。
            _newer = _latest_annual_filing(cik, after=fy)
            if _newer:
                _o, _c = _fcf_from_filing_xbrl(
                    cik, _newer,
                    (_IFRS_OCF if ns_used and ns_used[0] == "ifrs-full" else _OCF_TAGS),
                    (_IFRS_CAPEX if ns_used and ns_used[0] == "ifrs-full" else _CAPEX_TAGS))
                if _o and _c:
                    _common = sorted(set(_o) & set(_c))
                    if _common and _common[-1] > fy:
                        _f2 = _common[-1]
                        ocf, capex_map, fy = {_f2: _o[_f2]}, {_f2: _c[_f2]}, _f2
                        rec["src_note"] = "20-F XBRL 原文件(companyfacts 未收录该年)"
            rec.update(fy=fy, ocf=ocf[fy], capex=capex_map[fy],
                       fcf=ocf[fy] - capex_map[fy], status="ok")
            # 利息保障倍数(供合成评级用); 拿不到不影响 FCF
            ns, etags, itags = ns_used
            eb, itx = _pick(facts, etags, ns), _pick(facts, itags, ns)
            if fy in eb and fy in itx and itx[fy]:
                rec["ebit"], rec["interest"] = eb[fy], itx[fy]
                rec["coverage"] = round(eb[fy] / abs(itx[fy]), 2)
            ok_n += 1
            fys.append(fy)
            grp["members"].append(rec)
        out["groups"].append(grp)

    if ok_n:
        out["status"] = "ok"
        out["as_of"] = max(fys)
    out["ok_count"], out["total_count"] = ok_n, tot_n
    return out


def _latest_annual_filing(cik, after):
    """在 EDGAR submissions 里找【报告期晚于 after】的最新年报 accession。

    ★用途(Chao 2026-08-31 方案2): companyfacts 聚合库对 20-F 外国私人发行人收录滞后,
      本函数直接查申报索引, 确认是否存在更新的年报可读。
    after: "YYYY-MM-DD" 已有的最新财年末。无更新年报则返回 None。
    """
    import urllib.request
    try:
        u = f"https://data.sec.gov/submissions/CIK{int(cik):010d}.json"
        d = json.loads(urllib.request.urlopen(
            urllib.request.Request(u, headers=_SEC_UA),
            timeout=40).read().decode("utf-8", "replace"))
        r = d.get("filings", {}).get("recent", {})
        best = None
        for form, rep, accn in zip(r.get("form", []), r.get("reportDate", []),
                                   r.get("accessionNumber", [])):
            if not form.startswith(("10-K", "20-F")):
                continue
            if form.endswith("/A"):
                continue                       # 修订件先不用, 避免口径混乱
            if rep and rep > (after or "") and (best is None or rep > best[0]):
                best = (rep, accn)
        return best[1] if best else None
    except Exception:
        return None


def _fcf_from_filing_xbrl(cik, accn, ocf_tags, capex_tags, unit="usd"):
    """从【单份年报的 XBRL 原文件】取 OCF/CapEx —— companyfacts 缺该年时的补充源。

    ★为什么需要(Chao 2026-08-31 方案2): SEC companyfacts 聚合库对外国私人发行人
      (20-F) 收录严重滞后。实测 TSM 2025 年 20-F 已于 2026-04-16 申报,
      但 companyfacts 最新仍停在 2024-12-31 → 图上 TSM 落后整整一年。
      直接读那份 filing 的 XBRL 实例文件即可拿到, 数据仍是官方申报值, 口径不变。
    ★单位: 20-F 常同时申报本币(twd)与美元换算值, 必须按 unitRef 过滤,
      否则 dict 覆盖会混用两种币种 —— 实测混用会算出 FCF=-11999 億 的荒谬值。
    ★只取无 segment 的合并数(带 segment 的是分部/分项披露)。
    返回 {fy_end: value} 两个 dict (ocf, capex); 任一为空则返回 (None, None)。
    """
    import re as _re
    import urllib.request
    base = ("https://www.sec.gov/Archives/edgar/data/"
            f"{int(cik)}/{accn.replace('-', '')}/")
    try:
        idx = json.loads(urllib.request.urlopen(
            urllib.request.Request(base + "index.json", headers=_SEC_UA),
            timeout=40).read().decode("utf-8", "replace"))
        names = [it["name"] for it in idx["directory"]["item"]]
        inst = next((n for n in names if n.endswith("_htm.xml")), None)
        if not inst:
            return None, None
        raw = urllib.request.urlopen(
            urllib.request.Request(base + inst, headers=_SEC_UA),
            timeout=120).read().decode("utf-8", "ignore")
    except Exception:
        return None, None

    ctx = {}
    for m in _re.finditer(r'<context id="([^"]+)">(.*?)</context>', raw, _re.S):
        cid, body = m.group(1), m.group(2)
        s = _re.search(r"<startDate>([^<]+)", body)
        e = _re.search(r"<endDate>([^<]+)", body)
        if s and e:
            ctx[cid] = (s.group(1), e.group(1), "<segment>" in body)

    def _grab(tags):
        got = {}
        for tag in tags:
            t = tag.split(":")[-1]
            for m in _re.finditer(
                    r'<[a-z-]+:%s([^>]*)>(-?[0-9.]+)</[a-z-]+:%s>' % (t, t), raw):
                attrs, val = m.group(1), m.group(2)
                cid = _re.search(r'contextRef="([^"]+)"', attrs)
                un = _re.search(r'unitRef="([^"]+)"', attrs)
                if not cid or not un:
                    continue
                cid = cid.group(1)
                if un.group(1).lower() != unit:
                    continue                     # 币种不符 → 跳过(防混算)
                if cid not in ctx or ctx[cid][2]:
                    continue                     # 带 segment 的分部数据 → 跳过
                st, en, _ = ctx[cid]
                if st[:4] != en[:4]:
                    continue                     # 非完整自然年 → 跳过
                got[en] = float(val)
        return got

    ocf, capex = _grab(ocf_tags), _grab(capex_tags)
    if not ocf or not capex:
        return None, None
    return ocf, capex


def fetch_synthetic_spreads():
    """Damodaran 合成评级利差表(利息保障倍数 → 评级 → 违约利差)。

    返回 {"table":[{lo,hi,rating,spread_pct}], "status", "source", "caveat"}
    ★这是【模型推导】不是市场成交价: 免费源无单名企业债市场利差
      (FINRA TRACE 端点全 404 / Boerse-Frankfurt 403 / WSJ 401)。
    ★静态年度更新, 非日频。使用处必须显著标注"模型估算, 非市场报价"。
    """
    import re as _re
    import urllib.request
    out = {"table": [], "status": "未找到",
           "source": "Damodaran (NYU Stern) 合成评级违约利差表",
           "url": "https://pages.stern.nyu.edu/~adamodar/New_Home_Page/datafile/ratings.htm",
           "caveat": "★模型估算非市场报价: 由利息保障倍数映射评级再映射利差。"
                     "免费源无单名企业债市场成交利差(FINRA TRACE 已关闭公开端点)。"
                     "该表年度更新, 非日频。"}
    try:
        req = urllib.request.Request(out["url"], headers=_SEC_UA)
        with urllib.request.urlopen(req, timeout=35) as r:
            html = r.read().decode("utf-8", "replace")
    except Exception as e:
        out["status"] = f"抓取失败: {type(e).__name__}"
        return out

    txt = _re.sub(r"<[^>]+>", " ", html)
    txt = _re.sub(r"&nbsp;?", " ", txt)
    txt = _re.sub(r"\s+", " ", txt)
    pat = _re.compile(r"(-?\d+\.?\d*)\s+(-?\d+\.?\d*)\s+"
                      r"([A-Da-z0-9+\-/]{1,8})\s+(\d+\.\d+)%")
    seen = set()
    for lo, hi, rating, spr in pat.findall(txt):
        key = (rating, spr)
        if key in seen:
            continue
        seen.add(key)
        try:
            out["table"].append({"lo": float(lo), "hi": float(hi),
                                 "rating": rating, "spread_pct": float(spr)})
        except ValueError:
            continue
    if len(out["table"]) >= 10:      # 完整表 15 档; 少于 10 视为解析失败
        out["status"] = "ok"
    else:
        out["status"] = f"解析异常(仅 {len(out['table'])} 档)"
        out["table"] = []
    return out


def fetch_treasury_curve():
    """美债各期限收益率(利差基准腿), FRED 官方。返回 {tenor: (date, pct)}。"""
    out = {"points": {}, "status": "未找到", "as_of": None,
           "source": "FRED · U.S. Treasury Constant Maturity"}
    for tenor, sid in [("2Y", "DGS2"), ("5Y", "DGS5"), ("7Y", "DGS7"),
                       ("10Y", "DGS10"), ("20Y", "DGS20"), ("30Y", "DGS30")]:
        s = _fred_hist(sid, (datetime.date.today() -
                             datetime.timedelta(days=30)).strftime("%Y-%m-%d"))
        if s:
            out["points"][tenor] = {"date": s[-1][0], "pct": s[-1][1]}
    if out["points"]:
        out["status"] = "ok"
        out["as_of"] = max(v["date"] for v in out["points"].values())
    return out


# ══════════════════════════════════════════════════════════════════
#  中国 CIPS(人民币跨境支付系统) 使用量
# ══════════════════════════════════════════════════════════════════
#  官方一手源(权威, 优先):
#    · 当年月度 PDF —— 月度 笔数/金额/工作日/日均, 但【只有当年】
#    · 历年年度 PDF —— 2015 至今 年度 笔数/金额
#  第三方补充源(chinadata.live) 用于回补官方已下架的往年月度:
#    ★该源 CSV 列名写 amount_100m_rmb(亿元), 实际值是【十亿元】, 差 10 倍。
#      本函数用【官方重叠月】实测推断倍率, 绝不硬编码, 对不上就整体弃用。
CIPS_PDF_CUR = ("https://www.cips.com.cn/kjjqgs/articleFileDir/2025-12/05/"
                "c1ed77e02af749fb87cfb17484f9efb6.pdf")
CIPS_PDF_HIST = ("https://www.cips.com.cn/kjjqgs/articleFileDir/2025-12/05/"
                 "14c90b4ac33f4440b0085632f2c1ee73.pdf")
CIPS_3P_CSV = ("https://chinadata.live/api/v2/data/"
               "china-cips-payment-system-monthly?format=csv")
CIPS_STATS_PAGE = "https://www.cips.com.cn/kjjqgs/cipsfw/ywtj/index.shtml"


def _cips_pdf_text(url):
    """下载并抽取 CIPS PDF 文本。失败返回 None(绝不返空串冒充成功)。"""
    import urllib.request
    try:
        import pymupdf
    except ImportError:
        try:
            import fitz as pymupdf          # 旧版包名
        except ImportError:
            return None
    try:
        req = urllib.request.Request(
            url, headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"})
        with urllib.request.urlopen(req, timeout=60) as r:
            if r.status != 200:
                return None
            blob = r.read()
        if not blob.startswith(b"%PDF"):
            return None
        doc = pymupdf.open(stream=blob, filetype="pdf")
        return "\n".join(p.get_text() for p in doc)
    except Exception:
        return None


def fetch_cips(months=36):
    """中国 CIPS 跨境人民币支付系统 使用量(月度 + 年度)。

    返回 {"monthly":[{ym,count,amount_yi,workdays,avg_count,avg_amount_yi,
                       src}],
          "annual":[(year,count,amount_yi)],
          "as_of","status","source","note","official_months","third_months"}

    ★口径: amount 单位一律【亿元】。日均金额 = 官方直接披露值(非自算),
      仅第三方回补月自算(= 月度总额 / 工作日, 工作日缺失则留 None)。
    ★官方值优先; 第三方仅补官方缺失月, 且必须通过重叠月一致性校验。
    ★抓不到 → status 说明原因, 绝不编造/补零。
    """
    import csv as _csv
    import io as _io
    import re as _re
    import urllib.request

    out = {"monthly": [], "annual": [], "as_of": None, "status": "未找到",
           "source": "CIPS 跨境银行间支付清算有限责任公司 官方业务统计",
           "note": "", "official_months": 0, "third_months": 0,
           "page": CIPS_STATS_PAGE}

    # ---------- 1) 官方【当年】月度 ----------
    off = {}
    txt = _cips_pdf_text(CIPS_PDF_CUR)
    year_cur = None
    if txt:
        # 标题/正文里找年份; 找不到则不猜, 留 None 并跳过该源
        ym = _re.search(r"(20\d\d)\s*年", txt)
        if ym:
            year_cur = int(ym.group(1))
        else:
            # 当年表通常无年份字样 → 用历年表推断(见下), 先暂存原始行
            year_cur = -1
        rows = _re.findall(
            r"(\d{1,2})月\s+([\d,]+)\s+([\d,\.]+)\s+(\d+)\s+([\d,]+)\s+([\d,\.]+)",
            txt)
        for mo, cnt, amt, wd, dc, da in rows:
            off[int(mo)] = (int(cnt.replace(",", "")),
                            float(amt.replace(",", "")),
                            int(wd), int(dc.replace(",", "")),
                            float(da.replace(",", "")))

    # ---------- 2) 官方【历年】年度 ----------
    txt2 = _cips_pdf_text(CIPS_PDF_HIST)
    annual = []
    if txt2:
        for y, cnt, amt in _re.findall(r"(20\d\d)年\s+([\d,]+)\s+([\d,\.]+)", txt2):
            annual.append((int(y), int(cnt.replace(",", "")),
                           float(amt.replace(",", ""))))
        annual.sort()
    out["annual"] = annual

    # 当年月度表的年份 = 历年表最后一年 + 1 (历年表只收已完结年度)
    if off and year_cur == -1:
        year_cur = (annual[-1][0] + 1) if annual else None
    if off and not year_cur:
        off = {}
        out["note"] = "当年月度表年份无法确定, 已弃用该表(不猜年份)"

    monthly = {}
    if off and year_cur:
        for mo, (cnt, amt, wd, dc, da) in off.items():
            monthly[f"{year_cur}-{mo:02d}"] = {
                "ym": f"{year_cur}-{mo:02d}", "count": cnt, "amount_yi": amt,
                "workdays": wd, "avg_count": dc, "avg_amount_yi": da,
                "src": "official"}
    out["official_months"] = len(monthly)

    # ---------- 3) 第三方回补往年月度(需通过重叠月校验) ----------
    try:
        req = urllib.request.Request(
            CIPS_3P_CSV, headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)",
                                  "Accept": "text/csv,*/*"})
        with urllib.request.urlopen(req, timeout=45) as r:
            csv_txt = r.read().decode("utf-8", "replace") if r.status == 200 else None
    except Exception:
        csv_txt = None

    if csv_txt:
        third = {}
        for row in _csv.DictReader(_io.StringIO(csv_txt)):
            d = (row.get("date") or "").strip()
            if not _re.fullmatch(r"20\d\d-\d{2}", d):
                continue
            try:
                third[d] = (int(float(row["transactions"])),
                            float(row["amount_100m_rmb"]))
            except (KeyError, TypeError, ValueError):
                continue

        # ★倍率校验: 用官方重叠月推断第三方金额单位, 不硬编码
        ratios, cnt_ok, overlap = [], 0, 0
        for k, v in monthly.items():
            if k in third:
                overlap += 1
                if third[k][0] == v["count"]:
                    cnt_ok += 1
                if third[k][1]:
                    ratios.append(v["amount_yi"] / third[k][1])
        if overlap >= 3 and cnt_ok == overlap and ratios:
            mult = round(sum(ratios) / len(ratios), 4)
            spread = max(ratios) - min(ratios)
            # 倍率必须稳定(离散<1%)且接近整数倍, 否则不可信
            if spread < 0.01 * mult and abs(mult - round(mult)) < 0.02:
                mult = float(round(mult))
                added = 0
                for k, (c, a) in sorted(third.items()):
                    if k in monthly:
                        continue
                    monthly[k] = {"ym": k, "count": c,
                                  "amount_yi": round(a * mult, 1),
                                  "workdays": None, "avg_count": None,
                                  "avg_amount_yi": None, "src": "third_party"}
                    added += 1
                out["third_months"] = added
                out["source"] += (" ; 往年月度经第三方 chinadata.live 回补"
                                  f"(重叠 {overlap} 月与官方逐月一致, "
                                  f"金额单位换算 ×{mult:g})")
            else:
                out["note"] += (f" 第三方金额倍率不稳定(mult={mult:.3f}, "
                                f"离散={spread:.3f}), 已弃用第三方回补。")
        elif overlap:
            out["note"] += (f" 第三方与官方重叠 {overlap} 月中仅 {cnt_ok} 月笔数"
                            "一致, 未通过校验, 已弃用第三方回补。")

    if not monthly:
        out["status"] = "未找到"
        out["note"] = (out["note"] or "CIPS 官方 PDF 与第三方源均不可用").strip()
        return out

    ser = sorted(monthly.values(), key=lambda d: d["ym"])
    if months:
        ser = ser[-months:]
    # 第三方回补月无工作日数 → 日均金额留 None(不假装有), 前端标 n/a
    out["monthly"] = ser
    out["as_of"] = ser[-1]["ym"]
    out["status"] = "ok"
    return out


# ══════════════════════════════════════════════════════════════════════
# 中国模式断代监控 (Chao 2026-09-19)
# ══════════════════════════════════════════════════════════════════════
# ★需求背景: Ryan Perkins《China's Economic Crossroads》(2025-12-08) 判断 ——
#   中国已无法靠投资实现充分就业, 也无法在旧模式的失业问题解决前靠消费拉动增长;
#   结局二元: 软着陆 或 长期停滞+社会压力。他自己押"谨慎渐进的钢丝行走"。
#   本模块把他引用的判据变成可月度跟踪的指标。
#
# ★两个数据源(2026-09-19 实测):
#   ① 东方财富宏观 API   —— 月度总量, 最新 2026-08, 11 项
#   ② NBS 官方英文发布稿 —— FAI/地产<b>分项</b>(基建/制造/民间/分区域), 一手权威
#      NBS 的 easyquery JSON API 本机 403(数据中心 IP), 但英文发布稿 HTML 可达。
#
# ★口径铁律(开发时栽过两次, 务必遵守):
#   伯金斯全文讲的是<b>累计同比</b>(FAI 3.8%)。东财 BASE_SAME 是<b>单月同比</b>,
#   2026-08 为 -13.51%, 而官方累计口径是 -7.2% —— 差 6 个百分点。
#   两者都是真数但<b>不可混用</b>, 本模块每项都带 basis 字段标注口径。
#   同理社零必须用 RETAIL_ACCUMULATE_SAME(累计) 而非 RETAIL_TOTAL(绝对额)。
#
# ★诚实边界: 伯金斯引用但公开源确实拿不到的, 列在 unavailable 里显式呈现,
#   绝不用近似指标冒充(尤其"高技术/绿色制造 FAI 占比 28%"这个他的核心门槛)。

_CN_EM_SPECS = [
    # key, 名称, reportName, 字段, 单位, 口径, 引擎, 伯金斯参照
    ("fai_m",    "固定资产投资",   "RPT_ECONOMY_ASSET_INVEST",    "BASE_SAME",              "%", "单月同比", "old",    "3.8%"),
    ("boom",     "国房景气指数",   "RPT_ECONOMY_BOOM_INDEX",      "BOOM_INDEX",             "点", "季度",    "old",    None),
    ("retail_c", "社会消费品零售", "RPT_ECONOMY_TOTAL_RETAIL",    "RETAIL_ACCUMULATE_SAME", "%", "累计同比", "new",    None),
    ("income_c", "居民可支配收入", "RPT_ECONOMY_INCOME",          "ACCUMULATE_SAME",        "%", "累计同比", "new",    None),
    ("faith",    "消费者信心指数", "RPT_ECONOMY_FAITH_INDEX",     "CONSUMERS_FAITH_INDEX",  "点", "月度",    "new",    None),
    ("indus_m",  "规上工业增加值", "RPT_ECONOMY_INDUS_GROW",      "BASE_SAME",              "%", "单月同比", "new",    "5.2%"),
    ("cpi",      "CPI",            "RPT_ECONOMY_CPI",             "NATIONAL_SAME",          "%", "同比",    "price",  None),
    ("ppi",      "PPI",            "RPT_ECONOMY_PPI",             "BASE_SAME",              "%", "同比",    "price",  None),
    ("pmi",      "制造业 PMI",     "RPT_ECONOMY_PMI",             "MAKE_INDEX",             "点", "月度",    "policy", None),
    ("m2",       "M2",             "RPT_ECONOMY_CURRENCY_SUPPLY", "BASIC_CURRENCY_SAME",    "%", "同比",    "policy", None),
    ("tax",      "税收收入",       "RPT_ECONOMY_TAX",             "TAX_INCOME_SAME",        "%", "同比",    "policy", None),
]

# NBS 发布稿分项抽取规则
# ★坑: 间隔符不能用 [^.]*? —— 金额 "29,309.2 billion" 里的小数点会提前终止匹配,
#   导致最核心的 fai_total 永远抽不到。必须用 .{0,120}? 限长非贪婪。
_CN_NBS_RULES = [
    ("fai_total",   "固定资产投资 总",  r"national investment in fixed assets.{0,120}?year-on-year (decrease|increase) of ([\d.]+)%", "old",    "3.8%"),
    ("infra",       "基建投资",         r"investment in infrastructure.{0,120}?(decreased|increased) by ([\d.]+)% year on year", "old",    "+6.4%"),
    ("manuf",       "制造业投资",       r"investment in manufacturing (decreased|increased) by ([\d.]+)%", "new",    None),
    ("private",     "民间投资",         r"non-governmental investment in fixed assets (decreased|increased) by ([\d.]+)%", "old",    None),
    ("ip_products", "知识产权产品投资", r"investment in intellectual property products (increased|decreased) by ([\d.]+)%", "new",    "28% 高技术占比"),
    ("northeast",   "东北地区投资",     r"investment in the northeastern region (decreased|increased) by ([\d.]+)%", "social", None),
    ("estate_inv",  "房地产开发投资",   r"investment in real estate development was [\d,.]+ billion yuan, a year-on-year (decrease|increase) of ([\d.]+)%", "old", "-2.1%"),
    ("new_starts",  "房屋新开工面积",   r"floor space of buildings newly started was [\d,.]+ million square meters, (down|up) ([\d.]+)%", "old", None),
]

_CN_UNAVAILABLE = [
    {"name": "高技术/绿色制造 FAI 占比", "perkins_ref": "28%",
     "reason": "★伯金斯判断「新引擎是否接上」的唯一量化门槛。公开源无此分项；"
               "最接近的替代是 NBS「知识产权产品投资」增速，但那是增速不是占比，口径不同不可替代。"},
    {"name": "青年失业率 16-24 岁", "perkins_ref": ">25%（CEI 自家估计）",
     "reason": "World Bank 官方口径年度值约 15.8%，与 CEI 估计差约 10 个百分点；"
               "2023 年官方口径改版（不含在校生）。两者不可混用，故不取。"},
]

_CN_ENGINE_LABEL = {
    "old":    "旧引擎 · 投资/地产/信贷",
    "new":    "新引擎 · 消费/产业升级",
    "price":  "通缩漩涡 · 价格链条",
    "policy": "政策空间",
    "social": "社会压力 · 区域分化",
}


def _cn_em(report, size=40):
    import requests
    url = ("https://datacenter-web.eastmoney.com/api/data/v1/get?"
           f"reportName={report}&columns=ALL&pageSize={size}"
           "&sortColumns=REPORT_DATE&sortTypes=-1")
    r = requests.get(url, timeout=45, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0 Safari/537.36",
        "Referer": "https://data.eastmoney.com/"})
    r.raise_for_status()
    d = r.json()
    if not d.get("success"):
        raise RuntimeError("eastmoney success=false")
    return (d.get("result") or {}).get("data") or []


def _cn_plain(text):
    import re, html as _h
    t = re.sub(r"<script.*?</script>|<style.*?</style>", " ", text, flags=re.S)
    return re.sub(r"\s+", " ", _h.unescape(re.sub(r"<[^>]+>", " ", t)))


def fetch_china_regime_shift():
    """中国模式断代监控：旧引擎停摆 / 新引擎接续 的月度跟踪。"""
    import re, requests
    from datetime import timezone, timedelta
    _jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": _jst, "period": None,
           "monthly": [], "breakdown": [], "unavailable": list(_CN_UNAVAILABLE),
           "sources": {}, "errors": [], "engine_label": dict(_CN_ENGINE_LABEL)}

    # ① 东财月度总量
    for key, name, rpt, field, unit, basis, eng, pk in _CN_EM_SPECS:
        try:
            rows = _cn_em(rpt)
            s = [{"ym": (r.get("REPORT_DATE") or "")[:7], "v": float(r[field])}
                 for r in rows if r.get(field) is not None]
            s.reverse()
            if not s:
                out["errors"].append(f"东财/{name}: 字段 {field} 全空")
                continue
            last = s[-1]
            prev = s[-2] if len(s) > 1 else None
            yr = [x["v"] for x in s[-12:]]
            out["monthly"].append({
                "key": key, "name": name, "engine": eng, "unit": unit, "basis": basis,
                "perkins_ref": pk, "ym": last["ym"], "value": last["v"],
                "mom": round(last["v"] - prev["v"], 3) if prev else None,
                "yr_min": min(yr), "yr_max": max(yr),
                "yr_avg": round(sum(yr) / len(yr), 2),
                "series": s[-36:], "source": "eastmoney:" + rpt,
            })
        except Exception as e:
            out["errors"].append(f"东财/{name}: {type(e).__name__} {str(e)[:70]}")

    # ② NBS 官方分项
    try:
        base = "https://www.stats.gov.cn/english/PressRelease/"
        hdr = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                             "AppleWebKit/537.36 Chrome/120.0 Safari/537.36",
               "Accept-Language": "en-US,en;q=0.9"}
        lst = requests.get(base, timeout=50, headers=hdr)
        lst.raise_for_status()
        links = {}
        for href, title in re.findall(r'href="(\./\d+/t\d+_\d+\.html)"[^>]*>\s*([^<]{8,140})', lst.text):
            ti = re.sub(r"\s+", " ", title).strip()
            if "Investment in Fixed Assets" in ti and "fai" not in links:
                links["fai"] = (ti, base + href[2:])
            elif "Investment in Real Estate" in ti and "estate" not in links:
                links["estate"] = (ti, base + href[2:])
        merged = {}
        for which, (title, url) in links.items():
            out["sources"][which] = {"title": title, "url": url}
            page = requests.get(url, timeout=60, headers=hdr)
            page.raise_for_status()
            txt = _cn_plain(page.text)
            for key, name, pat, eng, pk in _CN_NBS_RULES:
                m = re.search(pat, txt, re.I)
                if m:
                    neg = m.group(1).lower() in ("decrease", "decreased", "down")
                    merged[key] = {
                        "key": key, "name": name, "engine": eng, "perkins_ref": pk,
                        "value": -float(m.group(2)) if neg else float(m.group(2)),
                        "unit": "%", "basis": "累计同比", "src": which,
                    }
            if not out["period"]:
                mm = re.search(r"from January to (\w+) (\d{4})", title, re.I)
                if mm:
                    out["period"] = f"{mm.group(2)} 年 1-{mm.group(1)}"
        out["breakdown"] = list(merged.values())
        if not out["breakdown"]:
            out["errors"].append("NBS: 发布稿解析出 0 项分项(页面格式可能已变)")
    except Exception as e:
        out["errors"].append(f"NBS: {type(e).__name__} {str(e)[:80]}")

    # ★★缓存回退(2026-09-21 修): NBS 偶发 RemoteDisconnected → breakdown 直接变 0 项,
    #   而 status 仍是 "ok(部分失败)", 页面照常渲染但少掉 8 项官方分项 —— 静默降级。
    #   纪律(项目既有): 抓不到时必须<b>保留旧真值</b>, 不能覆盖成空。
    import os as _os, json as _j
    # ★路径解析: __file__ 在被临时复制到别处执行时会失准(反测时踩到),
    #   故先试 __file__ 推导, 再回退到 cwd 下的 data/。
    _cands = []
    try:
        _cands.append(_os.path.join(
            _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))),
            "data", "china_regime_shift.json"))
    except Exception:
        pass
    _cands.append(_os.path.join(_os.getcwd(), "data", "china_regime_shift.json"))
    _cp = next((p for p in _cands if _os.path.exists(p)), _cands[0])
    if out["breakdown"]:
        # 抓到了 → 落盘供下次回退
        try:
            _os.makedirs(_os.path.dirname(_cp), exist_ok=True)
            _j.dump({"breakdown": out["breakdown"], "period": out.get("period"),
                     "sources": out.get("sources"),
                     "saved_at": datetime.datetime.now().isoformat()},
                    open(_cp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        except Exception as e:
            out["errors"].append(f"缓存写入失败(不影响本次): {type(e).__name__}")
    elif _os.path.exists(_cp):
        # 没抓到 → 读回缓存, 并<b>明确标注是缓存</b>(不假装是当期新数)
        try:
            _c = _j.load(open(_cp, encoding="utf-8"))
            if _c.get("breakdown"):
                out["breakdown"] = _c["breakdown"]
                out["period"] = out.get("period") or _c.get("period")
                out["sources"] = out.get("sources") or _c.get("sources") or {}
                _age = (datetime.datetime.now().timestamp() - _os.path.getmtime(_cp)) / 86400
                out["breakdown_cached"] = True
                out["breakdown_cache_age_days"] = round(_age, 1)
                out["errors"].append(
                    f"NBS 实时抓取失败 → 分项回退缓存({_age:.1f} 天前), 已在页面标注")
        except Exception as e:
            out["errors"].append(f"缓存读取失败: {type(e).__name__}")

    if not out["monthly"] and not out["breakdown"]:
        out["status"] = "error"
    elif out["errors"]:
        out["status"] = "ok(部分失败)"
    return out


# ══════════════════════════════════════════════════════════════════════
# WRS · Walmart 衰退信号 (Chao 2026-09-19, 甲方案)
# ══════════════════════════════════════════════════════════════════════
# ★口径来源: Jim Paulsen (Leuthold Group 前首席投资策略师) 提出的
#   WRS = Walmart 股价 / S&P Global Luxury Index (80 只奢侈品股)
#   逻辑: 经济走弱时消费从奢侈品转向折扣零售 → 比值上行 = 衰退压力上升。
#
# ★★分母是代理, 不是原指数 —— 这是甲方案的核心, 必须在图上标注:
#   S&P Global Luxury Index 是 S&P 专有指数, 免费源拿不到(实测 spglobal 403、
#   Yahoo ^SPGLUX 404)。本实现用 <b>LUXU.L = Amundi S&P Global Luxury UCITS ETF</b>,
#   它<b>追踪的正是该指数</b>且以 USD 计价(与 WMT 同币种, 无需汇率换算),
#   是最贴近原口径的公开代理。但 ETF 有管理费/跟踪误差/流动性折溢价,
#   <b>绝对水平与 Paulsen 原值不可直接比对</b>, 只看方向与分位。
#
# ★实测(2026-09-19): 共同交易日 2467 天(2016-09 至今), 当前 0.4858,
#   分位 95.9%, 历史峰值 0.6012 在 2026-05-19 —— <b>当前不是创新高</b>, 不许说成创新高。

_WRS_NUM = ("WMT", "Walmart")
_WRS_DEN = ("LUXU.L", "Amundi S&P Global Luxury UCITS ETF")


def _wrs_yahoo(sym, rng="10y"):
    import json as _j, urllib.request, datetime as _dt
    u = (f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}"
         f"?range={rng}&interval=1d")
    req = urllib.request.Request(u, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 Chrome/120.0 Safari/537.36"})
    with urllib.request.urlopen(req, timeout=45) as r:
        d = _j.loads(r.read())
    rr = (d.get("chart") or {}).get("result")
    if not rr:
        raise RuntimeError(f"{sym}: 无 result")
    rr = rr[0]
    ts = rr.get("timestamp") or []
    cl = ((rr.get("indicators") or {}).get("quote") or [{}])[0].get("close") or []
    out = {}
    for t, c in zip(ts, cl):
        if c is not None:
            out[_dt.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d")] = float(c)
    if not out:
        raise RuntimeError(f"{sym}: 无收盘价")
    return out, (rr.get("meta") or {}).get("currency", "?")


def fetch_wrs_recession_signal():
    """WRS · Walmart 衰退信号 = WMT / 奢侈品指数代理。"""
    from datetime import timezone, timedelta
    _jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": _jst,
           "numerator": {"sym": _WRS_NUM[0], "name": _WRS_NUM[1]},
           "denominator": {"sym": _WRS_DEN[0], "name": _WRS_DEN[1]},
           "is_proxy": True,
           "proxy_note": ("分母为代理：Paulsen 原式用 S&P Global Luxury Index（S&P 专有指数，"
                          "公开源不可得）。此处用追踪该指数的 Amundi UCITS ETF（LUXU.L，USD 计价）"
                          "替代，方向与分位可用，绝对水平与原式不可直接比对。"),
           "series": [], "errors": []}
    try:
        num, c1 = _wrs_yahoo(_WRS_NUM[0])
        den, c2 = _wrs_yahoo(_WRS_DEN[0])
        out["numerator"]["currency"] = c1
        out["denominator"]["currency"] = c2
        if c1 != c2:
            out["errors"].append(f"币种不一致 {c1} vs {c2}，比值含汇率噪声")
        common = sorted(set(num) & set(den))
        if len(common) < 100:
            out["status"] = "error"
            out["errors"].append(f"共同交易日仅 {len(common)} 天，不足以成图")
            return out
        ser = [{"d": d, "v": round(num[d] / den[d], 6)} for d in common]
        vals = [x["v"] for x in ser]
        cur = vals[-1]
        hi, lo = max(vals), min(vals)
        pct = round(sum(1 for v in vals if v <= cur) / len(vals) * 100, 1)
        # 峰值日期(诚实标注: 当前未必是新高)
        peak_i = max(range(len(vals)), key=lambda i: vals[i])
        # 近一年/近三月变化
        def _ago(n):
            return ser[-n]["v"] if len(ser) > n else None
        y1, m3 = _ago(252), _ago(63)
        out.update({
            "latest_date": ser[-1]["d"], "latest": cur,
            "hi": hi, "lo": lo, "pct_rank": pct,
            "peak_date": ser[peak_i]["d"], "peak": vals[peak_i],
            "is_new_high": cur >= hi - 1e-9,
            "chg_3m": round((cur / m3 - 1) * 100, 2) if m3 else None,
            "chg_1y": round((cur / y1 - 1) * 100, 2) if y1 else None,
            "n": len(ser), "start": ser[0]["d"],
            "series": ser[::5],          # 5 日抽样, 控体积
        })
        # 信号灯: 分位越高 = 衰退压力越大
        if pct >= 90:   out["level"], out["level_txt"] = "red", "高位"
        elif pct >= 70: out["level"], out["level_txt"] = "amber", "偏高"
        else:           out["level"], out["level_txt"] = "green", "中低"
    except Exception as e:
        out["status"] = "error"
        out["errors"].append(f"{type(e).__name__}: {str(e)[:100]}")
    return out


# ══════════════════════════════════════════════════════════════════════
# 国债市场微观结构 · Morgan Stanley Exhibit 3/4 对口 (Chao 2026-09-19, 乙方案)
# ══════════════════════════════════════════════════════════════════════
# ★需求: MS 报告 Exhibit 3(BrokerTec 2Y OTR 日内买卖价差) 与
#   Exhibit 4(2Y OTR 周度 DV01 成交量)。原 treasury_stress_panels 里用的是
#   期限溢价/IG OAS/曲线/HY OAS —— 那是<b>价格类风险补偿</b>, 不是流动性微观结构,
#   两者是不同维度。本模块补上真正的「成交量」与「交易成本」维度。
#
# ★Exhibit 4 = 真数据(不再是代理):
#   纽约联储一级交易商调查 FR2004, markets.newyorkfed.org 官方免费 API。
#   实测 702 个周度点(2013-04-03 起)。<b>按剩余期限分桶</b>:
#     PDTRGSC-L2(≤2年, 对应 MS 的 2Y) / -G2L3 / -G3L6 / -G6L7 / -G7L11
#   ⚠口径差: MS 用 <b>DV01 加权</b>成交量(久期风险敞口), NY Fed 是<b>名义金额</b>。
#     方向一致但不是同一个数 —— 必须在图上标注, 不许说成等同。
#   ⚠另: PDPOS* 是 POSITIONS(持仓), PDTR*/PDGSW* 才是 TRANSACTIONS(成交量)。
#     开发时特意区分过, 别拿持仓冒充成交量。
#
# ★Exhibit 3 = 仍是代理, 但比原来贴近:
#   BrokerTec 日内 bid-ask 是 CME 商业产品, 公开源确实没有。
#   本轮排查过并<b>排除</b>: 克利夫兰联储 FSI 流动性分项(LQTYSD678FRBCLE)
#     —— 看着对口, 但 metadata 显示 <b>2016-05-05 起 DISCONTINUED</b>, 不能用。
#   选定两条公开代理:
#     ① OFR FSI Funding 分项(官方, 日更, 2000 年起) = 融资/交易成本压力
#     ② 10Y 日内波幅(^TNX high-low, bp) = 价格发现成本, 价差走阔时同步放大
#   两者都不是 bid-ask 本身, 图上明确标注。

_UST_MICRO_BUCKETS = [
    ("PDTRGSC-L2",    "≤2 年", "★对应 MS Exhibit 4 的 2Y"),
    ("PDTRGSC-G2L3",  "2-3 年", ""),
    ("PDTRGSC-G3L6",  "3-6 年", ""),
    ("PDTRGSC-G6L7",  "6-7 年", ""),
    ("PDTRGSC-G7L11", "7-11 年", ""),
]
_UST_MICRO_TOTAL = ("PDGSWOEXTTOT", "全部合计")


def _nyfed_pd(keyid):
    import json as _j, urllib.request
    u = f"https://markets.newyorkfed.org/api/pd/get/{keyid}.json"
    req = urllib.request.Request(u, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 Chrome/120.0 Safari/537.36"})
    with urllib.request.urlopen(req, timeout=55) as r:
        d = _j.loads(r.read())
    rows = ((d.get("pd") or {}).get("timeseries")) or []
    out = []
    for x in rows:
        v = x.get("value")
        if v in (None, "", "*"):
            continue
        try:
            out.append({"d": x["asofdate"], "v": float(v)})
        except (ValueError, KeyError):
            continue
    return out


def _ofr_funding():
    import json as _j, urllib.request, datetime as _dt
    u = "https://www.financialresearch.gov/financial-stress-index/data/fsi.json"
    req = urllib.request.Request(u, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 Chrome/120.0 Safari/537.36"})
    with urllib.request.urlopen(req, timeout=55) as r:
        d = _j.loads(r.read())
    blk = d.get("Funding") or {}
    pts = blk.get("data") or []
    out = []
    for ms, v in pts:
        if v is None:
            continue
        out.append({"d": _dt.datetime.utcfromtimestamp(ms / 1000).strftime("%Y-%m-%d"),
                    "v": float(v)})
    return out


def _tnx_intraday_range(rng="3y"):
    """10Y 日内波幅(bp) = (high - low) * 100。价差走阔时同步放大。"""
    import json as _j, urllib.request, datetime as _dt
    u = (f"https://query1.finance.yahoo.com/v8/finance/chart/%5ETNX"
         f"?range={rng}&interval=1d")
    req = urllib.request.Request(u, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 Chrome/120.0 Safari/537.36"})
    with urllib.request.urlopen(req, timeout=50) as r:
        d = _j.loads(r.read())
    rr = d["chart"]["result"][0]
    ts = rr.get("timestamp") or []
    q = (rr.get("indicators") or {}).get("quote") or [{}]
    hi, lo = q[0].get("high") or [], q[0].get("low") or []
    out = []
    for t, h, l in zip(ts, hi, lo):
        if h is None or l is None:
            continue
        out.append({"d": _dt.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d"),
                    "v": round((h - l) * 100, 2)})
    return out


def fetch_ust_microstructure(years=3):
    """国债市场微观结构: 成交量(Exhibit 4 真数据) + 交易成本代理(Exhibit 3)。"""
    from datetime import timezone, timedelta
    _jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    cutoff = (datetime.date.today() - datetime.timedelta(days=int(years * 365.25))).isoformat()
    out = {"status": "ok", "as_of": _jst, "years": years,
           "volume": {"buckets": [], "total": None, "is_proxy": False,
                      "caliber_note": ("MS Exhibit 4 用 <b>DV01 加权</b>成交量（久期风险敞口），"
                                       "本图为纽约联储 FR2004 的<b>名义金额</b>成交量。"
                                       "方向一致，但不是同一个数，不可直接比对数值。"),
                      "source": "纽约联储一级交易商调查 FR2004（markets.newyorkfed.org 官方 API）"},
           "cost": {"series": [], "is_proxy": True,
                    "proxy_note": ("MS Exhibit 3 用 BrokerTec 日内 bid-ask 价差（CME 商业数据，"
                                   "公开源不可得）。此处用两条公开代理：OFR 金融压力指数 Funding 分项"
                                   "（官方日更）与 10Y 日内波幅（价格发现成本）。"
                                   "均非 bid-ask 本身，只反映交易成本方向。"),
                    "excluded_note": ("已排查并排除：克利夫兰联储 FSI 流动性利差分项"
                                      "（LQTYSD678FRBCLE）—— 看似对口，但 2016-05-05 起已停更。"),
                    "source": "OFR Financial Stress Index · Yahoo Finance ^TNX"},
           "errors": []}

    # ── Exhibit 4: 分桶成交量(真数据) ──
    try:
        for kid, name, tag in _UST_MICRO_BUCKETS:
            s = [p for p in _nyfed_pd(kid) if p["d"] >= cutoff]
            if not s:
                out["errors"].append(f"NYFed/{name}: 窗口内无数据")
                continue
            last = s[-1]["v"]
            vals = [p["v"] for p in s]
            yr = vals[-52:] if len(vals) >= 52 else vals
            out["volume"]["buckets"].append({
                "key": kid, "name": name, "tag": tag,
                "latest": last, "latest_date": s[-1]["d"],
                "avg_52w": round(sum(yr) / len(yr), 1),
                "vs_avg": round((last / (sum(yr) / len(yr)) - 1) * 100, 1),
                "max": max(vals), "min": min(vals), "n": len(s),
                "series": s,
            })
        tot = [p for p in _nyfed_pd(_UST_MICRO_TOTAL[0]) if p["d"] >= cutoff]
        if tot:
            vals = [p["v"] for p in tot]
            yr = vals[-52:] if len(vals) >= 52 else vals
            out["volume"]["total"] = {
                "name": _UST_MICRO_TOTAL[1], "latest": tot[-1]["v"],
                "latest_date": tot[-1]["d"], "avg_52w": round(sum(yr) / len(yr), 1),
                "vs_avg": round((tot[-1]["v"] / (sum(yr) / len(yr)) - 1) * 100, 1),
                "n": len(tot), "series": tot,
            }
    except Exception as e:
        out["errors"].append(f"NYFed 成交量: {type(e).__name__} {str(e)[:70]}")

    # ── Exhibit 3: 交易成本代理 ──
    for label, fn, unit, desc in [
        ("OFR Funding 压力", _ofr_funding, "", "OFR 金融压力指数融资分项（官方，日更）"),
        ("10Y 日内波幅", _tnx_intraday_range, "bp", "10 年期收益率日内高低差（价格发现成本）"),
    ]:
        try:
            s = [p for p in fn() if p["d"] >= cutoff]
            if not s:
                out["errors"].append(f"{label}: 窗口内无数据"); continue
            vals = [p["v"] for p in s]
            cur = vals[-1]
            pct = round(sum(1 for v in vals if v <= cur) / len(vals) * 100, 1)
            out["cost"]["series"].append({
                "name": label, "unit": unit, "desc": desc,
                "latest": cur, "latest_date": s[-1]["d"], "pct_rank": pct,
                "max": max(vals), "min": min(vals), "n": len(s),
                "series": s[::max(1, len(s) // 400)],
            })
        except Exception as e:
            out["errors"].append(f"{label}: {type(e).__name__} {str(e)[:60]}")

    if not out["volume"]["buckets"] and not out["cost"]["series"]:
        out["status"] = "error"
    elif out["errors"]:
        out["status"] = "ok(部分失败)"
    return out


# ══════════════════════════════════════════════════════════════════════
# 美元指数 vs 黄金 双轴图 + 大事件标注 (Chao 2026-09-19, 丙方案)
# ══════════════════════════════════════════════════════════════════════
# ★需求: 对齐中文财经视频里那张 TradingView 叠加图 —— 金价与 DXY 同图双轴,
#   并标注年初至今影响两者的大事件。
#
# ★事件从哪来(这是本模块最容易造假的部分, 纪律必须写死):
#   <b>绝不凭 LLM 记忆写 2026 年发生了什么。</b> 两条来源:
#   ① 日历事件: FOMC 议息日等已知日程 —— 联网核实后落盘 data/dxy_gold_events.json
#   ② 异动事件: 先用<b>真实价格</b>算出异动日(|日收益| >= 1.8σ, σ 取两年),
#      再对这些日子联网检索当天新闻, 抓到出处才写。
#   两者都带 url + confidence(high/medium/none), 前端只渲染有出处的,
#   confidence=none 的显示为「未找到明确触发事件」, 不硬凑。
#
# ★为什么是"价格先行": 让数据指出"这天有事", 而不是我想一个事件再去套价格。
#   这样标注的日期必定对应真实异动, 不会出现"标了个大事但当天其实没动"。
#
# ★更新节奏(Chao 定): 价格<b>日更</b>(随 dashboard 每日构建);
#   事件<b>周更</b>(联网检索成本高且事件不会天天有) —— 事件文件 mtime 超 8 天会在图上提示。

_DG_GOLD = ("GC=F", "COMEX 黄金期货")
_DG_DXY  = ("DX-Y.NYB", "ICE 美元指数 DXY")
_DG_EVENTS_FILE = "data/dxy_gold_events.json"
_DG_SIGMA_TH = 1.8


def _dg_yahoo(sym, rng="2y"):
    import json as _j, urllib.request, datetime as _dt
    u = f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}?range={rng}&interval=1d"
    req = urllib.request.Request(u, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 Chrome/120.0 Safari/537.36"})
    with urllib.request.urlopen(req, timeout=50) as r:
        d = _j.loads(r.read())
    rr = d["chart"]["result"][0]
    ts = rr.get("timestamp") or []
    cl = ((rr.get("indicators") or {}).get("quote") or [{}])[0].get("close") or []
    return {_dt.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d"): float(c)
            for t, c in zip(ts, cl) if c is not None}


def fetch_dxy_gold_events(start="2026-01-01", events_path=None):
    """美元指数 vs 黄金 双轴走势 + 已核实事件标注。"""
    import os as _os, json as _j, statistics as _stat
    from datetime import timezone, timedelta
    _jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": _jst, "start": start,
           "gold": {"sym": _DG_GOLD[0], "name": _DG_GOLD[1]},
           "dxy": {"sym": _DG_DXY[0], "name": _DG_DXY[1]},
           "series": [], "events": [], "anomalies": [],
           "events_stale": False, "events_updated": None,
           "sigma": {}, "errors": [],
           "method_note": ("事件标注不靠记忆：先用真实价格算出异动日（|日收益| ≥ "
                           f"{_DG_SIGMA_TH}σ，σ 取两年日波动），再对这些日子联网检索当天新闻，"
                           "抓到出处才标注。每条事件都带原文链接；核实不到的显示为"
                           "「未找到明确触发事件」，不硬凑。"),
           "cadence_note": "价格每日更新；事件每周重新联网核实一次。"}

    try:
        g = _dg_yahoo(_DG_GOLD[0])
        x = _dg_yahoo(_DG_DXY[0])
    except Exception as e:
        out["status"] = "error"
        out["errors"].append(f"价格抓取失败: {type(e).__name__} {str(e)[:70]}")
        return out

    days = sorted(set(g) & set(x))
    if len(days) < 60:
        out["status"] = "error"
        out["errors"].append(f"共同交易日仅 {len(days)} 天")
        return out

    # 两年 sigma(更稳), 只在窗口内标异动
    def _rets(s):
        return {days[i]: (s[days[i]] / s[days[i-1]] - 1) * 100 for i in range(1, len(days))}
    gr, xr = _rets(g), _rets(x)
    g_sd = _stat.pstdev(list(gr.values())) or 1.0
    x_sd = _stat.pstdev(list(xr.values())) or 1.0
    out["sigma"] = {"gold": round(g_sd, 3), "dxy": round(x_sd, 3), "th": _DG_SIGMA_TH}

    win = [d for d in days if d >= start]
    out["series"] = [{"d": d, "g": round(g[d], 2), "x": round(x[d], 3)} for d in win]

    for d in win[1:]:
        if d not in gr or d not in xr:
            continue
        zg, zx = gr[d] / g_sd, xr[d] / x_sd
        if abs(zg) >= _DG_SIGMA_TH or abs(zx) >= _DG_SIGMA_TH:
            out["anomalies"].append({
                "d": d, "gold_pct": round(gr[d], 2), "dxy_pct": round(xr[d], 2),
                "zg": round(zg, 2), "zx": round(zx, 2),
                "gold": round(g[d], 2), "dxy": round(x[d], 3),
                "score": round(max(abs(zg), abs(zx)), 2),
            })
    out["anomalies"].sort(key=lambda r: r["d"])

    # ── 读已核实事件(离线文件, 由周更任务写入) ──
    p = events_path or _os.path.join(_os.path.dirname(_os.path.dirname(
        _os.path.abspath(__file__))), _DG_EVENTS_FILE)
    if _os.path.exists(p):
        try:
            ev = _j.load(open(p, encoding="utf-8"))
            items = ev.get("events") if isinstance(ev, dict) else ev
            out["events"] = [e for e in (items or []) if e.get("date", "") >= start]
            out["events_updated"] = (ev.get("updated") if isinstance(ev, dict) else None)
            age_days = (datetime.datetime.now().timestamp() - _os.path.getmtime(p)) / 86400
            out["events_stale"] = age_days > 8
            out["events_age_days"] = round(age_days, 1)
        except Exception as e:
            out["errors"].append(f"事件文件解析失败: {type(e).__name__}")
    else:
        out["errors"].append("事件文件不存在（尚未跑过事件核实任务）")

    # 把事件挂到最近交易日 + 标出该日真实涨跌
    byday = {s["d"]: s for s in out["series"]}
    anod = {a["d"]: a for a in out["anomalies"]}
    for e in out["events"]:
        d = e.get("date")
        e["on_chart"] = d in byday
        if d in byday:
            e["gold"] = byday[d]["g"]
            e["dxy"] = byday[d]["x"]
        if d in anod:
            e["gold_pct"] = anod[d]["gold_pct"]
            e["dxy_pct"] = anod[d]["dxy_pct"]
            e["is_anomaly"] = True
        else:
            e["is_anomaly"] = False

    # 未被事件覆盖的异动日(诚实展示"有异动但没查到原因")
    covered = {e.get("date") for e in out["events"] if e.get("confidence") in ("high", "medium")}
    out["uncovered"] = [a for a in out["anomalies"] if a["d"] not in covered]

    if out["errors"]:
        out["status"] = "ok(部分失败)"
    return out


# ══════════════════════════════════════════════════════════════════════
# 黄金 ETF 分地区资金流 (WGC 官方 fsapi) — Chao 2026-09-19, 丙方案
# ══════════════════════════════════════════════════════════════════════
# ★数据源: https://fsapi.gold.org/api/v11/charts/etfv2/revised/flows-chart2
#   这是世界黄金协会(WGC) Goldhub「Gold ETF flows」图表的<b>官方数据后端</b>,
#   免费、无需 key。与 Chao 提供的财经视频截图中那张 WGC 图同源。
#   实测复核: 2026-06-12 = -24.7t、2026-06-26 = -38.2t, 与截图标注逐字吻合。
#
# ★提供: 北美/欧洲/亚洲/其他 四地区独立序列 + 金价(右轴), 双单位(吨/美元),
#   四种频率(周/月/季/年)。周度 454 点, 2018-01 至今。
#
# ★★白银为何没有同类数据(踩坑记录, 别再重复找):
#   1. WGC 只管黄金。`?metal=silver` <b>返回 HTTP 200 但内容与黄金一字不差</b>
#      (参数被静默忽略) —— 这是陷阱, 不验证就会把黄金数据当白银用。
#      判据: 比对两端点最新值完全相等 + 'Gold Price (rhs)' 仍是 ~4400/oz。
#   2. Silver Institute 无同类周度分地区 API(页面无数据文件/无 api 路径)。
#   3. iShares(SLV)/Sprott(PSLV)/abrdn(SIVR) 官方持仓接口均返回 HTML 壳或 404;
#      SPDR 只发 PDF; Yahoo quoteSummary 全线 401 拿不到 AUM。
#   结论: 白银的权威周度分地区 flows 基本只在彭博/路透终端。本模块<b>只做黄金</b>,
#   白银在 dashboard 上显式说明「无权威免费源」, 不用代理数据冒充。

_WGC_FLOWS = "https://fsapi.gold.org/api/v11/charts/etfv2/revised/flows-chart2"
_WGC_HOLD  = "https://fsapi.gold.org/api/v11/charts/etfv2/revised/holdings-chart2"
_WGC_REGIONS = ("North America", "Europe", "Asia", "Other")
_WGC_CN = {"North America": "北美", "Europe": "欧洲", "Asia": "亚洲",
           "Other": "其他", "Gold Price (rhs)": "金价"}


def _wgc_get(url):
    import json as _j, urllib.request, ssl as _ssl
    ctx = _ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = _ssl.CERT_NONE
    req = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 Chrome/120.0 Safari/537.36",
        "Accept": "application/json,*/*",
        "Referer": "https://www.gold.org/"})
    with urllib.request.urlopen(req, timeout=45, context=ctx) as r:
        return _j.loads(r.read())


def fetch_gold_etf_flows_by_region(weeks=52):
    """黄金 ETF 分地区周度资金流(WGC 官方)。"""
    import datetime as _dt
    from datetime import timezone, timedelta
    _jst = _dt.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": _jst, "regions": [], "weeks": [],
           "as_of_data": None, "totals": {}, "errors": [],
           "silver_note": ("白银无权威免费的分地区周度 flows：WGC 只覆盖黄金"
                           "（其 ?metal=silver 参数返回 200 但内容仍是黄金数据）；"
                           "Silver Institute 无同类 API；各发行商(iShares/Sprott/abrdn/SPDR)"
                           "持仓接口返回 HTML 壳、PDF 或 401。故本板块只做黄金，"
                           "不用代理数据冒充白银。"),
           "source": "World Gold Council (Goldhub) 官方 fsapi",
           "source_url": "https://www.gold.org/goldhub/data/gold-etfs-holdings-and-flows"}

    try:
        d = _wgc_get(_WGC_FLOWS)
    except Exception as e:
        out["status"] = "error"
        out["errors"].append(f"WGC flows 抓取失败: {type(e).__name__} {str(e)[:70]}")
        return out

    try:
        cd = d["chartData"]
        out["as_of_data"] = cd.get("asOfDate")
        ser = cd["data"]["Weekly"]["series"]["tonnes"]
    except Exception as e:
        out["status"] = "error"
        out["errors"].append(f"WGC flows 结构异常: {type(e).__name__}")
        return out

    def _d(ms):
        return _dt.datetime.utcfromtimestamp(ms / 1000).strftime("%Y-%m-%d")

    by = {}
    for s in ser:
        by[s.get("name")] = {_d(p[0]): p[1] for p in (s.get("data") or []) if p and p[0]}

    missing = [r for r in _WGC_REGIONS if r not in by]
    if missing:
        out["errors"].append(f"缺地区序列: {missing}")
    dates = sorted(by.get("North America") or by.get(list(by)[0]) or {})
    if not dates:
        out["status"] = "error"
        out["errors"].append("无日期序列")
        return out

    win = dates[-weeks:]
    px = by.get("Gold Price (rhs)", {})
    rows = []
    for dt in win:
        vals = {r: round(by.get(r, {}).get(dt, 0.0), 2) for r in _WGC_REGIONS}
        rows.append({"d": dt, **{_WGC_CN[r]: vals[r] for r in _WGC_REGIONS},
                     "total": round(sum(vals.values()), 2),
                     "px": round(px.get(dt, 0.0), 2)})
    out["weeks"] = rows
    out["regions"] = [_WGC_CN[r] for r in _WGC_REGIONS]

    # 汇总: 最新周 / 近4周 / 近13周 / YTD
    def _sum(n=None, ytd=False):
        sel = [r for r in rows if r["d"][:4] == rows[-1]["d"][:4]] if ytd else rows[-n:]
        return {k: round(sum(r[k] for r in sel), 1)
                for k in list(_WGC_CN[x] for x in _WGC_REGIONS) + ["total"]}
    out["totals"] = {"latest": {k: rows[-1][k] for k in
                                list(_WGC_CN[x] for x in _WGC_REGIONS) + ["total"]},
                     "w4": _sum(4), "w13": _sum(13), "ytd": _sum(ytd=True)}

    # 极值: 本年最大流入/流出周 + 全序列排名(用于判断"创纪录"这类说法)
    allr = [{"d": dt, "total": round(sum(by.get(r, {}).get(dt, 0.0)
                                         for r in _WGC_REGIONS), 2)} for dt in dates]
    srt = sorted(allr, key=lambda r: r["total"])
    out["extremes"] = {
        "worst_all": srt[:5], "best_all": srt[-5:][::-1],
        "latest_rank_out": sum(1 for r in allr if r["total"] < allr[-1]["total"]) + 1,
        "n_all": len(allr), "hist_start": dates[0],
    }

    # 持仓总量(另一端点, 拿不到不致命)
    try:
        h = _wgc_get(_WGC_HOLD)
        hs = h["chartData"]["data"]["Weekly"]
        ht = hs.get("tonnes") or (hs.get("series") or {}).get("tonnes")
        if isinstance(ht, list) and ht:
            tot = {}
            for s in ht:
                for p in (s.get("data") or []):
                    if p and p[0]:
                        tot[_d(p[0])] = tot.get(_d(p[0]), 0.0) + (p[1] or 0.0)
            if tot:
                k = max(tot)
                out["holdings_total_t"] = round(tot[k], 1)
                out["holdings_as_of"] = k
    except Exception as e:
        out["errors"].append(f"持仓端点跳过: {type(e).__name__}")

    if out["errors"]:
        out["status"] = "ok(部分失败)"
    return out


# ══════════════════════════════════════════════════════════════════════
# 贵金属期权波动率: 25D 偏斜 + ATM IV 期限结构 (Chao 2026-09-21)
# ══════════════════════════════════════════════════════════════════════
# ★数据源: CBOE cdn delayed quotes —— 逐合约含 iv + delta, 实测真数据。
#   https://cdn.cboe.com/api/global/delayed_quotes/options/{TICKER}.json
#
# ★两张图对应 Chao 提供的截图:
#   图A 25D 偏斜 = 25-delta put IV − 25-delta call IV
#       正值=看跌期权更贵=偏空; 负值=偏多。用 delta <b>线性插值</b>到 |delta|=0.25,
#       比取最近点准(合约 delta 不会正好等于 0.25)。
#   图B ATM IV 期限结构 = 各到期日 |delta|=0.5 处的 IV, 排成 1M/3M/6M/12M/18M/24M。
#       近月 > 远月 = 倒挂(市场预期近期震荡); 反之为正常上倾。
#
# ★★历史限制(重要, 别再重复找):
#   CBOE 只有<b>当下快照</b>。`?date=YYYY-MM-DD` 返回 200 但 timestamp 与现价
#   和不带参数时完全相同 —— 参数被静默忽略, 不是历史数据。
#   /historical_data/ 与 /historical_options/ 均 403。
#   Alpha Vantage: HISTORICAL_OPTIONS 是<b>付费端点</b>(免费 key 返回
#   "This is a premium endpoint"); REALTIME_OPTIONS 免费档返回的是
#   <b>占位假数据</b>(contractID=XXYYZZ..., expiration=2099-99-99, 价格全 100) —— 陷阱。
#   ORATS 401, Nasdaq Data Link 403。
#   → 结论: 偏斜与期限结构<b>无免费历史</b>, 只能每日快照自行积累(落 Notion DB_OPTVOL)。
#   → 可回填一年的只有: FRED GVZCLS(CBOE 黄金ETF波动率指数, 官方) 与
#     GLD/SLV 30日已实现波动率(自算, 字段名带 RV 标明非隐含波动率)。
#   ★白银官方 VXSLVCLS 已于 2022-02 DISCONTINUED(近一年 0 点), 不可用。

_PMV_TICKERS = ["GLD", "SLV", "IAU"]   # SGOL 到期日仅 4 个, 无法构建期限结构, 排除
_PMV_TENORS = [("1M", 30), ("3M", 91), ("6M", 182), ("12M", 365), ("18M", 548), ("24M", 730)]


def _pmv_chain(tk):
    import json as _j, re as _re, ssl as _ssl, urllib.request, datetime as _dt
    ctx = _ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = _ssl.CERT_NONE
    u = f"https://cdn.cboe.com/api/global/delayed_quotes/options/{tk}.json"
    req = urllib.request.Request(u, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 Chrome/120.0 Safari/537.36"})
    with urllib.request.urlopen(req, timeout=50, context=ctx) as r:
        d = _j.loads(r.read())
    dd = d.get("data", {})
    rows = []
    for o in dd.get("options", []):
        m = _re.match(r"^[A-Z]+(\d{6})([CP])(\d{8})$", o.get("option", ""))
        if not m:
            continue
        e = m.group(1)
        try:
            exp = _dt.date(2000 + int(e[:2]), int(e[2:4]), int(e[4:6]))
        except Exception:
            continue
        iv, dl = o.get("iv"), o.get("delta")
        if not iv or dl is None:
            continue
        rows.append({"exp": exp, "cp": m.group(2), "k": int(m.group(3)) / 1000.0,
                     "iv": float(iv), "d": float(dl), "oi": o.get("open_interest", 0)})
    return dd.get("current_price"), d.get("timestamp"), rows


def _pmv_interp(arr, target):
    """在 |delta|=target 处线性插值 IV。"""
    pts = sorted([(abs(x["d"]), x["iv"]) for x in arr if 0.01 < abs(x["d"]) < 0.99])
    if len(pts) < 2:
        return None
    for i in range(len(pts) - 1):
        (d0, v0), (d1, v1) = pts[i], pts[i + 1]
        if d0 <= target <= d1:
            return v0 if d1 == d0 else v0 + (v1 - v0) * (target - d0) / (d1 - d0)
    return min(pts, key=lambda p: abs(p[0] - target))[1]


def fetch_pm_option_vol():
    """贵金属期权: 25D 偏斜 + ATM IV 期限结构(CBOE 真数据, 当下快照)。"""
    import datetime as _dt
    from datetime import timezone, timedelta
    from collections import defaultdict
    _jst = _dt.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    today = _dt.date.today()
    out = {"status": "ok", "as_of": _jst, "date": today.isoformat(),
           "tickers": {}, "errors": [],
           "source": "CBOE delayed quotes (逐合约 IV + delta)",
           "hist_note": ("偏斜与期限结构<b>无免费历史源</b>：CBOE 只提供当下快照"
                         "（?date= 参数返回 200 但内容不变，实为忽略）；"
                         "Alpha Vantage 历史期权为付费端点，其免费实时接口返回占位假数据。"
                         "故这两项自 2026-09-21 起每日快照累积；"
                         "可回溯一年的只有 GVZ（FRED 官方）与 30 日已实现波动率（自算）。"),
           "method_note": ("25D 偏斜 = 25-delta put IV − 25-delta call IV，"
                           "正值＝看跌期权更贵＝偏空。ATM 取 |delta|=0.5。"
                           "两者均按 delta <b>线性插值</b>求得，非取最近合约。")}

    for tk in _PMV_TICKERS:
        try:
            spot, ts, rows = _pmv_chain(tk)
        except Exception as e:
            out["errors"].append(f"{tk}: {type(e).__name__} {str(e)[:60]}")
            continue
        if not rows:
            out["errors"].append(f"{tk}: 无有效合约")
            continue
        by = defaultdict(list)
        for r in rows:
            by[r["exp"]].append(r)
        term, skew = [], []
        for lab, dte in _PMV_TENORS:
            cand = [(abs((e - today).days - dte), e) for e in by if (e - today).days >= 5]
            if not cand:
                continue
            cand.sort()
            e = cand[0][1]
            arr = by[e]
            d = (e - today).days
            cs = [x for x in arr if x["cp"] == "C"]
            ps = [x for x in arr if x["cp"] == "P"]
            if len(cs) < 3 or len(ps) < 3:
                continue
            ac, ap = _pmv_interp(cs, 0.50), _pmv_interp(ps, 0.50)
            c25, p25 = _pmv_interp(cs, 0.25), _pmv_interp(ps, 0.25)
            atm = (ac + ap) / 2 * 100 if (ac and ap) else None
            sk = (p25 - c25) * 100 if (p25 and c25) else None
            term.append({"t": lab, "dte": d, "exp": e.isoformat(),
                         "atm": round(atm, 2) if atm else None})
            skew.append({"t": lab, "dte": d,
                         "put": round(p25 * 100, 2) if p25 else None,
                         "call": round(c25 * 100, 2) if c25 else None,
                         "skew": round(sk, 2) if sk is not None else None})
        # ★去重: 到期日不足时多个 tenor 会落到同一到期 —— 不可当成不同期限展示
        seen, term2, skew2 = set(), [], []
        for a, b in zip(term, skew):
            if a["exp"] in seen:
                continue
            seen.add(a["exp"])
            term2.append(a)
            skew2.append(b)
        dropped = len(term) - len(term2)
        vals = [x["atm"] for x in term2 if x["atm"]]
        shape = None
        if len(vals) >= 3:
            shape = "倒挂" if vals[0] > vals[-1] else "正常上倾"
        out["tickers"][tk] = {
            "spot": spot, "quote_ts": ts, "term": term2, "skew": skew2,
            "shape": shape, "n_exp": len(by), "n_contract": len(rows),
            "dup_dropped": dropped,
            "dup_note": (f"有 {dropped} 个期限因到期日不足而落在同一到期，已剔除"
                         if dropped else ""),
        }

    if not out["tickers"]:
        out["status"] = "error"
    elif out["errors"]:
        out["status"] = "ok(部分失败)"
    return out


# ══════════════════════════════════════════════════════════════════════
# COMEX 期货期权持仓 (Chao 2026-09-21) —— 注意: 这不是 IV
# ══════════════════════════════════════════════════════════════════════
# ★★为什么没有 IV(踩坑记录, 别再重复找):
#   COMEX 期货期权的隐含波动率是 <b>CME 的商业数据产品</b>, 免费源全封:
#     CME 官网 API 403 / CME ftp settlement 403 / CME datamine 需订阅 /
#     Barchart 403 / Yahoo 期货期权链 401(v7+v6 × 三种 symbol 全试) /
#     CBOE 只覆盖股票与 ETF 期权, 不含期货期权 /
#     ScraperAPI 能绕 IP 封锁但 CME 把请求导向中文营销页, 不是 API 响应。
#   → 绝不拿 GLD/SLV 的 ETF 期权 IV 冒充 "COMEX 期货期权 IV"(口径完全不同)。
#
# ★本模块给的是<b>持仓方向</b>, 回答「Managed Money 在期权上加多还是加空」:
#   期权部分 = Disagg(Futures AND Options) − Disagg(Futures Only)
#   两者同为 CFTC 官方 Disaggregated 口径, 含相同的 m_money 字段, 可相减。
#   ★已在 <b>170 周(2023-06 至今, 3年+)</b> 长样本验证: 金银两市场
#     期权 OI 差均 <b>0 次为负</b> —— 若口径不匹配必然出现负数, 故口径成立。
#   (注: 6dca-aqww 是 Legacy 口径 noncomm/comm, 不含 m_money, 不能与 Disagg 混减)
#
# ★单位换算: COMEX 黄金 1 张 = 100 金衡盎司; 白银 1 张 = 5000 盎司;
#   1 吨 = 32150.7 金衡盎司。吨口径用于对齐市面上 WGC/行研图表的「吨」表述。

_CXO_MKTS = {"gold": ("GOLD - COMMODITY EXCHANGE INC.", 100.0, "金"),
             "silver": ("SILVER - COMMODITY EXCHANGE INC.", 5000.0, "银")}
_CXO_FAO = "kh3c-gbw2"   # Disaggregated Futures AND Options
_CXO_FO  = "72hh-3qpy"   # Disaggregated Futures Only
_CXO_OZ_PER_TONNE = 32150.7


def _cxo_pull(ds, mkt, limit=170):
    import json as _j, ssl as _ssl, urllib.request, urllib.parse
    ctx = _ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = _ssl.CERT_NONE
    u = (f"https://publicreporting.cftc.gov/resource/{ds}.json?$limit={limit}"
         f"&$order=report_date_as_yyyy_mm_dd%20DESC"
         f"&market_and_exchange_names={urllib.parse.quote(mkt)}")
    req = urllib.request.Request(u, headers={
        "User-Agent": "Mozilla/5.0 Chrome/120.0", "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=70, context=ctx) as r:
        rows = _j.loads(r.read())
    out = {}
    for x in rows:
        d = (x.get("report_date_as_yyyy_mm_dd") or "")[:10]
        try:
            out[d] = {
                "l": float(x["m_money_positions_long_all"]),
                "s": float(x["m_money_positions_short_all"]),
                "oi": float(x["open_interest_all"]),
            }
        except Exception:
            pass
    return out


def fetch_comex_option_positions(weeks=104):
    """COMEX 金/银 期货期权持仓(Managed Money 方向)。"""
    import datetime as _dt
    from datetime import timezone, timedelta
    _jst = _dt.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": _jst, "metals": {}, "errors": [],
           "not_iv_note": ("本板块是<b>持仓方向</b>不是隐含波动率。COMEX 期货期权的 IV 属 "
                           "CME 商业数据，免费源全部封闭（CME API/ftp 403、Barchart 403、"
                           "Yahoo 期货期权链 401、CBOE 不覆盖期货期权），"
                           "故不以 ETF 期权 IV 冒充。"),
           "method_note": ("期权部分 = CFTC Disaggregated（期货+期权）− （仅期货），"
                           "两者同口径同字段。已在 170 周（3 年+）样本验证："
                           "金银期权 OI 差均 0 次为负，口径成立。"),
           "source": "CFTC 官方 Socrata（publicreporting.cftc.gov）"}

    for key, (mkt, oz, cn) in _CXO_MKTS.items():
        try:
            fao = _cxo_pull(_CXO_FAO, mkt, weeks)
            fo = _cxo_pull(_CXO_FO, mkt, weeks)
        except Exception as e:
            out["errors"].append(f"{key}: {type(e).__name__} {str(e)[:60]}")
            continue
        common = sorted(set(fao) & set(fo))
        if len(common) < 8:
            out["errors"].append(f"{key}: 共同周仅 {len(common)}")
            continue
        series, neg = [], 0
        for d in common:
            a, b = fao[d], fo[d]
            net_ao = a["l"] - a["s"]
            net_fo = b["l"] - b["s"]
            oi_opt = a["oi"] - b["oi"]
            if oi_opt < 0:
                neg += 1
            series.append({
                "d": d,
                "net_fut": round(net_fo),
                "net_opt": round(net_ao - net_fo),
                "oi_opt": round(oi_opt),
                "oi_all": round(a["oi"]),
                "tonnes": round(net_fo * oz / _CXO_OZ_PER_TONNE, 1),
            })
        last = series[-1]
        opts = [x["net_opt"] for x in series]
        out["metals"][key] = {
            "cn": cn, "series": series, "latest": last,
            "neg_weeks": neg, "n": len(series),
            "dir": "净多" if last["net_opt"] > 0 else ("净空" if last["net_opt"] < 0 else "中性"),
            "opt_min": min(opts), "opt_max": max(opts),
            "pct": round(sum(1 for v in opts if v < last["net_opt"]) / len(opts) * 100, 1),
            "caveat": (f"期权 OI 差有 {neg} 周为负，口径可能不匹配" if neg else ""),
        }

    if not out["metals"]:
        out["status"] = "error"
    elif out["errors"]:
        out["status"] = "ok(部分失败)"

    # ★2026-10-06 (Chao「在这两个图上覆盖一下金银的 COMEX」): 每个品种叠加 COMEX 连续合约日线收盘价(Yahoo GC=F / SI=F)。
    #   抓不到读上次缓存并标「缓存」; 缓存也没有则不叠加并标「未获取」, 绝不编数。价格失败不影响持仓数据。
    import os as _os
    _cache = _os.path.join(_os.path.dirname(__file__), "..", "data", "comex_optpos_price_cache.json")

    def _load_px_cache():
        try:
            with open(_cache, encoding="utf-8") as f:
                return json.load(f)
        except Exception:  # noqa: BLE001
            return {}

    def _yahoo_px(sym):
        import urllib.request as _ur
        u = f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}?range=3y&interval=1d"
        try:
            r = _ur.urlopen(_ur.Request(u, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                                                   "AppleWebKit/537.36 Chrome/120 Safari/537.36"}), timeout=45)
            res = json.loads(r.read().decode("utf-8", "ignore"))["chart"]["result"][0]
            cl = res["indicators"]["quote"][0]["close"]
            return [[_dt.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d"), round(float(c), 2)]
                    for t, c in zip(res.get("timestamp") or [], cl) if c is not None]
        except Exception:  # noqa: BLE001
            return []

    _pc = _load_px_cache()
    _fresh = {}
    for key, sym in (("gold", "GC=F"), ("silver", "SI=F")):
        if key not in out["metals"]:
            continue
        px = _yahoo_px(sym)
        if len(px) >= 20:
            _fresh[key] = px
            out["metals"][key]["price"] = px
            out["metals"][key]["price_status"] = "ok"
        elif (_pc.get(key) or []):
            out["metals"][key]["price"] = _pc[key]
            out["metals"][key]["price_status"] = "缓存"
            out["errors"].append(f"{key} 价格抓取失败, 用缓存(最新 {_pc[key][-1][0]})")
        else:
            out["metals"][key]["price"] = []
            out["metals"][key]["price_status"] = "未获取"
            out["errors"].append(f"{key} 价格未获取, 图上不叠加")
        out["metals"][key]["price_source"] = f"Yahoo Finance · {sym} COMEX 连续合约 日收盘(美元/盎司)"
    if _fresh:
        try:
            _merged = dict(_pc); _merged.update(_fresh)
            with open(_cache, "w", encoding="utf-8") as f:
                json.dump(_merged, f, ensure_ascii=False)
        except Exception:  # noqa: BLE001
            pass
    return out


# ══════════════════════════════════════════════════════════════════════
# 全球货币供应量(五大经济体聚合) + 加速度  (Chao 2026-09-21)
# ══════════════════════════════════════════════════════════════════════
# ★★口径声明(务必随图展示):「全球货币供应」<b>没有官方单一口径</b>。
#   IMF/BIS/世界银行都不发布这个数。本模块是**自行聚合五大经济体、
#   按当期汇率折美元**的构造值, 不是任何机构的官方发布。
#
# ★数据源(均已实测可用):
#   美国 M2    FRED M2SL                      月  1959-01~  十亿美元
#   欧元区 M3  ECB SDMX BSI                   月  1996-01~  百万欧元
#   日本 M2    BOJ MD02 第9列(存量列)          月  2003-04~  亿日元
#   英国 M4    BoE IADB LPMAUYM               月  1996-01~  百万英镑
#   中国 M2    东方财富 RPT_ECONOMY_CURRENCY   月  2008-01~  亿元
#   汇率       FRED DEXUSEU/DEXJPUS/DEXUSUK/DEXCHUS
#
# ★★三个历史缺口是<b>数据本身的边界</b>, 不是抓漏(已逐一查证):
#   ① 欧元区 1996-1998 无数据 —— <b>欧元 1999-01 才诞生</b>, 之前不存在 EUR 汇率。
#      任何号称"1996 年全球 M2"的图, 这一段要么用马克/法郎折算要么是估的。
#   ② 日本 1996-2003 —— BOJ 该表 1980 起的列是<b>同比增速%</b>,
#      <b>存量列 2003-04 才开始</b>。★第一次解析抓到的 "2"、"1.2" 就是增速,
#      靠数量级检查(M2 应在 1.2e7 亿円)才发现, 否则会把增速当存量画成图。
#   ③ 中国 1996-2007 —— 东财只到 2008; 世界银行有年度值(1996: 76,095 亿元)可插值。
#   → 图上按时段标明"本段包含哪几国", 绝不假装 30 年都完整。
#
# ★加速度 = 同比增速的变化率(二阶导): accel_t = yoy_t − yoy_{t-12}
#   正值 = 放水在加速; 负值 = 放水在减速(即便 yoy 仍为正)。

_GMS_START = "1996-01"


def _gms_fred(sid, key, start="1996-01-01"):
    import json as _j, ssl as _ssl, urllib.request
    ctx = _ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = _ssl.CERT_NONE
    u = (f"https://api.stlouisfed.org/fred/series/observations?series_id={sid}"
         f"&api_key={key}&file_type=json&observation_start={start}")
    with urllib.request.urlopen(u, timeout=60, context=ctx) as r:
        d = _j.loads(r.read())
    return {o["date"][:7]: float(o["value"])
            for o in d.get("observations", []) if o["value"] not in (".", "")}


def _gms_get(u, t=70):
    import ssl as _ssl, urllib.request
    ctx = _ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = _ssl.CERT_NONE
    req = urllib.request.Request(u, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 Chrome/120.0 Safari/537.36"})
    with urllib.request.urlopen(req, timeout=t, context=ctx) as r:
        return r.read()


def fetch_global_money_supply():
    """五大经济体 M2/M3/M4 聚合(折美元) + 同比 + 加速度。"""
    import os as _os, re as _re, csv as _csv, io as _io, json as _j
    import html as _html, datetime as _dt
    from datetime import timezone, timedelta
    _jst = _dt.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")

    out = {"status": "ok", "as_of": _jst, "countries": {}, "series": [],
           "errors": [], "coverage": [],
           "caliber_note": ("「全球货币供应」无官方单一口径 —— IMF/BIS/世界银行均不发布该数。"
                            "本图为<b>自行聚合五大经济体、按当期汇率折美元</b>的构造值，"
                            "非任何机构官方发布。"),
           "gap_note": ("三段历史缺口属数据本身边界：① 欧元区 1996-1998 —— "
                        "<b>欧元 1999 年才诞生</b>，此前无 EUR 汇率；"
                        "② 日本 1996-2003 —— BOJ 表中 1980 年起的是同比增速列，"
                        "M2 <b>存量列 2003-04 才开始</b>；"
                        "③ 中国 1996-2007 —— 月度源起于 2008。"
                        "图中按时段标明覆盖国数，未做无依据外推。"),
           "accel_note": ("加速度 = 同比增速的年度变化（yoyₜ − yoyₜ₋₁₂）。"
                          "正值＝放水在提速；负值＝放水在减速（即便同比仍为正）。")}

    key = ""
    for _p in ("/home/user/Projects/Eco-and-Volatility-Checker/.env", ".env"):
        try:
            for ln in open(_p):
                if ln.startswith("FRED_API_KEY"):
                    key = ln.split("=", 1)[1].strip().strip('"\'')
            if key:
                break
        except Exception:
            pass
    if not key:
        out["status"] = "error"
        out["errors"].append("FRED_API_KEY 缺失")
        return out

    src = {}
    # ① 美国 M2 (十亿美元 → 万亿美元 ÷1000)
    try:
        src["US"] = (_gms_fred("M2SL", key), 1e-3, None, "美国 M2", "FRED M2SL")
    except Exception as e:
        out["errors"].append(f"US: {type(e).__name__}")

    # ② 欧元区 M3 (百万欧元 → 万亿欧元 ÷1e6)
    try:
        b = _gms_get("https://data-api.ecb.europa.eu/service/data/BSI/"
                     "M.U2.Y.V.M30.X.1.U2.2300.Z01.E?format=csvdata&startPeriod=1996-01")
        ez = {r["TIME_PERIOD"]: float(r["OBS_VALUE"])
              for r in _csv.DictReader(_io.StringIO(b.decode("utf-8", "ignore")))
              if r.get("OBS_VALUE")}
        src["EZ"] = (ez, 1e-6, ("EUR", False), "欧元区 M3", "ECB SDMX")
    except Exception as e:
        out["errors"].append(f"EZ: {type(e).__name__}")

    # ③ 日本 M2 (亿日元 → 万亿日元 ÷1e4) —— ★第9列是存量, 前面几列是增速%
    try:
        t = _gms_get("https://www.stat-search.boj.or.jp/ssi/mtshtml/"
                     "md02_m_1_en.html").decode("utf-8", "ignore")
        jp = {}
        for row in _re.findall(r"<tr[^>]*>(.*?)</tr>", t, _re.S):
            cells = [_re.sub(r"<[^>]+>", "", _html.unescape(c)).strip()
                     for c in _re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", row, _re.S)]
            if cells and _re.match(r"^\d{4}/\s*\d{1,2}$", cells[0]) and len(cells) > 9:
                y, m = _re.match(r"^(\d{4})/\s*(\d{1,2})$", cells[0]).groups()
                v = cells[9].replace(",", "")
                try:
                    fv = float(v)
                except Exception:
                    continue
                # ★数量级护栏: M2 存量应 > 1e6 亿円; 小于说明抓到了增速列
                if fv > 1e6:
                    jp[f"{y}-{int(m):02d}"] = fv
        if not jp:
            raise ValueError("BOJ 第9列未解析出存量(可能页面结构变化)")
        src["JP"] = (jp, 1e-4, ("JPY", True), "日本 M2", "BOJ MD02")
    except Exception as e:
        out["errors"].append(f"JP: {type(e).__name__} {str(e)[:50]}")

    # ④ 英国 M4 (百万英镑 → 万亿英镑 ÷1e6)
    try:
        b = _gms_get("https://www.bankofengland.co.uk/boeapps/iadb/fromshowcolumns.asp?"
                     "csv.x=yes&Datefrom=01/Jan/1996&Dateto=01/Dec/2026"
                     "&SeriesCodes=LPMAUYM&CSVF=TN&UsingCodes=Y&VPD=Y&VFD=N")
        gb = {}
        for r in _csv.DictReader(_io.StringIO(b.decode("utf-8", "ignore"))):
            try:
                d0 = _dt.datetime.strptime((r.get("DATE") or "").strip(), "%d %b %Y")
                gb[d0.strftime("%Y-%m")] = float(r["LPMAUYM"])
            except Exception:
                pass
        if not gb:
            raise ValueError("BoE CSV 解析出 0 行")
        src["GB"] = (gb, 1e-6, ("GBP", False), "英国 M4", "BoE LPMAUYM")
    except Exception as e:
        out["errors"].append(f"GB: {type(e).__name__} {str(e)[:50]}")

    # ⑤ 中国 M2 (亿元 → 万亿元 ÷1e4)
    try:
        b = _gms_get("https://datacenter-web.eastmoney.com/api/data/v1/get?"
                     "reportName=RPT_ECONOMY_CURRENCY_SUPPLY&columns=ALL"
                     "&pageSize=400&sortColumns=REPORT_DATE&sortTypes=1")
        cn = {}
        for r in (_j.loads(b).get("result") or {}).get("data") or []:
            d0 = (r.get("REPORT_DATE") or "")[:7]
            v = r.get("BASIC_CURRENCY")
            if d0 and v:
                cn[d0] = float(v)
        if not cn:
            raise ValueError("东财返回 0 行")
        src["CN"] = (cn, 1e-4, ("CNY", True), "中国 M2", "东方财富")
    except Exception as e:
        out["errors"].append(f"CN: {type(e).__name__} {str(e)[:50]}")

    if not src:
        out["status"] = "error"
        return out

    # 汇率
    fx = {}
    for nm, sid in (("EUR", "DEXUSEU"), ("JPY", "DEXJPUS"),
                    ("GBP", "DEXUSUK"), ("CNY", "DEXCHUS")):
        try:
            fx[nm] = _gms_fred(sid, key, "1990-01-01")
        except Exception as e:
            out["errors"].append(f"FX {nm}: {type(e).__name__}")

    # 元数据
    for cc, (d, sc, fxk, cn_name, srcname) in src.items():
        out["countries"][cc] = {"name": cn_name, "src": srcname,
                                "n": len(d), "first": min(d) if d else None,
                                "last": max(d) if d else None}

    # 逐月折美元
    months = sorted({m for d, *_ in src.values() for m in d if m >= _GMS_START})
    rows = []
    for m in months:
        tot, parts, miss = 0.0, {}, []
        for cc, (d, sc, fxk, *_r) in src.items():
            if m not in d:
                miss.append(cc)
                continue
            local = d[m] * sc                       # → 万亿本币
            if fxk is None:
                usd = local                          # 美国本来就是美元
            else:
                nm, inv = fxk
                rk = [k for k in fx.get(nm, {}) if k <= m]
                if not rk:
                    miss.append(cc)
                    continue
                rate = fx[nm][max(rk)]
                usd = local / rate if inv else local * rate
            parts[cc] = round(usd, 3)
            tot += usd
        rows.append({"m": m, "total": round(tot, 3), "parts": parts,
                     "n_cty": len(parts), "missing": miss})

    # 同比 + 加速度
    by_m = {r["m"]: r for r in rows}
    for r in rows:
        y, mm = r["m"].split("-")
        prev = by_m.get(f"{int(y)-1}-{mm}")
        # ★只在<b>覆盖国家集合相同</b>时算同比, 否则国家数变化会伪装成货币暴增
        if prev and prev["total"] > 0 and set(prev["parts"]) == set(r["parts"]):
            r["yoy"] = round((r["total"] / prev["total"] - 1) * 100, 2)
        else:
            r["yoy"] = None
    for r in rows:
        y, mm = r["m"].split("-")
        prev = by_m.get(f"{int(y)-1}-{mm}")
        r["accel"] = (round(r["yoy"] - prev["yoy"], 2)
                      if (r.get("yoy") is not None and prev
                          and prev.get("yoy") is not None) else None)

    # ★★剔除尾部「未满员」月份(2026-08 只有中日发布 → 61 万亿, 画上去是断崖假象)。
    #   判据: 从尾部往前, 只要该月覆盖国数 < 历史最大覆盖国数, 就丢弃;
    #   一旦遇到满员月即停(只砍尾巴, 不动中间的真实历史缺口)。
    _maxn = max((r["n_cty"] for r in rows), default=0)
    _dropped = []
    while rows and rows[-1]["n_cty"] < _maxn:
        _dropped.append(rows[-1]["m"])
        rows.pop()
    if _dropped:
        out["tail_dropped"] = list(reversed(_dropped))
        out["tail_note"] = (f"末尾 {len(_dropped)} 个月（{_dropped[-1]}"
                            f"{'~' + _dropped[0] if len(_dropped) > 1 else ''}）"
                            f"因部分经济体尚未发布当月数据而暂不显示——"
                            f"若强行纳入会呈现虚假暴跌。")

    out["series"] = rows
    # 覆盖度分段
    cur, seg = None, None
    for r in rows:
        k = tuple(sorted(r["parts"]))
        if k != cur:
            if seg:
                out["coverage"].append(seg)
            cur = k
            seg = {"from": r["m"], "to": r["m"], "cty": list(k)}
        else:
            seg["to"] = r["m"]
    if seg:
        out["coverage"].append(seg)

    last = rows[-1] if rows else {}
    out["latest"] = last
    if out["errors"]:
        out["status"] = "ok(部分失败)"
    return out


if __name__ == "__main__":
    import json
    print("=== Credit Impulse 三国(信贷脉冲) ===")
    _ci = fetch_credit_impulse()
    for cc, d in _ci.items():
        if d["points"]:
            tail = ", ".join(f"{p['date'][:7]}:{p['ci']:+.2f}" for p in d["points"][-4:])
            print(f"  {d['flag']} {d['name']}: latest {d['latest']:+.2f} ({d['latest_date']}) [{d['signal']}]  近4季: {tail}")
        else:
            print(f"  {d['flag']} {d['name']}: {d['status']}")
    print("\n=== KOL 状态变化(近10天) ===")
    for ch in kol_stance_changes()[:8]:
        print(f"  {ch['date']} {ch['kol']}({ch['sector']}): {ch['prev_dir']} → {ch['new_dir']}")
    print("\n=== 流动性关键点 ===")
    print(json.dumps(fetch_liquidity_points(), ensure_ascii=False, indent=2, default=str)[:1500])
    print("\n=== 外国官方托管美债 (Fed H.4.1) ===")
    print(json.dumps(fetch_foreign_custody_ust(), ensure_ascii=False, indent=2, default=str))
    print("\n=== 美国国债拍卖 ===")
    print(json.dumps(fetch_treasury_auctions(), ensure_ascii=False, indent=2, default=str))


# ══════════════════════════════════════════════════════════════════
# WGC Weekly Markets Monitor —— ETF / COMEX 期货 ATM IV 期限结构(周更)
# ══════════════════════════════════════════════════════════════════
_WMM_LIST_URL = "https://www.gold.org/goldhub/gold-focus"

# ★IGLN/SGOL 长端为平推值(2026-09-21 实测: 12M/18M/24M 的 PDF y 坐标完全相同,
#   不是三个独立报价)。这两个标的流动性差(SGOL 本项目自算侧也因"只有 4 个到期日"
#   被排除), Bloomberg 侧多半是用末个有效值前向填充。页面必须标注,
#   否则读者会把它当成四条同等真实的曲线。


def _wmm_cache_path():
    """双候选路径(照 fetch_china_regime_shift 的做法)。
    ★为什么不只用 __file__: 反测/工具脚本可能从别处导入本模块,
      __file__ 推导会算到错误目录 —— 今天在中国断代那里已经栽过一次。"""
    import os as _os
    c = [_os.path.abspath(_os.path.join(_os.path.dirname(__file__), "..",
                                        "data", "wgc_weekly_iv.json")),
         _os.path.join(_os.getcwd(), "data", "wgc_weekly_iv.json")]
    for p in c:
        if _os.path.exists(p):
            return p
    return c[0]


def _wmm_latest_pdf_url():
    """从 gold-focus 列表页找最新一期 WMM 的 PDF 直链。
    ★跳过 'no-release' 期(WGC 确实会停发, 实测有 weekly-markets-monitor-no-release)。"""
    import re as _re
    import requests
    try:
        r = requests.get(_WMM_LIST_URL, timeout=55,
                         headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"})
        if r.status_code != 200:
            return None, None
        html = r.text
    except Exception:
        return None, None
    _errs = []
    seen, uniq = set(), []
    for m in _re.finditer(
            r'href="(/goldhub/gold-focus/\d{4}/\d{2}/weekly-markets-monitor[^"]*)"', html):
        u = m.group(1)
        if u not in seen:
            seen.add(u)
            uniq.append(u)
    for u in uniq[:5]:
        if "no-release" in u:
            continue
        try:
            p = requests.get("https://www.gold.org" + u, timeout=55,
                             headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"})
            if p.status_code != 200:
                continue
            m = _re.search(r'href="(/download/file/\d+/WMM_[^"]+\.pdf)"', p.text)
            if m:
                return "https://www.gold.org" + m.group(1), u
        except Exception as _e:
            # ★不吞: 记下真实异常类型, 否则只看到"未找到链接"会往错方向排查
            #   (2026-09-21 实测: NameError 被吞掉, 我查了三轮正则)
            _errs.append("%s: %s" % (type(_e).__name__, _e))
            continue
    if _errs:
        print("[external_data] WMM 详情页异常: %s" % _errs[:3])
    return None, None


def _wmm_parse(blob):
    """解析 WMM PDF 里的 ATM IV 期限结构图(ETF 左 / 期货 右)。

    ★方法: 轴刻度文字 bbox **中心**做 y→% 线性标定; 折线按颜色分组;
      x 吸附到期限刻度; 图例色块反查 ticker。
    ★坑(2026-09-21 实测): 文字 bbox 的 x 是左边缘, 而数据点落在文字中心,
      差约半个字宽(6.6pt)。用左边缘吸附会全部落空 —— 必须用 (x0+x1)/2。
    ★硬校验: 反推值必须落在该图 y 轴范围内, 否则整条丢弃(宁缺毋错)。
    """
    import re as _re
    from collections import defaultdict
    try:
        import pymupdf as _fz
    except ImportError:
        try:
            import fitz as _fz
        except ImportError:
            return {"error": "pymupdf 未安装"}
    try:
        doc = _fz.open(stream=blob, filetype="pdf")
    except Exception as e:
        return {"error": f"PDF 打开失败 {type(e).__name__}"}

    tgt = None
    for i, pg in enumerate(doc):
        if "ATM IV term structure" in pg.get_text():
            tgt = i
            break
    if tgt is None:
        return {"error": "未找到 ATM IV term structure 页"}

    pg = doc[tgt]
    ylab, xlab, legends = [], [], []
    for b in pg.get_text("dict")["blocks"]:
        for l in b.get("lines", []):
            for s in l.get("spans", []):
                t = s["text"].strip()
                x0, y0, x1, y1 = s["bbox"]
                cx, cy = (x0 + x1) / 2.0, (y0 + y1) / 2.0
                if _re.fullmatch(r"\d{2}", t) and cy > 200:
                    ylab.append((cx, cy, int(t)))
                elif t in ("1M", "3M", "6M", "12M", "18M", "24M"):
                    xlab.append((cx, cy, t))
                elif "Equity" in t or "COMDTY" in t:
                    legends.append({"tk": t.split()[0], "x0": x0, "cy": cy})

    draws = pg.get_drawings()
    legend_map = {}
    for g in legends:
        best = None
        for d in draws:
            col = d.get("fill") or d.get("color")
            if col is None:
                continue
            if max(col) - min(col) < 0.02 and col[0] > 0.7:
                continue
            rc = d["rect"]
            if (g["x0"] - 30) <= rc.x1 <= (g["x0"] - 1) and abs((rc.y0 + rc.y1) / 2 - g["cy"]) <= 7:
                dist = g["x0"] - rc.x1
                if best is None or dist < best[0]:
                    best = (dist, tuple(round(c, 3) for c in col))
        if best:
            legend_map[best[1]] = g["tk"]

    out = {"page": tgt + 1, "series": {}, "flat_tail": [],
           "legend_matched": len(legend_map), "legend_total": len(legends)}
    for grp, lo, hi in (("etf", 60, 480), ("futures", 490, 900)):
        ys = sorted([(cy, v) for cx, cy, v in ylab if lo <= cx <= hi], key=lambda z: z[0])
        xs = sorted([(cx, t) for cx, cy, t in xlab if lo <= cx <= hi], key=lambda z: z[0])
        if len(ys) < 2 or len(xs) < 4:
            continue
        (ya, va), (yb, vb) = ys[0], ys[-1]
        alo, ahi = min(v for _, v in ys), max(v for _, v in ys)
        grouped = defaultdict(list)
        for d in draws:
            col = d.get("color")
            if col is None:
                continue
            if max(col) - min(col) < 0.02 and col[0] > 0.7:
                continue
            pts = []
            for it in d["items"]:
                if it[0] == "l":
                    for p in (it[1], it[2]):
                        if lo <= p.x <= hi and 260 <= p.y <= 445:
                            pts.append((round(p.x, 1), round(p.y, 2)))
            if len(pts) >= 6:
                grouped[tuple(round(c, 3) for c in col)].extend(pts)
        order = [t for _, t in xs]
        for col, pts in grouped.items():
            tk = legend_map.get(col)
            if not tk:
                continue
            snap = {}
            for px, py in sorted(set(pts)):
                bx, bt = min(xs, key=lambda z: abs(z[0] - px))
                if abs(bx - px) <= 9:
                    snap.setdefault(bt, []).append(py)
            if len(snap) < 5:
                continue
            vals = {}
            for t in order:
                if t in snap:
                    yy = sum(snap[t]) / len(snap[t])
                    vals[t] = round(va + (yy - ya) * (vb - va) / (yb - ya), 2)
            if not all(alo - 1.5 <= v <= ahi + 1.5 for v in vals.values()):
                continue
            tail = [snap[t][0] for t in ("12M", "18M", "24M") if t in snap]
            if len(tail) >= 3 and len(set(tail)) == 1:
                out["flat_tail"].append(tk)
            out["series"][tk] = {"group": grp, "iv": vals}
    return out


def fetch_wgc_weekly_iv(cache_path=None):
    """WGC Weekly Markets Monitor 的 ATM IV 期限结构(ETF + COMEX 期货)。

    ★为什么做(2026-09-21 Chao 指出这图出自 WGC 周报): 本项目的 IV 是 CBOE
      逐合约自算, WGC 周报里有同类图(源 = Bloomberg)。两者并排 = 交叉验证 ——
      对得上说明自算算法可信, 对不上就是算法有问题。实测四个点差 <1pt。
    ★额外价值: 周报含 GC1 COMDTY(COMEX 黄金期货 IV)。本项目此前排查八条路径
      全封(CME API/ftp 403、Barchart 403、Yahoo 401、CBOE 不覆盖期货期权),
      判定 CME 的 IV 是商业数据拿不到 —— WGC 买来印在公开 PDF 里, 每周免费更新。
    ★频率差异: WGC 周更(周一), 本项目 CBOE 日更。页面必须分别标明, 不能混为一谈。
    """
    import os as _os
    out = {"status": "ok", "source": "WGC Weekly Markets Monitor(数据源 Bloomberg)",
           "freq": "weekly", "series": {}, "flat_tail": [], "warnings": [],
           "flat_tail_note": "长端多期相同 = 该标的流动性不足, 源侧前向填充, 非独立报价"}
    cp = cache_path or _wmm_cache_path()
    try:
        import re as _re
        pdf_url, page_url = _wmm_latest_pdf_url()
        if not pdf_url:
            raise RuntimeError("未找到 WMM PDF 链接")
        out["pdf_url"] = pdf_url
        out["page_url"] = ("https://www.gold.org" + page_url) if page_url else ""
        m = _re.search(r"WMM_(\d{1,2})_([A-Za-z]+)_(\d{4})", pdf_url)
        if m:
            out["issue"] = "%s %s %s" % (m.group(1), m.group(2), m.group(3))
        import urllib.request as _ur
        req = _ur.Request(pdf_url, headers={"User-Agent": "Mozilla/5.0 (EcoVolChecker research)"})
        with _ur.urlopen(req, timeout=120) as r:
            if r.status != 200:
                raise RuntimeError("PDF HTTP %s" % r.status)
            blob = r.read()
        if not blob.startswith(b"%PDF"):
            raise RuntimeError("返回的不是 PDF(可能是反爬页)")
        out["pdf_bytes"] = len(blob)
        pa = _wmm_parse(blob)
        if pa.get("error"):
            raise RuntimeError(pa["error"])
        if not pa.get("series"):
            raise RuntimeError("解析到 0 条曲线")
        out.update({k: pa[k] for k in ("series", "flat_tail", "page",
                                       "legend_matched", "legend_total")})
        out["as_of"] = datetime.date.today().isoformat()
        try:
            _os.makedirs(_os.path.dirname(cp), exist_ok=True)
            json.dump(out, open(cp, "w"), ensure_ascii=False, indent=1)
        except Exception:
            pass
    except Exception as e:
        try:
            cached = json.load(open(cp))
        except Exception:
            cached = None
        if cached and cached.get("series"):
            cached["status"] = "ok(缓存)"
            cached["cached"] = True
            cached["warnings"] = ["本次抓取失败(%s: %s), 显示缓存" % (type(e).__name__, e)]
            return cached
        out["status"] = "error"
        out["warnings"].append("%s: %s" % (type(e).__name__, e))
    return out



# ════════════════════════════════════════════════════════════════════
# 道指 ÷ 黄金 · 标普500 ÷ 黄金  50 年比值（Chao 2026-09-30）
# ════════════════════════════════════════════════════════════════════
# 月度序列 = 历史底表 data/stock_gold_ratio_base.json（1976-01 起，已完结月份，
#   道指/标普月末收盘，黄金伦敦金月均价）+ Yahoo 日线追加底表之后的月份：
#   股指取月末(当月则取最新)收盘，黄金取 GC=F 当月日收盘均值（与底表"月均价"同口径）。
# 另给"今日比值"= 最新股指收盘 ÷ 最新 GC=F 收盘（日报/卡片用）。
# Yahoo 失败 → 用上次成功的缓存 data/stock_gold_ratio_cache.json，并标 warnings；绝不编。

def _sgr_yahoo(sym, rng="1y"):
    import json as _j, urllib.request, datetime as _dt
    u = (f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}?range={rng}&interval=1d")
    req = urllib.request.Request(u, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124"})
    with urllib.request.urlopen(req, timeout=40) as r:
        d = _j.loads(r.read())
    rr = ((d.get("chart") or {}).get("result") or [None])[0]
    if not rr:
        raise RuntimeError(f"{sym}: 无 result")
    out = {}
    for t, c in zip(rr.get("timestamp") or [], ((rr.get("indicators") or {}).get("quote") or [{}])[0].get("close") or []):
        if c is not None:
            out[_dt.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d")] = float(c)
    if not out:
        raise RuntimeError(f"{sym}: 无收盘价")
    return out


def fetch_stock_gold_ratio(base_path=None, cache_path=None):
    from datetime import timezone, timedelta
    root = os.path.join(os.path.dirname(__file__), "..", "data")
    base_path = base_path or os.path.join(root, "stock_gold_ratio_base.json")
    cache_path = cache_path or os.path.join(root, "stock_gold_ratio_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "rows": []}
    try:
        base = json.load(open(base_path))
    except Exception as e:
        out["status"] = "error"; out["warnings"].append(f"底表缺失: {e}")
        return out
    rows = {r["m"]: dict(r, src="base") for r in base["rows"]}
    last_base = max(rows)
    try:
        dj, sp, gc = _sgr_yahoo("%5EDJI"), _sgr_yahoo("%5EGSPC"), _sgr_yahoo("GC%3DF")
        months = sorted({d[:7] for d in gc if d[:7] > last_base})
        for m in months:
            gd = [v for d, v in gc.items() if d[:7] == m]
            dd = [dj[d] for d in sorted(dj) if d[:7] == m]
            ss = [sp[d] for d in sorted(sp) if d[:7] == m]
            if gd and dd and ss:
                rows[m] = {"m": m, "dow": round(dd[-1], 2), "spx": round(ss[-1], 2),
                           "gold": round(sum(gd) / len(gd), 2), "src": "yahoo", "gold_days": len(gd)}
        ld = max(d for d in dj if d in gc)
        ls = max(d for d in sp if d in gc)
        out["today"] = {"dow_date": ld, "dow": round(dj[ld], 2), "gold_dow_day": round(gc[ld], 2),
                        "dow_gold": round(dj[ld] / gc[ld], 3),
                        "spx_date": ls, "spx": round(sp[ls], 2), "gold_spx_day": round(gc[ls], 2),
                        "spx_gold": round(sp[ls] / gc[ls], 4)}
    except Exception as e:
        try:
            c = json.load(open(cache_path))
            c.setdefault("warnings", []).append(f"本次 Yahoo 抓取失败({type(e).__name__}: {e})，显示 {c.get('as_of')} 缓存")
            c["status"] = "ok(cache)"
            return c
        except Exception:
            out["warnings"].append(f"Yahoo 失败且无缓存: {type(e).__name__}: {e}；仅显示底表")
    ser = []
    for m in sorted(rows):
        r = rows[m]
        ser.append({"m": m, "dow": r["dow"], "spx": r["spx"], "gold": r["gold"],
                    "dg": round(r["dow"] / r["gold"], 4), "sg": round(r["spx"] / r["gold"], 5), "src": r["src"]})
    out["rows"] = ser
    for k in ("dg", "sg"):
        v = [x[k] for x in ser]
        imax = max(range(len(v)), key=lambda i: v[i]); imin = min(range(len(v)), key=lambda i: v[i])
        cur = out.get("today", {}).get("dow_gold" if k == "dg" else "spx_gold", v[-1])
        out[k] = {"max": v[imax], "max_m": ser[imax]["m"], "min": v[imin], "min_m": ser[imin]["m"],
                  "median": sorted(v)[len(v) // 2], "latest_month": v[-1],
                  "pct_rank": round(sum(1 for x in v if x <= cur) / len(v) * 100, 1)}
    if out["status"] == "ok" and "today" in out:
        try:
            json.dump(out, open(cache_path, "w"), ensure_ascii=False)
        except Exception:
            pass
    return out



# ════════════════════════════════════════════════════════════════════
# 密歇根消费者信心 vs 美国全要素生产率相对趋势（复刻 BofA Chart 10，Chao 2026-09-30）
# ════════════════════════════════════════════════════════════════════
# TFP: 旧金山联储 Fernald 季度 TFP（quarterly_tfp.xlsx，商业部门），两个口径都算：
#   dtfp（普通）与 dtfp_util（剔除开工率），年化对数增速 /400 累加成对数水平。
# 趋势: 每季用过去 40 季对数水平线性拟合，偏离＝实际−拟合（无未来信息）。BofA 未公布算法，图上注明。
# 信心: FRED UMCSENT 月度。
# 每日运行: xlsx 每季更新一次，本地缓存 24h 内不重下；下载失败用缓存文件并标「缓存」。

_TFP_URL = "https://www.frbsf.org/wp-content/uploads/quarterly_tfp.xlsx"


def fetch_tfp_sentiment(cache_dir=None):
    import time, urllib.request
    import numpy as _np
    from datetime import timezone, timedelta
    root = cache_dir or os.path.join(os.path.dirname(__file__), "..", "data")
    xlsx = os.path.join(root, "fernald_quarterly_tfp.xlsx")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "tfp": [], "umcsent": []}
    fresh = os.path.exists(xlsx) and time.time() - os.path.getmtime(xlsx) < 86400
    if not fresh:
        try:
            req = urllib.request.Request(_TFP_URL, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124"})
            with urllib.request.urlopen(req, timeout=60) as r:
                blob = r.read()
            if blob[:2] != b"PK" or len(blob) < 100000:
                raise RuntimeError(f"下载内容不是 xlsx（{len(blob)} 字节）")
            tmp = xlsx + ".tmp"
            open(tmp, "wb").write(blob); os.replace(tmp, xlsx)
        except Exception as e:
            if os.path.exists(xlsx):
                out["warnings"].append(f"旧金山联储 TFP 下载失败（{type(e).__name__}），用缓存文件（"
                                       + datetime.datetime.fromtimestamp(os.path.getmtime(xlsx)).strftime("%Y-%m-%d") + "）")
            else:
                out["status"] = "error"; out["warnings"].append(f"TFP 下载失败且无缓存: {e}")
    if os.path.exists(xlsx):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
            rows = list(wb["quarterly"].iter_rows(values_only=True))
            hdr = [str(c) for c in rows[1]]
            ci, cu = hdr.index("dtfp"), hdr.index("dtfp_util")
            prod = next((str(r[0]) for r in list(wb["readme"].iter_rows(values_only=True))[:6]
                         if r and r[1] and "Produced on" in str(r[1])), None)
            vint = next((str(r[1]) for r in list(wb["readme"].iter_rows(values_only=True))[:6]
                         if r and len(r) > 1 and r[1] and "Produced on" in str(r[1])), "")
            q = [(r[0], r[ci], r[cu]) for r in rows[2:] if r[0] and ":Q" in str(r[0])
                 and r[ci] is not None and r[cu] is not None]
            d = [x[0] for x in q]
            res = {}
            for k, col in (("raw", 1), ("util", 2)):
                lv = _np.cumsum([x[col] / 400 for x in q])
                dev = [None] * len(lv)
                for i in range(39, len(lv)):
                    w = lv[i - 39:i + 1]
                    dev[i] = float(w[-1] - _np.polyval(_np.polyfit(_np.arange(40), w, 1), 39))
                res[k] = dev
            out["tfp"] = [{"q": d[i], "raw": round(res["raw"][i], 5), "util": round(res["util"][i], 5)}
                          for i in range(len(d)) if res["raw"][i] is not None and d[i] >= "1978:Q1"]
            import re as _re
            m = _re.search(r"Produced on ([A-Za-z]+ \d{1,2}, \d{4})", vint)
            out["tfp_vintage"] = m.group(1) if m else ""
        except Exception as e:
            out["status"] = "error"; out["warnings"].append(f"TFP 解析失败: {type(e).__name__}: {e}")
    try:
        um = _fred_series_hist("UMCSENT", "1978-01-01")
        out["umcsent"] = [{"m": dd[:7], "v": v} for dd, v in um]
        if len(out["umcsent"]) < 100:
            out["warnings"].append(f"UMCSENT 点数异常少: {len(out['umcsent'])}")
    except Exception as e:
        out["warnings"].append(f"UMCSENT 失败: {e}")
    if out["tfp"] and out["umcsent"]:
        um = {x["m"]: x["v"] for x in out["umcsent"]}
        def qmean(qq):
            y, n = qq.split(":Q"); v = [um.get(f"{y}-{(int(n)-1)*3+k:02d}") for k in (1, 2, 3)]
            v = [x for x in v if x is not None]; return sum(v) / len(v) if v else None
        for k in ("raw", "util"):
            pr = [(x[k], qmean(x["q"])) for x in out["tfp"] if qmean(x["q"]) is not None]
            out[f"corr_{k}"] = round(float(_np.corrcoef([a for a, _ in pr], [b for _, b in pr])[0, 1]), 2)
    return out



# ════════════════════════════════════════════════════════════════════
# LLM Token 支出指数（Silicon Data SDLLMTK，Chao 2026-09-30）
# ════════════════════════════════════════════════════════════════════
# 官方 portal 只给滚动 7 天、无历史下载 → 本项目每天抓一次追加进 data/llm_token_index.json 自建历史。
# 种子：2026-08-22 起第三方每日存档(与 portal 重叠日逐日一致)。抓取失败 → 用已存历史并标 warnings。
_SDL_URL = "https://portal.silicondata.com/token-indexes-chart"
_SDL_TOK = {"expenditure": "all", "open_expenditure": "open", "closed_expenditure": "prop"}


def fetch_llm_token_index(path=None):
    import re, time, requests
    from datetime import timezone, timedelta
    path = path or os.path.join(os.path.dirname(__file__), "..", "data", "llm_token_index.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    try:
        store = json.load(open(path))
    except Exception:
        store = {"rows": {}}
    out = {"status": "ok", "as_of": jst, "warnings": [], "added": 0, "revised": 0}
    Q = r'\\*"'
    re_idx = re.compile(Q + "indexes" + Q + r"\s*:\s*\{(.*?)\}", re.S)
    re_pair = re.compile(Q + r"(\d{4}-\d{2}-\d{2})" + Q + r"\s*:\s*" + Q + r"([0-9]+\.[0-9]+)" + Q)
    got = {}
    s = requests.Session()
    s.headers["User-Agent"] = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0 Safari/537.36"
    for tok, col in _SDL_TOK.items():
        try:
            r = s.get(_SDL_URL, params={"token": tok, "_": str(int(time.time() * 1000))},
                      headers={"Cache-Control": "no-cache", "Pragma": "no-cache"}, timeout=45)
            r.raise_for_status()
            if tok not in r.text:
                raise RuntimeError("返回页未回显所请求的 token")
            m = re_idx.search(r.text)
            pairs = re_pair.findall(m.group(1)) if m else []
            if not pairs:
                raise RuntimeError("页面里没解析到 indexes 数据(格式可能变了)")
            for d, v in pairs:
                got.setdefault(d, {})[col] = float(v)
        except Exception as e:
            out["warnings"].append(f"{col}: {type(e).__name__} {str(e)[:60]}")
    rows = store.setdefault("rows", {})
    for d, vals in got.items():
        if set(vals) != {"all", "open", "prop"}:
            continue
        old = rows.get(d)
        if old is None:
            rows[d] = dict(vals, src="portal"); out["added"] += 1
        elif any(abs(old.get(k, -1) - vals[k]) > 1e-9 for k in vals):
            rows[d] = dict(vals, src="portal(revised)"); out["revised"] += 1
    if got and (out["added"] or out["revised"]):
        try:
            json.dump(store, open(path, "w"), ensure_ascii=False, indent=0)
        except Exception as e:
            out["warnings"].append(f"写回失败: {e}")
    if not got:
        out["warnings"].insert(0, "本次 Silicon Data 抓取失败，显示已存历史")
    ser = [dict(d=d, **{k: rows[d][k] for k in ("all", "open", "prop")}) for d in sorted(rows)]
    out["series"] = ser
    out["source_note"] = store.get("seed_note", "")
    if ser:
        v = [x["all"] for x in ser]
        out["latest"] = ser[-1]
        out["lo"] = min(ser, key=lambda x: x["all"]); out["hi"] = max(ser, key=lambda x: x["all"])
        out["chg_7d"] = round((v[-1] / v[-8] - 1) * 100, 1) if len(v) >= 8 else None
    else:
        out["status"] = "error"
    return out



# ════════════════════════════════════════════════════════════════════
# 中美 AI 资本开支（季度）+ 美国数据中心建设支出（月度）  Chao 2026-09-30 选方案 C
# ════════════════════════════════════════════════════════════════════
# 美国五家: SEC companyfacts XBRL(官方), 单季=直接季度事实, 否则 YTD 差分; 与第三方逐季比对 0 差异。
# 中国三家: data/cn_ai_capex_seed.json(阿里/腾讯四季加总与年报逐一相符) + 每日尝试从官方原文追加新季度:
#   阿里/百度=EDGAR 6-K ex99-1 正文/表; 腾讯=PRNewswire 季报「Capital expenditures (d)」表首列。
# 汇率: FRED DEXCHUS 季度均值。数据中心: 美国统计局 C30「Value of Private Construction Put in Place, SA」Data center 列。
_AIC_US = {"亚马逊": (1018724, "PaymentsToAcquireProductiveAssets"),
           "微软": (789019, "PaymentsToAcquirePropertyPlantAndEquipment"),
           "谷歌": (1652044, "PaymentsToAcquirePropertyPlantAndEquipment"),
           "Meta": (1326801, "PaymentsToAcquirePropertyPlantAndEquipment"),
           "甲骨文": (1341439, "PaymentsToAcquirePropertyPlantAndEquipment")}
_SEC_H = {"User-Agent": "Eco-Volatility-Checker research contact@example.com"}


def _aic_cq(end):
    d = datetime.date.fromisoformat(end) - datetime.timedelta(days=40)
    return f"{d.year}Q{(d.month - 1) // 3 + 1}"


def _aic_us_quarters(cik, concept):
    import urllib.request
    u = json.load(urllib.request.urlopen(urllib.request.Request(
        f"https://data.sec.gov/api/xbrl/companyconcept/CIK{cik:010d}/us-gaap/{concept}.json", headers=_SEC_H), timeout=40))["units"]["USD"]
    facts = {}
    for x in u:
        if x.get("form") in ("10-Q", "10-K", "10-Q/A", "10-K/A") and "start" in x:
            facts[(x["start"], x["end"])] = x["val"]
    dd = lambda a, b: (datetime.date.fromisoformat(b) - datetime.date.fromisoformat(a)).days
    q = {e: v for (s, e), v in facts.items() if 80 <= dd(s, e) <= 100}
    for (s, e), v in facts.items():
        if dd(s, e) > 100 and e not in q:
            prev = [(e2, v2) for (s2, e2), v2 in facts.items() if s2 == s and e2 < e]
            if prev:
                e2, v2 = max(prev)
                if 80 <= dd(e2, e) <= 100:
                    q[e] = v - v2
    return {_aic_cq(e): (v / 1e9, e) for e, v in q.items()}


def _aic_cn_update(seed, warns):
    """尝试从官方原文追加中国三家的新季度。只追加 seed 里没有的季度; 失败只记 warning。"""
    import urllib.request, re as _re, html as _h, subprocess, time as _t
    def T(b): return _re.sub(r"\s+", " ", _h.unescape(_re.sub(r"<[^>]+>", " ", b.decode("utf-8", "ignore"))))
    def sec(u):
        _t.sleep(0.2); return urllib.request.urlopen(urllib.request.Request(u, headers=_SEC_H), timeout=40).read()
    M = {"March": 1, "June": 2, "September": 3, "December": 4}
    added = []
    for name, cik in (("alibaba", 1577552), ("baidu", 1329099)):
        try:
            r = json.loads(sec(f"https://data.sec.gov/submissions/CIK{cik:010d}.json"))["filings"]["recent"]
            recent = [(r["filingDate"][i], r["accessionNumber"][i]) for i in range(len(r["form"])) if r["form"][i] == "6-K"][:6]
            for fd, acc in recent:
                a = acc.replace("-", "")
                items = json.loads(sec(f"https://www.sec.gov/Archives/edgar/data/{cik}/{a}/index.json"))["directory"]["item"]
                for d in [x["name"] for x in items if x["name"].lower().endswith(".htm") and "99" in x["name"]][:1]:
                    t = T(sec(f"https://www.sec.gov/Archives/edgar/data/{cik}/{a}/{d}"))
                    if name == "alibaba":
                        m = _re.search(r"quarter ended (March|June|September|December) 3[01], (\d{4})[^.]{0,200}?capital expenditures? (?:of|were|was) RMB\s?([\d,]+) million", t)
                        if m:
                            k = f"{m.group(2)}Q{M[m.group(1)]}"
                            if k not in seed[name]:
                                seed[name][k] = float(m.group(3).replace(",", "")) / 1000; added.append(f"阿里 {k}")
                    else:
                        mq = _re.search(r"Announces (First|Second|Third|Fourth) Quarter(?: and Full Year)? (\d{4}) Results", t)
                        mc = _re.search(r"Less: Capital expenditures((?: \([\d,]+ \)){9,12})", t)
                        if mq and mc:
                            k = f"{mq.group(2)}Q{ {'First': 1, 'Second': 2, 'Third': 3, 'Fourth': 4}[mq.group(1)] }"
                            nums = [float(x.replace(",", "")) / 1000 for x in _re.findall(r"\(([\d,]+) \)", mc.group(1))]
                            if k not in seed[name]:
                                seed[name][k] = nums[8]; added.append(f"百度 {k}")
        except Exception as e:
            warns.append(f"{name} 新季度检查失败: {type(e).__name__}")
    try:
        lst = subprocess.run(["curl", "-sL", "-m", "30", "-A", "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124",
                              "https://www.prnewswire.com/news/tencent/"], capture_output=True, timeout=40).stdout.decode("utf-8", "ignore")
        for slug in sorted(set(_re.findall(r"tencent-announces-(\d{4})-(?:annual-and-)?(first|second|third|fourth)-quarter[a-z-]*results-\d+\.html", lst))):
            pass
        for m in _re.finditer(r"(tencent-announces-(\d{4})-(?:annual-and-)?(first|second|third|fourth)-quarter[a-z-]*results-\d+\.html)", lst):
            k = f"{m.group(2)}Q{ {'first': 1, 'second': 2, 'third': 3, 'fourth': 4}[m.group(3)] }"
            if k in seed["tencent"]:
                continue
            for base in ("https://www.prnewswire.com/apac/news-releases/", "https://www.prnewswire.com/news-releases/"):
                h = subprocess.run(["curl", "-sL", "-m", "30", "-A", "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124", base + m.group(1)],
                                   capture_output=True, timeout=40).stdout
                mc = _re.search(r"Capital expenditures \(d\) ([\d,]+) ", T(h))
                if mc:
                    seed["tencent"][k] = float(mc.group(1).replace(",", "")) / 1000; added.append(f"腾讯 {k}"); break
    except Exception as e:
        warns.append(f"腾讯 新季度检查失败: {type(e).__name__}")
    return added


def fetch_ai_capex(cache_dir=None):
    import time, urllib.request
    from datetime import timezone, timedelta
    root = cache_dir or os.path.join(os.path.dirname(__file__), "..", "data")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "us": {}, "cn": {}, "fx": {}, "dc": []}
    cache = os.path.join(root, "ai_capex_cache.json")
    # 美国
    for nm, (cik, con) in _AIC_US.items():
        try:
            out["us"][nm] = {k: v[0] for k, v in _aic_us_quarters(cik, con).items() if k >= "2023Q1"}
        except Exception as e:
            out["warnings"].append(f"{nm} SEC 失败: {type(e).__name__}")
    if len(out["us"]) < 5:
        try:
            c = json.load(open(cache))
            for nm in _AIC_US:
                if nm not in out["us"] and nm in c.get("us", {}):
                    out["us"][nm] = c["us"][nm]; out["warnings"].append(f"{nm} 用缓存（{c.get('as_of')}）")
        except Exception:
            pass
    # 中国
    seed_p = os.path.join(root, "cn_ai_capex_seed.json")
    try:
        seed = json.load(open(seed_p))
        added = _aic_cn_update(seed, out["warnings"])
        if added:
            json.dump(seed, open(seed_p, "w"), ensure_ascii=False, indent=1)
            out["cn_added"] = added
        out["cn"] = {"阿里": seed["alibaba"], "腾讯": seed["tencent"], "百度": seed["baidu"]}
    except Exception as e:
        out["warnings"].append(f"中国种子数据缺失: {e}")
    # 汇率
    try:
        fx = _fred_series_hist("DEXCHUS", "2023-01-01")
        acc = {}
        for d, v in fx:
            k = f"{d[:4]}Q{(int(d[5:7]) - 1) // 3 + 1}"
            acc.setdefault(k, []).append(v)
        out["fx"] = {k: sum(v) / len(v) for k, v in acc.items()}
    except Exception as e:
        out["warnings"].append(f"汇率失败: {e}")
    # 数据中心(月度)
    xlsx = os.path.join(root, "census_private_construction_sa.xlsx")
    if not (os.path.exists(xlsx) and time.time() - os.path.getmtime(xlsx) < 86400):
        try:
            b = urllib.request.urlopen(urllib.request.Request("https://www.census.gov/construction/c30/xlsx/privsatime.xlsx",
                                                               headers={"User-Agent": "Mozilla/5.0"}), timeout=60).read()
            if b[:2] != b"PK":
                raise RuntimeError("不是 xlsx")
            open(xlsx + ".tmp", "wb").write(b); os.replace(xlsx + ".tmp", xlsx)
        except Exception as e:
            out["warnings"].append(f"美国统计局下载失败（{type(e).__name__}），用缓存文件")
    try:
        import openpyxl
        ws = openpyxl.load_workbook(xlsx, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == "Date")
        ci = list(rows[hi]).index("Data center")
        mon = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6, "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12}
        dc = []
        for r in rows[hi + 1:]:
            if not r or not r[0] or not isinstance(r[ci], (int, float)):
                continue
            s = str(r[0]); mm = mon.get(s[:3])
            if not mm:
                continue
            yy = 2000 + int(s[4:6]); flag = s[6:].strip()
            dc.append({"m": f"{yy}-{mm:02d}", "v": r[ci] / 1000, "flag": {"p": "初值", "r": "修订"}.get(flag, "")})
        out["dc"] = sorted(dc, key=lambda x: x["m"])
    except Exception as e:
        out["warnings"].append(f"数据中心数据解析失败: {e}")
    # 汇总
    qs = sorted({k for d in out["us"].values() for k in d} | {k for d in out["cn"].values() for k in d})
    tbl = []
    for k in qs:
        us = [out["us"][n].get(k) for n in _AIC_US]
        cn = [out["cn"].get(n, {}).get(k) for n in ("阿里", "腾讯", "百度")]
        fx = out["fx"].get(k)
        row = {"q": k, "us_parts": dict(zip(_AIC_US, us)), "cn_parts_rmb": dict(zip(("阿里", "腾讯", "百度"), cn)), "fx": fx}
        row["us"] = sum(us) if all(x is not None for x in us) else None
        row["cn_rmb"] = sum(cn) if all(x is not None for x in cn) else None
        row["cn"] = row["cn_rmb"] / fx if (row["cn_rmb"] is not None and fx) else None
        if row["us"] is not None or row["cn"] is not None:
            tbl.append(row)
    out["table"] = tbl
    if out["us"] and len(out["us"]) == 5:
        try:
            json.dump({"as_of": jst, "us": out["us"]}, open(cache, "w"), ensure_ascii=False)
        except Exception:
            pass
    return out



# ════════════════════════════════════════════════════════════════════
# 金油比（黄金 ÷ WTI 原油）1971 起  Chao 2026-10-01
# ════════════════════════════════════════════════════════════════════
# 底表 data/gold_oil_ratio_base.json: 黄金=伦敦金月均价(与股金比同源), 原油=FRED WTISPLC 月均。
# 底表之后的月份: 黄金=Yahoo GC=F 当月日收盘均值, 原油=FRED DCOILWTICO 当月日均(同为月均口径)。
# 今日比值 = 最新 GC=F ÷ 最新 DCOILWTICO（取两者都有的最近一天）。失败用缓存并标注, 绝不编。
def fetch_gold_oil_ratio(base_path=None, cache_path=None):
    from datetime import timezone, timedelta
    import statistics as _st
    root = os.path.join(os.path.dirname(__file__), "..", "data")
    base_path = base_path or os.path.join(root, "gold_oil_ratio_base.json")
    cache_path = cache_path or os.path.join(root, "gold_oil_ratio_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "rows": []}
    try:
        base = json.load(open(base_path))
    except Exception as e:
        out.update(status="error"); out["warnings"].append(f"底表缺失: {e}"); return out
    rows = {r["m"]: dict(r, src="base") for r in base["rows"]}
    last_base = max(rows)
    try:
        gc = _sgr_yahoo("GC%3DF", "1y")
        oil = {d: v for d, v in _fred_series_hist("DCOILWTICO", f"{last_base[:4]}-01-01") if v is not None and v > 0}
        if not oil:
            raise RuntimeError("FRED DCOILWTICO 无数据")
        for m in sorted({d[:7] for d in gc if d[:7] > last_base}):
            gd = [v for d, v in gc.items() if d[:7] == m]
            od = [v for d, v in oil.items() if d[:7] == m]
            if gd and od:
                rows[m] = {"m": m, "gold": round(sum(gd) / len(gd), 2), "oil": round(sum(od) / len(od), 3),
                           "src": "live", "days": f"金{len(gd)}天/油{len(od)}天"}
        both = sorted(d for d in oil if d in gc)
        if both:
            d = both[-1]
            out["today"] = {"date": d, "gold": round(gc[d], 2), "oil": round(oil[d], 2), "ratio": round(gc[d] / oil[d], 2)}
    except Exception as e:
        try:
            c = json.load(open(cache_path))
            c.setdefault("warnings", []).append(f"本次抓取失败({type(e).__name__}: {str(e)[:60]})，显示 {c.get('as_of')} 缓存")
            c["status"] = "ok(cache)"
            return c
        except Exception:
            out["warnings"].append(f"实时数据失败且无缓存: {type(e).__name__}；仅显示底表")
    ser = [{"m": m, "gold": rows[m]["gold"], "oil": rows[m]["oil"], "r": round(rows[m]["gold"] / rows[m]["oil"], 3),
            "src": rows[m]["src"]} for m in sorted(rows)]
    out["rows"] = ser
    v = [x["r"] for x in ser]
    cur = out.get("today", {}).get("ratio", v[-1])
    imax = max(range(len(v)), key=lambda i: v[i]); imin = min(range(len(v)), key=lambda i: v[i])
    out["stats"] = {"mean": round(_st.mean(v), 2), "median": round(_st.median(v), 2),
                    "max": v[imax], "max_m": ser[imax]["m"], "min": v[imin], "min_m": ser[imin]["m"],
                    "pct_rank": round(sum(1 for x in v if x <= cur) / len(v) * 100, 1),
                    "mean_ex2020": round(_st.mean([x["r"] for x in ser if not ("2020-03" <= x["m"] <= "2020-05")]), 2)}
    if out["status"] == "ok" and "today" in out:
        try:
            json.dump(out, open(cache_path, "w"), ensure_ascii=False)
        except Exception:
            pass
    return out



# ════════════════════════════════════════════════════════════════════
# 美国货运: 铁路(AAR 周度, 10 类商品+联运) + 公路(ATA 卡车吨位/Cass 发货量 月度)  Chao 2026-10-01
# ════════════════════════════════════════════════════════════════════
def fetch_us_freight():
    import sys as _sys, urllib.request
    from datetime import timezone, timedelta
    root = os.path.join(os.path.dirname(__file__), "..")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "rail": {}, "road": {}}
    # 铁路: 增量抓最新周报
    try:
        _sys.path.insert(0, os.path.join(root, "tools"))
        import aar_rail_weekly as _a
        p = _a.OUT
        store = json.load(open(p)) if os.path.exists(p) else {"weeks": {}, "src": {}}
        have = set(store.get("src", {}).values())
        new = 0
        for u, d in _a.list_pdfs()[-6:]:
            if u in have:
                continue
            try:
                b = urllib.request.urlopen(urllib.request.Request(u, headers=_a.UA), timeout=60).read()
                wk, rec = _a.parse(b)
                store["weeks"][wk] = rec; store.setdefault("src", {})[wk] = u; new += 1
            except Exception as e:
                out["warnings"].append(f"AAR {d} 周报解析失败: {str(e)[:50]}")
        if new:
            json.dump(store, open(p, "w"), ensure_ascii=False, indent=0)
        out["rail_new"] = new
    except Exception as e:
        out["warnings"].append(f"AAR 周报检查失败（{type(e).__name__}），用已存数据")
        try:
            store = json.load(open(os.path.join(root, "data", "aar_rail_weekly.json")))
        except Exception:
            store = {"weeks": {}}
    wk = {k: v for k, v in store.get("weeks", {}).items() if k >= "2021-07-01"}
    out["rail"] = {"weeks": [dict(d=k, **wk[k]) for k in sorted(wk)]}
    # 公路(月度): ATA 卡车吨位指数 + Cass 发货量指数
    for sid, key in (("TRUCKD11", "ata"), ("FRGSHPUSM649NCIS", "cass")):
        try:
            out["road"][key] = [{"m": d[:7], "v": v} for d, v in _fred_series_hist(sid, "2021-01-01")]
        except Exception as e:
            out["warnings"].append(f"FRED {sid} 失败: {type(e).__name__}")
    # 成本侧: 柴油(EIA 周度, FRED GASDESW) + 整车长途运价 PPI + Cass 运费支出(月度)
    try:
        out["diesel"] = [{"d": d, "v": v} for d, v in _fred_series_hist("GASDESW", "2021-01-01")]
    except Exception as e:
        out["warnings"].append(f"柴油价格失败: {type(e).__name__}")
    for sid, key in (("PCU484121484121", "ppi_tl"), ("FRGEXPUSM649NCIS", "cass_exp")):
        try:
            out["road"][key] = [{"m": d[:7], "v": v} for d, v in _fred_series_hist(sid, "2021-01-01")]
        except Exception as e:
            out["warnings"].append(f"FRED {sid} 失败: {type(e).__name__}")
    # SONAR 公开快照(拒载率/现货运价): 官方不给历史, 每天抓一次公开页面的当日读数自建序列; 被拦截就沿用已存
    out["sonar"] = _sonar_snapshot(os.path.join(root, "data", "sonar_snapshots.json"))
    if not out["rail"]["weeks"]:
        out["status"] = "error"
    return out


def _sonar_snapshot(path):
    import re as _re, html as _h, urllib.request
    try:
        store = json.load(open(path))
    except Exception:
        store = {"rows": {}, "note": "SONAR Market Monitor 公开页当日读数, 本站每日抓取自建; 页面不标数据日期, 以抓取日(JST)记"}
    today = datetime.datetime.utcnow() + datetime.timedelta(hours=9)
    day = today.strftime("%Y-%m-%d")
    res = {"ok": False, "err": ""}
    if day not in store["rows"] and today.weekday() < 5:
        try:
            req = urllib.request.Request("https://gosonar.com/freightwaves-sonar-market-monitor",
                                         headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124"})
            b = urllib.request.urlopen(req, timeout=40).read().decode("utf-8", "ignore")
            if "Attention Required" in b or "you have been blocked" in b:
                raise RuntimeError("被 Cloudflare 拦截")
            t = _re.sub(r"\s+", " ", _h.unescape(_re.sub(r"<[^>]+>", " ", _re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", b, flags=_re.S))))
            def g(pat):
                m = _re.search(pat, t)
                return float(m.group(1).replace(",", "")) if m else None
            row = {"stri": g(r"STRI\.USA ([\d.]+)%"), "otri": g(r"OTRI ([\d.]+)%"), "nti": g(r"NTI\.USA \$([\d.]+)/mi"),
                   "vcrpm1": g(r"VCRPM1\.USA \$([\d.]+)/mi"), "stvi": g(r"STVI\.USA ([\d,]+\.?\d*)")}
            if row["stri"] is None or row["nti"] is None:
                raise RuntimeError("页面格式变了, 没解析到 STRI/NTI")
            store["rows"][day] = row
            json.dump(store, open(path, "w"), ensure_ascii=False, indent=0)
            res["ok"] = True
        except Exception as e:
            res["err"] = f"{type(e).__name__}: {str(e)[:50]}"
    res["rows"] = [dict(d=k, **store["rows"][k]) for k in sorted(store["rows"])]
    return res


# ════════════════════════════════════════════════════════
# 比特币 ÷ 黄金（2010-08 起，月度）+ 暴涨暴跌事件标注   Chao 2026-10-02
# ════════════════════════════════════════════════════════
# 比特币: blockchain.com 市场均价(美元, 日度, 2010-08-18 起首次有价)，与 FRED CBBTCUSD(Coinbase, 2014-12 起)
#         重叠 4288 天中位偏差 1.5%（2014-12~2015-01 个别日因早期交易所价差较大），仅作来源可信度参考。
# 黄金:   与金油比同一底表 data/gold_oil_ratio_base.json(伦敦金月均价)；底表之后用 Yahoo GC=F 当月日收盘均值。
# 比值 = 月均 BTC 美元价 ÷ 月均黄金美元价（1 盎司黄金 ≈ 几个比特币的倒数；这里直接给「1 个比特币 = 几盎司黄金」）。
# 事件: data/btc_events.json（价格先行、联网检索、逐条验证 URL，见文件内 method）。
def fetch_btc_gold_ratio(base_path=None, events_path=None, cache_path=None):
    from datetime import timezone, timedelta
    import statistics as _st
    root = os.path.join(os.path.dirname(__file__), "..", "data")
    base_path = base_path or os.path.join(root, "gold_oil_ratio_base.json")
    events_path = events_path or os.path.join(root, "btc_events.json")
    cache_path = cache_path or os.path.join(root, "btc_gold_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "rows": [], "events": [], "daily": []}
    try:
        import requests as _rq
        j = _rq.get("https://api.blockchain.info/charts/market-price?timespan=all&format=json&sampled=false",
                    headers={"User-Agent": "Mozilla/5.0"}, timeout=60).json()
        btc = {}
        for x in j.get("values", []):
            d = datetime.datetime.utcfromtimestamp(x["x"]).strftime("%Y-%m-%d")
            if x["y"] and x["y"] > 0 and d >= "2010-08-18":
                btc[d] = float(x["y"])
        if len(btc) < 3000:
            raise RuntimeError(f"blockchain.com 返回仅 {len(btc)} 天")
        base = json.load(open(base_path))
        gold = {r["m"]: float(r["gold"]) for r in base["rows"]}
        last_base = max(gold)
        try:
            gc = _sgr_yahoo("GC%3DF", "1y")
            for m in sorted({d[:7] for d in gc if d[:7] > last_base}):
                gd = [v for d, v in gc.items() if d[:7] == m]
                if gd:
                    gold[m] = round(sum(gd) / len(gd), 2)
            _both = sorted(d for d in gc if d in btc)
            gold_today_d = _both[-1] if _both else None   # 取两者都有报价的最近一天, 避免日期错位
            gold_today = gc[gold_today_d] if gold_today_d else None
        except Exception as e:
            out["warnings"].append(f"当月黄金实时失败({type(e).__name__})，月度只到 {last_base}")
            gold_today = None; gold_today_d = None
        mb = {}
        for d, v in btc.items():
            mb.setdefault(d[:7], []).append(v)
        rows = []
        for m in sorted(mb):
            if m in gold:
                b = sum(mb[m]) / len(mb[m])
                rows.append({"m": m, "btc": round(b, 4 if b < 1 else 2), "gold": gold[m], "r": round(b / gold[m], 5),
                             "days": len(mb[m])})
        out["rows"] = rows
        # 日度比特币(用于标注日、悬停): 只保留 2014 起每周一个点以控制体积; 标注日单独保留
        ev = []
        try:
            ev = json.load(open(events_path)).get("events", [])
            ev_age = (datetime.datetime.now() - datetime.datetime.fromtimestamp(os.path.getmtime(events_path))).days
            if ev_age > 30:
                out["warnings"].append(f"事件清单已 {ev_age} 天未更新")
        except Exception as e:
            out["warnings"].append(f"事件清单缺失: {type(e).__name__}")
        out["events"] = ev
        # 叠加线(用户 2026-10-02 加): ① 全球货币供应量=五大经济体 M2/M3 折美元聚合(万亿美元, 复用既有 fetcher)
        #                            ② 美国联邦政府债务总额=FRED GFDEBTN(季度, 百万美元→万亿美元)
        out["overlay"] = {"m2": [], "debt": [], "warnings": []}
        try:
            _g = fetch_global_money_supply()
            out["overlay"]["m2"] = [{"m": x["m"], "v": x["total"], "n": x.get("n_cty")} for x in _g.get("series", []) if x.get("total") is not None]
            if _g.get("status") != "ok":
                out["overlay"]["warnings"].append("全球 M2 来源状态: " + str(_g.get("status")))
        except Exception as e:
            out["overlay"]["warnings"].append(f"全球 M2 失败: {type(e).__name__}")
        try:
            _d = [(d, v) for d, v in _fred_series_hist("GFDEBTN", "1996-01-01") if v is not None]
            out["overlay"]["debt"] = [{"m": d[:7], "v": round(v / 1e6, 4)} for d, v in _d]   # 百万美元 → 万亿美元
        except Exception as e:
            out["overlay"]["warnings"].append(f"美国联邦债务失败: {type(e).__name__}")
        td = gold_today_d or max(btc)
        out["today"] = {"date": td, "btc": round(btc[td], 2), "gold": round(gold_today, 2) if gold_today else None,
                        "ratio": round(btc[td] / gold_today, 3) if gold_today else None}
        v = [r["r"] for r in rows]
        cur = rows[-1]["r"]
        imax = max(range(len(v)), key=lambda i: v[i])
        out["stats"] = {"max": v[imax], "max_m": rows[imax]["m"], "cur_m": rows[-1]["m"], "cur": cur,
                        "pct_rank": round(sum(1 for x in v if x <= cur) / len(v) * 100, 1),
                        "n": len(rows)}
        # 逐次暴涨: 从周期低点到下一个高点 (zigzag, 阈值 50%/回撤 50%), 供「用了多久涨几百倍」
        s = sorted(btc.items())
        out["btc_daily_n"] = len(s)
        out["btc_first"] = s[0]
        out["btc_peak"] = max(s, key=lambda x: x[1])
        # 标注日附近的日度价(用于图上定位)
        dd = dict(s)
        for e in ev:
            e["btc_px"] = dd.get(e["date"])
    except Exception as e:
        try:
            c = json.load(open(cache_path))
            c.setdefault("warnings", []).append(f"本次抓取失败({type(e).__name__}: {str(e)[:60]})，显示 {c.get('as_of')} 缓存")
            c["status"] = "ok(cache)"
            return c
        except Exception:
            out["status"] = "error"; out["warnings"].append(f"实时数据失败且无缓存: {type(e).__name__}: {str(e)[:60]}")
            return out
    if out["status"] == "ok":
        try:
            json.dump(out, open(cache_path, "w"), ensure_ascii=False)
        except Exception:
            pass
    return out


GS_TECH_SYMBOLS = [
    # (代码, 名称, 分组, 说明)
    ("GC=F", "COMEX 黄金期货", "金属本身", "COMEX 连续合约(Yahoo), 非现货/LBMA 定盘价; 换月会有跳点"),
    ("SI=F", "COMEX 白银期货", "金属本身", "COMEX 连续合约(Yahoo), 非现货/LBMA 定盘价; 换月会有跳点"),
    ("GLD", "GLD 黄金 ETF", "金属本身", "SPDR Gold Shares, 美元"),
    ("SLV", "SLV 白银 ETF", "金属本身", "iShares Silver Trust, 美元"),
    ("GDX", "GDX 美国金矿 ETF", "美国矿企指数", "VanEck Gold Miners ETF, 美元"),
    ("GDXJ", "GDXJ 美国小型金矿 ETF", "美国矿企指数", "VanEck Junior Gold Miners ETF, 美元"),
    ("SIL", "SIL 美国银矿 ETF", "美国矿企指数", "Global X Silver Miners ETF, 美元"),
    ("SILJ", "SILJ 美国小型银矿 ETF", "美国矿企指数", "Amplify Junior Silver Miners ETF, 美元"),
    ("159562.SZ", "中国黄金股 ETF 159562", "中国矿企", "深交所 ETF(代表中国黄金矿企指数), 人民币"),
    ("600547.SS", "山东黄金", "中国矿企", "上交所 A 股, 人民币"),
    ("600489.SS", "中金黄金", "中国矿企", "上交所 A 股, 人民币"),
    ("601899.SS", "紫金矿业", "中国矿企", "上交所 A 股, 人民币"),
]


def _gs_sma(vals, n):
    out = [None] * len(vals)
    if len(vals) < n:
        return out
    s = sum(vals[:n]); out[n - 1] = s / n
    for i in range(n, len(vals)):
        s += vals[i] - vals[i - n]; out[i] = s / n
    return out


def _gs_rsi(vals, n=14):
    """Wilder RSI(StockCharts 口径): 前 n 根简单平均起步, 之后 Wilder 平滑。"""
    out = [None] * len(vals)
    if len(vals) <= n:
        return out
    gains = [max(vals[i] - vals[i - 1], 0.0) for i in range(1, len(vals))]
    losses = [max(vals[i - 1] - vals[i], 0.0) for i in range(1, len(vals))]
    ag = sum(gains[:n]) / n; al = sum(losses[:n]) / n
    out[n] = 100.0 if al == 0 else 100.0 - 100.0 / (1.0 + ag / al)
    for i in range(n, len(gains)):
        ag = (ag * (n - 1) + gains[i]) / n; al = (al * (n - 1) + losses[i]) / n
        out[i + 1] = 100.0 if al == 0 else 100.0 - 100.0 / (1.0 + ag / al)
    return out


def fetch_gold_silver_tech(cache_path=None, keep=260):
    """金银短线图: 12 个标的日线 OHLC + SMA20/50/200 + RSI14(Wilder)。
    数据: Yahoo Finance 日线(约 3 年历史, 保证 1 年窗口内 200 日线完整); 指标在全历史上算, 只缓存最近 keep 根。
    单个标的失败 → 用该标的旧缓存并标「缓存」, 绝不用空覆盖旧真值。"""
    import requests as _rq, json as _json, time
    from datetime import timezone, timedelta
    root = os.path.join(os.path.dirname(__file__), "..", "data")
    cache_path = cache_path or os.path.join(root, "gold_silver_tech_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    try:
        old = {x["sym"]: x for x in _json.load(open(cache_path, encoding="utf-8")).get("items", [])}
    except Exception:
        old = {}
    out = {"status": "ok", "as_of": jst, "warnings": [], "items": []}
    H = {"User-Agent": "Mozilla/5.0"}
    t1 = int(datetime.datetime.utcnow().timestamp()) + 86400; t0 = t1 - 1250 * 86400
    ok_n = 0
    for sym, name, grp, note in GS_TECH_SYMBOLS:
        item = None
        try:
            r = None
            for _ in range(2):
                r = _rq.get(f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}", params={"period1": t0, "period2": t1, "interval": "1d"}, headers=H, timeout=40)
                if r.status_code == 200:
                    break
                time.sleep(1.5)
            if r.status_code != 200:
                raise RuntimeError(f"HTTP {r.status_code}")
            j = r.json()["chart"]["result"][0]; q = j["indicators"]["quote"][0]
            rows = []
            for k, ts in enumerate(j["timestamp"]):
                o_, h_, l_, c_ = q["open"][k], q["high"][k], q["low"][k], q["close"][k]
                if None in (o_, h_, l_, c_):
                    continue
                rows.append((datetime.datetime.utcfromtimestamp(ts).strftime("%Y-%m-%d"), float(o_), float(h_), float(l_), float(c_), float(q["volume"][k] or 0)))
            dd = {}
            for x in rows:
                dd[x[0]] = x
            rows = [dd[k] for k in sorted(dd)]
            if len(rows) < 260:
                raise RuntimeError(f"仅 {len(rows)} 根日线")
            cl = [x[4] for x in rows]
            m20, m50, m200, rs = _gs_sma(cl, 20), _gs_sma(cl, 50), _gs_sma(cl, 200), _gs_rsi(cl, 14)
            sl = slice(len(rows) - keep, len(rows))
            rd = lambda v, k=4: None if v is None else round(v, k)
            item = {"sym": sym, "name": name, "grp": grp, "note": note, "src": "live",
                    "d": [x[0] for x in rows[sl]], "o": [rd(x[1]) for x in rows[sl]], "h": [rd(x[2]) for x in rows[sl]],
                    "l": [rd(x[3]) for x in rows[sl]], "c": [rd(x[4]) for x in rows[sl]], "v": [int(x[5]) for x in rows[sl]],
                    "ma20": [rd(v) for v in m20[sl]], "ma50": [rd(v) for v in m50[sl]], "ma200": [rd(v) for v in m200[sl]], "rsi": [rd(v, 2) for v in rs[sl]],
                    "n_hist": len(rows), "contract": (j.get("meta", {}).get("longName") or j.get("meta", {}).get("shortName") or "")}
            ok_n += 1
        except Exception as e:
            out["warnings"].append(f"{sym} 抓取失败({type(e).__name__}: {str(e)[:50]}), " + ("使用该标的旧缓存" if sym in old else "且无缓存, 本图缺失"))
            if sym in old:
                item = dict(old[sym]); item["src"] = "cache"
        if item:
            out["items"].append(item)
        time.sleep(0.4)
    if ok_n == 0:
        out["status"] = "error" if not out["items"] else "cached"
        return out
    if out["warnings"]:
        out["status"] = "partial"
    _json.dump(out, open(cache_path, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    return out


URANIUM_TECH_SYMBOLS = [
    # (代码, 名称, 分组, 说明)
    ("SRUUF", "SPUT 实物铀信托(美元)", "铀本身", "Sprott Physical Uranium Trust 美股 OTC 价; 是信托份额二级市场价, 与 U3O8 现货有溢价/折价, 非现货价"),
    ("YCA.L", "Yellow Cake 实物铀(伦敦)", "铀本身", "Yellow Cake plc 伦敦主板, 便士计价; 持有实物铀, 二级市场价, 非 U3O8 现货价"),
    ("URA", "URA 铀与核能 ETF", "铀 ETF", "Global X Uranium ETF, 美元"),
    ("URNM", "URNM 铀矿 ETF", "铀 ETF", "Sprott Uranium Miners ETF, 美元"),
    ("NLR", "NLR 铀与核电 ETF", "铀 ETF", "VanEck Uranium and Nuclear ETF, 美元"),
    ("CCJ", "Cameco", "矿企", "加拿大最大铀生产商, 纽交所美元价"),
    ("NXE", "NexGen", "矿企", "Athabasca 盆地开发商(Rook I), 纽交所美元价"),
    ("KAP.IL", "Kazatomprom", "矿企", "哈萨克斯坦国家原子能公司, 伦敦 GDR 美元价"),
    ("PDN.AX", "Paladin", "矿企", "澳交所澳元价"),
    ("DNN", "Denison", "矿企", "美股美元价"),
    ("UUUU", "Energy Fuels", "矿企", "美股美元价"),
    ("UEC", "Uranium Energy", "矿企", "美股美元价"),
    ("BOE.AX", "Boss Energy", "矿企", "澳交所澳元价"),
    ("1164.HK", "中广核矿业", "矿企", "港股港元价"),
]


def fetch_uranium_tech(cache_path=None, keep=260):
    """铀短线图: 14 个标的日线 + SMA20/50/200 + RSI14(Wilder), 与金银短线图同口径。
    铀现货(U3O8)无免费可靠日线(Yahoo UX=F 仅 1 根), 用两只实物铀信托作代理并在页面注明。
    单个标的失败 → 用旧缓存并标「缓存」, 绝不用空覆盖旧真值。"""
    import requests as _rq, json as _json, time
    from datetime import timezone, timedelta
    root = os.path.join(os.path.dirname(__file__), "..", "data")
    cache_path = cache_path or os.path.join(root, "uranium_tech_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    try:
        old = {x["sym"]: x for x in _json.load(open(cache_path, encoding="utf-8")).get("items", [])}
    except Exception:
        old = {}
    out = {"status": "ok", "as_of": jst, "warnings": [], "items": []}
    H = {"User-Agent": "Mozilla/5.0"}
    t1 = int(datetime.datetime.utcnow().timestamp()) + 86400; t0 = t1 - 1250 * 86400
    ok_n = 0
    for sym, name, grp, note in URANIUM_TECH_SYMBOLS:
        item = None
        try:
            r = None
            for _ in range(2):
                r = _rq.get(f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}", params={"period1": t0, "period2": t1, "interval": "1d"}, headers=H, timeout=30)
                if r.status_code == 200:
                    break
                time.sleep(1.5)
            if r.status_code != 200:
                raise RuntimeError(f"HTTP {r.status_code}")
            j = r.json()["chart"]["result"][0]; q = j["indicators"]["quote"][0]
            rows = []
            for k, ts in enumerate(j["timestamp"]):
                o_, h_, l_, c_ = q["open"][k], q["high"][k], q["low"][k], q["close"][k]
                if None in (o_, h_, l_, c_):
                    continue
                rows.append((datetime.datetime.utcfromtimestamp(ts).strftime("%Y-%m-%d"), float(o_), float(h_), float(l_), float(c_), float(q["volume"][k] or 0)))
            dd = {}
            for x in rows:
                dd[x[0]] = x
            rows = [dd[k] for k in sorted(dd)]
            if len(rows) < 260:
                raise RuntimeError(f"仅 {len(rows)} 根日线")
            cl = [x[4] for x in rows]
            m20, m50, m200, rs = _gs_sma(cl, 20), _gs_sma(cl, 50), _gs_sma(cl, 200), _gs_rsi(cl, 14)
            sl = slice(len(rows) - keep, len(rows))
            rd = lambda v, k=4: None if v is None else round(v, k)
            item = {"sym": sym, "name": name, "grp": grp, "note": note, "src": "live",
                    "d": [x[0] for x in rows[sl]], "o": [rd(x[1]) for x in rows[sl]], "h": [rd(x[2]) for x in rows[sl]],
                    "l": [rd(x[3]) for x in rows[sl]], "c": [rd(x[4]) for x in rows[sl]], "v": [int(x[5]) for x in rows[sl]],
                    "ma20": [rd(v) for v in m20[sl]], "ma50": [rd(v) for v in m50[sl]], "ma200": [rd(v) for v in m200[sl]], "rsi": [rd(v, 2) for v in rs[sl]],
                    "n_hist": len(rows), "contract": (j.get("meta", {}).get("longName") or j.get("meta", {}).get("shortName") or "")}
            ok_n += 1
        except Exception as e:
            out["warnings"].append(f"{sym} 抓取失败({type(e).__name__}: {str(e)[:50]}), " + ("使用该标的旧缓存" if sym in old else "且无旧缓存, 该标的本次缺失"))
            if sym in old:
                item = dict(old[sym]); item["src"] = "cache"
        if item:
            out["items"].append(item)
        time.sleep(0.4)
    if ok_n == 0:
        out["status"] = "error" if not out["items"] else "cached"
        return out
    if out["warnings"]:
        out["status"] = "partial"
    _json.dump(out, open(cache_path, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    return out


def fetch_fed_gold_oil(cache_path=None):
    """联邦基金利率 × 金价 × 油价(月度, 1970 起)。
    利率: FRED FEDFUNDS(有效利率月均), 最新未发布月用 DFF 日度当月均值补。
    金价: datahub gold-prices(世界银行 Pink Sheet 月均, 与项目「金油比/股金比」底表同源, 逐月对过 668 个月 0 差异), 其后月份用 Yahoo GC=F 当月日均补。
    油价: FRED WTISPLC(WTI 现货月均; 1973 年前为管制标价, 页面注明), 最新未发布月用 DCOILWTICO 日度当月均值补。
    失败回退缓存, 绝不用空覆盖旧真值。"""
    import requests as _rq, json as _json, csv as _csv, io as _io
    from datetime import timezone, timedelta
    root = os.path.join(os.path.dirname(__file__), "..", "data")
    cache_path = cache_path or os.path.join(root, "fed_gold_oil_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "m": [], "ff": [], "gold": [], "oil": [], "src": {}}
    H = {"User-Agent": "Mozilla/5.0"}
    def _monthly_mean(daily):
        acc = {}
        for d, v in daily:
            acc.setdefault(d[:7], []).append(v)
        return {m: sum(x) / len(x) for m, x in acc.items()}
    try:
        ff = {d[:7]: v for d, v in _fred_series("FEDFUNDS", "1970-01-01")}
        if len(ff) < 600:
            raise RuntimeError(f"FEDFUNDS 仅 {len(ff)} 月")
        dff = _monthly_mean(_fred_series("DFF", "2026-01-01"))
        for m, v in dff.items():
            if m not in ff:
                ff[m] = v; out["src"]["ff_" + m] = "DFF 日度当月均值(FEDFUNDS 月值未发布)"
        oil = {d[:7]: v for d, v in _fred_series("WTISPLC", "1970-01-01")}
        if len(oil) < 600:
            raise RuntimeError(f"WTISPLC 仅 {len(oil)} 月")
        wti = _monthly_mean(_fred_series("DCOILWTICO", "2026-01-01"))
        for m, v in wti.items():
            if m not in oil:
                oil[m] = v; out["src"]["oil_" + m] = "DCOILWTICO 日度当月均值(WTISPLC 月值未发布)"
        t = _rq.get("https://raw.githubusercontent.com/datasets/gold-prices/master/data/monthly.csv", headers=H, timeout=40).text
        gold = {r["Date"]: float(r["Price"]) for r in _csv.DictReader(_io.StringIO(t)) if r["Date"] >= "1970-01"}
        if len(gold) < 600:
            raise RuntimeError(f"金价仅 {len(gold)} 月")
        last_g = max(gold)
        t0 = int(datetime.datetime(2026, 1, 1, tzinfo=timezone.utc).timestamp()); t1 = int(datetime.datetime.utcnow().timestamp()) + 86400
        r = _rq.get("https://query1.finance.yahoo.com/v8/finance/chart/GC=F", params={"period1": t0, "period2": t1, "interval": "1d"}, headers=H, timeout=40)
        j = r.json()["chart"]["result"][0]
        gd = [(datetime.datetime.utcfromtimestamp(x).strftime("%Y-%m-%d"), float(c)) for x, c in zip(j["timestamp"], j["indicators"]["quote"][0]["close"]) if c]
        for m, v in _monthly_mean(gd).items():
            if m > last_g:
                gold[m] = v; out["src"]["gold_" + m] = "Yahoo GC=F 当月日均(底表未发布)"
        months = sorted(set(ff) & set(gold) & set(oil))
        months = [m for m in months if m >= "1970-01"]
        if len(months) < 600:
            raise RuntimeError(f"三序列共同月仅 {len(months)}")
        out["m"] = months
        out["ff"] = [round(ff[m], 3) for m in months]
        out["gold"] = [round(gold[m], 2) for m in months]
        out["oil"] = [round(oil[m], 2) for m in months]
        _json.dump(out, open(cache_path, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    except Exception as e:
        try:
            c = _json.load(open(cache_path, encoding="utf-8"))
            c["status"] = "cached"; c["warnings"] = (c.get("warnings") or []) + [f"在线抓取失败({type(e).__name__}: {str(e)[:60]}), 使用缓存"]
            return c
        except Exception:
            out["status"] = "error"; out["warnings"].append(f"抓取失败且无缓存: {type(e).__name__}: {str(e)[:80]}")
    return out


def fetch_gold_volatility(cache_path=None):
    """黄金波动率(周线): Yahoo COMEX 黄金期货 GC=F 周线 OHLC + FRED GVZCLS(CBOE 黄金波动率指数, 日度→取每周最后值)。
    指标全部自算: 5 条简单均线(82/105/127/148/170 周)、布林带宽度(40,2)/(80,2)=(上轨-下轨)/中轨×100(总体σ)、ATR(14) Wilder。
    2026-10-09 实测对账(同一周): GVZ 与图上读数完全一致; 均线偏差一致 +0.5%(期货 vs 现货基差); BB 宽度差 1~3%; ATR 同量级。
    价格是期货不是现货, 页面写明。失败回退缓存, 绝不用空覆盖旧真值。"""
    import requests as _rq, json as _json, math as _m
    from datetime import timezone, timedelta
    root = os.path.join(os.path.dirname(__file__), "..", "data")
    cache_path = cache_path or os.path.join(root, "gold_vol_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "w": {}}
    try:
        t0 = int(datetime.datetime(2005, 1, 1, tzinfo=timezone.utc).timestamp()); t1 = int(datetime.datetime.utcnow().timestamp()) + 86400
        r = _rq.get("https://query1.finance.yahoo.com/v8/finance/chart/GC=F", params={"period1": t0, "period2": t1, "interval": "1wk"},
                    headers={"User-Agent": "Mozilla/5.0"}, timeout=40)
        r.raise_for_status()
        j = r.json()["chart"]["result"][0]; q = j["indicators"]["quote"][0]
        rows = [(datetime.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d"), q["open"][i], q["high"][i], q["low"][i], q["close"][i])
                for i, t in enumerate(j["timestamp"]) if q["close"][i] and q["high"][i] and q["low"][i] and q["open"][i]]
        if len(rows) < 600:
            raise RuntimeError(f"GC=F 周线仅 {len(rows)} 根")
        d = [x[0] for x in rows]; o = [x[1] for x in rows]; h = [x[2] for x in rows]; l = [x[3] for x in rows]; c = [x[4] for x in rows]
        n = len(c)
        def sma(k):
            return [None if i < k - 1 else sum(c[i - k + 1:i + 1]) / k for i in range(n)]
        def bbw(k, mult=2.0):
            res = []
            for i in range(n):
                if i < k - 1: res.append(None); continue
                seg = c[i - k + 1:i + 1]; m = sum(seg) / k; sd = _m.sqrt(sum((x - m) ** 2 for x in seg) / k)
                res.append(2 * mult * sd / m * 100)
            return res
        tr = [h[0] - l[0]] + [max(h[i] - l[i], abs(h[i] - c[i - 1]), abs(l[i] - c[i - 1])) for i in range(1, n)]
        atr = [None] * n; a = sum(tr[:14]) / 14; atr[13] = a
        for i in range(14, n):
            a = (a * 13 + tr[i]) / 14; atr[i] = a
        gv = {}
        for dd, v in _fred_series("GVZCLS", "2005-01-01"):
            gv[dd] = v
        gk = sorted(gv); gvz = []; p = 0; last = None
        for dd in d:                                   # 每根周线取"该周标记日之前(含)最近一个 GVZ 值", 不向前借用未来值
            while p < len(gk) and gk[p] <= dd:
                last = gv[gk[p]]; p += 1
            gvz.append(last)
        out["w"] = {"d": d, "o": o, "h": h, "l": l, "c": c, "ma": {str(k): sma(k) for k in (82, 105, 127, 148, 170)},
                    "bbw40": bbw(40), "bbw80": bbw(80), "atr": atr, "gvz": gvz, "last_bar": d[-1]}
        rd = lambda L: [None if v is None else round(v, 3) for v in L]
        out["w"]["ma"] = {k: rd(v) for k, v in out["w"]["ma"].items()}
        for k in ("bbw40", "bbw80", "atr"): out["w"][k] = rd(out["w"][k])
        out["w"]["gvz"] = rd(out["w"]["gvz"])
        out["w"]["o"], out["w"]["h"], out["w"]["l"], out["w"]["c"] = rd(o), rd(h), rd(l), rd(c)
        _json.dump(out, open(cache_path, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    except Exception as e:
        try:
            cc = _json.load(open(cache_path, encoding="utf-8"))
            cc["status"] = "cached"; cc["warnings"] = (cc.get("warnings") or []) + [f"在线抓取失败({type(e).__name__}: {str(e)[:60]}), 使用缓存"]
            return cc
        except Exception:
            out["status"] = "error"; out["warnings"].append(f"抓取失败且无缓存: {type(e).__name__}: {str(e)[:80]}")
    return out


def fetch_analog_corrections(cache_path=None):
    """金银「类比路径」: 把历史年份的日价格路径, 按【从该年 1 月 28 日起的相同日历天数偏移】叠到 2026 年 1/28 起点上。
    数据: Yahoo COMEX 主力连续合约(GC=F/SI=F)日收盘, 另用 datahub 月度金价补 1973(日度缺, 如实标月度)。
    不做归一化, 全部是原始美元价; 图上同一对齐规则写明(无法确认原图的对齐日, 是本看板的口径)。
    另: GDX÷SPY(Yahoo 日线相除)。失败回退缓存, 绝不用空覆盖旧真值。"""
    import requests as _rq, json as _json, csv as _csv, io as _io
    from datetime import timezone, timedelta
    root = os.path.join(os.path.dirname(__file__), "..", "data")
    cache_path = cache_path or os.path.join(root, "analog_corrections_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "gold": {}, "silver": {}, "gdx_spy": {}, "gold_1973_monthly": []}
    H = {"User-Agent": "Mozilla/5.0"}
    def _ts(d):
        return int(datetime.datetime.strptime(d, "%Y-%m-%d").replace(tzinfo=timezone.utc).timestamp())
    def _yh(sym, a, b):
        r = _rq.get(f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}",
                    params={"period1": _ts(a), "period2": _ts(b), "interval": "1d"}, headers=H, timeout=40)
        r.raise_for_status()
        j = r.json()["chart"]["result"][0]
        return [(datetime.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d"), float(v))
                for t, v in zip(j.get("timestamp") or [], j["indicators"]["quote"][0]["close"]) if v and v > 0]
    START = (1, 28)
    try:
        today = datetime.datetime.utcnow().strftime("%Y-%m-%d")
        spec = {"gold": [("GC=F", 2026, "2026"), ("GC=F", 2006, "2006"), ("GC=F", 2008, "2008")],
                "silver": [("SI=F", 2026, "2026"), ("SI=F", 2008, "2008"), ("SI=F", 2011, "2011")]}
        for metal, items in spec.items():
            for sym, yr, label in items:
                rows = _yh(sym, f"{yr}-01-01", (today if yr == 2026 else f"{yr + 1}-02-15"))
                base = datetime.date(yr, *START)
                ser = []
                for d, v in rows:
                    off = (datetime.date.fromisoformat(d) - base).days
                    if off >= 0:
                        ser.append([off, round(v, 3)])
                if len(ser) < 50:
                    raise RuntimeError(f"{sym} {yr} 仅 {len(ser)} 点")
                out[metal][label] = {"sym": sym, "start": str(base), "pts": ser, "last_date": rows[-1][0]}
        # GDX ÷ SPY(2006-06 起, 日线)
        g = dict(_yh("GDX", "2006-05-01", today)); sp = dict(_yh("SPY", "2006-05-01", today))
        days = sorted(set(g) & set(sp))
        out["gdx_spy"] = {"d": days, "v": [round(g[d] / sp[d], 5) for d in days]}
        if len(days) < 1000:
            raise RuntimeError(f"GDX/SPY 共同日仅 {len(days)}")
        # 1973 月度金价(datahub 公开 CSV, 月度)
        t = _rq.get("https://raw.githubusercontent.com/datasets/gold-prices/master/data/monthly.csv", headers=H, timeout=40).text
        mm = [[r["Date"], float(r["Price"])] for r in _csv.DictReader(_io.StringIO(t)) if r["Date"].startswith("1973")]
        if len(mm) == 12:
            out["gold_1973_monthly"] = mm
        else:
            out["warnings"].append(f"1973 月度金价仅 {len(mm)} 个月, 不画")
        _json.dump(out, open(cache_path, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    except Exception as e:
        try:
            c = _json.load(open(cache_path, encoding="utf-8"))
            c["status"] = "cached"; c["warnings"] = (c.get("warnings") or []) + [f"在线抓取失败({type(e).__name__}: {str(e)[:60]}), 使用缓存"]
            return c
        except Exception:
            out["status"] = "error"; out["warnings"].append(f"抓取失败且无缓存: {type(e).__name__}: {str(e)[:80]}")
    return out


def fetch_silver_curve_spread(cache_path=None):
    """白银期货曲线价差(Yahoo COMEX 各月合约收盘价, 非 CME 官方结算价, 非 EFP)。
    各远月合约相对"最近一个仍交易的月份"的升水(美元/盎司与百分比), 并按 3 个月 SOFR/国债利率估算持有成本(理论套利上限)对比。
    合约到期会从 Yahoo 消失(404) → 近月基准按"当日仍有报价的最近月份"滚动, 并在 rows 里记录 base 合约。"""
    import requests as _rq
    from datetime import timezone, timedelta
    root = os.path.join(os.path.dirname(__file__), "..", "data")
    cache_path = cache_path or os.path.join(root, "silver_curve_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "rows": [], "contracts": [], "today": {}, "carry": {}}
    MC = {"F": 1, "G": 2, "H": 3, "J": 4, "K": 5, "M": 6, "N": 7, "Q": 8, "U": 9, "V": 10, "X": 11, "Z": 12}
    now = datetime.datetime.utcnow()
    cands = []
    for yr in (now.year, now.year + 1, now.year + 2):
        for code, mo in MC.items():
            if (yr, mo) >= (now.year, now.month):
                cands.append((f"SI{code}{str(yr)[-2:]}.CMX", yr, mo, code))
    cands = cands[:14]
    try:
        series = {}
        for sym, yr, mo, code in cands:
            try:
                r = _rq.get(f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}?range=1y&interval=1d",
                            headers={"User-Agent": "Mozilla/5.0"}, timeout=25)
                if r.status_code != 200:
                    continue
                d = r.json()["chart"]["result"][0]
                ser = {datetime.datetime.utcfromtimestamp(t).strftime("%Y-%m-%d"): float(v)
                       for t, v in zip(d["timestamp"], d["indicators"]["quote"][0]["close"]) if v and v > 0}
                if len(ser) >= 60:
                    series[sym] = {"ym": f"{yr}-{mo:02d}", "label": f"{yr}年{mo}月", "ser": ser, "code": code, "yr": yr, "mo": mo}
            except Exception:
                continue
        if len(series) < 4:
            raise RuntimeError(f"仅取到 {len(series)} 个合约")
        syms = sorted(series, key=lambda k: series[k]["ym"])
        keep = syms[:9]                              # 最近 9 个有报价的月份合约
        days = sorted(set.intersection(*[set(series[k]["ser"]) for k in keep]))
        if len(days) < 60:
            raise RuntimeError(f"共同交易日仅 {len(days)}")
        # 近月基准 = 每日"仍交易的最近月份"; 近月到期后基准自动换下一个(记录在 base)
        rows = []
        for d in days:
            base = next(k for k in keep if d in series[k]["ser"])
            b = series[base]["ser"][d]
            rows.append({"d": d, "base": series[base]["ym"], "b": round(b, 3),
                         "px": {series[k]["ym"]: round(series[k]["ser"][d], 3) for k in keep}})
        # 利率(持有成本): 3M 国债收益率(FRED DGS3MO), 缺失用最近值
        rate = {}
        try:
            for dd, v in _fred_series("DGS3MO", days[0]):
                rate[dd] = v / 100.0
        except Exception as e:
            out["warnings"].append(f"利率(DGS3MO)获取失败({type(e).__name__}), 持有成本线不画")
        last_r = None; rr = []
        for dd in days:
            if dd in rate:
                last_r = rate[dd]
            rr.append(last_r)
        for i, row in enumerate(rows):
            row["r"] = rr[i]
        out["rows"] = rows
        out["contracts"] = [{"ym": series[k]["ym"], "label": series[k]["label"]} for k in keep]
        last = rows[-1]; prev30 = rows[max(0, len(rows) - 22)]
        out["today"] = {"d": last["d"], "base": last["base"], "b": last["b"], "px": last["px"], "r": last["r"],
                        "prev_d": prev30["d"], "prev_px": prev30["px"], "prev_base": prev30["base"], "prev_b": prev30["b"]}
        json.dump(out, open(cache_path, "w"), ensure_ascii=False)
    except Exception as e:
        out["warnings"].append(f"实时获取失败({type(e).__name__}: {str(e)[:60]})")
        try:
            c = json.load(open(cache_path))
            c["status"] = "ok(缓存)"; c["warnings"] = list(c.get("warnings", [])) + out["warnings"]
            return c
        except Exception:
            out["status"] = "error"
            return out
    return out


# ══════════════════════════════════════════════════════════════════
# 标普500 ÷ 黄金 · 1929 起长历史 + 均线 + 回撤(Chao 2026-10-16 看到 StockCharts 原图, 要求做)
#   ★独立于 fetch_stock_gold_ratio(1976+ 已验收面板), 不改它。
#   标普: datasets/s-and-p-500 (Shiller 月度, 本口径为【月均值】, 与 1976+ 面板的【月末收盘】不同 → 页面写明);
#   黄金: datasets/gold-prices 月度; 1934-01 前 20.67 美元、1934-1971 官价 35 美元 → 此段比值≈指数本身, 图上标灰。
#   最新月: 用 1976+ 面板的 Yahoo 月末收盘/当月金均值续接(标注), 不混用口径做均线之外的统计。
#   均线: 月度 40/78 期(≈原图周线 170/340 的月度近似, 不等同); 「跌破两均线」自动判定:
#     之前 12 个月内曾同时站上两线 → 此后连续 ≥12 个月同时低于两线 → 取起点, 间隔 ≥24 个月。
#   回撤: 标普月度序列的「峰→谷」跌幅(月均值口径, 比日内口径浅), 三段固定窗口。
def fetch_spx_gold_long(cache_path=None):
    import csv as _csv, io as _io, json as _j, urllib.request as _ur, os as _os
    from datetime import timezone, timedelta
    root = _os.path.join(_os.path.dirname(__file__), "..", "data")
    cache_path = cache_path or _os.path.join(root, "spx_gold_long_cache.json")
    jst = datetime.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "rows": []}
    def _get(u):
        req = _ur.Request(u, headers={"User-Agent": "Mozilla/5.0"})
        with _ur.urlopen(req, timeout=40) as r:
            return list(_csv.reader(_io.StringIO(r.read().decode("utf-8"))))[1:]
    try:
        sp = {r[0][:7]: float(r[1]) for r in _get("https://raw.githubusercontent.com/datasets/s-and-p-500/main/data/data.csv")}
        gd = {r[0][:7]: float(r[1]) for r in _get("https://raw.githubusercontent.com/datasets/gold-prices/main/data/monthly.csv")}
    except Exception as e:
        try:
            c = _j.load(open(cache_path)); c.setdefault("warnings", []).append("本次抓取失败(%s), 显示 %s 缓存" % (type(e).__name__, c.get("as_of")))
            c["status"] = "ok(cache)"; return c
        except Exception:
            out["status"] = "error"; out["warnings"].append("长历史抓取失败且无缓存: %s" % type(e).__name__); return out
    ms = sorted(m for m in sp if m in gd and m >= "1929-01")
    if len(ms) < 200:
        out["status"] = "error"; out["warnings"].append("重叠月份过少(%d)" % len(ms)); return out
    ser = [sp[m] / gd[m] for m in ms]
    def _sma(n): return [None] * (n - 1) + [round(sum(ser[i - n + 1:i + 1]) / n, 4) for i in range(n - 1, len(ser))]
    m40, m78 = _sma(40), _sma(78)
    ev = []; i = 78
    while i < len(ms) - 12:
        prior = any(ser[j] > m40[j] and ser[j] > m78[j] for j in range(max(78, i - 12), i))
        now = all(ser[j] < m40[j] and ser[j] < m78[j] for j in range(i, i + 12))
        edge = ser[i] < m40[i] and ser[i] < m78[i] and not (ser[i - 1] < m40[i - 1] and ser[i - 1] < m78[i - 1])
        if prior and now and edge:
            ev.append(ms[i]); i += 24
        else:
            i += 1
    streak = 0
    for j in range(len(ms) - 1, 77, -1):
        if ser[j] < m40[j] and ser[j] < m78[j]: streak += 1
        else: break
    def _dd(a, b):
        sub = [(m, sp[m]) for m in sorted(sp) if a <= m <= b]
        pk = max(sub, key=lambda x: x[1]); aft = [x for x in sub if x[0] >= pk[0]]; tr = min(aft, key=lambda x: x[1])
        return {"peak_m": pk[0], "trough_m": tr[0], "pct": round((tr[1] / pk[1] - 1) * 100, 1)}
    dds = [dict(_dd(a, b), label=l) for a, b, l in (("1929-01", "1933-12", "1929–32"), ("1972-06", "1975-06", "1973–74"), ("2000-01", "2003-06", "2000–02"))]
    out["rows"] = [{"m": m, "sg": round(ser[k], 5), "spx": sp[m], "gold": gd[m], "ma40": m40[k], "ma78": m78[k]} for k, m in enumerate(ms)]
    out["events"] = ev; out["streak"] = streak; out["drawdowns"] = dds
    out["last_m"] = ms[-1]; out["latest"] = round(ser[-1], 3)
    out["ma_latest"] = {"ma40": m40[-1], "ma78": m78[-1]}
    try:
        _j.dump(out, open(cache_path, "w"), ensure_ascii=False)
    except Exception:
        pass
    return out


# ══════════════════════════════════════════════════════════════════
# 美债三图(Chao 2026-10-16 看到 MacroMicro/Investing 拼图, 要求做 1/2/3; 第 4 张「美国主权 CDS」为付费数据未做)
#   ① 2Y/10Y/30Y 收益率 2019 起(FRED DGS2/DGS10/DGS30 日频→周度末值);
#   ② 10Y 期限溢价(NY Fed ACM, THREEFYTP10) + MOVE(Yahoo ^MOVE) + 10Y 收益率, 近 2 年: MOVE 用 agg="peak" 保尖峰;
#   ③ 彭博美国综合债券指数的【ETF 代理】: Yahoo AGG(主)与 BND(对照)日线 → 价格, 不是指数点位, 页面写明。
#   ACM 期限溢价每日发布但滞后数日(末值日期可能早于其他序列), 页面写各序列各自的末日。
def fetch_ust_trio(cache_path=None):
    import json as _j, os as _os, datetime as _dt, requests as _rq
    from datetime import timezone, timedelta
    root = _os.path.join(_os.path.dirname(__file__), "..", "data")
    cache_path = cache_path or _os.path.join(root, "ust_trio_cache.json")
    jst = _dt.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "yields": {}, "tpm": {}, "agg": {}}
    today = _dt.date.today()
    def _pk(pts, weekly=True, agg="last"):
        if not pts: return []
        wp = _weekly_resample(pts, agg=agg) if weekly else pts
        return [[d, round(v, 4)] for d, v in wp]
    # ① 收益率
    y_start = "2019-01-01"
    ys = {k: _fred_series_hist(sid, y_start) for k, sid in (("y2", "DGS2"), ("y10", "DGS10"), ("y30", "DGS30"))}
    for k, v in ys.items():
        if len(v) < 200: out["warnings"].append("收益率 %s 点数过少(%d)" % (k, len(v)))
    out["yields"] = {"y2": _pk(ys["y2"]), "y10": _pk(ys["y10"]), "y30": _pk(ys["y30"]),
                     "last": {k: ({"d": v[-1][0], "v": v[-1][1]} if v else None) for k, v in ys.items()}}
    # ② 期限溢价 + MOVE + 10Y (近 2 年)
    t_start = (today - _dt.timedelta(days=740)).isoformat()
    tp = _fred_series_hist("THREEFYTP10", t_start)
    mv = _fetch_move_history(t_start)
    y10 = [(d, v) for d, v in ys["y10"] if d >= t_start]
    out["tpm"] = {"tp": _pk(tp), "move": _pk(mv, agg="peak"), "y10": _pk(y10),
                  "last": {"tp": ({"d": tp[-1][0], "v": tp[-1][1]} if tp else None),
                           "move": ({"d": mv[-1][0], "v": mv[-1][1]} if mv else None),
                           "y10": ({"d": y10[-1][0], "v": y10[-1][1]} if y10 else None)}}
    if not tp: out["warnings"].append("期限溢价 THREEFYTP10 无数据")
    if not mv: out["warnings"].append("MOVE 无数据")
    # ③ AGG / BND 日线(Yahoo chart v8)
    def _yh(sym):
        try:
            r = _rq.get("https://query1.finance.yahoo.com/v8/finance/chart/" + sym, params={"interval": "1d", "range": "2y"},
                        headers={"User-Agent": "Mozilla/5.0"}, timeout=30)
            if r.status_code != 200: return []
            j = r.json()["chart"]["result"][0]; c = j["indicators"]["quote"][0]["close"]
            return [(_dt.datetime.fromtimestamp(t, timezone(timedelta(hours=9))).strftime("%Y-%m-%d"), v) for t, v in zip(j["timestamp"], c) if v is not None]
        except Exception:
            return []
    agg, bnd = _yh("AGG"), _yh("BND")
    if len(agg) < 200: out["warnings"].append("AGG 日线点数过少(%d)" % len(agg))
    out["agg"] = {"agg": [[d, round(v, 3)] for d, v in agg], "bnd": [[d, round(v, 3)] for d, v in bnd],
                  "last": {"agg": ({"d": agg[-1][0], "v": round(agg[-1][1], 3)} if agg else None),
                           "bnd": ({"d": bnd[-1][0], "v": round(bnd[-1][1], 3)} if bnd else None)}}
    ok = len(ys["y10"]) > 200 and len(agg) > 200
    if not ok:
        try:
            c = _j.load(open(cache_path)); c.setdefault("warnings", []).append("本次抓取不完整(%s), 显示 %s 缓存" % ("; ".join(out["warnings"][:2]), c.get("as_of")))
            c["status"] = "ok(cache)"; return c
        except Exception:
            out["status"] = "error"; return out
    try:
        _j.dump(out, open(cache_path, "w"), ensure_ascii=False)
    except Exception:
        pass
    return out


# ══════════════════════════════════════════════════════════════════
# 日元套息(Chao 2026-10-18 贴 MacroMicro 四宫格 + 周报文字摘要, 要求「做 + 每天类似的推导结论」, 接受用已实现波动率代替隐含波动率)
def fetch_jpy_carry(cache_path=None):
    import json as _j, os as _os, datetime as _dt, math as _m, requests as _rq
    from datetime import timezone, timedelta
    root = _os.path.join(_os.path.dirname(__file__), "..", "data")
    cache_path = cache_path or _os.path.join(root, "jpy_carry_cache.json")
    jst = _dt.datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M JST")
    out = {"status": "ok", "as_of": jst, "warnings": [], "cftc": {}, "spread": {}, "kline": {}, "facts": []}
    today = _dt.date.today(); H = {"User-Agent": "Mozilla/5.0"}
    # ── A. CFTC 日元(只取 CME JAPANESE YEN 主合约) ──
    def _cftc(ds, fields):
        rows, off = [], 0
        while off < 4000:
            r = _rq.get("https://publicreporting.cftc.gov/resource/%s.json" % ds, params={
                "$where": "market_and_exchange_names='JAPANESE YEN - CHICAGO MERCANTILE EXCHANGE' AND report_date_as_yyyy_mm_dd>='%s'" % ((today - _dt.timedelta(days=1100)).isoformat()),
                "$order": "report_date_as_yyyy_mm_dd ASC", "$limit": 1000, "$offset": off}, headers=H, timeout=45)
            if r.status_code != 200: break
            b = r.json()
            if not b: break
            rows += b; off += len(b)
            if len(b) < 1000: break
        out_rows = []
        for x in rows:
            try:
                d = x["report_date_as_yyyy_mm_dd"][:10]
                out_rows.append([d] + [float(x[f]) for f in fields])
            except Exception:
                continue
        return out_rows
    try:
        leg = _cftc("6dca-aqww", ["noncomm_positions_long_all", "noncomm_positions_short_all"])
        tff = _cftc("gpe5-46if", ["lev_money_positions_long", "lev_money_positions_short", "asset_mgr_positions_long", "asset_mgr_positions_short", "open_interest_all"])
    except Exception as e:
        leg, tff = [], []; out["warnings"].append("CFTC 日元取数失败: %s" % type(e).__name__)
    nc = [[r[0], r[1] - r[2]] for r in leg]
    lv = [[r[0], r[1] - r[2]] for r in tff]
    am = [[r[0], r[3] - r[4]] for r in tff]
    out["cftc"] = {"noncomm": nc, "lev": lv, "assetmgr": am,
                   "last": {"noncomm": ({"d": nc[-1][0], "v": nc[-1][1], "long": leg[-1][1], "short": leg[-1][2]} if nc else None),
                            "lev": ({"d": lv[-1][0], "v": lv[-1][1]} if lv else None)}}
    if len(nc) < 100: out["warnings"].append("CFTC 非商业点数过少(%d)" % len(nc))
    # ── B. 美日 10Y 利差 + USDJPY ──
    start = (today - _dt.timedelta(days=760)).isoformat()
    us10 = dict(_fred_series_hist("DGS10", start))
    try: jp10 = {k: v for k, v in _mof_jgb_yields(10).items() if k >= start}
    except Exception as e: jp10 = {}; out["warnings"].append("日债10Y(MOF)取数失败: %s" % type(e).__name__)
    def _yh(sym, rng="5y"):
        try:
            r = _rq.get("https://query1.finance.yahoo.com/v8/finance/chart/" + sym, params={"interval": "1d", "range": rng}, headers=H, timeout=40)
            if r.status_code != 200: return []
            j = r.json()["chart"]["result"][0]; q = j["indicators"]["quote"][0]
            res = []
            for i, t in enumerate(j["timestamp"]):
                o, h, l, c = q["open"][i], q["high"][i], q["low"][i], q["close"][i]
                if c is None or o is None: continue
                res.append((_dt.datetime.fromtimestamp(t, timezone.utc).strftime("%Y-%m-%d"), o, h, l, c))
            return res
        except Exception:
            return []
    fx = _yh("USDJPY=X")
    if len(fx) < 500: out["warnings"].append("USDJPY 日线点数过少(%d)" % len(fx))
    fxc = {d: c for d, o, h, l, c in fx}
    dts = sorted(set(us10) & set(jp10))
    spread = [(d, round(us10[d] - jp10[d], 3)) for d in dts]
    sp_w = _weekly_resample(spread) if spread else []
    fx_w = _weekly_resample([(d, fxc[d]) for d in sorted(fxc) if d >= start]) if fxc else []
    out["spread"] = {"spread": [[d, round(v, 3)] for d, v in sp_w], "usdjpy": [[d, round(v, 3)] for d, v in fx_w],
                     "us10": [[d, v] for d, v in _weekly_resample(sorted((d, v) for d, v in us10.items()))],
                     "jp10": [[d, v] for d, v in _weekly_resample(sorted(jp10.items()))],
                     "last": {"spread": ({"d": spread[-1][0], "v": spread[-1][1]} if spread else None),
                              "usdjpy": ({"d": fx[-1][0], "v": round(fx[-1][4], 3)} if fx else None),
                              "jp10": ({"d": max(jp10), "v": jp10[max(jp10)]} if jp10 else None),
                              "us10": ({"d": max(us10), "v": us10[max(us10)]} if us10 else None)}}
    if not spread: out["warnings"].append("美日利差无共同交易日")
    # ── C. USDJPY 周K + 均线 + 已实现波动率(20 日年化, 代理) ──
    wk = {}
    for d, o, h, l, c in fx:
        dt = _dt.date.fromisoformat(d); k = (dt - _dt.timedelta(days=dt.weekday())).isoformat()   # 周一为周标
        w = wk.setdefault(k, {"o": o, "h": h, "l": l, "c": c, "d": d})
        w["h"] = max(w["h"], h); w["l"] = min(w["l"], l); w["c"] = c; w["d"] = d
    wks = sorted(wk)[-130:]
    allw = sorted(wk)
    closes = [wk[k]["c"] for k in allw]
    def _ma(n, i): return round(sum(closes[i - n + 1:i + 1]) / n, 3) if i >= n - 1 else None
    kl = []
    for k in wks:
        i = allw.index(k)
        kl.append({"d": wk[k]["d"], "o": round(wk[k]["o"], 3), "h": round(wk[k]["h"], 3), "l": round(wk[k]["l"], 3), "c": round(wk[k]["c"], 3),
                   "ma5": _ma(5, i), "ma10": _ma(10, i), "ma20": _ma(20, i), "ma30": _ma(30, i)})
    rv = []
    cl = [(d, c) for d, o, h, l, c in fx]
    for i in range(20, len(cl)):
        rets = [_m.log(cl[j][1] / cl[j - 1][1]) for j in range(i - 19, i + 1)]
        mu = sum(rets) / 20; var = sum((x - mu) ** 2 for x in rets) / 19
        rv.append([cl[i][0], round(_m.sqrt(var) * _m.sqrt(252) * 100, 2)])
    rv_w = _weekly_resample(rv) if rv else []
    out["kline"] = {"k": kl, "rv": [[d, round(v, 2)] for d, v in rv_w if d >= (today - _dt.timedelta(days=760)).isoformat()],
                    "last_rv": ({"d": rv[-1][0], "v": rv[-1][1]} if rv else None)}
    # ── facts: 规则型事实, 每条带口径与日期; 不预测不建议 ──
    F = []
    def add(tag, txt): F.append({"tag": tag, "txt": txt})
    def _lag(d):
        """数据日 d(YYYY-MM-DD) 距今日(JST)的日历天数文案。"""
        try: n = (_dt.datetime.now(timezone(timedelta(hours=9))).date() - _dt.date.fromisoformat(str(d)[:10])).days
        except Exception: return ""
        return "滞后 %d 天" % n if n > 0 else "当日"
    if len(nc) >= 2:
        a, b = nc[-1], nc[-2]
        chg = (a[1] - b[1]) / abs(b[1]) * 100 if b[1] else None
        add("CFTC", "日元非商业净头寸（数据日 %s，%s）：%+d 手（上周 %+d，周变动 %+d 手%s）" % (a[0], _lag(a[0]), a[1], b[1], a[1] - b[1], ("，%+.0f%%" % chg) if chg is not None and abs(b[1]) >= 5000 else "，上周基数过小不算百分比"))
    if len(lv) >= 2:
        _t = [r for r in tff if len(r) >= 6]
        if len(_t) >= 2:
            _a, _b = _t[-2], _t[-1]; _oi = _b[5]
            add("CFTC", "杠杆基金（TFF）日元（数据日 %s，%s）：多头 %d→%d（%+d），空头 %d→%d（%+d），净 %+d 手（周变动 %+d）；空头占全市场持仓量 %.1f%%（上周 %.1f%%）；空头变动占持仓量 %+.2f%%"
                % (_b[0], _lag(_b[0]), _a[1], _b[1], _b[1] - _a[1], _a[2], _b[2], _b[2] - _a[2], lv[-1][1], lv[-1][1] - lv[-2][1],
                   100 * _b[2] / _oi if _oi else 0, 100 * _a[2] / _a[5] if _a[5] else 0, 100 * (_b[2] - _a[2]) / _oi if _oi else 0))
        else:
            add("CFTC", "杠杆基金（TFF）日元净头寸（数据日 %s，%s）：%+d 手（周变动 %+d）" % (lv[-1][0], _lag(lv[-1][0]), lv[-1][1], lv[-1][1] - lv[-2][1]))
    if jp10:
        dk = sorted(jp10); v, p = jp10[dk[-1]], jp10[dk[-2]]
        hi = max(jp10[k] for k in dk[-30:])
        add("日债", "日本 10 年国债收益率（数据日 %s，%s）：%.3f%%（前一交易日 %.3f%%；近 30 个交易日最高 %.3f%%）%s" % (dk[-1], _lag(dk[-1]), v, p, hi, ("，当前" + ("高于" if v > 3 else "低于") + " 3%")))
    if us10:
        dk = sorted(us10); v = us10[dk[-1]]
        add("美债", "美国 10 年国债收益率（数据日 %s，%s）：%.2f%%，%s 5%%" % (dk[-1], _lag(dk[-1]), v, "高于" if v > 5 else "低于"))
    if spread:
        a = spread[-1]; b = spread[-6] if len(spread) > 6 else spread[0]
        add("利差", "美日 10 年利差（数据日 %s，%s）：%.2f 个百分点（5 个交易日前 %.2f，%s）" % (a[0], _lag(a[0]), a[1], b[1], "走阔" if a[1] > b[1] else "收窄"))
    if fx:
        last = fx[-1][4]
        add("汇率", "美元/日元（数据日 %s，%s）：%.2f（距 160 关口 %.2f 日元；近 1 年高点 %.2f）" % (fx[-1][0], _lag(fx[-1][0]), last, 160 - last, max(x[2] for x in fx[-250:])))
    if rv:
        add("波动", "美元/日元 20 日已实现波动率（数据日 %s，%s）：%.1f%%（年化；近 1 年中位 %.1f%%）——是已实现波动率，不是隐含波动率" % (rv[-1][0], _lag(rv[-1][0]), rv[-1][1], sorted(x[1] for x in rv[-250:])[len(rv[-250:]) // 2]))
    out["facts"] = F
    ok = len(nc) >= 100 and len(fx) >= 500 and bool(spread)
    if not ok:
        try:
            c = _j.load(open(cache_path)); c.setdefault("warnings", []).append("本次抓取不完整(%s), 显示 %s 缓存" % ("; ".join(out["warnings"][:2]), c.get("as_of")))
            c["status"] = "ok(cache)"; return c
        except Exception:
            out["status"] = "error"; return out
    try: _j.dump(out, open(cache_path, "w"), ensure_ascii=False)
    except Exception: pass
    return out


def fetch_bbb_oas_3y(cache_path=None):
    """BBB 档公司债 OAS 近 3 年(FRED BAMLC0A4CBBB, 日频→周度末值, 百分点×100=bp)。Oracle CDS 图的行业背景线; 不是 Oracle 自己的利差。失败回退缓存, 绝不编。"""
    import json as _j, os as _os, datetime as _dt
    cache_path = cache_path or _os.path.join(_os.path.dirname(__file__), "..", "data", "bbb_oas_3y_cache.json")
    start = (_dt.date.today() - _dt.timedelta(days=3 * 365 + 20)).isoformat()
    pts = _fred_series_hist("BAMLC0A4CBBB", start)
    if len(pts) < 300:
        try:
            c = _j.load(open(cache_path)); c["status"] = "ok(cache)"; c.setdefault("warnings", []).append("本次抓取点数过少(%d), 显示缓存" % len(pts)); return c
        except Exception:
            return {"status": "error", "warnings": ["BBB OAS 取数失败且无缓存"], "series": []}
    wk = _weekly_resample(pts)
    out = {"status": "ok", "as_of": pts[-1][0], "warnings": [], "series": [[d, round(v * 100, 1)] for d, v in wk],
           "last": {"d": pts[-1][0], "bp": round(pts[-1][1] * 100, 1)}, "start": start}
    try: _j.dump(out, open(cache_path, "w"), ensure_ascii=False)
    except Exception: pass
    return out


def fetch_mm_ratio(cache_path=None, weeks_back=110):
    """白银/黄金 Managed Money 多空比 = 多头手数 / 空头手数(CFTC Disaggregated, 周度, COMEX, 排除 MICRO)。
    两个口径并存: Combined(期货+期权, 数据集 kh3c-gbw2) 与 FuturesOnly(72hh-3qpy)。价格取 data/cot_managed_money.json 里已有的日频连续合约(不重复抓)。
    Chao 2026-10-18「做」(对标 MacroMicro 白银多空比 vs 价格)。失败回退缓存, 绝不编。"""
    import json as _j, os as _os, datetime as _dt, requests as _rq
    root = _os.path.join(_os.path.dirname(__file__), "..")
    cache_path = cache_path or _os.path.join(root, "data", "mm_ratio_cache.json")
    base = "https://publicreporting.cftc.gov/resource/%s.json"
    start = (_dt.date.today() - _dt.timedelta(days=weeks_back * 7)).isoformat()
    names = {"silver": ("SILVER - COMMODITY EXCHANGE", "白银"), "gold": ("GOLD - COMMODITY EXCHANGE", "黄金")}
    def pull(ds, like):
        q = {"$where": "market_and_exchange_names like '%s%%' AND report_date_as_yyyy_mm_dd >= '%sT00:00:00.000'" % (like, start),
             "$order": "report_date_as_yyyy_mm_dd ASC", "$limit": 400}
        r = _rq.get(base % ds, params=q, headers={"User-Agent": "Mozilla/5.0"}, timeout=60); r.raise_for_status()
        out = {}
        for x in r.json():
            nm = x.get("market_and_exchange_names") or ""
            if "MICRO" in nm.upper(): continue
            d = x["report_date_as_yyyy_mm_dd"][:10]
            L = int(float(x.get("m_money_positions_long_all") or 0)); S = int(float(x.get("m_money_positions_short_all") or 0))
            out[d] = (L, S)
        return out
    res = {"status": "ok", "warnings": [], "metals": {}}
    try:
        px = _j.load(open(_os.path.join(root, "data", "cot_managed_money.json")))
    except Exception:
        px = {}
    for key, (like, zh) in names.items():
        try:
            cb = pull("kh3c-gbw2", like); fo = pull("72hh-3qpy", like)
        except Exception as ex:
            res["warnings"].append("%s CFTC 取数失败: %s" % (zh, type(ex).__name__)); continue
        ds = sorted(set(cb) & set(fo))
        if len(ds) < 20:
            res["warnings"].append("%s 周数过少(%d)" % (zh, len(ds))); continue
        rows = []
        for d in ds:
            Lc, Sc = cb[d]; Lf, Sf = fo[d]
            if Sc <= 0 or Sf <= 0: continue
            rows.append({"d": d, "cb_long": Lc, "cb_short": Sc, "cb_ratio": round(Lc / Sc, 3), "fo_long": Lf, "fo_short": Sf, "fo_ratio": round(Lf / Sf, 3)})
        gaps = [(rows[i]["d"], rows[i + 1]["d"]) for i in range(len(rows) - 1) if (_dt.date.fromisoformat(rows[i + 1]["d"]) - _dt.date.fromisoformat(rows[i]["d"])).days != 7]
        price = [p for p in ((px.get(key) or {}).get("price") or []) if p[0] >= rows[0]["d"]]
        res["metals"][key] = {"name": zh, "rows": rows, "price": price, "gaps": gaps,
                              "exchange": "COMEX", "price_sym": (px.get(key) or {}).get("symbol")}
    if not res["metals"]:
        try:
            c = _j.load(open(cache_path)); c["status"] = "ok(cache)"; c.setdefault("warnings", []).extend(res["warnings"]); return c
        except Exception:
            res["status"] = "error"; return res
    res["as_of"] = max(m["rows"][-1]["d"] for m in res["metals"].values())
    try: _j.dump(res, open(cache_path, "w"), ensure_ascii=False)
    except Exception: pass
    return res


def fetch_h8_deposits(cache_path=None, weeks_back=110):
    """美联储 H.8 周度存款(FRED): 总/大型国内银行/小型国内银行/外国相关机构 × SA/NSA, 十亿美元; 另 M2SL 月度。
    ID: DPSACBW027(总) DPSLCBW027(大) DPSSCBW027(小) DPSFRIW027(外国相关机构) + SBOG(SA)/NBOG(NSA)。H.8 周三发布(周三收盘余额, 以周三为周序列日期)。
    ★H.8 只含商业银行存款, 不含货币市场基金/国债等, 不等于「市场上的钱」。失败回退缓存, 绝不编。"""
    import json as _j, os as _os, datetime as _dt
    root = _os.path.join(_os.path.dirname(__file__), "..")
    cache_path = cache_path or _os.path.join(root, "data", "h8_deposits_cache.json")
    start = (_dt.date.today() - _dt.timedelta(days=weeks_back * 7 + 14)).isoformat()
    ids = {"tot": "DPSACBW027", "lg": "DPSLCBW027", "sm": "DPSSCBW027", "fr": "DPSFRIW027"}
    S = {}
    warns = []
    for k, base in ids.items():
        for sfx, tag in (("SBOG", "sa"), ("NBOG", "nsa")):
            try:
                S[(k, tag)] = dict(_fred_series_hist(base + sfx, start))
            except Exception as ex:
                warns.append("%s%s 取数失败: %s" % (base, sfx, type(ex).__name__)); S[(k, tag)] = {}
    dates = sorted(set(S[("tot", "sa")]) & set(S[("tot", "nsa")]))
    if len(dates) < 20 or any(len(v) < 20 for v in S.values()):
        try:
            c = _j.load(open(cache_path)); c["status"] = "ok(cache)"; c.setdefault("warnings", []).extend(warns + ["本次抓取不完整, 显示缓存"]); return c
        except Exception:
            return {"status": "error", "warnings": warns + ["H.8 取数失败且无缓存"], "rows": []}
    rows = []
    for d in dates:
        r = {"d": d}
        ok = True
        for k in ids:
            for tag in ("sa", "nsa"):
                v = S[(k, tag)].get(d)
                if v is None: ok = False; break
                r["%s_%s" % (k, tag)] = round(v, 1)
            if not ok: break
        if ok: rows.append(r)
    # 自洽检查: 大+小+外 == 总 (SA), 容差 0.5; 超差记警告而不改数
    bad = [r["d"] for r in rows if abs(r["lg_sa"] + r["sm_sa"] + r["fr_sa"] - r["tot_sa"]) > 0.5]
    if bad: warns.append("大+小+外≠总 的周(容差0.5): %s" % bad[:5])
    m2 = []
    try:
        m2 = [[d, round(v, 1)] for d, v in _fred_series_hist("M2SL", (_dt.date.today() - _dt.timedelta(days=240)).isoformat())][-6:]
    except Exception as ex:
        warns.append("M2SL 取数失败: %s" % type(ex).__name__)
    out = {"status": "ok", "as_of": rows[-1]["d"], "warnings": warns, "rows": rows, "m2": m2, "unit": "十亿美元"}
    try: _j.dump(out, open(cache_path, "w"), ensure_ascii=False)
    except Exception: pass
    return out


def fetch_hf_domiciles_ust(cache_path=None):
    """TIC SLT Table 3: 英国/开曼群岛/巴哈马 持有美国国债(百万美元, 月度, 2020-01 起)。
    ★口径: 按托管/登记地统计, 不是『对冲基金持有』; 三地合计≠全体对冲基金。只用 Holdings, 不用 n.a. 的净买卖。失败回退缓存, 绝不编。"""
    import json as _j, os as _os, re as _re, requests as _rq
    root = _os.path.join(_os.path.dirname(__file__), "..")
    cache_path = cache_path or _os.path.join(root, "data", "hf_domiciles_ust_cache.json")
    names = {"United Kingdom": "uk", "Cayman Islands": "cy", "Bahamas": "bs"}
    warns = []
    try:
        r = _rq.get(TIC_SLT_T3_URL, headers={"User-Agent": "Mozilla/5.0"}, timeout=60)
        if r.status_code != 200: raise RuntimeError("HTTP %s" % r.status_code)
        D = {"uk": {}, "cy": {}, "bs": {}}
        for l in r.text.splitlines():
            c = [x.strip() for x in l.split("\t")]
            if len(c) > 3 and c[0] in names and _re.match(r"\d{4}-\d{2}$", c[2]):
                try: D[names[c[0]]][c[2]] = int(c[3])
                except Exception: pass
        months = sorted(set(D["uk"]) & set(D["cy"]) & set(D["bs"]))
        months = [m for m in months if m >= "2020-01"]
        if len(months) < 24: raise RuntimeError("三国共有月份仅 %d" % len(months))
        rows = [{"m": m, "uk": D["uk"][m], "cy": D["cy"][m], "bs": D["bs"][m], "tot": D["uk"][m] + D["cy"][m] + D["bs"][m]} for m in months]
        out = {"status": "ok", "as_of": months[-1], "warnings": warns, "rows": rows, "unit": "百万美元"}
        try: _j.dump(out, open(cache_path, "w"), ensure_ascii=False)
        except Exception: pass
        return out
    except Exception as ex:
        try:
            c = _j.load(open(cache_path)); c["status"] = "ok(cache)"; c.setdefault("warnings", []).append("本次抓取失败(%s), 显示缓存" % type(ex).__name__); return c
        except Exception:
            return {"status": "error", "warnings": ["TIC Table 3 取数失败且无缓存: %s" % ex], "rows": []}
